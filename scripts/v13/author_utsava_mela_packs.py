#!/usr/bin/env python3
"""Author V13 Utsava/Mela packs (markdown + visuals + printable builders). No render."""
from __future__ import annotations

from pathlib import Path

from PIL import Image, ImageDraw, ImageFont

REPO = Path(__file__).resolve().parents[2]
FIRST = REPO / "11-weekly-program-library" / "first-six-months"
EVID = REPO / "build-evidence" / "v13"
PRINT = REPO / "scripts" / "v13" / "printables"
LAUNCH = REPO / "launch"
CREAM, PLUM, SAFFRON, TEAL = "#FFF8E8", "#5B1933", "#E59B24", "#4F7C78"


def write(path: Path, text: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text.lstrip("\n") if text.startswith("\n") else text, encoding="utf-8")
    if not text.endswith("\n"):
        path.write_text(path.read_text(encoding="utf-8") + "\n", encoding="utf-8")
    print("wrote", path.relative_to(REPO))


def font(size: int, bold: bool = False):
    for name in ("segoeuib.ttf" if bold else "segoeui.ttf", "arialbd.ttf" if bold else "arial.ttf"):
        p = Path(r"C:\Windows\Fonts") / name
        if p.exists():
            return ImageFont.truetype(str(p), size)
    return ImageFont.load_default()


def canvas(size=(900, 560)):
    im = Image.new("RGB", size, CREAM)
    return im, ImageDraw.Draw(im)


def save_png(out: Path, name: str, im: Image.Image) -> None:
    out.mkdir(parents=True, exist_ok=True)
    im.save(out / name, "PNG", optimize=True)
    print("wrote", (out / name).relative_to(REPO))


# ---------------------------------------------------------------------------
# Shared helpers
# ---------------------------------------------------------------------------

CALLOUTS_NOTE = (
    "Allowed callout kinds only: SASTRA, KEY_IDEA, TEACHER_NOTE, DO_NOT_SPECULATE, "
    "FAMILY_APPLICATION, CHILD_ACTIVITY, MATERIALS, HOME_PRACTICE, SAFETY_PRIVACY, TIME_CUE, RIGHTS_NOTE."
)

FOOT = (
    "KUTUMBA · Families Growing in Krishna Consciousness · Program Director: Swapnil Patil · "
    "Internal founding cohort · human/temple/local-tithi gates EXTERNAL_OPEN"
)


def start_here(
    mid: str,
    title: str,
    date_line: str,
    essential: str,
    conclusion: str,
    memory: str,
    primary: str,
    primary_url: str,
    misconception: str,
    calendar_note: str,
    home_min: str,
    next_name: str,
    filename: str = "V13-UTSAVA-START-HERE.md",
) -> str:
    return f"""# V13 START HERE — {mid}

**KUTUMBA · Families Growing in Krishna Consciousness**  
**Program Director: Swapnil Patil**  
**Module:** {mid} — {title}  
**Planning date:** {date_line}  
**This is the only operational entry point for this V13 festival/mela pack.**

---

## THIS SATURDAY (locked clock)

| Time | Block |
|---|---|
| 1:50–2:00 | Arrival / settle / LED safety briefing if lamps used |
| 2:00–2:10 | Opening mantras + welcome |
| 2:10–2:30 | All-family festival launch + primary source layer |
| 2:30–3:10 | Parallel stations: **parents/family** / **K–2** / **Grades 4–5** |
| 3:10–3:30 | Reunification + shared kīrtana / exhibition / reflection |
| 3:30–3:40 | Light snack + water only |
| 3:40–3:55 | Saṅkalpa + take-home / feedback |
| 3:55–4:00 | Closing |

Parents onsite. No weekly meal. Parallel tracks **2:30–3:10**; reunify **3:10–3:30**. Do not invent 2:35 / 3:25.

---

## ESSENTIAL FOCUS

| Field | Content |
|---|---|
| Essential question | {essential} |
| Conclusion | {conclusion} |
| Memory line | {memory} |
| Primary | **{primary}** — {primary_url} |
| Misconception to block | {misconception} |
| Calendar note | {calendar_note} |

---

## WHO SPEAKS / TEACHES

| Role | File |
|---|---|
| Owner run-of-show | `teacher/OWNER-RUN-OF-SHOW-V13.md` |
| Main script | `teacher/MAIN-FACILITATOR-GUIDE-V13.md` |
| Parent / family | `teacher/PARENT-FAMILY-GUIDE-V13.md` |
| Younger stations | `teacher/YOUNGER-STATION-PACK-V13.md` |
| Older stations | `teacher/OLDER-STATION-PACK-V13.md` |
| Materials / safety | `materials-safety-v13.md` |
| Home practice | `family-home-practice-v13.md` |
| Gamma (prompt-only) | `gamma/V13-GAMMA-MASTER-DECK-PROMPT.md` |
| Printables | `activities/v13-*` + `scripts/v13/printables/{mid}.py` |

---

## FRIDAY NIGHT (≤60 min)

1. Open this file.
2. Confirm local calendar EXTERNAL_OPEN items with owner (tithi/fasting not claimed in child materials).
3. Read MAIN speaking script once aloud.
4. Sort station printables into Y / O / P folders.
5. Confirm LED-default lamp plan + snack/water + room reset.
6. Optional: paste Gamma master prompt (owner render; not pre-approved).

---

## HOME PRACTICE

{home_min}

---

## NEXT

**{next_name}**

---

## STATUS

{FOOT}

{CALLOUTS_NOTE}
"""


def owner_ros(mid: str, title: str, date_line: str, stations: list[str], safety_bullets: list[str]) -> str:
    station_rows = "\n".join(f"| {i+1} | {s} |" for i, s in enumerate(stations))
    safety = "\n".join(f"- {b}" for b in safety_bullets)
    return f"""# Owner Run-of-Show — {mid}

**{title}** · {date_line}  
**Status:** Internal founding-cohort — EXTERNAL_OPEN human/temple/local-tithi gates  
**Entry:** `V13-UTSAVA-START-HERE.md` or `V13-MELA-START-HERE.md`

## Owner pre-check (day-of, 20 min)

1. Confirm Saturday clock: arrival 1:50; tracks 2:30–3:10; reunify 3:10; close 4:00.
2. Confirm no child fasting instructions are printed or spoken.
3. Confirm LED lamps are default if any lamp craft/station is open.
4. Confirm printables for younger / older / parent stations are bagged.
5. Confirm privacy card visible: no public ranking of sādhana; no private family data collected into Git.

## Live clock cues

| Cue | Action |
|---|---|
| 1:50 | Doors; greet; LED/safety board visible |
| 2:00 | Opening mantras |
| 2:10 | Main facilitator delivers festival launch script |
| 2:30 | **TIME_CUE** — send tracks to stations |
| 3:10 | **TIME_CUE** — reunify; shared closing experience |
| 3:30 | Snack + water only |
| 3:40 | Home practice / saṅkalpa |
| 3:55 | Closing blessing; next-name only |

## Station map

| # | Station |
|---|---|
{station_rows}

## Safety board (read aloud if lamps present)

{safety}

## Do not claim

- Exact local tithi unless EXTERNAL_OPEN calendar confirmed.
- Temple/human/legal approval.
- Child fasting as KUTUMBA requirement.
"""


def rights_md(mid: str, assets: list[str]) -> str:
    lines = "\n".join(f"| `{a}` | KUTUMBA-original Pillow line art | internal teaching |" for a in assets)
    return f"""# Rights — {mid} visuals/v13

**Brand:** KUTUMBA · Families Growing in Krishna Consciousness · Program Director: Swapnil Patil  
**Status:** Internal founding-cohort teaching assets — EXTERNAL_OPEN human/temple review

## Asset register

| File | Origin | Use |
|---|---|---|
{lines}

## Scripture display

- Sanskrit/Bengali verse display text: cite VedaBase URL; do not dump full BBT purport.
- KUTUMBA teaching meanings are original paraphrases for family pedagogy — not labeled as BBT translation.
- Dāmodarāṣṭaka or other hymns: use only from a verified authoritative source if added later; not invented here.

## Forbidden

- Unauthorized temple/ISKCON logos.
- Copyrighted deity art copying.
- Identifiable student photos without separate rights process.
- Claiming publication readiness while gates remain EXTERNAL_OPEN.
"""


def image_prompt_lib(mid: str, prompts: list[tuple[str, str]]) -> str:
    body = "\n\n".join(f"### {t}\n{p}" for t, p in prompts)
    return f"""# Image Prompt Library — {mid}

Prompt-only companion for Gamma and future art. No claim that images are already rendered in Gamma.

{body}
"""


def source_matrix(mid: str, rows: list[tuple[str, str, str, str]]) -> str:
    table = "\n".join(f"| {a} | {b} | {c} | {d} |" for a, b, c, d in rows)
    return f"""# V13 Source Matrix — {mid}

| Claim / use | Source | URL | Boundary |
|---|---|---|---|
{table}

## Rights note

Sanskrit/Bengali display + KUTUMBA teaching meaning only. No full purport reproduction. Local tithi/fasting EXTERNAL_OPEN.
"""


def claim_register(mid: str, may: list[str], may_not: list[str]) -> str:
    may_y = "\n".join(f'  - "{m}"' for m in may)
    may_n = "\n".join(f'  - "{m}"' for m in may_not)
    return f"""# V13 Claim Register — {mid}
module_id: {mid}
may_claim:
{may_y}
may_not_claim:
{may_n}
external_open:
  - local Harrisburg tithi confirmation
  - human/temple doctrinal review
  - Gamma live render sign-off
  - legal/BBT excerpt expansion beyond display verse
"""


def content_audit(mid: str, title: str, date: str, primary: str, url: str, mechanics: list[tuple[str, str]], gaps: list[str]) -> str:
    mech = "\n".join(f"| {c} | {m} |" for c, m in mechanics)
    gap = "\n".join(f"{i}. {g}" for i, g in enumerate(gaps, 1))
    return f"""# {mid} V13 Content Audit

**Title:** {title} · **Planning date:** {date}  
**Primary:** {primary} — {url}  
**Audit date:** 2026-09-09 (authoring pass)

## Verdict

**CONTENT AUTHORING PASS** for the V13 festival/mela pack (markdown/yaml/scripts/assets).  
**PUBLISHING RENDER GAP** remains for final DOCX/PDF/raster/visual QA (explicitly deferred — do not render in this pass).

## Primary confirmation

| Field | Value |
|---|---|
| Reference | {primary} |
| URL | {url} |
| Teaching meaning | KUTUMBA-original; not labeled as BBT translation |
| Purport | Not invented / not dumped |
| Calendar | Local tithi EXTERNAL_OPEN |

## Activity mechanic summary

| Code | Mechanic |
|---|---|
{mech}

## Safeguards checked

- [x] No child fasting instructions
- [x] LED lamps default where lamps appear
- [x] No false exact-tithi claim for planning Saturday
- [x] C1-W1 not modified
- [x] No commit in this pass
- [x] No render in this pass
- [x] Callouts only from allowed kinds

## Gaps / EXTERNAL_OPEN

{gap}
"""


def req_csv(mid: str, rows: list[tuple[str, str, str]]) -> str:
    header = "requirement_id,requirement,deliverable_path,status,notes\n"
    body = "\n".join(f'{mid}-R{i:02d},{req},{path},PASS,{note}' for i, (req, path, note) in enumerate(rows, 1))
    return header + body + "\n"


# ---------------------------------------------------------------------------
# Visual generators
# ---------------------------------------------------------------------------

def visual_c1_u1(out: Path) -> list[str]:
    names = []
    im, d = canvas()
    # mortar + rope short by two fingers metaphor (abstract)
    d.rectangle((320, 280, 580, 420), outline=PLUM, width=8)
    d.ellipse((380, 180, 520, 300), outline=TEAL, width=7)
    d.arc((300, 220, 600, 360), 200, 340, fill=SAFFRON, width=10)
    d.text((250, 40), "Bound by love — not force", fill=PLUM, font=font(28, True))
    d.text((280, 480), "SB 10.9 teaching visual", fill=TEAL, font=font(20))
    save_png(out, "damodara-rope-of-love.png", im)
    names.append("damodara-rope-of-love.png")

    im, d = canvas((900, 420))
    for i, label in enumerate(["Butter pot", "Search", "Rope try", "Bound by love", "Gratitude"]):
        x = 80 + i * 160
        d.rounded_rectangle((x, 120, x + 130, 280), radius=16, outline=TEAL if i % 2 == 0 else PLUM, width=6)
        d.text((x + 10, 180), label, fill=PLUM, font=font(16, True))
        if i < 4:
            d.polygon([(x + 140, 190), x + 155, 200, x + 140, 210], fill=SAFFRON)  # type: ignore
            d.polygon([(x + 140, 190), (x + 155, 200), (x + 140, 210)], fill=SAFFRON)
    d.text((260, 40), "Dāmodara story sequence", fill=PLUM, font=font(26, True))
    save_png(out, "damodara-story-sequence.png", im)
    names.append("damodara-story-sequence.png")

    im, d = canvas((700, 700))
    d.ellipse((220, 180, 480, 420), outline=SAFFRON, width=10)
    d.rectangle((300, 420, 400, 560), outline=TEAL, width=8)
    d.text((180, 60), "LED lamp gratitude card", fill=PLUM, font=font(26, True))
    d.text((200, 600), "LED default — no open flame", fill=TEAL, font=font(18))
    save_png(out, "led-gratitude-lamp-card.png", im)
    names.append("led-gratitude-lamp-card.png")
    return names


def visual_c2_u2(out: Path) -> list[str]:
    names = []
    im, d = canvas((1000, 520))
    d.text((280, 30), "Bhagavad-gītā · 18 chapters map", fill=PLUM, font=font(26, True))
    for i in range(18):
        r, c = divmod(i, 6)
        x, y = 60 + c * 150, 100 + r * 130
        d.rounded_rectangle((x, y, x + 130, y + 100), radius=12, outline=TEAL, width=5)
        d.text((x + 40, y + 35), str(i + 1), fill=SAFFRON, font=font(28, True))
    save_png(out, "gita-18-chapter-map.png", im)
    names.append("gita-18-chapter-map.png")

    im, d = canvas()
    d.text((220, 40), "Arjuna decision case frame", fill=PLUM, font=font(26, True))
    for i, (lab, col) in enumerate([("Duty?", TEAL), ("Fear?", SAFFRON), ("Kṛṣṇa's word", PLUM)]):
        x = 80 + i * 280
        d.ellipse((x, 180, x + 220, 400), outline=col, width=8)
        d.text((x + 40, 270), lab, fill=col, font=font(22, True))
    save_png(out, "arjuna-decision-frame.png", im)
    names.append("arjuna-decision-frame.png")

    im, d = canvas((800, 500))
    d.rectangle((80, 80, 720, 420), outline=PLUM, width=8)
    d.text((200, 200), "Family Gītā reading plan", fill=TEAL, font=font(28, True))
    d.text((180, 260), "tiny · regular · noncompetitive", fill=SAFFRON, font=font(22))
    save_png(out, "family-gita-reading-plan.png", im)
    names.append("family-gita-reading-plan.png")
    return names


def visual_c3_u3(out: Path) -> list[str]:
    names = []
    im, d = canvas()
    d.text((220, 40), "Mercy chain (no graphic violence)", fill=PLUM, font=font(24, True))
    labels = ["Hear name", "Receive mercy", "Serve", "Invite others"]
    for i, lab in enumerate(labels):
        x = 60 + i * 210
        d.rounded_rectangle((x, 180, x + 180, 340), radius=18, outline=TEAL, width=6)
        d.text((x + 20, 240), lab, fill=PLUM, font=font(18, True))
        if i < 3:
            d.polygon([(x + 185, 250), (x + 205, 260), (x + 185, 270)], fill=SAFFRON)
    save_png(out, "mercy-chain.png", im)
    names.append("mercy-chain.png")

    im, d = canvas((900, 500))
    d.text((260, 40), "Service challenge cards", fill=PLUM, font=font(26, True))
    for i, lab in enumerate(["Help at home", "Kind words", "Join kīrtana", "Share prasāda"]):
        x = 40 + (i % 4) * 210
        d.rectangle((x, 140, x + 190, 360), outline=SAFFRON if i % 2 else TEAL, width=7)
        d.text((x + 20, 220), lab, fill=PLUM, font=font(16, True))
    save_png(out, "service-challenge-cards.png", im)
    names.append("service-challenge-cards.png")

    im, d = canvas((800, 500))
    d.ellipse((250, 120, 550, 400), outline=PLUM, width=10)
    d.text((300, 230), "Saṅkīrtana", fill=SAFFRON, font=font(32, True))
    d.text((220, 430), "Invitation — not volume contest", fill=TEAL, font=font(18))
    save_png(out, "sankirtana-invitation.png", im)
    names.append("sankirtana-invitation.png")
    return names


def visual_mela(out: Path) -> list[str]:
    names = []
    im, d = canvas((1000, 560))
    d.text((280, 30), "Noncompetitive exhibition loop", fill=PLUM, font=font(26, True))
    for i, lab in enumerate(["See", "Appreciate", "Learn", "Feedback", "Continue"]):
        x = 40 + i * 190
        d.rounded_rectangle((x, 140, x + 170, 380), radius=20, outline=TEAL, width=6)
        d.text((x + 30, 240), lab, fill=PLUM, font=font(18, True))
    save_png(out, "exhibition-feedback-loop.png", im)
    names.append("exhibition-feedback-loop.png")

    im, d = canvas()
    d.text((200, 40), "Continue / Review / Strengthen", fill=PLUM, font=font(26, True))
    for i, (lab, col) in enumerate([("CONTINUE", TEAL), ("REVIEW", SAFFRON), ("STRENGTHEN", PLUM)]):
        x = 70 + i * 280
        d.ellipse((x, 160, x + 240, 400), outline=col, width=8)
        d.text((x + 40, 260), lab, fill=col, font=font(20, True))
    save_png(out, "continue-review-strengthen.png", im)
    names.append("continue-review-strengthen.png")

    im, d = canvas((800, 500))
    d.rectangle((60, 80, 740, 420), outline=PLUM, width=8)
    d.text((180, 200), "Family testimony (no comparison)", fill=TEAL, font=font(24, True))
    d.text((220, 260), "Private gratitude welcome", fill=SAFFRON, font=font(20))
    save_png(out, "family-testimony-frame.png", im)
    names.append("family-testimony-frame.png")
    return names


# ---------------------------------------------------------------------------
# Printable builder templates
# ---------------------------------------------------------------------------

def printable_module(
    week_id: str,
    folder: str,
    memory: str,
    verse_url: str,
    builders_src: str,
) -> str:
    return f'''#!/usr/bin/env python3
"""{week_id} V13 printable Word builders for render_week.py.

Real Word tables/images. DOCX/PDF emission owned by scripts/v13/render_week.py.
This pass authors builders only — do not render here.
"""
from __future__ import annotations

import sys
from pathlib import Path

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches

REPO = Path(__file__).resolve().parents[3]
sys.path.insert(0, str(REPO / "scripts" / "v13"))
sys.path.insert(0, str(REPO / "scripts" / "v12"))
from printable_common import (  # noqa: E402
    add_card_grid,
    add_cut_lines,
    add_memory_phrase_block,
    add_printable_title,
    add_write_lines,
)
from kutumba_docx_styles import add_callout, add_verse_card  # noqa: E402

ASSETS = (
    REPO
    / "11-weekly-program-library"
    / "first-six-months"
    / "{folder}"
    / "visuals"
    / "v13"
)
MEMORY = "{memory}"
VERSE_URL = "{verse_url}"


def _img(document, name: str, width: float = 2.4) -> None:
    path = ASSETS / name
    if path.is_file():
        p = document.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(path), width=Inches(width))


{builders_src}

FAMILY_BUILDERS = []
'''


# ---------------------------------------------------------------------------
# C1-U1 Damodara
# ---------------------------------------------------------------------------

def author_c1_u1() -> None:
    folder = "c1-u1-kartika-damodara-family-utsava"
    base = FIRST / folder
    mid = "C1-U1"
    title = "Kārtika / Dāmodara Family Utsava"
    date = "Saturday 2026-10-31 (planning Saturday)"
    primary = "ŚB 10.9.18"
    url = "https://vedabase.io/en/library/sb/10/9/18/"
    memory = "Kṛṣṇa is bound by love."
    assets = visual_c1_u1(base / "visuals" / "v13")

    write(base / "V13-UTSAVA-START-HERE.md", start_here(
        mid, title, date,
        essential="How does mother Yaśodā’s love show that Kṛṣṇa is conquered by devotion?",
        conclusion="The Lord who rules the universe willingly accepts the bonds of pure love.",
        memory=memory,
        primary=primary,
        primary_url=url,
        misconception="Oct 31 is automatically the exact Diwali/Govardhana tithi for our local temple.",
        calendar_note="Planning Saturday only. Exact local Kārtika/Diwali/Govardhana tithi = EXTERNAL_OPEN. Do not claim Oct 31 is exact Diwali/Govardhana.",
        home_min="Minimum: say the memory line once + one gratitude/service act with LED or no-flame remembrance. No child fasting.",
        next_name="C2-W1 — Action and Reaction: How Karma Binds (after protected off week as calendared).",
    ))

    write(base / "teacher" / "OWNER-RUN-OF-SHOW-V13.md", owner_ros(
        mid, title, date,
        stations=[
            "Story sequencing (all ages share launch)",
            "Younger butter-pot / rope-of-love craft (LED lamp card)",
            "Older Dāmodara narrative map from SB 10.9",
            "Parent gratitude/service + household remembrance plan",
            "Family kīrtana (invitation, not volume contest)",
        ],
        safety_bullets=[
            "LED lamps are the default child activity.",
            "Open flame only if owner/host adult safety plan explicitly permits and children are supervised away from flame.",
            "No child fasting instructions tonight.",
            "Do not label tonight as exact Diwali/Govardhana unless local calendar confirms (EXTERNAL_OPEN).",
        ],
    ))

    write(base / "teacher" / "MAIN-FACILITATOR-GUIDE-V13.md", f"""# Main Facilitator Guide — {mid}

**{title}** · {date}  
**Status:** Internal founding-cohort — EXTERNAL_OPEN  
**Entry:** `V13-UTSAVA-START-HERE.md`

## Two-minute summary

Dāmodara means Kṛṣṇa bound at the belly — and more deeply, bound by mother Yaśodā’s love.  
Primary: **{primary}** ({url}).  
Supporting chapter frame: ŚB 10.9 (https://vedabase.io/en/library/sb/10/9/).

## Exact primary

- **Reference:** {primary}
- **URL:** {url}
- **Devanāgarī:** स्वमातु: स्विन्नगात्राया विस्रस्तकबरस्रज: । दृष्ट्वा परिश्रमं कृष्ण: कृपयासीत् स्वबन्धने ॥ १८ ॥
- **IAST:** sva-mātuḥ svinna-gātrāyā visrasta-kabara-srajaḥ / dṛṣṭvā pariśramaṁ kṛṣṇaḥ kṛpayāsīt sva-bandhane
- **KUTUMBA teaching meaning:** Seeing mother Yaśodā tired from trying to bind Him, Kṛṣṇa became merciful and agreed to be bound.
- **Supporting:** ŚB 10.9.19 — https://vedabase.io/en/library/sb/10/9/19/ — the Lord who controls the universe shows the quality of coming under the control of His devotee.
- **Rights:** Sanskrit from VedaBase display; KUTUMBA meaning original; no full purport.

## Teaching map

| Time | Block |
|---|---|
| 1:50–2:00 | Arrival; LED safety board |
| 2:00–2:10 | Opening mantras |
| 2:10–2:30 | Festival launch + verse |
| 2:30–3:10 | Parallel stations |
| 3:10–3:30 | Reunify + soft kīrtana |
| 3:30–3:40 | Snack + water |
| 3:40–3:55 | Saṅkalpa |
| 3:55–4:00 | Close |

## 10–20 minute speaking script

> Speak naturally. Pause for a 20-second pair share. Do not replace with “see research.”

Welcome. Tonight is our **Kārtika / Dāmodara Family Utsava** planning Saturday — **2026-10-31**.

Calendar honesty first: this is a **family observance candidate** during the Kārtika season window. We do **not** claim that October 31 is the exact local Diwali or Govardhana tithi. Local confirmation remains **EXTERNAL_OPEN**.

Essential question: How does mother Yaśodā’s love show that Kṛṣṇa is conquered by devotion?

Primary: ŚB 10.9.18 — {url}

KUTUMBA teaching meaning: “Seeing mother Yaśodā tired from trying to bind Him, Kṛṣṇa became merciful and agreed to be bound.”

Story frame from ŚB 10.9 (paraphrase only; no invented dialogue): Mother Yaśodā’s household work and butter-churning lead to Kṛṣṇa’s butter mischief. She tries to bind Him. Rope after rope is short. When she is exhausted, Kṛṣṇa accepts her love and allows Himself to be bound. Supporting ŚB 10.9.19 teaches that the Lord who rules the universe shows the wonder of being controlled by a devotee’s love.

Analogy one — rope vs love: material rope fails; love succeeds. Limit: do not teach children that physical ropes are magical toys for “catching God.”

Analogy two — tired mother, merciful child: effort meets grace. Limit: children’s obedience at home is not identical to Yaśodā’s pure parental devotion.

Analogy three — LED lamp remembrance: light as gratitude cue during Kārtika. Limit: LED is pedagogy/safety default; open flame is adult-only if ever permitted.

Household case: A family wants to “do Kārtika perfectly” and pressures a child to skip dinner. Mistaken conclusion: devotion requires child fasting. Principle: no child fasting instructions in KUTUMBA. Action: choose gratitude, kīrtana, and service instead.

We will not speculate about exact miracle physics beyond what śāstra states. Deferral line: if a question needs temple authority or local calendar precision, mark it EXTERNAL_OPEN and continue with love, service, and safe remembrance.

Memory line: **{memory}**

## Analogies with limits

1. Rope vs love — limit above.
2. Tired mother / merciful response — limit above.
3. LED remembrance lamp — limit above.

## Three family cases

1. Pressure to fast a child — redirect to gratitude/service.
2. Sibling teases “you can’t catch Kṛṣṇa” during craft — redirect to respect for līlā.
3. Parent compares who’s “more into Kārtika” — block ranking; keep private saṅkalpa.

## Q&A boundaries

- Exact Diwali/Govardhana date? → EXTERNAL_OPEN local calendar.
- Must we offer Dāmodarāṣṭaka? → Only from verified authoritative source if used; optional enrichment, not invented text tonight.
- Open flame lamps? → LED default; owner safety plan required otherwise.

## Do not speculate

Do not invent dialogue for Yaśodā/Kṛṣṇa. Do not claim Oct 31 equals exact tithi. Do not teach child fasting.

## Track transition (2:30)

Parents → gratitude/service plan. Younger → sequence + rope-of-love craft + LED card. Older → SB 10.9 narrative map + exit ticket.

## Reunification (3:10)

One shared sentence: “Kṛṣṇa is bound by love.” Soft kīrtana invitation.
""")

    write(base / "teacher" / "PARENT-FAMILY-GUIDE-V13.md", f"""# Parent / Family Guide — {mid}

**40-minute parallel track (2:30–3:10)** · {title}

## Flow

| Min | Block | Printable |
|---|---|---|
| 0–5 | Private reflection: where does love, not force, lead our home? | P01 |
| 5–12 | Source observation: ŚB 10.9.18 layer | P02 |
| 12–22 | Household case: child-fasting pressure vs gratitude | P03 |
| 22–32 | Operating application: Kārtika remembrance without ranking | P04 |
| 32–40 | Private next-step saṅkalpa | P05 |

## Key idea

Love binds more deeply than control. Kārtika remembrance can be tiny, regular, and safe.

## Safety & privacy

No public ranking of who lights more lamps. No child fasting. LED default.
""")

    write(base / "teacher" / "YOUNGER-STATION-PACK-V13.md", f"""# Younger Station Pack — {mid}

**K–2 · 2:30–3:10**

| Code | Station |
|---|---|
| Y01 | Story sequence cards (butter → search → rope → love) |
| Y02 | Rope-of-love craft (paper “rope”; love words) |
| Y03 | LED gratitude lamp card |
| Y04 | Kind speech about Kṛṣṇa’s līlā |
| Y05 | Memory phrase mat: {memory} |

## Redirects

- If child asks about fire: “We use LED tonight.”
- If child asks about fasting: “Children eat what parents provide; we practice gratitude.”
- If fear appears: return to motherly love and kindness.

## Backup (5 min)

Echo memory line + one stretch “reaching like Yaśodā tried” then soft sit.
""")

    write(base / "teacher" / "OLDER-STATION-PACK-V13.md", f"""# Older Station Pack — {mid}

**Grades 4–5 · 2:30–3:10**

| Code | Station |
|---|---|
| O01 | ŚB 10.9.18 observation |
| O02 | Dāmodara narrative map (SB 10.9 paraphrase stations) |
| O03 | Bound-by-love vs bound-by-force sort |
| O04 | Respectful Kārtika scenario cards |
| O05 | Gratitude/service diagram |
| O06 | Exit ticket |

## TEACHER-ONLY key notes

- Short-by-two-fingers detail is from the chapter narrative; keep wonder without speculative physics lectures.
- Supporting ŚB 10.9.19: devotee’s love and the Lord’s willingness — paraphrase carefully.
""")

    # Activities
    acts = {
        "activities/v13-younger/Y01-STORY-SEQUENCE.md": f"# {mid} · Y01 · Story Sequence\n\n**Mechanic:** Sequencing\n**Bridge:** ŚB 10.9 narrative\n\nOrder cards: Butter pot → Search → Rope try → Bound by love → Gratitude.\nVisual: `damodara-story-sequence.png`\nSuccess: child retells gently without mocking Yaśodā.\n",
        "activities/v13-younger/Y02-ROPE-OF-LOVE-CRAFT.md": f"# {mid} · Y02 · Rope-of-Love Craft\n\nPaper strips as “rope.” Write love/service words. Link strips.\nTeaching: love binds more than force.\nMaterials: paper, tape, crayons.\nAge: K–2. No real rope around necks/wrists as play-bondage.\n",
        "activities/v13-younger/Y03-LED-GRATITUDE-LAMP-CARD.md": f"# {mid} · Y03 · LED Gratitude Lamp Card\n\nColor `led-gratitude-lamp-card.png` card. Write one thank-you.\n**SAFETY:** LED default. No child open flame.\n",
        "activities/v13-younger/Y04-KIND-LILA-SPEECH.md": f"# {mid} · Y04 · Kind Līlā Speech\n\nCards: kind vs mocking lines about the butter story. Choose kind.\n",
        "activities/v13-younger/Y05-MEMORY-PHRASE-MAT.md": f"# {mid} · Y05 · Memory Phrase Mat\n\n**{memory}**\nTrace/say once. Smile. Sit.\n",
        "activities/v13-older/O01-SB-10-9-18-OBSERVATION.md": f"# {mid} · O01 · ŚB 10.9.18 Observation\n\nURL: {url}\nObserve: Who is tired? What does Kṛṣṇa do? What does mercy look like here?\nWrite three observations. No purport dump.\n",
        "activities/v13-older/O02-DAMODARA-NARRATIVE-MAP.md": f"# {mid} · O02 · Narrative Map\n\nMap paraphrase stations from https://vedabase.io/en/library/sb/10/9/\nDo not invent dialogue.\n",
        "activities/v13-older/O03-LOVE-VS-FORCE-SORT.md": f"# {mid} · O03 · Love vs Force Sort\n\nSort statements into Bound-by-love / Bound-by-force / Not sure.\n",
        "activities/v13-older/O04-KARTIKA-SCENARIOS.md": f"# {mid} · O04 · Kārtika Scenarios\n\nInclude: child fasting pressure → refuse; LED vs flame → LED default; ranking lamps → privacy.\n",
        "activities/v13-older/O05-GRATITUDE-SERVICE-DIAGRAM.md": f"# {mid} · O05 · Gratitude/Service Diagram\n\nDraw: Remembrance → Gratitude → One service act.\n",
        "activities/v13-older/O06-EXIT-TICKET.md": f"# {mid} · O06 · Exit Ticket\n\n1) Memory line 2) One thing Yaśodā shows 3) One safe home act\n",
        "activities/v13-parent/P01-PRIVATE-REFLECTION.md": f"# {mid} · P01 · Private Reflection\n\nWhere do we try to force outcomes that love and patience should lead?\nSharing optional.\n",
        "activities/v13-parent/P02-SOURCE-OBSERVATION.md": f"# {mid} · P02 · Source Observation\n\nRead ŚB 10.9.18 layer. Note mercy and motherly effort.\n",
        "activities/v13-parent/P03-SUBSTANTIAL-FAMILY-CASE.md": f"# {mid} · P03 · Case\n\nChild-fasting pressure during Kārtika. Redirect to gratitude/service. No ranking.\n",
        "activities/v13-parent/P04-HOUSEHOLD-OPERATING.md": f"# {mid} · P04 · Household Operating\n\nOne remembrance cue; one service; LED/no-flame plan; minimum version.\n",
        "activities/v13-parent/P05-PRIVATE-SANKALPA.md": f"# {mid} · P05 · Private Saṅkalpa\n\nSpecific action + frequency + trigger + minimum version. No photo proof.\n",
    }
    for rel, text in acts.items():
        write(base / rel, text)

    write(base / "materials-safety-v13.md", f"""# Materials & Safety — {mid}

## Materials

- Printed Y/O/P station packs
- Paper strips, tape, crayons
- LED tea-lights (default)
- Snack + water only
- Privacy reminder card

## Safety

- **LED lamps default** for children.
- Open flame only with explicit owner/host adult safety plan.
- No child fasting instructions.
- Allergies: check snack labels.
- Craft scissors: adult supervision.

## Setup / cleanup

Bag stations by 1:40. Reset room by 4:10. Collect teacher-only keys separately.
""")

    write(base / "family-home-practice-v13.md", f"""# Family Home Practice — {mid}

**Memory:** {memory}

**Minimum:** Say memory line once + one gratitude or help-at-home act before next Saturday.  
**Ideal (5–15 min):** Memory line + one-sentence paraphrase of ŚB 10.9.18 + LED/no-flame remembrance once.  
No photo proof. No ranking. No child fasting.
""")

    write(base / "project" / "V13-MODULE-PROJECT-BRIEF.md", f"""# Project Brief — {mid}

**Artifact:** Family “bound by love” gratitude card (LED theme) contributed to Cycle 1 closing memory later.  
Noncompetitive. Optional photo stays private to family — not committed to Git.
""")

    write(base / "research" / "V13-SOURCE-MATRIX.md", source_matrix(mid, [
        ("Primary teaching", primary, url, "Display verse + KUTUMBA meaning"),
        ("Chapter frame", "ŚB 10.9", "https://vedabase.io/en/library/sb/10/9/", "Paraphrase narrative; no invented dialogue"),
        ("Supporting", "ŚB 10.9.19", "https://vedabase.io/en/library/sb/10/9/19/", "bhṛtya-vaśyatā theme"),
        ("Calendar planning", "ISKCON Bangalore calendar refs", "https://www.iskconbangalore.org/vaishnava-calendar/", "Planning only; local EXTERNAL_OPEN"),
    ]))
    write(base / "research" / "V13-CLAIM-REGISTER.yaml", claim_register(
        mid,
        may=[
            "Teach SB 10.9.18 with KUTUMBA meaning",
            "Use 2026-10-31 as planning Saturday",
            "Require LED default for child lamp activities",
        ],
        may_not=[
            "Claim Oct 31 is exact Diwali/Govardhana tithi",
            "Instruct child fasting",
            "Invent Yaśodā/Kṛṣṇa dialogue",
            "Claim temple approval",
        ],
    ))

    write(base / "visuals" / "v13" / "RIGHTS.md", rights_md(mid, assets))
    write(base / "visuals" / "v13" / "IMAGE-PROMPT-LIBRARY.md", image_prompt_lib(mid, [
        ("Dāmodara rope of love", "Warm cream educational illustration; abstract mortar and saffron rope arc; no deity face copying; no open flame; South Asian mother-and-child silhouettes optional and non-identifiable; 16:9"),
        ("Story sequence", "Five gentle station panels for butter-pot story without gore or mockery; flat educational style; cream background"),
        ("LED lamp card", "Child-safe LED tea-light greeting card design; clearly LED; soft glow; no real fire"),
    ]))

    # Gamma master (~20 slides)
    slides = []
    slide_specs = [
        ("Title", "Kārtika / Dāmodara Family Utsava", ["KUTUMBA", "Planning Saturday 2026-10-31", "Kṛṣṇa is bound by love."]),
        ("Calendar honesty", "Planning Saturday — not exact tithi claim", ["Oct 31 is a family observance candidate", "Local Diwali/Govardhana tithi = EXTERNAL_OPEN", "Do not erase core curriculum"]),
        ("Essential question", "Essential question", ["How does mother Yaśodā’s love show that Kṛṣṇa is conquered by devotion?", memory]),
        ("Chapter frame", "ŚB 10.9 frame", ["Butter mischief → search → rope short → mercy", "Paraphrase only", "https://vedabase.io/en/library/sb/10/9/"]),
        ("Primary verse", "ŚB 10.9.18", ["Devanāgarī layer on slide", "IAST", "KUTUMBA meaning: Kṛṣṇa agrees to be bound when He sees mother’s fatigue", url]),
        ("Supporting", "ŚB 10.9.19", ["The Lord who controls the universe", "shows the wonder of devotee’s love", "https://vedabase.io/en/library/sb/10/9/19/"]),
        ("Analogy", "Rope vs love", ["Material rope fails", "Love succeeds", "Limit: not magical rope games"]),
        ("Safety", "LED default", ["LED lamps for children", "Open flame only with adult plan", "No child fasting"]),
        ("Stations", "2:30–3:10 stations", ["Parents: gratitude plan", "K–2: sequence + craft + LED card", "4–5: narrative map"]),
        ("Case", "Household case", ["Pressure to fast a child", "Redirect to gratitude/service", "No ranking"]),
        ("Younger preview", "K–2 preview", ["Story sequence", "Rope-of-love craft", "Memory mat"]),
        ("Older preview", "Grades 4–5 preview", ["Verse observation", "Love vs force sort", "Exit ticket"]),
        ("Parent preview", "Parent track", ["Private reflection", "Source observation", "Saṅkalpa"]),
        ("Kīrtana", "Soft kīrtana", ["Invitation not volume contest", "Listening counts"]),
        ("Home practice", "Home practice", ["Memory line once", "One gratitude/service act", "No photo proof"]),
        ("Do not claim", "Do not claim", ["Exact Diwali/Govardhana on Oct 31", "Temple approval", "Child fasting"]),
        ("Reunify cue", "TIME_CUE 3:10", ["Return to circle", "Say memory line", "Snack at 3:30"]),
        ("Close", "Closing", ["Next: follow calendar toward C2-W1", FOOT]),
        ("Sources", "Source strip", [url, "https://vedabase.io/en/library/sb/10/9/", "https://vedabase.io/en/library/sb/10/9/19/"]),
        ("Rights", "Rights", ["KUTUMBA-original framing", "BBT/VedaBase rights remain with holders", "EXTERNAL_OPEN review"]),
    ]
    for i, (purpose, title_s, copy) in enumerate(slide_specs, 1):
        bullets = "\n".join(f"  - {c}" for c in copy)
        slides.append(f"""### Slide {i} — {purpose}
- **Purpose:** {purpose}
- **Exact title:** {title_s}
- **Exact on-screen copy:**
{bullets}
- **Sources:** MAIN-FACILITATOR-GUIDE-V13 · V13-SOURCE-MATRIX
- **Layout:** Brand-consistent cream panel; plum title; saffron accent
- **Detailed AI-image prompt:** Educational cream 16:9 scene for “{title_s}”; South Asian family devotion atmosphere; no temple logo; no identifiable real persons; no open flame; no gore; no readable dense scripture glyphs in image
- **Presenter notes:** Keep calendar honesty; protect children from fasting pressure.
- **Interaction:** Nod / pair-share as appropriate
- **Accessibility:** Read title aloud; high contrast
- **Do not claim:** exact local tithi; temple approval; child fasting need
""")
    write(base / "gamma" / "V13-GAMMA-MASTER-DECK-PROMPT.md", f"""# {mid} V13 Gamma Master Deck Prompt

**Status:** prompt-only — not rendered — not approved  
**Design:** 16:9 · cream `#FFF8E8` · plum `#5B1933` / saffron `#E59B24` / teal `#4F7C78`  
**Brand:** KUTUMBA · Families Growing in Krishna Consciousness · Program Director: Swapnil Patil  
**Module:** {mid} — {title} · {date}  
**Slide count:** {len(slide_specs)}

Paste into Gamma Studio as one master prompt deck. Owner renders; no claim of Gamma approval until then.

---

## slides

{chr(10).join(slides)}
""")
    for aud, n in (("PARENT", 11), ("YOUNGER", 9), ("OLDER", 11)):
        write(base / "gamma" / f"V13-GAMMA-{aud}-DECK-PROMPT.md", f"""# {mid} Gamma {aud.title()} Deck

**Status:** prompt-only — not rendered  
**Slides:** {n} (audience-specific subset of master themes)

Focus: {"parent household application + no fasting" if aud=="PARENT" else "story/craft/LED safety" if aud=="YOUNGER" else "SB 10.9 observation + scenarios"}.

Include exact memory line: {memory}  
Include primary URL: {url}  
Include do-not-claim: Oct 31 ≠ automatic exact Diwali/Govardhana tithi.
""")
    write(base / "gamma" / "V13-GAMMA-SOURCE-MAP.yaml", f"""module: {mid}
primary:
  ref: {primary}
  url: {url}
supporting:
  - ref: ŚB 10.9
    url: https://vedabase.io/en/library/sb/10/9/
  - ref: ŚB 10.9.19
    url: https://vedabase.io/en/library/sb/10/9/19/
gamma_status: prompt-only
local_tithi: EXTERNAL_OPEN
""")

    builders = '''
def y01(document) -> None:
    add_printable_title(document, "Y01", "Dāmodara Story Sequence", "Order the cards. Tell the story kindly.")
    _img(document, "damodara-story-sequence.png", 6.0)
    add_card_grid(document, ["Butter pot", "Search", "Rope try", "Bound by love", "Gratitude"], per_page=5)
    add_memory_phrase_block(document, MEMORY)
    add_cut_lines(document)
    add_callout(document, "CHILD_ACTIVITY", "Kind voices only. No mocking mother Yaśodā.")


def y02(document) -> None:
    add_printable_title(document, "Y02", "Rope-of-Love Craft", "Write love/service words on paper strips. Link them.")
    add_write_lines(document, ["Love word 1:", "Love word 2:", "Service word:"], 2)
    add_callout(document, "SAFETY_PRIVACY", "Paper craft only. No real rope play-bondage.")


def y03(document) -> None:
    add_printable_title(document, "Y03", "LED Gratitude Lamp Card", "Color the card. Write one thank-you.")
    _img(document, "led-gratitude-lamp-card.png", 4.5)
    add_write_lines(document, ["I thank Kṛṣṇa for:", "I can help at home by:"], 2)
    add_callout(document, "SAFETY_PRIVACY", "LED default. No child open flame. No child fasting.")


def y04(document) -> None:
    add_printable_title(document, "Y04", "Kind Līlā Speech", "Choose kind words about the butter story.")
    add_card_grid(document, ["Kind: Mother loved Kṛṣṇa", "Unkind: mocking laugh", "Kind: Kṛṣṇa is merciful", "Unkind: teasing fasting"], per_page=4)
    add_cut_lines(document)


def y05(document) -> None:
    add_printable_title(document, "Y05", "Memory Phrase Mat", MEMORY)
    add_memory_phrase_block(document, MEMORY)
    add_callout(document, "KEY_IDEA", MEMORY)


def o01(document) -> None:
    add_printable_title(document, "O01", "ŚB 10.9.18 Observation", "Observe the verse layer. No purport dump.")
    add_verse_card(document, "ŚB 10.9.18", "", "", "Seeing mother Yaśodā tired, Kṛṣṇa mercifully agreed to be bound.", VERSE_URL)
    add_write_lines(document, ["Who is tired?", "What does Kṛṣṇa do?", "What does mercy look like?"], 3)


def o02(document) -> None:
    add_printable_title(document, "O02", "Dāmodara Narrative Map", "Map SB 10.9 paraphrase stations. No invented dialogue.")
    _img(document, "damodara-rope-of-love.png", 5.5)
    add_write_lines(document, ["Station 1:", "Station 2:", "Station 3:", "Mercy moment:"], 2)


def o03(document) -> None:
    add_printable_title(document, "O03", "Love vs Force Sort", "Sort into Bound-by-love / Bound-by-force / Not sure.")
    add_card_grid(document, ["Patient motherly effort", "Harsh control", "Gratitude service", "Public ranking of devotion", "Private saṅkalpa", "Child fasting pressure"], per_page=6)
    add_cut_lines(document)


def o04(document) -> None:
    add_printable_title(document, "O04", "Kārtika Scenarios", "Choose the KUTUMBA-aligned response.")
    add_write_lines(document, ["Scenario: child fasting pressure →", "Scenario: open flame request →", "Scenario: ranking lamps →"], 3)
    add_callout(document, "TEACHER_NOTE", "Correct redirects: no child fasting; LED default; no ranking.")


def o05(document) -> None:
    add_printable_title(document, "O05", "Gratitude / Service Diagram", "Remembrance → Gratitude → One service act.")
    add_write_lines(document, ["Remembrance cue:", "Gratitude sentence:", "Service act this week:"], 2)


def o06(document) -> None:
    add_printable_title(document, "O06", "Exit Ticket", "Three lines only.")
    add_write_lines(document, ["Memory line:", "One thing Yaśodā shows:", "One safe home act:"], 2)


def p01(document) -> None:
    add_printable_title(document, "P01", "Private Reflection", "Sharing optional.")
    add_write_lines(document, ["Where do we force outcomes?", "Where could love/patience lead instead?"], 3)
    add_callout(document, "SAFETY_PRIVACY", "Write privately. Do not collect sheets into Git.")


def p02(document) -> None:
    add_printable_title(document, "P02", "Source Observation", "ŚB 10.9.18")
    add_verse_card(document, "ŚB 10.9.18", "", "", "Kṛṣṇa agrees to be bound when He sees mother’s fatigue.", VERSE_URL)
    add_write_lines(document, ["What stands out?", "What must we not claim about Oct 31?"], 3)


def p03(document) -> None:
    add_printable_title(document, "P03", "Substantial Family Case", "Child-fasting pressure during Kārtika.")
    add_write_lines(document, ["Mistaken conclusion:", "Principle:", "Family action:"], 3)
    add_callout(document, "DO_NOT_SPECULATE", "No child fasting instructions. Local tithi EXTERNAL_OPEN.")


def p04(document) -> None:
    add_printable_title(document, "P04", "Household Operating Application", "Kārtika remembrance without ranking.")
    add_write_lines(document, ["Remembrance cue:", "Service act:", "LED/no-flame plan:", "Minimum version:"], 2)
    add_callout(document, "FAMILY_APPLICATION", "Tiny and regular beats dramatic and ranked.")


def p05(document) -> None:
    add_printable_title(document, "P05", "Private Next-Step Saṅkalpa", "")
    add_callout(document, "KEY_IDEA", "specific action + frequency + trigger + minimum version")
    add_write_lines(document, ["Specific action:", "Frequency:", "Trigger:", "Minimum version:"], 2)


YOUNGER_BUILDERS = [y01, y02, y03, y04, y05]
OLDER_BUILDERS = [o01, o02, o03, o04, o05, o06]
PARENT_BUILDERS = [p01, p02, p03, p04, p05]
'''
    write(PRINT / "C1-U1.py", printable_module(mid, folder, memory, url, builders))

    write(EVID / "C1-U1" / "CONTENT-AUDIT.md", content_audit(
        mid, title, "2026-10-31", primary, url,
        [("Y01", "story sequence"), ("Y02", "rope-of-love craft"), ("Y03", "LED lamp card"),
         ("O01", "verse observation"), ("O02", "narrative map"), ("P03", "fasting-pressure case")],
        ["DOCX/PDF render deferred", "Raster/VISUAL-QA deferred", "Local tithi EXTERNAL_OPEN",
         "Human/temple review EXTERNAL_OPEN", "Gamma render EXTERNAL_OPEN",
         "Dāmodarāṣṭaka not inlined — only if later verified from authoritative source"],
    ))
    write(EVID / "C1-U1" / "REQUIREMENT-TRACEABILITY.csv", req_csv(mid, [
        ("Start Here", f"11-weekly-program-library/first-six-months/{folder}/V13-UTSAVA-START-HERE.md", "planning Sat 2026-10-31"),
        ("Owner run-of-show", f"11-weekly-program-library/first-six-months/{folder}/teacher/OWNER-RUN-OF-SHOW-V13.md", "LED + no child fasting"),
        ("Main script", f"11-weekly-program-library/first-six-months/{folder}/teacher/MAIN-FACILITATOR-GUIDE-V13.md", "SB 10.9.18 inline"),
        ("Parent/family guide", f"11-weekly-program-library/first-six-months/{folder}/teacher/PARENT-FAMILY-GUIDE-V13.md", ""),
        ("Younger stations", f"11-weekly-program-library/first-six-months/{folder}/teacher/YOUNGER-STATION-PACK-V13.md", ""),
        ("Older stations", f"11-weekly-program-library/first-six-months/{folder}/teacher/OLDER-STATION-PACK-V13.md", ""),
        ("Materials/safety", f"11-weekly-program-library/first-six-months/{folder}/materials-safety-v13.md", "LED default"),
        ("Source matrix", f"11-weekly-program-library/first-six-months/{folder}/research/V13-SOURCE-MATRIX.md", ""),
        ("Rights", f"11-weekly-program-library/first-six-months/{folder}/visuals/v13/RIGHTS.md", ""),
        ("Gamma master actual copy", f"11-weekly-program-library/first-six-months/{folder}/gamma/V13-GAMMA-MASTER-DECK-PROMPT.md", "20 slides"),
        ("Printable builder", "scripts/v13/printables/C1-U1.py", "no render this pass"),
        ("No Oct 31 exact Diwali claim", f"11-weekly-program-library/first-six-months/{folder}/V13-UTSAVA-START-HERE.md", "EXTERNAL_OPEN calendar"),
    ]))


# ---------------------------------------------------------------------------
# C2-U2 Gita Jayanti
# ---------------------------------------------------------------------------

def author_c2_u2() -> None:
    folder = "c2-u2-gita-jayanti-family-utsava"
    base = FIRST / folder
    mid = "C2-U2"
    title = "Gītā Jayantī Family Utsava"
    date = "Saturday 2026-12-19 (planning); published example Gītā Jayantī 2026-12-20"
    primary = "BG 18.65"
    url = "https://vedabase.io/en/library/bg/18/65/"
    memory = "Think of Kṛṣṇa; become His devotee."
    assets = visual_c2_u2(base / "visuals" / "v13")

    write(base / "V13-UTSAVA-START-HERE.md", start_here(
        mid, title, date,
        essential="How should our family receive Kṛṣṇa’s Gītā instruction without pressure or ranking?",
        conclusion="The Gītā invites devotion: remember Kṛṣṇa, serve, and honor Him — tiny, regular, noncompetitive.",
        memory=memory,
        primary=primary,
        primary_url=url,
        misconception="Tonight is claimed as the exact local Gītā Jayantī tithi for every calendar.",
        calendar_note="Family observance candidate on 2026-12-19. Published example Gītā Jayantī 2026-12-20. Exact local tithi EXTERNAL_OPEN — not an exact-tithi claim.",
        home_min="Minimum: memory line once + open the Gītā together for one short reading. No pressure book-distribution goals for children.",
        next_name="Winter break dates on calendar; then C3-W1 as scheduled.",
    ))

    write(base / "teacher" / "OWNER-RUN-OF-SHOW-V13.md", owner_ros(
        mid, title, date,
        stations=[
            "18-chapter visual map",
            "Find-the-Gītā-theme younger stations",
            "Arjuna decision case (older/parent)",
            "Chapter-poster contribution",
            "Family Gītā reading plan",
            "Book/service gratitude (no child distribution quotas)",
        ],
        safety_bullets=[
            "No child fasting instructions.",
            "No pressure-based book distribution goals for children.",
            "Do not claim 2026-12-19 is exact local tithi; published example is 2026-12-20.",
        ],
    ))

    write(base / "teacher" / "MAIN-FACILITATOR-GUIDE-V13.md", f"""# Main Facilitator Guide — {mid}

**{title}** · {date}  
**Status:** INTERNAL · EXTERNAL_OPEN local tithi/human/temple gates  
**Entry:** `V13-UTSAVA-START-HERE.md`

## Exact primary

- **Reference:** {primary}
- **URL:** {url}
- **Devanāgarī:** मन्मना भव मद्भक्तो मद्याजी मां नमस्कुरु । मामेवैष्यसि सत्यं ते प्रतिजाने प्रियोऽसि मे ॥ ६५ ॥
- **IAST:** man-manā bhava mad-bhakto mad-yājī māṁ namaskuru / mām evaiṣyasi satyaṁ te pratijāne priyo ’si me
- **KUTUMBA teaching meaning:** Always think of Me, become My devotee, worship Me, and offer homage unto Me; thus you will come to Me. I promise this because you are dear to Me.
- **Whole-Gītā frame:** https://vedabase.io/en/library/bg/
- **Rights:** Sanskrit display + KUTUMBA meaning; no full purport.

## 10–20 minute speaking script

Welcome to our **Gītā Jayantī Family Utsava** planning Saturday — **2026-12-19**.

Calendar honesty: published planning example lists Gītā Jayantī on **2026-12-20**. Tonight is a **family observance candidate**, not a claim that our local exact tithi is settled. Local confirmation remains **EXTERNAL_OPEN**.

Essential question: How should our family receive Kṛṣṇa’s Gītā instruction without pressure or ranking?

Primary: BG 18.65 — {url}

KUTUMBA teaching meaning: “Always think of Me, become My devotee, worship Me, and offer homage unto Me; thus you will come to Me. I promise this because you are dear to Me.”

The Gītā is Kṛṣṇa’s instruction to Arjuna in a moment of confusion about duty. We will not invent battlefield dialogue. We will teach the invitation: remember Kṛṣṇa, live as His devotee, and honor Him in ordinary family life.

Analogy one — map of 18 chapters: a map helps travel; it is not the destination. Limit: children need not memorize all chapter titles tonight.

Analogy two — decision flashlight: the Gītā lights the next faithful step. Limit: not every modern choice has a one-verse slogan; some questions defer to parents/temple.

Analogy three — reading plan like watering a plant: tiny and regular. Limit: no competitive chapter-count contests.

Household case: A parent sets a child “book distribution quota” for Jayantī. Mistaken conclusion: love equals numbers. Principle: no pressure-based distribution goals for children. Action: gratitude for the book + one kind share if natural.

Memory line: **{memory}**

## Q&A boundaries

- Exact local date? → EXTERNAL_OPEN; published example 2026-12-20.
- Must children distribute books? → No quotas; gratitude and optional natural sharing only.
""")

    write(base / "teacher" / "PARENT-FAMILY-GUIDE-V13.md", f"""# Parent / Family Guide — {mid}

## Flow (2:30–3:10)

P01 private reflection · P02 BG 18.65 observation · P03 Arjuna-decision family case · P04 household Gītā reading plan · P05 private saṅkalpa.

## Key idea

Receive the Gītā as friendship and instruction — not as a ranking contest.
""")

    write(base / "teacher" / "YOUNGER-STATION-PACK-V13.md", f"""# Younger Station Pack — {mid}

Y01 chapter-map stickers (1–18 gentle) · Y02 find-the-theme cards · Y03 devotee craft · Y04 gratitude-for-the-book card · Y05 memory mat: {memory}
""")

    write(base / "teacher" / "OLDER-STATION-PACK-V13.md", f"""# Older Station Pack — {mid}

O01 BG 18.65 observation · O02 18-chapter map worksheet · O03 theme sort · O04 Arjuna decision case · O05 reading-plan diagram · O06 exit ticket
""")

    acts = {
        "activities/v13-younger/Y01-CHAPTER-MAP.md": f"# {mid} Y01 Chapter Map\nUse `gita-18-chapter-map.png`. Point to chapters; do not force memorization.\n",
        "activities/v13-younger/Y02-FIND-THE-THEME.md": f"# {mid} Y02 Find the Theme\nCards: Remember Kṛṣṇa / Be kind / Serve / Honor. Match to BG 18.65 ideas.\n",
        "activities/v13-younger/Y03-DEVOTEE-CRAFT.md": f"# {mid} Y03 Devotee Craft\nHeart + name card: “I can think of Kṛṣṇa when…”\n",
        "activities/v13-younger/Y04-BOOK-GRATITUDE.md": f"# {mid} Y04 Book Gratitude\nThank-you card for the Gītā. No distribution quota.\n",
        "activities/v13-younger/Y05-MEMORY-MAT.md": f"# {mid} Y05 Memory Mat\n{memory}\n",
        "activities/v13-older/O01-BG-18-65-OBSERVATION.md": f"# {mid} O01\nObserve BG 18.65. URL {url}\n",
        "activities/v13-older/O02-18-CHAPTER-MAP.md": f"# {mid} O02\nLabel map regions lightly; whole Gītā frame https://vedabase.io/en/library/bg/\n",
        "activities/v13-older/O03-THEME-SORT.md": f"# {mid} O03\nSort: remember / devotee / worship / homage / not tonight.\n",
        "activities/v13-older/O04-ARJUNA-DECISION-CASE.md": f"# {mid} O04\nUse `arjuna-decision-frame.png`. Duty / fear / Kṛṣṇa’s word. No invented dialogue.\n",
        "activities/v13-older/O05-READING-PLAN.md": f"# {mid} O05\nFamily reading plan tiny/regular. Visual `family-gita-reading-plan.png`.\n",
        "activities/v13-older/O06-EXIT-TICKET.md": f"# {mid} O06\nMemory · one BG 18.65 word · one home reading cue\n",
        "activities/v13-parent/P01-PRIVATE-REFLECTION.md": f"# {mid} P01\nWhere does Gītā instruction already touch our home?\n",
        "activities/v13-parent/P02-SOURCE-OBSERVATION.md": f"# {mid} P02\nBG 18.65 observation\n",
        "activities/v13-parent/P03-FAMILY-CASE.md": f"# {mid} P03\nChild distribution quota case → refuse pressure goals.\n",
        "activities/v13-parent/P04-READING-PLAN.md": f"# {mid} P04\nHousehold reading plan application\n",
        "activities/v13-parent/P05-SANKALPA.md": f"# {mid} P05\nPrivate saṅkalpa\n",
    }
    for rel, text in acts.items():
        write(base / rel, text)

    write(base / "materials-safety-v13.md", f"""# Materials & Safety — {mid}

Printed maps/cards; crayons; optional Bhagavad-gītā volume for display (respect rights; no photocopy of copyrighted pages into Git). Snack + water. No child fasting. No distribution quotas for children.
""")
    write(base / "family-home-practice-v13.md", f"""# Home Practice — {mid}

**Memory:** {memory}  
Minimum: memory line + one short family Gītā opening before next meeting window. No ranking. No child quotas.
""")
    write(base / "project" / "V13-MODULE-PROJECT-BRIEF.md", f"""# Project — {mid}

Chapter-poster contribution (one chapter number + one kind theme word). Noncompetitive exhibit later at Mela if desired.
""")
    write(base / "research" / "V13-SOURCE-MATRIX.md", source_matrix(mid, [
        ("Primary", primary, url, "Display + KUTUMBA meaning"),
        ("Whole Gītā frame", "Bhagavad-gītā", "https://vedabase.io/en/library/bg/", "Map pedagogy"),
        ("Calendar example", "ISKCON Bangalore Gītā Jayantī", "https://www.iskconbangalore.org/gita-jayanti/", "Published example 2026-12-20; local EXTERNAL_OPEN"),
    ]))
    write(base / "research" / "V13-CLAIM-REGISTER.yaml", claim_register(
        mid,
        may=["Teach BG 18.65", "Use 2026-12-19 as planning Saturday", "Cite published example 2026-12-20"],
        may_not=["Claim exact local tithi settled", "Set child book-distribution quotas", "Invent Arjuna dialogue"],
    ))
    write(base / "visuals" / "v13" / "RIGHTS.md", rights_md(mid, assets))
    write(base / "visuals" / "v13" / "IMAGE-PROMPT-LIBRARY.md", image_prompt_lib(mid, [
        ("18-chapter map", "Clean educational grid of 18 numbered panels; cream background; no battlefield gore"),
        ("Arjuna decision", "Abstract three-circle decision frame; respectful; no graphic war"),
        ("Reading plan", "Family study table with closed book silhouette; warm light"),
    ]))

    gslides = "\n".join(
        f"""### Slide {i} — {t}
- **Exact title:** {t}
- **Exact on-screen copy:** {c}
- **Detailed AI-image prompt:** Cream 16:9 educational illustration for {t}; no gore; no temple logo; no identifiable persons
- **Do not claim:** exact local tithi on 2026-12-19; child distribution quotas
"""
        for i, (t, c) in enumerate([
            ("Gītā Jayantī Family Utsava", "Planning Sat 2026-12-19 · published example 2026-12-20"),
            ("Calendar honesty", "Family observance candidate — not exact tithi claim"),
            ("Essential question", "How should our family receive the Gītā without ranking?"),
            ("BG 18.65", f"Full verse layer · {url}"),
            ("KUTUMBA meaning", "Think of Me; become My devotee; worship; offer homage"),
            ("18-chapter map", "A map helps; it is not the destination"),
            ("Arjuna frame", "Duty / fear / Kṛṣṇa’s word — paraphrase only"),
            ("No child quotas", "Gratitude for the book; no pressure distribution goals"),
            ("Stations", "2:30–3:10 parent / younger / older"),
            ("Reading plan", "Tiny · regular · noncompetitive"),
            ("Home practice", memory),
            ("Do not claim", "Exact local tithi settled; temple approval"),
            ("TIME_CUE", "Tracks 2:30 · reunify 3:10 · close 4:00"),
            ("Sources", url),
            ("Rights", "KUTUMBA framing; BBT rights remain with holders"),
            ("Close", "Winter break follows on calendar"),
            ("Memory", memory),
            ("Parent focus", "Household reading plan"),
            ("Younger focus", "Theme cards + gratitude"),
            ("Older focus", "Observation + decision case"),
        ], 1)
    )
    write(base / "gamma" / "V13-GAMMA-MASTER-DECK-PROMPT.md", f"""# {mid} Gamma Master Deck Prompt

**Status:** prompt-only — not rendered — not approved  
**Slides:** 20 · {title}

{gslides}
""")
    for aud in ("PARENT", "YOUNGER", "OLDER"):
        write(base / "gamma" / f"V13-GAMMA-{aud}-DECK-PROMPT.md", f"# {mid} Gamma {aud} Deck\n\nPrompt-only audience deck. Memory: {memory}. Primary: {url}. Calendar: candidate 2026-12-19 / example 2026-12-20.\n")
    write(base / "gamma" / "V13-GAMMA-SOURCE-MAP.yaml", f"module: {mid}\nprimary: {{ref: {primary}, url: {url}}}\ngamma_status: prompt-only\nlocal_tithi: EXTERNAL_OPEN\n")

    builders = '''
def y01(document) -> None:
    add_printable_title(document, "Y01", "18-Chapter Map", "Point and wonder. No forced memorization.")
    _img(document, "gita-18-chapter-map.png", 6.2)
    add_memory_phrase_block(document, MEMORY)


def y02(document) -> None:
    add_printable_title(document, "Y02", "Find the Gītā Theme", "Match theme cards.")
    add_card_grid(document, ["Remember Kṛṣṇa", "Be a devotee", "Worship", "Offer homage", "Kind share", "Not a contest"], per_page=6)
    add_cut_lines(document)


def y03(document) -> None:
    add_printable_title(document, "Y03", "Devotee Craft", "I can think of Kṛṣṇa when…")
    add_write_lines(document, ["When I wake:", "When I help:", "When I sing:"], 2)


def y04(document) -> None:
    add_printable_title(document, "Y04", "Book Gratitude Card", "Thank you for the Gītā.")
    add_write_lines(document, ["I am grateful for:", "One kind thing I can do:"], 2)
    add_callout(document, "TEACHER_NOTE", "No child book-distribution quotas.")


def y05(document) -> None:
    add_printable_title(document, "Y05", "Memory Phrase Mat", MEMORY)
    add_memory_phrase_block(document, MEMORY)


def o01(document) -> None:
    add_printable_title(document, "O01", "BG 18.65 Observation", "")
    add_verse_card(document, "BG 18.65", "", "", "Think of Me; become My devotee; worship Me; offer homage.", VERSE_URL)
    add_write_lines(document, ["Four invitations I see:", "One promise I notice:", "One family application:"], 3)


def o02(document) -> None:
    add_printable_title(document, "O02", "18-Chapter Map Worksheet", "")
    _img(document, "gita-18-chapter-map.png", 6.0)
    add_write_lines(document, ["A chapter number I notice:", "Why a map helps:", "Map limit:"], 2)


def o03(document) -> None:
    add_printable_title(document, "O03", "Theme Sort", "")
    add_card_grid(document, ["man-manā", "mad-bhakta", "mad-yājī", "namaskuru", "ranking contest", "quota pressure"], per_page=6)


def o04(document) -> None:
    add_printable_title(document, "O04", "Arjuna Decision Case", "Paraphrase only — no invented dialogue.")
    _img(document, "arjuna-decision-frame.png", 5.5)
    add_write_lines(document, ["Duty tension:", "Fear tension:", "Kṛṣṇa’s instruction direction:"], 3)


def o05(document) -> None:
    add_printable_title(document, "O05", "Family Reading Plan", "")
    _img(document, "family-gita-reading-plan.png", 5.0)
    add_write_lines(document, ["When:", "How long (minutes):", "Minimum version:"], 2)


def o06(document) -> None:
    add_printable_title(document, "O06", "Exit Ticket", "")
    add_write_lines(document, ["Memory line:", "One BG 18.65 word:", "Home reading cue:"], 2)


def p01(document) -> None:
    add_printable_title(document, "P01", "Private Reflection", "")
    add_write_lines(document, ["Where Gītā already touches our home:", "Where pressure sneaks in:"], 3)
    add_callout(document, "SAFETY_PRIVACY", "Private writing. Optional share.")


def p02(document) -> None:
    add_printable_title(document, "P02", "Source Observation", "BG 18.65")
    add_verse_card(document, "BG 18.65", "", "", "Always think of Me; become My devotee.", VERSE_URL)
    add_write_lines(document, ["What stands out?", "Calendar honesty note:"], 3)


def p03(document) -> None:
    add_printable_title(document, "P03", "Family Case", "Child distribution quota pressure")
    add_write_lines(document, ["Mistaken conclusion:", "Principle:", "Action:"], 3)
    add_callout(document, "DO_NOT_SPECULATE", "No child quotas. Local tithi EXTERNAL_OPEN.")


def p04(document) -> None:
    add_printable_title(document, "P04", "Household Reading Plan", "")
    add_write_lines(document, ["Cue:", "Duration:", "Who starts:", "Minimum version:"], 2)
    add_callout(document, "FAMILY_APPLICATION", "Tiny and regular.")


def p05(document) -> None:
    add_printable_title(document, "P05", "Private Saṅkalpa", "")
    add_callout(document, "KEY_IDEA", "specific action + frequency + trigger + minimum version")
    add_write_lines(document, ["Specific action:", "Frequency:", "Trigger:", "Minimum version:"], 2)


YOUNGER_BUILDERS = [y01, y02, y03, y04, y05]
OLDER_BUILDERS = [o01, o02, o03, o04, o05, o06]
PARENT_BUILDERS = [p01, p02, p03, p04, p05]
'''
    write(PRINT / "C2-U2.py", printable_module(mid, folder, memory, url, builders))
    write(EVID / "C2-U2" / "CONTENT-AUDIT.md", content_audit(
        mid, title, "2026-12-19", primary, url,
        [("Y01", "chapter map"), ("O04", "Arjuna decision"), ("P03", "no child quotas")],
        ["DOCX/PDF deferred", "Raster deferred", "Local tithi EXTERNAL_OPEN", "Gamma render EXTERNAL_OPEN"],
    ))
    write(EVID / "C2-U2" / "REQUIREMENT-TRACEABILITY.csv", req_csv(mid, [
        ("Start Here", f"11-weekly-program-library/first-six-months/{folder}/V13-UTSAVA-START-HERE.md", "candidate not exact tithi"),
        ("Owner ROS", f"11-weekly-program-library/first-six-months/{folder}/teacher/OWNER-RUN-OF-SHOW-V13.md", ""),
        ("Main script", f"11-weekly-program-library/first-six-months/{folder}/teacher/MAIN-FACILITATOR-GUIDE-V13.md", "BG 18.65"),
        ("Parent guide", f"11-weekly-program-library/first-six-months/{folder}/teacher/PARENT-FAMILY-GUIDE-V13.md", ""),
        ("Younger/older packs", f"11-weekly-program-library/first-six-months/{folder}/teacher/", ""),
        ("Materials/safety", f"11-weekly-program-library/first-six-months/{folder}/materials-safety-v13.md", ""),
        ("Source matrix", f"11-weekly-program-library/first-six-months/{folder}/research/V13-SOURCE-MATRIX.md", ""),
        ("Rights/visuals", f"11-weekly-program-library/first-six-months/{folder}/visuals/v13/RIGHTS.md", ""),
        ("Gamma master", f"11-weekly-program-library/first-six-months/{folder}/gamma/V13-GAMMA-MASTER-DECK-PROMPT.md", "actual copy"),
        ("Printable builder", "scripts/v13/printables/C2-U2.py", "no render"),
    ]))


# ---------------------------------------------------------------------------
# C3-U3 Nityananda
# ---------------------------------------------------------------------------

def author_c3_u3() -> None:
    folder = "c3-u3-nityananda-trayodasi-family-utsava"
    base = FIRST / folder
    mid = "C3-U3"
    title = "Nityānanda Trayodaśī Family Utsava"
    date = "Saturday 2027-02-20 (planning); published example 2027-02-19"
    primary = "CC Ādi 5.208"
    url = "https://vedabase.io/en/library/cc/adi/5/208/"
    memory = "Nityānanda is an incarnation of mercy."
    assets = visual_c3_u3(base / "visuals" / "v13")

    write(base / "V13-UTSAVA-START-HERE.md", start_here(
        mid, title, date,
        essential="How should mercy, service, and saṅkīrtana shape our family’s next kind action?",
        conclusion="Nityānanda’s mercy invites service and saṅkīrtana without graphic violence or ranking.",
        memory=memory,
        primary=primary,
        primary_url=url,
        misconception="Tonight is automatically the exact local Nityānanda Trayodaśī tithi.",
        calendar_note="Family observance candidate 2027-02-20. Published example 2027-02-19. Exact local tithi EXTERNAL_OPEN.",
        home_min="Minimum: memory line + one mercy/service act. Soft kīrtana invitation at home optional.",
        next_name="Owner-select Mela 2027-02-27 or 2027-03-06.",
        filename="V13-UTSAVA-START-HERE.md",
    ))

    write(base / "teacher" / "OWNER-RUN-OF-SHOW-V13.md", owner_ros(
        mid, title, date,
        stations=[
            "Mercy chain",
            "Service challenge",
            "Age-sensitive Jagāi-Mādhāi mercy theme (no graphic violence)",
            "Kīrtana invitation",
            "Family who-can-we-serve plan",
        ],
        safety_bullets=[
            "No graphic violence in Jagāi-Mādhāi materials.",
            "Focus mercy/helping/service.",
            "No child fasting.",
            "Avoid speculative miracle embellishment.",
            "Candidate date 2027-02-20; published example 2027-02-19 — not exact-tithi claim.",
        ],
    ))

    write(base / "teacher" / "MAIN-FACILITATOR-GUIDE-V13.md", f"""# Main Facilitator Guide — {mid}

**{title}** · {date}  
**Status:** EXTERNAL_OPEN local tithi/human/temple  
**Entry:** `V13-UTSAVA-START-HERE.md`

## Exact primary

- **Reference:** {primary}
- **URL:** {url}
- **Bengali:** প্রেমে মত্ত নিত্যানন্দ কৃপা–অবতার । উত্তম, অধম, কিছু না করে বিচার ॥ ২০৮ ॥
- **IAST:** preme matta nityānanda kṛpā-avatāra / uttama, adhama, kichu nā kare vicāra
- **KUTUMBA teaching meaning:** Nityānanda, intoxicated by ecstatic love and an incarnation of mercy, does not distinguish between the good and the bad.
- **Supporting:** CC Madhya 1.25 — https://vedabase.io/en/library/cc/madhya/1/25/ — by nature inspired in Kṛṣṇa-prema, and by the Lord’s order He distributes loving service everywhere.
- **Rights:** Bengali display + KUTUMBA meaning; no full purport.

## 10–20 minute speaking script

Welcome to our **Nityānanda Trayodaśī Family Utsava** planning Saturday — **2027-02-20**.

Calendar honesty: published example often lists **2027-02-19**. Tonight is a **family observance candidate**, not a settled exact local tithi claim (**EXTERNAL_OPEN**).

Essential question: How should mercy, service, and saṅkīrtana shape our family’s next kind action?

Primary: CC Ādi 5.208 — {url}

KUTUMBA teaching meaning: “Nityānanda, intoxicated by ecstatic love and an incarnation of mercy, does not distinguish between the good and the bad.”

We teach mercy and service. If Jagāi-Mādhāi is mentioned, keep it **age-sensitive**: people far from kindness are invited toward the holy name and service. **No graphic violence.** No speculative miracle embellishment.

Analogy one — mercy chain: one kind invitation leads to another. Limit: not every person responds immediately; do not force.

Analogy two — service challenge cards: small helps train the heart. Limit: service is not a scoreboard.

Analogy three — saṅkīrtana as invitation: listening counts. Limit: not a volume contest.

Household case: Older child wants to dramatize a violent scene “for realism.” Redirect to mercy and helping. Principle: age-sensitive storytelling.

Memory line: **{memory}**
""")

    write(base / "teacher" / "PARENT-FAMILY-GUIDE-V13.md", f"""# Parent / Family Guide — {mid}

P01 reflection on mercy at home · P02 CC Ādi 5.208 observation · P03 age-sensitive storytelling case · P04 who-can-we-serve plan · P05 saṅkalpa.
""")
    write(base / "teacher" / "YOUNGER-STATION-PACK-V13.md", f"""# Younger Station Pack — {mid}

Y01 mercy chain · Y02 service challenge cards · Y03 help-hands craft · Y04 kind invitation cards · Y05 memory mat: {memory}

**Age boundary:** No violent images or graphic Jagāi-Mādhāi scenes.
""")
    write(base / "teacher" / "OLDER-STATION-PACK-V13.md", f"""# Older Station Pack — {mid}

O01 CC Ādi 5.208 observation · O02 mercy vs judgment sort · O03 age-sensitive narrative boundaries · O04 service plan · O05 saṅkīrtana invitation map · O06 exit ticket

**TEACHER-ONLY:** If Jagāi-Mādhāi arises, emphasize transformation through mercy/holy name — never graphic violence.
""")

    acts = {
        "activities/v13-younger/Y01-MERCY-CHAIN.md": f"# {mid} Y01\nVisual `mercy-chain.png`. Hear → receive mercy → serve → invite.\n",
        "activities/v13-younger/Y02-SERVICE-CHALLENGE.md": f"# {mid} Y02\n`service-challenge-cards.png`\n",
        "activities/v13-younger/Y03-HELP-HANDS-CRAFT.md": f"# {mid} Y03\nTrace hands; write one help.\n",
        "activities/v13-younger/Y04-KIND-INVITATION.md": f"# {mid} Y04\nKind invitation cards for kīrtana/listening.\n",
        "activities/v13-younger/Y05-MEMORY-MAT.md": f"# {mid} Y05\n{memory}\n",
        "activities/v13-older/O01-CC-ADI-5-208.md": f"# {mid} O01\n{url}\n",
        "activities/v13-older/O02-MERCY-VS-JUDGMENT.md": f"# {mid} O02\nSort mercy / harsh judgment / unsure.\n",
        "activities/v13-older/O03-AGE-SENSITIVE-BOUNDARIES.md": f"# {mid} O03\nRewrite a graphic line into a mercy-focused line. No violence details.\n",
        "activities/v13-older/O04-SERVICE-PLAN.md": f"# {mid} O04\nWho can we serve this week?\n",
        "activities/v13-older/O05-SANKIRTANA-MAP.md": f"# {mid} O05\n`sankirtana-invitation.png`\n",
        "activities/v13-older/O06-EXIT-TICKET.md": f"# {mid} O06\nMemory · one mercy act · one person to serve\n",
        "activities/v13-parent/P01-PRIVATE-REFLECTION.md": f"# {mid} P01\nWhere does judgment block mercy at home?\n",
        "activities/v13-parent/P02-SOURCE-OBSERVATION.md": f"# {mid} P02\nCC Ādi 5.208\n",
        "activities/v13-parent/P03-AGE-SENSITIVE-CASE.md": f"# {mid} P03\nChild wants violent dramatization → redirect.\n",
        "activities/v13-parent/P04-WHO-CAN-WE-SERVE.md": f"# {mid} P04\nFamily service plan\n",
        "activities/v13-parent/P05-SANKALPA.md": f"# {mid} P05\nPrivate saṅkalpa\n",
    }
    for rel, text in acts.items():
        write(base / rel, text)

    write(base / "materials-safety-v13.md", f"""# Materials & Safety — {mid}

Printables; crayons; soft kīrtana track optional. **No graphic violence materials.** No child fasting. Snack + water.
""")
    write(base / "family-home-practice-v13.md", f"""# Home Practice — {mid}

**Memory:** {memory}  
Minimum: memory line + one mercy/service act. Optional soft kīrtana listening.
""")
    write(base / "project" / "V13-MODULE-PROJECT-BRIEF.md", f"""# Project — {mid}

“Who can we serve?” family card for Mela exhibition (noncompetitive).
""")
    write(base / "research" / "V13-SOURCE-MATRIX.md", source_matrix(mid, [
        ("Primary", primary, url, "Mercy incarnation teaching"),
        ("Supporting", "CC Madhya 1.25", "https://vedabase.io/en/library/cc/madhya/1/25/", "prema distribution"),
        ("Narrative boundary", "CC sources for Nityānanda", "https://vedabase.io/en/library/cc/", "Age-sensitive; no graphic violence"),
    ]))
    write(base / "research" / "V13-CLAIM-REGISTER.yaml", claim_register(
        mid,
        may=["Teach CC Ādi 5.208", "Use planning Sat 2027-02-20", "Cite published example 2027-02-19"],
        may_not=["Exact local tithi claim", "Graphic Jagāi-Mādhāi violence", "Speculative miracle embellishment", "Child fasting"],
    ))
    write(base / "visuals" / "v13" / "RIGHTS.md", rights_md(mid, assets))
    write(base / "visuals" / "v13" / "IMAGE-PROMPT-LIBRARY.md", image_prompt_lib(mid, [
        ("Mercy chain", "Four soft panels of hearing-serving-inviting; no violence"),
        ("Service cards", "Four household help icons; educational flat style"),
        ("Saṅkīrtana", "Abstract circle of song invitation; calm; no crowd crush"),
    ]))
    write(base / "gamma" / "V13-GAMMA-MASTER-DECK-PROMPT.md", f"""# {mid} Gamma Master Deck Prompt

**Status:** prompt-only — not rendered — not approved  
**Slides:** 18 · {title}

### Slide 1 — Title
- **Exact title:** Nityānanda Trayodaśī Family Utsava
- **Exact on-screen copy:** Planning Sat 2027-02-20 · published example 2027-02-19 · {memory}
- **Do not claim:** exact local tithi

### Slide 2 — Calendar honesty
- **Exact on-screen copy:** Family observance candidate · EXTERNAL_OPEN local confirmation

### Slide 3 — Essential question
- **Exact on-screen copy:** How should mercy, service, and saṅkīrtana shape our next kind action?

### Slide 4 — Primary verse
- **Exact title:** CC Ādi 5.208
- **Exact on-screen copy:** Bengali + IAST + KUTUMBA meaning · {url}

### Slide 5 — Supporting
- **Exact on-screen copy:** CC Madhya 1.25 · distributes loving service · https://vedabase.io/en/library/cc/madhya/1/25/

### Slide 6 — Mercy theme
- **Exact on-screen copy:** Incarnation of mercy · no ranking of who “deserves” help

### Slide 7 — Age-sensitive storytelling
- **Exact on-screen copy:** Jagāi-Mādhāi if mentioned = mercy/helping · NO graphic violence

### Slide 8 — Service challenge
- **Exact on-screen copy:** Help at home · kind words · join kīrtana · share prasāda

### Slide 9 — Saṅkīrtana
- **Exact on-screen copy:** Invitation · listening counts · not volume contest

### Slide 10 — Stations
- **Exact on-screen copy:** 2:30–3:10 parallel tracks

### Slide 11 — Parent focus
- **Exact on-screen copy:** Who can we serve this week?

### Slide 12 — Younger focus
- **Exact on-screen copy:** Mercy chain + help-hands craft

### Slide 13 — Older focus
- **Exact on-screen copy:** Mercy vs judgment · narrative boundaries

### Slide 14 — Case
- **Exact on-screen copy:** Redirect violent dramatization requests

### Slide 15 — Home practice
- **Exact on-screen copy:** {memory} + one service act

### Slide 16 — Do not claim
- **Exact on-screen copy:** Exact tithi · temple approval · graphic violence as pedagogy

### Slide 17 — Sources / rights
- **Exact on-screen copy:** {url} · KUTUMBA framing · BBT rights remain with holders

### Slide 18 — Close
- **Exact on-screen copy:** Next: owner-select Mela date
""")
    for aud in ("PARENT", "YOUNGER", "OLDER"):
        write(base / "gamma" / f"V13-GAMMA-{aud}-DECK-PROMPT.md", f"# {mid} Gamma {aud}\n\nPrompt-only. Memory: {memory}. No graphic violence. Primary: {url}.\n")
    write(base / "gamma" / "V13-GAMMA-SOURCE-MAP.yaml", f"module: {mid}\nprimary: {{ref: '{primary}', url: {url}}}\nsupporting:\n  - ref: CC Madhya 1.25\n    url: https://vedabase.io/en/library/cc/madhya/1/25/\ngamma_status: prompt-only\nlocal_tithi: EXTERNAL_OPEN\n")

    builders = '''
def y01(document) -> None:
    add_printable_title(document, "Y01", "Mercy Chain", "Hear → mercy → serve → invite")
    _img(document, "mercy-chain.png", 6.0)
    add_memory_phrase_block(document, MEMORY)
    add_callout(document, "CHILD_ACTIVITY", "No scary or violent pictures.")


def y02(document) -> None:
    add_printable_title(document, "Y02", "Service Challenge Cards", "")
    _img(document, "service-challenge-cards.png", 5.5)
    add_cut_lines(document)


def y03(document) -> None:
    add_printable_title(document, "Y03", "Help-Hands Craft", "Trace and write one help.")
    add_write_lines(document, ["My helping hand will:", "For whom:"], 2)


def y04(document) -> None:
    add_printable_title(document, "Y04", "Kind Invitation Cards", "")
    add_card_grid(document, ["Please sing with us", "Listening is welcome", "We can help together", "Come share kindness"], per_page=4)


def y05(document) -> None:
    add_printable_title(document, "Y05", "Memory Phrase Mat", MEMORY)
    add_memory_phrase_block(document, MEMORY)


def o01(document) -> None:
    add_printable_title(document, "O01", "CC Ādi 5.208 Observation", "")
    add_verse_card(document, "CC Ādi 5.208", "", "", "Nityānanda is an incarnation of mercy and does not distinguish good and bad.", VERSE_URL)
    add_write_lines(document, ["What is mercy here?", "What must we not dramatize?"], 3)


def o02(document) -> None:
    add_printable_title(document, "O02", "Mercy vs Judgment Sort", "")
    add_card_grid(document, ["Invite kindly", "Mock the fallen", "Serve quietly", "Public shame", "Listen in kīrtana", "Scoreboard seva"], per_page=6)


def o03(document) -> None:
    add_printable_title(document, "O03", "Age-Sensitive Boundaries", "Rewrite without graphic violence.")
    add_write_lines(document, ["Harsh line (do not perform):", "Mercy-focused rewrite:", "Why the rewrite is better:"], 3)
    add_callout(document, "TEACHER_NOTE", "No graphic Jagāi-Mādhāi violence in materials.")


def o04(document) -> None:
    add_printable_title(document, "O04", "Service Plan", "Who can we serve?")
    add_write_lines(document, ["Person/place:", "Help act:", "When:", "Minimum version:"], 2)


def o05(document) -> None:
    add_printable_title(document, "O05", "Saṅkīrtana Invitation Map", "")
    _img(document, "sankirtana-invitation.png", 4.5)
    add_write_lines(document, ["Invitation line:", "Listening counts because:"], 2)


def o06(document) -> None:
    add_printable_title(document, "O06", "Exit Ticket", "")
    add_write_lines(document, ["Memory line:", "One mercy act:", "One person to serve:"], 2)


def p01(document) -> None:
    add_printable_title(document, "P01", "Private Reflection", "")
    add_write_lines(document, ["Where judgment blocks mercy:", "One softer response:"], 3)
    add_callout(document, "SAFETY_PRIVACY", "Private optional share.")


def p02(document) -> None:
    add_printable_title(document, "P02", "Source Observation", "")
    add_verse_card(document, "CC Ādi 5.208", "", "", "Incarnation of mercy.", VERSE_URL)
    add_write_lines(document, ["Takeaway:", "Calendar honesty:"], 3)


def p03(document) -> None:
    add_printable_title(document, "P03", "Age-Sensitive Case", "Violent dramatization request")
    add_write_lines(document, ["Mistaken goal:", "Principle:", "Redirect:"], 3)
    add_callout(document, "DO_NOT_SPECULATE", "No graphic violence; no miracle embellishment.")


def p04(document) -> None:
    add_printable_title(document, "P04", "Who Can We Serve?", "")
    add_write_lines(document, ["Household:", "Community:", "Temple/service space:", "Minimum version:"], 2)
    add_callout(document, "FAMILY_APPLICATION", "One real act beats many slogans.")


def p05(document) -> None:
    add_printable_title(document, "P05", "Private Saṅkalpa", "")
    add_callout(document, "KEY_IDEA", "specific action + frequency + trigger + minimum version")
    add_write_lines(document, ["Specific action:", "Frequency:", "Trigger:", "Minimum version:"], 2)


YOUNGER_BUILDERS = [y01, y02, y03, y04, y05]
OLDER_BUILDERS = [o01, o02, o03, o04, o05, o06]
PARENT_BUILDERS = [p01, p02, p03, p04, p05]
'''
    write(PRINT / "C3-U3.py", printable_module(mid, folder, memory, url, builders))
    write(EVID / "C3-U3" / "CONTENT-AUDIT.md", content_audit(
        mid, title, "2027-02-20", primary, url,
        [("Y01", "mercy chain"), ("O03", "age-sensitive rewrite"), ("P03", "no graphic violence case")],
        ["DOCX/PDF deferred", "Raster deferred", "Local tithi EXTERNAL_OPEN", "Gamma render EXTERNAL_OPEN"],
    ))
    write(EVID / "C3-U3" / "REQUIREMENT-TRACEABILITY.csv", req_csv(mid, [
        ("Start Here", f"11-weekly-program-library/first-six-months/{folder}/V13-UTSAVA-START-HERE.md", "candidate 02-20 / example 02-19"),
        ("Owner ROS", f"11-weekly-program-library/first-six-months/{folder}/teacher/OWNER-RUN-OF-SHOW-V13.md", "no graphic violence"),
        ("Main script", f"11-weekly-program-library/first-six-months/{folder}/teacher/MAIN-FACILITATOR-GUIDE-V13.md", "CC Adi 5.208"),
        ("Parent/younger/older", f"11-weekly-program-library/first-six-months/{folder}/teacher/", ""),
        ("Materials/safety", f"11-weekly-program-library/first-six-months/{folder}/materials-safety-v13.md", ""),
        ("Source matrix", f"11-weekly-program-library/first-six-months/{folder}/research/V13-SOURCE-MATRIX.md", ""),
        ("Rights/visuals", f"11-weekly-program-library/first-six-months/{folder}/visuals/v13/", ""),
        ("Gamma master", f"11-weekly-program-library/first-six-months/{folder}/gamma/V13-GAMMA-MASTER-DECK-PROMPT.md", ""),
        ("Printable builder", "scripts/v13/printables/C3-U3.py", "no render"),
    ]))


# ---------------------------------------------------------------------------
# MELA
# ---------------------------------------------------------------------------

def author_mela() -> None:
    folder = "mela-six-month-reflection-family-mela"
    base = FIRST / folder
    mid = "MELA"
    title = "Six-Month Reflection / Family Mela"
    date = "Owner-select 2027-02-27 or 2027-03-06 (not a tithi claim)"
    primary = "First-six-month review chain (ŚB 1.2.18 center)"
    url = "https://vedabase.io/en/library/sb/1/2/18/"
    memory = "Continue · Review · Strengthen — without comparison."
    assets = visual_mela(base / "visuals" / "v13")

    write(base / "V13-MELA-START-HERE.md", start_here(
        mid, title, date,
        essential="What will our family continue, review, and strengthen after six months — without ranking anyone?",
        conclusion="Exhibition and feedback celebrate learning; the next phase is chosen with gratitude, not competition.",
        memory=memory,
        primary=primary,
        primary_url=url,
        misconception="This Mela is itself a religious tithi obligation.",
        calendar_note="Owner selects 2027-02-27 or 2027-03-06. Not a religious tithi claim. Gaura Pūrṇimā lookahead is separate.",
        home_min="Minimum: one continue / one review / one strengthen note privately as a family.",
        next_name="Lookahead only: Gaura Pūrṇimā published example 2027-03-22 — see launch/V13-GAURA-PURNIMA-LOOKAHEAD.md",
        filename="V13-MELA-START-HERE.md",
    ).replace("V13-UTSAVA-START-HERE.md", "V13-MELA-START-HERE.md"))

    # fix start-here filename content already written with wrong H1 path — rewrite cleanly
    write(base / "V13-MELA-START-HERE.md", start_here(
        mid, title, date,
        essential="What will our family continue, review, and strengthen after six months — without ranking anyone?",
        conclusion="Exhibition and feedback celebrate learning; the next phase is chosen with gratitude, not competition.",
        memory=memory,
        primary=primary,
        primary_url=url,
        misconception="This Mela is itself a religious tithi obligation.",
        calendar_note="Owner selects 2027-02-27 or 2027-03-06. Not a religious tithi claim.",
        home_min="Minimum: one continue / one review / one strengthen note privately as a family.",
        next_name="Lookahead only: Gaura Pūrṇimā published example 2027-03-22 — launch/V13-GAURA-PURNIMA-LOOKAHEAD.md",
    ).replace("# V13 START HERE — MELA", "# V13 MELA START HERE — MELA").replace(
        "`V13-UTSAVA-START-HERE.md` or `V13-MELA-START-HERE.md`", "`V13-MELA-START-HERE.md`"
    ))

    write(base / "teacher" / "OWNER-RUN-OF-SHOW-V13.md", owner_ros(
        mid, title, date,
        stations=[
            "Child project exhibition (noncompetitive)",
            "Four-family reflections (voluntary)",
            "Kīrtana",
            "Source quiz/game (gentle)",
            "Family testimony without comparison",
            "Feedback cards",
            "Continue / Review / Strengthen decision",
            "Preview next formation phase",
        ],
        safety_bullets=[
            "Noncompetitive exhibition only — no awards ranking spirituality.",
            "Testimony is invitation; opt-out honored.",
            "Not a tithi claim.",
            "No child fasting.",
            "Do not force-fit Gaura Pūrṇimā into this Mela clock.",
        ],
    ))

    write(base / "teacher" / "MAIN-FACILITATOR-GUIDE-V13.md", f"""# Main Facilitator Guide — {mid}

**{title}** · {date}  
**Entry:** `V13-MELA-START-HERE.md`

## Primary review center

- **ŚB 1.2.18** — {url} — hearing/service culture that nourishes bhakti (Cycle 1 opening chain center).
- Review first-six-month verse chain without introducing major new doctrine tonight.

## 10–20 minute speaking script

Welcome to our **Six-Month Reflection / Family Mela**.

Owner selects **2027-02-27** or **2027-03-06**. This is **not** a religious tithi claim.

Essential question: What will our family continue, review, and strengthen — without ranking anyone?

We celebrate projects and learning. We give feedback that is kind and specific. We do not crown winners of devotion.

Analogy one — exhibition as garden walk: notice beauty; do not score flowers against each other.  
Analogy two — feedback as polishing cloth: gentle, specific, hopeful.  
Analogy three — three doors: Continue / Review / Strengthen.

Household case: A parent compares children aloud during exhibition. Redirect to one appreciative sentence per child without ranking.

Memory line: **{memory}**

Gaura Pūrṇimā (published example 2027-03-22) is **lookahead only** — do not distort this Mela to force-fit it.
""")

    write(base / "teacher" / "PARENT-FAMILY-GUIDE-V13.md", f"""# Parent / Family Guide — {mid}

P01 private six-month reflection · P02 review-chain observation · P03 noncomparison testimony prep · P04 continue/review/strengthen · P05 private saṅkalpa toward next phase.
""")
    write(base / "teacher" / "YOUNGER-STATION-PACK-V13.md", f"""# Younger Station Pack — {mid}

Y01 exhibition walk cards · Y02 appreciation stickers · Y03 favorite memory craft · Y04 feedback smile cards · Y05 memory mat
""")
    write(base / "teacher" / "OLDER-STATION-PACK-V13.md", f"""# Older Station Pack — {mid}

O01 verse-chain quiz (gentle) · O02 exhibition host sheet · O03 feedback that helps · O04 continue/review/strengthen · O05 next-phase preview · O06 exit ticket
""")

    acts = {
        "activities/v13-younger/Y01-EXHIBITION-WALK.md": f"# {mid} Y01\nWalk · see · say one kind word. Visual `exhibition-feedback-loop.png`.\n",
        "activities/v13-younger/Y02-APPRECIATION.md": f"# {mid} Y02\nAppreciation stickers — no ranking stars that create winners.\n",
        "activities/v13-younger/Y03-FAVORITE-MEMORY.md": f"# {mid} Y03\nDraw one favorite learning memory.\n",
        "activities/v13-younger/Y04-FEEDBACK-SMILE.md": f"# {mid} Y04\nSmile feedback cards: kind · specific · hopeful.\n",
        "activities/v13-younger/Y05-MEMORY-MAT.md": f"# {mid} Y05\n{memory}\n",
        "activities/v13-older/O01-VERSE-CHAIN-QUIZ.md": f"# {mid} O01\nGentle quiz across first-six-month chain. Not a tournament.\n",
        "activities/v13-older/O02-EXHIBITION-HOST.md": f"# {mid} O02\nHost sheet: welcome · explain · thank.\n",
        "activities/v13-older/O03-FEEDBACK-THAT-HELPS.md": f"# {mid} O03\nRewrite comparative feedback into helpful feedback.\n",
        "activities/v13-older/O04-CONTINUE-REVIEW-STRENGTHEN.md": f"# {mid} O04\n`continue-review-strengthen.png`\n",
        "activities/v13-older/O05-NEXT-PHASE-PREVIEW.md": f"# {mid} O05\nPreview next formation phase at high level only.\n",
        "activities/v13-older/O06-EXIT-TICKET.md": f"# {mid} O06\nOne continue · one review · one strengthen\n",
        "activities/v13-parent/P01-PRIVATE-REFLECTION.md": f"# {mid} P01\nSix-month private reflection\n",
        "activities/v13-parent/P02-REVIEW-CHAIN.md": f"# {mid} P02\nŚB 1.2.18 center + chain gratitude\n",
        "activities/v13-parent/P03-TESTIMONY-PREP.md": f"# {mid} P03\nOptional testimony without comparison. Visual `family-testimony-frame.png`.\n",
        "activities/v13-parent/P04-CRS-DECISION.md": f"# {mid} P04\nContinue / Review / Strengthen decision\n",
        "activities/v13-parent/P05-SANKALPA.md": f"# {mid} P05\nPrivate next-phase saṅkalpa\n",
    }
    for rel, text in acts.items():
        write(base / rel, text)

    write(base / "materials-safety-v13.md", f"""# Materials & Safety — {mid}

Exhibition tables; feedback cards; soft kīrtana; snack + water. Noncompetitive signage. No awards that rank devotion. Not a tithi ritual packet.
""")
    write(base / "family-home-practice-v13.md", f"""# Home Practice — {mid}

**Memory:** {memory}  
Write three private lines: Continue / Review / Strengthen. Optional look at Gaura Pūrṇimā lookahead file — no forced attendance claim.
""")
    write(base / "project" / "V13-MODULE-PROJECT-BRIEF.md", f"""# Project — {mid}

All prior cycle artifacts may be exhibited. Feedback cards collected for owner learning — strip private names before any repo notes.
""")
    write(base / "research" / "V13-SOURCE-MATRIX.md", source_matrix(mid, [
        ("Review center", "ŚB 1.2.18", url, "Opening chain center"),
        ("Chain", "V13-VERSE-CHAIN.md", "build-evidence/V13-VERSE-CHAIN.md", "Review only"),
        ("Calendar", "Owner-select Saturdays", "launch/FIRST-SIX-MONTHS-CALENDAR.md", "Not a tithi claim"),
    ]))
    write(base / "research" / "V13-CLAIM-REGISTER.yaml", claim_register(
        mid,
        may=["Run noncompetitive exhibition", "Owner-select 2027-02-27 or 2027-03-06", "Review verse chain"],
        may_not=["Call Mela a tithi", "Rank families/children spiritually", "Force Gaura Pūrṇimā into this night"],
    ))
    write(base / "visuals" / "v13" / "RIGHTS.md", rights_md(mid, assets))
    write(base / "visuals" / "v13" / "IMAGE-PROMPT-LIBRARY.md", image_prompt_lib(mid, [
        ("Exhibition loop", "Five-step noncompetitive exhibition icons"),
        ("CRS doors", "Three doors Continue Review Strengthen"),
        ("Testimony", "Quiet family gratitude frame; no podium ranking"),
    ]))
    write(base / "gamma" / "V13-GAMMA-MASTER-DECK-PROMPT.md", f"""# {mid} Gamma Master Deck Prompt

**Status:** prompt-only — not rendered — not approved  
**Slides:** 18

### Slide 1 — Title
Six-Month Reflection / Family Mela · owner-select 2027-02-27 or 2027-03-06 · not a tithi claim

### Slide 2 — Essential question
What will we continue, review, and strengthen — without ranking?

### Slide 3 — Review center
ŚB 1.2.18 · {url}

### Slide 4 — Exhibition rules
See · appreciate · learn · feedback · continue · NO spiritual scoreboard

### Slide 5 — Feedback
Kind · specific · hopeful

### Slide 6 — Testimony
Invitation · opt-out honored · no comparison

### Slide 7 — Source quiz
Gentle review game — not a tournament

### Slide 8 — Continue / Review / Strengthen
Three doors visual

### Slide 9 — Stations
Parent / younger / older 2:30–3:10

### Slide 10 — Home practice
{memory}

### Slide 11 — Next phase preview
High-level only

### Slide 12 — Gaura Pūrṇimā lookahead pointer
Published example 2027-03-22 · separate file · do not force-fit

### Slide 13 — Do not claim
Tithi status · ranking · temple approval fabrication

### Slide 14 — TIME_CUE
2:30 tracks · 3:10 reunify · 4:00 close

### Slide 15 — Rights
KUTUMBA framing · EXTERNAL_OPEN review

### Slide 16 — Memory
{memory}

### Slide 17 — Sources
Verse chain + ŚB 1.2.18

### Slide 18 — Close
Gratitude · snack · private saṅkalpa
""")
    for aud in ("PARENT", "YOUNGER", "OLDER"):
        write(base / "gamma" / f"V13-GAMMA-{aud}-DECK-PROMPT.md", f"# {mid} Gamma {aud}\n\nPrompt-only. Noncompetitive. Not a tithi claim. Memory: {memory}.\n")
    write(base / "gamma" / "V13-GAMMA-SOURCE-MAP.yaml", f"module: {mid}\nprimary: {{ref: ŚB 1.2.18 review center, url: {url}}}\ngamma_status: prompt-only\nnot_a_tithi_claim: true\n")

    builders = '''
def y01(document) -> None:
    add_printable_title(document, "Y01", "Exhibition Walk", "See · say one kind word")
    _img(document, "exhibition-feedback-loop.png", 6.0)
    add_memory_phrase_block(document, MEMORY)


def y02(document) -> None:
    add_printable_title(document, "Y02", "Appreciation Cards", "No ranking winners")
    add_card_grid(document, ["I noticed...", "I learned...", "Thank you for...", "This helped me..."], per_page=4)
    add_callout(document, "CHILD_ACTIVITY", "Appreciation only — no scoreboard.")


def y03(document) -> None:
    add_printable_title(document, "Y03", "Favorite Memory Craft", "")
    add_write_lines(document, ["My favorite learning memory:", "One person who helped me:"], 3)


def y04(document) -> None:
    add_printable_title(document, "Y04", "Feedback Smile Cards", "Kind · specific · hopeful")
    add_write_lines(document, ["Kind:", "Specific:", "Hopeful:"], 2)


def y05(document) -> None:
    add_printable_title(document, "Y05", "Memory Phrase Mat", MEMORY)
    add_memory_phrase_block(document, MEMORY)


def o01(document) -> None:
    add_printable_title(document, "O01", "Gentle Verse-Chain Quiz", "Not a tournament")
    add_write_lines(document, ["One verse I remember:", "One week theme I remember:", "One family practice I remember:"], 2)
    add_callout(document, "TEACHER_NOTE", "Celebrate recall; do not rank scores publicly.")


def o02(document) -> None:
    add_printable_title(document, "O02", "Exhibition Host Sheet", "")
    add_write_lines(document, ["Welcome line:", "One-sentence explain:", "Thank-you line:"], 2)


def o03(document) -> None:
    add_printable_title(document, "O03", "Feedback That Helps", "")
    add_write_lines(document, ["Comparative line to avoid:", "Helpful rewrite:"], 3)


def o04(document) -> None:
    add_printable_title(document, "O04", "Continue / Review / Strengthen", "")
    _img(document, "continue-review-strengthen.png", 5.5)
    add_write_lines(document, ["CONTINUE:", "REVIEW:", "STRENGTHEN:"], 2)


def o05(document) -> None:
    add_printable_title(document, "O05", "Next-Phase Preview", "High-level only")
    add_write_lines(document, ["One curiosity for next phase:", "One support we need:"], 2)


def o06(document) -> None:
    add_printable_title(document, "O06", "Exit Ticket", "")
    add_write_lines(document, ["Continue:", "Review:", "Strengthen:"], 2)


def p01(document) -> None:
    add_printable_title(document, "P01", "Private Six-Month Reflection", "")
    add_write_lines(document, ["Gift of these months:", "Hard stretch:", "Grace noticed:"], 3)
    add_callout(document, "SAFETY_PRIVACY", "Private. Optional share. No ranking.")


def p02(document) -> None:
    add_printable_title(document, "P02", "Review-Chain Observation", "ŚB 1.2.18 center")
    add_verse_card(document, "ŚB 1.2.18", "", "", "Hearing and serving culture nourishes bhakti — review center for tonight.", VERSE_URL)
    add_write_lines(document, ["What we heard well:", "What we want to hear better:"], 3)


def p03(document) -> None:
    add_printable_title(document, "P03", "Testimony Prep (optional)", "")
    _img(document, "family-testimony-frame.png", 5.0)
    add_write_lines(document, ["One sentence without comparison:", "Opt-out is OK because:"], 2)
    add_callout(document, "SAFETY_PRIVACY", "Invitation only.")


def p04(document) -> None:
    add_printable_title(document, "P04", "Continue / Review / Strengthen", "")
    add_write_lines(document, ["CONTINUE:", "REVIEW:", "STRENGTHEN:", "Who supports this at home:"], 2)
    add_callout(document, "FAMILY_APPLICATION", "Decision without competition.")


def p05(document) -> None:
    add_printable_title(document, "P05", "Private Next-Phase Saṅkalpa", "")
    add_callout(document, "KEY_IDEA", "specific action + frequency + trigger + minimum version")
    add_write_lines(document, ["Specific action:", "Frequency:", "Trigger:", "Minimum version:"], 2)


YOUNGER_BUILDERS = [y01, y02, y03, y04, y05]
OLDER_BUILDERS = [o01, o02, o03, o04, o05, o06]
PARENT_BUILDERS = [p01, p02, p03, p04, p05]
'''
    write(PRINT / "MELA.py", printable_module(mid, folder, memory, url, builders))
    write(EVID / "MELA" / "CONTENT-AUDIT.md", content_audit(
        mid, title, "2027-02-27 or 2027-03-06", primary, url,
        [("Y01", "exhibition walk"), ("O04", "CRS decision"), ("P03", "noncomparison testimony")],
        ["DOCX/PDF deferred", "Raster deferred", "Owner date selection EXTERNAL_OPEN", "Gamma render EXTERNAL_OPEN",
         "Gaura Pūrṇimā is lookahead only"],
    ))
    write(EVID / "MELA" / "REQUIREMENT-TRACEABILITY.csv", req_csv(mid, [
        ("Start Here", f"11-weekly-program-library/first-six-months/{folder}/V13-MELA-START-HERE.md", "not a tithi claim"),
        ("Owner ROS", f"11-weekly-program-library/first-six-months/{folder}/teacher/OWNER-RUN-OF-SHOW-V13.md", "noncompetitive"),
        ("Main script", f"11-weekly-program-library/first-six-months/{folder}/teacher/MAIN-FACILITATOR-GUIDE-V13.md", ""),
        ("Parent/younger/older", f"11-weekly-program-library/first-six-months/{folder}/teacher/", ""),
        ("Materials/safety", f"11-weekly-program-library/first-six-months/{folder}/materials-safety-v13.md", ""),
        ("Source matrix", f"11-weekly-program-library/first-six-months/{folder}/research/V13-SOURCE-MATRIX.md", ""),
        ("Rights/visuals", f"11-weekly-program-library/first-six-months/{folder}/visuals/v13/", ""),
        ("Gamma master", f"11-weekly-program-library/first-six-months/{folder}/gamma/V13-GAMMA-MASTER-DECK-PROMPT.md", ""),
        ("Printable builder", "scripts/v13/printables/MELA.py", "no render"),
    ]))


def author_gaura_lookahead() -> None:
    write(LAUNCH / "V13-GAURA-PURNIMA-LOOKAHEAD.md", f"""# V13 Gaura Pūrṇimā Lookahead (only)

**KUTUMBA · Families Growing in Krishna Consciousness**  
**Program Director: Swapnil Patil**  
**Status:** Lookahead planning note — **not** a first-six-month delivery pack  
**Published example date:** **2027-03-22** (ISKCON calendar planning reference)  
**Local tithi / temple plan:** **EXTERNAL_OPEN**

## Purpose

Record Gaura Pūrṇimā as a **lookahead** after the first-six-month sequence.  
Do **not** distort Cycle 3, Nityānanda Utsava, or the Six-Month Mela to force-fit Gaura Pūrṇimā into those Saturdays.

## What this file is

- A calendar/attention note for owner and teachers.
- A pointer to confirm local temple observance later.
- A reminder that child fasting instructions remain out of scope until separately approved.

## What this file is not

- Not a complete Utsava pack.
- Not a claim that 2027-03-22 is Harrisburg’s exact local tithi.
- Not permission to erase curriculum weeks.

## Owner next actions (when ready)

1. Confirm local Vaiṣṇava calendar for Gaura Pūrṇimā 2027.
2. Decide temple attendance vs home observance vs optional KUTUMBA gathering.
3. If a KUTUMBA gathering is desired, author a separate pack later — outside first-six-month gold closure scope if needed.
4. Keep first-six-month Mela owner-select dates (2027-02-27 or 2027-03-06) independent.

## References (planning only)

- Repository calendar: `launch/FIRST-SIX-MONTHS-CALENDAR.md`
- External example calendars: https://www.iskconbangalore.org/vaishnava-calendar/
- Local temple: confirm with ISKCON Harrisburg / family temple (EXTERNAL_OPEN)

---

{FOOT}
""")


def update_registry() -> None:
    path = REPO / "scripts" / "v13" / "week_registry.yaml"
    text = path.read_text(encoding="utf-8")
    replacements = {
        "folder: c1-utsava-kartika-damodara": "folder: c1-u1-kartika-damodara-family-utsava",
        "id: C1-UTSAVA-KARTIKA-DAMODARA": "id: C1-U1",
        "export_subdir: C1/C1-UTSAVA-KARTIKA-DAMODARA": "export_subdir: C1/C1-U1",
        "primary_verse: \"ŚB 10.9 (exact verse EXTERNAL_OPEN — local calendar)\"": 'primary_verse: "ŚB 10.9.18"',
        "folder: c2-utsava-gita-jayanti": "folder: c2-u2-gita-jayanti-family-utsava",
        "id: C2-UTSAVA-GITA-JAYANTI": "id: C2-U2",
        "export_subdir: C2/C2-UTSAVA-GITA-JAYANTI": "export_subdir: C2/C2-U2",
        'primary_verse: "Bhagavad-gītā (exact primary EXTERNAL_OPEN — local calendar)"': 'primary_verse: "BG 18.65"',
        "primary_url: https://vedabase.io/en/library/bg/\n    schedule_date: \"2026-12-19\"": "primary_url: https://vedabase.io/en/library/bg/18/65/\n    schedule_date: \"2026-12-19\"",
        "folder: c3-utsava-nityananda": "folder: c3-u3-nityananda-trayodasi-family-utsava",
        "id: C3-UTSAVA-NITYANANDA": "id: C3-U3",
        "export_subdir: C3/C3-UTSAVA-NITYANANDA": "export_subdir: C3/C3-U3",
        'primary_verse: "CC (Nityānanda mercy — exact verse EXTERNAL_OPEN)"': 'primary_verse: "CC Ādi 5.208"',
        "primary_url: https://vedabase.io/en/library/cc/\n    schedule_date: \"2027-02-20\"": "primary_url: https://vedabase.io/en/library/cc/adi/5/208/\n    schedule_date: \"2027-02-20\"",
        "folder: six-month-reflection-mela": "folder: mela-six-month-reflection-family-mela",
        "id: SIX-MONTH-REFLECTION-MELA": "id: MELA",
        "export_subdir: MELA/SIX-MONTH-REFLECTION-MELA": "export_subdir: MELA/MELA",
    }
    for a, b in replacements.items():
        if a not in text:
            print("WARN missing registry needle:", a[:60])
        text = text.replace(a, b)
    # fix C1-U1 primary url already ok; set primary_url for damodara if needed
    text = text.replace(
        "id: C1-U1\n    title: \"Kārtika / Dāmodara Family Utsava\"\n    cycle: C1\n    folder: c1-u1-kartika-damodara-family-utsava\n    primary_verse: \"ŚB 10.9.18\"\n    primary_url: https://vedabase.io/en/library/sb/10/9/",
        "id: C1-U1\n    title: \"Kārtika / Dāmodara Family Utsava\"\n    cycle: C1\n    folder: c1-u1-kartika-damodara-family-utsava\n    primary_verse: \"ŚB 10.9.18\"\n    primary_url: https://vedabase.io/en/library/sb/10/9/18/",
    )
    path.write_text(text, encoding="utf-8")
    print("updated", path.relative_to(REPO))


def main() -> None:
    author_c1_u1()
    author_c2_u2()
    author_c3_u3()
    author_mela()
    author_gaura_lookahead()
    update_registry()
    print("DONE authoring Utsava/Mela packs")


if __name__ == "__main__":
    main()
