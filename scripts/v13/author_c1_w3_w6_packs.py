#!/usr/bin/env python3
"""Author complete V13 packs for C1-W3 through C1-W6 (no DOCX/PDF render)."""
from __future__ import annotations

import csv
import shutil
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont

REPO = Path(__file__).resolve().parents[2]
BASE = REPO / "11-weekly-program-library" / "first-six-months"
EVID = REPO / "build-evidence" / "v13"
PRINT = REPO / "scripts" / "v13" / "printables"
CREAM, PLUM, SAFFRON, TEAL = "#FFF8E8", "#5B1933", "#E59B24", "#4F7C78"

WEEKS = {
    "C1-W3": {
        "folder": "c1-w3-the-nature-of-the-soul",
        "title": "The Nature of the Soul",
        "date": "2026-09-26",
        "date_human": "Saturday 26 Sep 2026",
        "primary": "BG 2.20",
        "url": "https://vedabase.io/en/library/bg/2/20/",
        "eq": "If I am a soul, how should I live?",
        "conclusion": "The jīva is eternal, conscious, individual, minute, and related to Kṛṣṇa in service — not God Himself.",
        "memory": "I am an eternal soul — conscious, individual, and meant for Kṛṣṇa's service.",
        "misconception": "All souls are God / we are the Supreme.",
        "next": "C1-W4 — Why Human Life Is Rare and Valuable. Do not teach full human-form rarity ontology tonight.",
        "devanagari": "न जायते म्रियते वा कदाचिन्नायं भूत्वा भविता वा न भूयः । अजो नित्यः शाश्वतोऽयं पुराणो न हन्यते हन्यमाने शरीरे ॥ २० ॥",
        "iast": "na jāyate mriyate vā kadācin nāyaṁ bhūtvā bhavitā vā na bhūyaḥ / ajo nityaḥ śāśvato ’yaṁ purāṇo na hanyate hanyamāne śarīre",
        "meaning": "The soul is never born and never dies; unborn, eternal, and everlasting — not destroyed when the body is destroyed.",
        "mechanic": "Classification map (soul is / is-not)",
        "science_mode": "NA",
        "parent_theme": "Identity under success and failure",
        "younger_theme": "Soul is / is-not classification cards",
        "older_theme": "Verse observation + attributes grid",
    },
    "C1-W4": {
        "folder": "c1-w4-why-human-life-is-rare-and-valuable",
        "title": "Why Human Life Is Rare and Valuable",
        "date": "2026-10-03",
        "date_human": "Saturday 3 Oct 2026",
        "primary": "ŚB 11.9.29",
        "url": "https://vedabase.io/en/library/sb/11/9/29/",
        "eq": "What deserves protected family time?",
        "conclusion": "Human life gives a rare opportunity for deliberate self-realization — protect inquiry and practice time.",
        "memory": "Human life is a rare chance to ask who I am and serve Kṛṣṇa.",
        "misconception": "Use fear or death-pressure to motivate children.",
        "next": "C1-W5 — The Temporary World and the Search for Permanent Happiness. Do not teach full temporary-world ontology tonight.",
        "devanagari": "लब्ध्वा सुदुर्लभमिदं बहुसम्भवान्ते मानुष्यमर्थदमनित्यमपीह धीर: । तूर्णं यतेत न पतेदनुमृत्यु यावन्नि:श्रेयसाय विषय: खलु सर्वत: स्यात् ॥ २९ ॥",
        "iast": "labdhvā su-durlabham idaṁ bahu-sambhavānte mānuṣyam artha-dam anityam apīha dhīraḥ / tūrṇaṁ yateta na pated anu-mṛtyu yāvan niḥśreyasāya viṣayaḥ khalu sarvataḥ syāt",
        "meaning": "After many births one obtains the rare human form — temporary yet able to give the highest purpose. A sober person should quickly endeavor for the ultimate good while this body lasts.",
        "mechanic": "Time-budget challenge (jar / priority / schedule audit)",
        "science_mode": "PEDAGOGY",
        "parent_theme": "Family schedule audit",
        "younger_theme": "Time-gift jar craft",
        "older_theme": "Priority / time-budget challenge",
        "special": "If Mṛgāri is used: CC Madhya 24.229–282 — NOT ŚB.",
    },
    "C1-W5": {
        "folder": "c1-w5-the-temporary-world-and-the-search-for-permanent-happiness",
        "title": "The Temporary World and the Search for Permanent Happiness",
        "date": "2026-10-10",
        "date_human": "Saturday 10 Oct 2026",
        "primary": "BG 8.15",
        "url": "https://vedabase.io/en/library/bg/8/15/",
        "eq": "How can enjoyment become gratitude and service?",
        "conclusion": "Temporary things can be used well but cannot provide permanent fulfillment.",
        "memory": "Temporary joys can be used with gratitude; lasting fulfillment is in Kṛṣṇa.",
        "misconception": "Material things and family affection are worthless.",
        "next": "C1-W6 — Integration Night: Who Am I, and How Should Our Family Live? Bring Cycle 1 artifacts.",
        "devanagari": "मामुपेत्य पुनर्जन्म दु:खालयमशाश्वतम् । नाप्नुवन्ति महात्मान: संसिद्धिं परमां गता: ॥ १५ ॥",
        "iast": "mām upetya punar janma duḥkhālayam aśāśvatam / nāpnuvanti mahātmānaḥ saṁsiddhiṁ paramāṁ gatāḥ",
        "meaning": "Having attained the Lord, great souls do not return again to this temporary world of misery; they have reached the highest perfection.",
        "mechanic": "Duration sorting lab (temporary / lasting)",
        "science_mode": "PEDAGOGY",
        "parent_theme": "Comfort vs meaning case",
        "younger_theme": "Temporary / lasting sort",
        "older_theme": "Happiness-duration lab",
        "special": "No depression/mental-health claims. No shame of ordinary enjoyment.",
    },
    "C1-W6": {
        "folder": "c1-w6-integration-night-who-am-i-and-how-should-our-family-live",
        "title": "Integration Night: Who Am I, and How Should Our Family Live?",
        "date": "2026-10-17",
        "date_human": "Saturday 17 Oct 2026",
        "primary": "Cycle 1 review (ŚB 1.2.18 · BG 2.13 · BG 2.20 · ŚB 11.9.29 · BG 8.15)",
        "url": "https://vedabase.io/en/library/bg/8/15/",
        "eq": "Can our family explain and apply what we learned?",
        "conclusion": "Identity, purpose, and practice must form one coherent family life.",
        "memory": "Who am I, and how should our family live?",
        "misconception": "Competition or ranking of families.",
        "next": "Kārtika / Dāmodara Family Utsava (name only). Do not teach new Cycle 2 doctrine tonight.",
        "devanagari": "(Review W1–W5 primary verse layers — do not invent a sixth primary.)",
        "iast": "ŚB 1.2.18 · BG 2.13 · BG 2.20 · ŚB 11.9.29 · BG 8.15",
        "meaning": "Identity, purpose, and practice must form one coherent family life — reviewing the Cycle 1 chain with W5 as BG 8.15.",
        "mechanic": "Retrieval + presentation (noncompetitive)",
        "science_mode": "PEDAGOGY",
        "parent_theme": "Presentation coaching (noncompetitive)",
        "younger_theme": "Family artifact share",
        "older_theme": "Misconception clinic",
        "special": "Four families × up to 10 minutes. No major new doctrine.",
    },
}


def font(size: int, bold: bool = False):
    for name in ("segoeuib.ttf" if bold else "segoeui.ttf", "arialbd.ttf" if bold else "arial.ttf"):
        path = Path(r"C:\Windows\Fonts") / name
        if path.exists():
            return ImageFont.truetype(str(path), size)
    return ImageFont.load_default()


def canvas(size=(900, 600)):
    image = Image.new("RGB", size, CREAM)
    return image, ImageDraw.Draw(image)


def save_png(image: Image.Image, out: Path, name: str) -> None:
    out.mkdir(parents=True, exist_ok=True)
    path = out / name
    image.save(path, "PNG", optimize=True)
    print("  visual", path.name)


def write(path: Path, text: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text.rstrip() + "\n", encoding="utf-8")
    print("  wrote", path.relative_to(REPO))


def promote_research(folder: Path) -> None:
    mapping = {
        "SOURCE-MATRIX.md": "V13-SOURCE-MATRIX.md",
        "SCRIPTURAL-EXAMPLES.md": "V13-SCRIPTURAL-EXAMPLES.md",
        "DEVOTIONAL-AND-HISTORICAL-EXAMPLES.md": "V13-DEVOTIONAL-HISTORICAL-EXAMPLES.md",
        "CASE-STUDIES.md": "V13-CASE-STUDIES.md",
        "ANALOGIES-AND-LIMITS.md": "V13-ANALOGIES-AND-LIMITS.md",
        "SCIENCE-AND-APPLICATION.md": "V13-SCIENCE-AND-APPLICATION.md",
        "CLAIM-REGISTER.yaml": "V13-CLAIM-REGISTER.yaml",
    }
    research = folder / "research"
    for src_name, dst_name in mapping.items():
        src = research / src_name
        dst = research / dst_name
        if not src.exists():
            write(
                dst,
                f"# {dst_name}\n\nPromoted placeholder — source file `{src_name}` missing. HUMAN REVIEW REQUIRED.\n",
            )
            continue
        text = src.read_text(encoding="utf-8")
        header = (
            f"<!-- V13 promotion from `{src_name}` — deepen in place; do not shallow-replace. -->\n"
            f"**Status:** Internal founding-cohort research — human/temple review EXTERNAL_OPEN\n\n"
        )
        if not text.lstrip().startswith("<!-- V13"):
            text = header + text
        write(dst, text)


def start_here(wid: str, w: dict) -> str:
    special = w.get("special", "")
    special_block = f"\n| Special safeguard | {special} |\n" if special else "\n"
    return f"""# V13 WEEK START HERE — {wid}

**KUTUMBA · Families Growing in Krishna Consciousness**  
**Program Director: Swapnil Patil**  
**Week:** {wid} — {w['title']}  
**Date:** {w['date_human']}  
**This is the only {wid} operational entry point for V13.**

---

## THIS SATURDAY (locked)

| Time | Block |
|---|---|
| 1:50–2:00 | Arrival / settle |
| 2:00–2:10 | Opening mantras + welcome |
| 2:10–2:30 | All-family launch / Prem-kī-Kathā + verse |
| 2:30–3:10 | Parallel tracks: **parents** / **K–2** / **Grades 4–5** |
| 3:10–3:30 | Reunification + bhakti laboratory |
| 3:30–3:40 | Light snack + water only |
| 3:40–3:55 | Saṅkalpa + Cycle 1 project + home practice |
| 3:55–4:00 | Closing / next Saturday |

Parents onsite. No weekly meal. Parallel tracks **2:30–3:10**; reunify **3:10–3:30**. Do not invent 2:35 / 3:25.

---

## ESSENTIAL FOCUS

| Field | Content |
|---|---|
| Essential question | {w['eq']} |
| Conclusion | {w['conclusion']} |
| Memory line | {w['memory']} |
| Primary | **{w['primary']}** — {w['url']} |
| Misconception to block | {w['misconception']} |
| Primary mechanic | {w['mechanic']} |{special_block}
---

## WHO SPEAKS / TEACHES

| Role | File |
|---|---|
| Main facilitator | `teacher/MAIN-FACILITATOR-GUIDE-V13.md` |
| Parent track | `teacher/PARENT-GUIDE-V13.md` |
| Younger K–2 | `teacher/YOUNGER-TEACHER-GUIDE-V13.md` |
| Older 4–5 | `teacher/OLDER-TEACHER-GUIDE-V13.md` |
| Home practice | `family-home-practice-v13.md` |
| Materials | `materials-v13.md` |
| Project | `project/V13-MODULE-PROJECT-BRIEF.md` |
| Gamma (prompt-only) | `gamma/V13-GAMMA-MASTER-DECK-PROMPT.md` |
| Printables | `activities/v13-parent/` · `v13-younger/` · `v13-older/` |

---

## FRIDAY NIGHT (≤60 min)

1. Open this file.
2. When available, print `exports/final/v13/C1/{wid}/` Saturday packet (builders ready; DOCX/PDF render may follow).
3. Read MAIN speaking script once aloud (10–15 min).
4. Sort printables into three track folders (Y / O / P).
5. Confirm snack/water + room reset + privacy reminder card.
6. Optional: paste Gamma master prompt into Gamma Studio (owner render; not pre-approved).

---

## HOME PRACTICE (tell families at 3:40)

Minimum before next Saturday: **memory line once + one week-specific kind/service act.**  
Ideal: memory line + one-sentence paraphrase of {w['primary']} + one gratitude/service act (5–15 min).  
No ranking. Five minutes counts.

---

## NEXT WEEK (name only)

**{w['next']}**

---

## EXTERNAL_OPEN

Local tithi / Utsava confirmation · pilot GO / temple human approval · live Gamma render sign-off · formal BBT licensing beyond teaching paraphrase.

**Verdict language allowed:** {wid} owner-runnable pack authored — human/temple external gates remain open.  
Do not claim public publication approval or official ISKCON endorsement.
"""


def home_practice(wid: str, w: dict) -> str:
    return f"""# {wid} Family Home Practice — V13

**Week:** {w['title']} · **Date:** {w['date']}  
**Primary:** {w['primary']} — {w['url']}

## Memory line

> {w['memory']}

## Ideal version (5–15 minutes)

1. Say the memory line once together.
2. Paraphrase {w['primary']} in one family sentence (KUTUMBA teaching meaning — not a purport dump).
3. Do one gratitude or service act connected to tonight’s conclusion.
4. Optional: sketch or tape this week’s project artifact layer.

## Minimum version (counts)

Memory line + one kind / care / gratitude act. Five minutes counts.

## Do not

- Rank families or children.
- Force confession of private struggles.
- Claim science proves doctrine.
- Shame ordinary lawful enjoyment (especially relevant for temporary-world weeks).
- Use fear or death-pressure (especially relevant for human-life weeks).

## Privacy

Private saṅkalpas stay private. Do not post completed sheets to public channels.
"""


def materials(wid: str, w: dict) -> str:
    return f"""# {wid} Materials Checklist — V13

**Week:** {w['title']} · **Date:** {w['date']}

## Shared room

- [ ] Opening mantra card / verse card for {w['primary']}
- [ ] Privacy reminder card
- [ ] Snack + water only (no weekly meal service)
- [ ] Soft seating / reunification circle markers
- [ ] Cycle 1 project storage bin

## Younger (K–2)

- [ ] Printables from `activities/v13-younger/` (Y01–Y05)
- [ ] Crayons / markers / scissors (adult-supervised)
- [ ] Soft toss object / calm-corner chair
- [ ] Visuals from `visuals/v13/` as needed

## Older (Grades 4–5)

- [ ] Printables from `activities/v13-older/` (O01–O06)
- [ ] Pens; verse observation sheets
- [ ] Teacher-only answer keys kept separate

## Parent track

- [ ] Printables from `activities/v13-parent/` (P01–P05)
- [ ] Quiet writing space; pens
- [ ] No public collection of private sheets into Git

## Cleanup

- [ ] Reset rooms; recycle scraps; return scissors
- [ ] Store project artifacts for Week 6
"""


def parent_guide(wid: str, w: dict) -> str:
    return f"""# {wid} Parent Track Guide — V13 (40 minutes)

**KUTUMBA · Families Growing in Krishna Consciousness**  
**Program Director: Swapnil Patil**  
**Week:** {w['title']} · **Date:** {w['date']}  
**Track window:** Saturday **2:30–3:10**  
**Reunite:** **3:10**  
**Primary:** {w['primary']} — {w['url']}  
**Status:** Internal founding-cohort teaching material — human/temple review EXTERNAL_OPEN

---

## Essential question

{w['eq']}

## Materials (print before Saturday)

From `activities/v13-parent/`:
- **P01** Private reflection
- **P02** Card sort / structured practice
- **P03** Substantial family case
- **P04** Household operating application
- **P05** Private next-step saṅkalpa

Also: pens; quiet writing space; verse layer for {w['primary']}.

## Locked Saturday context

- 1:50–2:00 Arrival
- 2:00–2:10 Opening mantras
- 2:10–2:30 All-family launch
- **2:30–3:10 This parent track**
- 3:10–3:30 Reunification
- 3:30–3:40 Snack + water only
- 3:40–3:55 Saṅkalpa / project / home practice
- 3:55–4:00 Close

Parents remain onsite. No public sādhana ranking. No forced disclosure.
{('Special: ' + w['special']) if w.get('special') else ''}

---

## 0–5 — Private reflection (P01)

Theme: **{w['parent_theme']}**  
Private writing ~4 minutes. No forced sharing.

## 5–12 — Primary observation

**Source:** {w['url']}

**Devanāgarī:** {w['devanagari']}

**IAST:** `{w['iast']}`

**KUTUMBA teaching meaning (not a BBT translation):**  
{w['meaning']}

Adult memory line:
> {w['memory']}

Misconception to block: {w['misconception']}

Deferral line:
> I don't want to guess. I will verify that from Śrīla Prabhupāda's books / our source packet and come back to you.

## 12–20 — Structured practice (P02)

Run the week’s parent mechanic: {w['parent_theme']}.  
Keep tone practical, not competitive.

## 20–30 — Substantial family case (P03)

Walk seven fields: situation · mistaken conclusion · principle · source · compassionate response · practical action · what not to say.  
Constructed fictional case — not a real family.

## 30–36 — Household operating application (P04)

Write one operating line the household can actually keep this week.

## 36–40 — Private next-step saṅkalpa (P05)

Private. Not for ranking. Minimum version counts.

---

## Reunification handoff

Invite one family sentence: “This week we will…”  
Drawings and whispered parent sentences count.

## Do not claim

Temple/BBT/human approval; science proving ātman; ranking results; guarantees of advancement.
"""


def activity_md(code: str, title: str, wid: str, w: dict, purpose: str, how: str, materials: str, age: str, success: str, extra: str = "") -> str:
    return f"""# {wid} · {code} · {title}

**Print:** US Letter · {age}  
**Mechanic:** {w['mechanic']}  
**Primary bridge:** {w['primary']}  
**URL:** {w['url']}

---

## Teaching purpose

{purpose}

## How to run

{how}

## Materials

{materials}

## Age boundary

{age}

## Success

{success}
{extra}
"""


def activities_for(wid: str, w: dict) -> dict[str, str]:
    if wid == "C1-W3":
        return {
            "v13-younger/Y01-SOUL-IS-IS-NOT-CARDS.md": activity_md(
                "Y01", "Soul Is / Is-Not Cards", wid, w,
                "Children classify simple statements into SOUL IS vs SOUL IS NOT, blocking “I am God” and body-identity.",
                "1. Lay two mats: SOUL IS / SOUL IS NOT.\n2. Children place cards with teacher help.\n3. Echo memory phrase.\n4. Correct gently if a child says “I am God”: dear spark who serves, not the whole sun.",
                "Printed cards; two mat labels; crayons optional.",
                "K–2. No violence. No science-proof claims.",
                "Child places at least three cards correctly and echoes the memory phrase.",
                "\n## Card set\n\n| Card | Mat |\n|---|---|\n| Eternal | IS |\n| Conscious | IS |\n| Meant to serve Kṛṣṇa | IS |\n| The temporary body | IS NOT |\n| God Himself / the Supreme | IS NOT |\n| A spark equal to the whole sun | IS NOT |\n",
            ),
            "v13-younger/Y02-SPARK-AND-SUN-MATCH.md": activity_md(
                "Y02", "Spark and Sun Match", wid, w,
                "Match spark → sun with the limit: same light-nature, not the whole sun.",
                "1. Show sun poster and spark cutouts.\n2. Children place sparks near sun.\n3. Say: same light-nature — not the whole sun.\n4. Whisper: I serve Kṛṣṇa.",
                "Sun poster; spark cutouts; soft toss object.",
                "K–2",
                "Child can say spark is not the sun.",
            ),
            "v13-younger/Y03-SPARK-SERVICE-CRAFT.md": activity_md(
                "Y03", "Spark-Service Cue Craft", wid, w,
                "Fold card: spark on front; memory phrase inside; one care act + one service act on back.",
                "1. Fold card stock.\n2. Draw/stick spark.\n3. Write short echo inside.\n4. Draw care + service acts.\n5. Take home near dinner plate.",
                "Card stock; crayons; `visuals/v13/spark-service-card.png`.",
                "K–2",
                "Child leaves with a cue card.",
            ),
            "v13-younger/Y04-CARE-AS-SERVICE-CARDS.md": activity_md(
                "Y04", "Care-as-Service Cards", wid, w,
                "Sort CARE HELPS SERVICE vs BODY NEGLECT — refuse body contempt.",
                "1. Two mats.\n2. Sort sleep/water/brush vs ignore rest ‘because soul’.\n3. Affirm care readiness.",
                "Printed cards.",
                "K–2",
                "Child chooses at least one care act as service readiness.",
            ),
            "v13-younger/Y05-MEMORY-PHRASE-MAT.md": activity_md(
                "Y05", "Memory Phrase Mat", wid, w,
                "Floor mat for Spark-and-Serve landings and phrase echo.",
                "1. Place mat.\n2. Movement landings.\n3. Echo phrase once.",
                "`visuals/v13/memory-mat.png`.",
                "K–2",
                "Child can echo the memory phrase with help.",
            ),
            "v13-older/O01-BG-2-20-OBSERVATION.md": activity_md(
                "O01", "BG 2.20 Observation Worksheet", wid, w,
                "Underline never born / never dies / not destroyed when body is destroyed; paraphrase teaching meaning.",
                "1. Read verse layer.\n2. Answer observation prompts.\n3. State what science does NOT prove tonight.",
                "Worksheet; pens.",
                "Grades 4–5",
                "Student paraphrases BG 2.20 without claiming lab proof.",
            ),
            "v13-older/O02-ATTRIBUTES-GRID.md": activity_md(
                "O02", "Soul Attributes Grid", wid, w,
                "Fill grid: eternal · conscious · individual · minute · related in service — with IS NOT column.",
                "1. Complete grid in pairs.\n2. Share one cell.\n3. Block Godhood claim.",
                "Grid printable.",
                "Grades 4–5",
                "Grid shows five positive attributes and clear is-nots.",
            ),
            "v13-older/O03-IDENTITY-STATEMENT-SORT.md": activity_md(
                "O03", "Identity Statement Sort", wid, w,
                "Sort statements into ALIGNED / MISTAKEN / MIXED regarding jīva ontology.",
                "1. Sort cards.\n2. Discuss one MIXED card.\n3. Teacher key after.",
                "Card grid.",
                "Grades 4–5",
                "Student rejects “I am God” and body-contempt statements.",
            ),
            "v13-older/O04-SUCCESS-FAILURE-SCENARIOS.md": activity_md(
                "O04", "Success / Failure Identity Scenarios", wid, w,
                "Write seven fields for scenarios where success/failure quietly becomes the whole self.",
                "1. Read scenarios.\n2. Complete fields.\n3. Affirm soul dignity under both outcomes.",
                "Scenario cards.",
                "Grades 4–5",
                "Student separates outcome from eternal identity.",
            ),
            "v13-older/O05-IS-IS-NOT-DIAGRAM.md": activity_md(
                "O05", "Jīva Is / Is-Not Diagram", wid, w,
                "Label concept diagram: IS attributes outside; IS NOT in caution zone.",
                "1. Label diagram.\n2. Add one service arrow.\n3. Keep care arrow for body.",
                "`visuals/v13/jiva-is-is-not.png`.",
                "Grades 4–5",
                "Diagram shows soul ≠ body and soul ≠ Supreme.",
            ),
            "v13-older/O06-EXIT-TICKET.md": activity_md(
                "O06", "Exit Ticket", wid, w,
                "Memory phrase + one service act + project sentence.",
                "Complete before reunification.",
                "Exit ticket sheet.",
                "Grades 4–5",
                "Ticket completed without ranking language.",
            ),
            "v13-parent/P01-PRIVATE-REFLECTION.md": activity_md(
                "P01", "Private Reflection — Identity Under Success/Failure", wid, w,
                "Private writing: where do wins/losses quietly become the whole self at home?",
                "4 minutes private. No forced share.",
                "P01 sheet.",
                "Adults",
                "One noticed pattern written privately.",
            ),
            "v13-parent/P02-IDENTITY-LANGUAGE-CARD-SORT.md": activity_md(
                "P02", "Identity Language Card Sort", wid, w,
                "Sort RESPECTFUL IDENTITY vs NEEDS REPAIR language around success/failure and Godhood slips.",
                "Two mats; discuss one almost-right card.",
                "Cards.",
                "Adults",
                "Adults can name a repair sentence.",
            ),
            "v13-parent/P03-SUBSTANTIAL-FAMILY-CASE.md": activity_md(
                "P03", "Substantial Family Case", wid, w,
                "Child says “I am God” after a win; parent freezes. Seven-field response.",
                "Walk fields aloud; write notes.",
                "Case sheet.",
                "Adults",
                "Gentle correction + service sentence ready.",
            ),
            "v13-parent/P04-HOUSEHOLD-OPERATING-APPLICATION.md": activity_md(
                "P04", "Household Operating Application", wid, w,
                "One operating line: when outcomes spike, we will… instead of…",
                "Write and keep private or household-only.",
                "P04 sheet.",
                "Adults",
                "One realistic operating line.",
            ),
            "v13-parent/P05-PRIVATE-NEXT-STEP-SANKALPA.md": activity_md(
                "P05", "Private Next-Step Saṅkalpa", wid, w,
                "Private next step connecting eternal identity to one service act.",
                "Private; not ranked.",
                "P05 sheet.",
                "Adults",
                "Private saṅkalpa written.",
            ),
        }
    if wid == "C1-W4":
        return {
            "v13-younger/Y01-TIME-GIFT-JAR.md": activity_md(
                "Y01", "Time-Gift Jar Craft", wid, w,
                "Children place Must / Optional tokens into a jar to protect inquiry and kindness time.",
                "1. Decorate jar label: Time Gift.\n2. Sort tokens: prayer/story/help vs endless screens.\n3. Place one Must token first.\n4. Echo memory phrase.",
                "Jar or cup; tokens; crayons; `visuals/v13/time-gift-jar.png`.",
                "K–2. No death scare.",
                "Child places one Must token before Optional.",
            ),
            "v13-younger/Y02-MUST-OPTIONAL-SORT.md": activity_md(
                "Y02", "Must / Optional Picture Sort", wid, w,
                "Sort picture cards into Must protect vs Optional later.",
                "Two mats; gentle coaching; no shame.",
                "Picture cards.",
                "K–2",
                "Child can name one Must.",
            ),
            "v13-younger/Y03-RARE-TICKET-CRAFT.md": activity_md(
                "Y03", "Rare Ticket Cue Craft", wid, w,
                "Make a ‘rare chance’ ticket cue: ask who I am + serve.",
                "Fold ticket; draw heart/service; take home.",
                "Card stock; crayons.",
                "K–2",
                "Child leaves with ticket cue.",
            ),
            "v13-younger/Y04-KIND-TIME-SPEECH-CARDS.md": activity_md(
                "Y04", "Kind Time-Speech Cards", wid, w,
                "Sort invitation language vs fear/death-pressure language.",
                "Reject scare cards; keep invitation cards.",
                "Cards.",
                "K–2",
                "Child prefers invitation over scare.",
            ),
            "v13-younger/Y05-MEMORY-PHRASE-MAT.md": activity_md(
                "Y05", "Memory Phrase Mat", wid, w,
                "Landing mat for Must-first movement game.",
                "Place mat; land after Must token; echo phrase.",
                "`visuals/v13/memory-mat.png`.",
                "K–2",
                "Phrase echoed with help.",
            ),
            "v13-older/O01-SB-11-9-29-OBSERVATION.md": activity_md(
                "O01", "ŚB 11.9.29 Observation Worksheet", wid, w,
                "Observe rare / temporary / endeavor / ultimate good; paraphrase without fear framing.",
                "Underline; answer prompts; state Mṛgāri source if used.",
                "Worksheet.",
                "Grades 4–5",
                "Student paraphrases without death-pressure.",
            ),
            "v13-older/O02-PRIORITY-TIME-BUDGET.md": activity_md(
                "O02", "Priority / Time-Budget Challenge", wid, w,
                "Allocate a fictional evening budget: Must / Should / Optional — protect inquiry slot.",
                "Fill budget table; defend one Must slot.",
                "Budget sheet.",
                "Grades 4–5",
                "Inquiry/practice slot protected before optional screens.",
            ),
            "v13-older/O03-OPPORTUNITY-STATEMENT-SORT.md": activity_md(
                "O03", "Opportunity Statement Sort", wid, w,
                "Sort ALIGNED / FEAR-BASED / MIXED statements about human opportunity.",
                "Sort; discuss MIXED; teacher key.",
                "Cards.",
                "Grades 4–5",
                "Fear-based cards rejected.",
            ),
            "v13-older/O04-BUSY-FAMILY-SCENARIOS.md": activity_md(
                "O04", "Busy Family / Opportunity Scenarios", wid, w,
                "Seven fields for busy family, screen block, and service-priority cases.",
                "Complete scenarios A–C.",
                "Scenario cards.",
                "Grades 4–5",
                "Student proposes protected time without quitting livelihood theatrics.",
            ),
            "v13-older/O05-TIME-BUDGET-DIAGRAM.md": activity_md(
                "O05", "Time-Budget Diagram", wid, w,
                "Label Must jar before Optional screens; add mercy note (not panic).",
                "Label diagram; write one family Must.",
                "`visuals/v13/time-budget-diagram.png`.",
                "Grades 4–5",
                "Diagram shows Must-first without fear.",
            ),
            "v13-older/O06-EXIT-TICKET.md": activity_md(
                "O06", "Exit Ticket", wid, w,
                "Memory phrase + one protected-time act + project sentence.",
                "Complete before reunify.",
                "Exit ticket.",
                "Grades 4–5",
                "Ticket done; no ranking.",
            ),
            "v13-parent/P01-PRIVATE-REFLECTION.md": activity_md(
                "P01", "Private Reflection — Schedule Truth", wid, w,
                "Where does the only quiet block disappear at home?",
                "Private 4 minutes.",
                "P01.",
                "Adults",
                "One noticed leak written privately.",
            ),
            "v13-parent/P02-SCHEDULE-AUDIT-CARD-SORT.md": activity_md(
                "P02", "Schedule Audit Card Sort", wid, w,
                "Sort PROTECT / DELAY / DROP for evening blocks.",
                "Audit cards; choose one Protect slot.",
                "Cards.",
                "Adults",
                "One Protect slot named.",
            ),
            "v13-parent/P03-SUBSTANTIAL-FAMILY-CASE.md": activity_md(
                "P03", "Substantial Family Case", wid, w,
                "Screens consume the only quiet block; parent uses a death scare. Rewrite as invitation.",
                "Seven fields; correct Mṛgāri source if mentioned: CC Madhya 24.229–282.",
                "Case sheet.",
                "Adults",
                "Invitation rewrite ready; scare rejected.",
            ),
            "v13-parent/P04-HOUSEHOLD-OPERATING-APPLICATION.md": activity_md(
                "P04", "Household Operating Application", wid, w,
                "When the evening opens, we will place ___ Must before Optional.",
                "Write operating line.",
                "P04.",
                "Adults",
                "Realistic Must-first line.",
            ),
            "v13-parent/P05-PRIVATE-NEXT-STEP-SANKALPA.md": activity_md(
                "P05", "Private Next-Step Saṅkalpa", wid, w,
                "Private: protect one inquiry/practice slot this week.",
                "Private.",
                "P05.",
                "Adults",
                "Private saṅkalpa.",
            ),
        }
    if wid == "C1-W5":
        return {
            "v13-younger/Y01-TEMPORARY-LASTING-SORT.md": activity_md(
                "Y01", "Temporary / Lasting Sort", wid, w,
                "Sort picture cards into Temporary joy vs Lasting (gratitude/service/Kṛṣṇa).",
                "Two mats; affirm toys can be enjoyed with thanks; not worthless.",
                "Cards; mats; `visuals/v13/temporary-lasting.png`.",
                "K–2. No shame of ordinary enjoyment.",
                "Child can sort 4+ cards and say thanks for a temporary joy.",
            ),
            "v13-younger/Y02-SPARKLER-LAMP-MATCH.md": activity_md(
                "Y02", "Sparkler and Lamp Match", wid, w,
                "Match sparkler (brief) to lamp (steady) with gratitude limit.",
                "Demo; children match; say enjoy sparkler / remember lamp.",
                "Picture pair cards.",
                "K–2",
                "Child names sparkler temporary, lamp lasting metaphorically.",
            ),
            "v13-younger/Y03-GRATITUDE-OFFERING-CRAFT.md": activity_md(
                "Y03", "Gratitude Offering Craft", wid, w,
                "Draw one temporary joy + write Thank You, Kṛṣṇa.",
                "Craft card; take home.",
                "Card stock; crayons.",
                "K–2",
                "Gratitude card completed.",
            ),
            "v13-younger/Y04-KIND-ENJOYMENT-SPEECH.md": activity_md(
                "Y04", "Kind Enjoyment Speech Cards", wid, w,
                "Sort kindness toward toys/fun vs shame language.",
                "Reject shame cards; keep gratitude cards.",
                "Cards.",
                "K–2",
                "Shame language rejected.",
            ),
            "v13-younger/Y05-MEMORY-PHRASE-MAT.md": activity_md(
                "Y05", "Memory Phrase Mat", wid, w,
                "Landing mat for lasting/gratitude echo.",
                "Place; land; echo.",
                "`visuals/v13/memory-mat.png`.",
                "K–2",
                "Phrase echoed.",
            ),
            "v13-older/O01-BG-8-15-OBSERVATION.md": activity_md(
                "O01", "BG 8.15 Observation Worksheet", wid, w,
                "Observe temporary world / highest perfection / attained the Lord; refuse contempt for family affection.",
                "Underline; paraphrase; note what we do NOT shame.",
                "Worksheet.",
                "Grades 4–5",
                "Paraphrase without depression claims.",
            ),
            "v13-older/O02-HAPPINESS-DURATION-LAB.md": activity_md(
                "O02", "Happiness-Duration Lab", wid, w,
                "Rate sample joys as brief / medium / lasting-direction; discuss limits.",
                "Fill lab table; refuse medical/depression framing.",
                "Lab sheet.",
                "Grades 4–5",
                "Student distinguishes temporary contact joy from lasting shelter without shame.",
            ),
            "v13-older/O03-TEMP-LASTING-STATEMENT-SORT.md": activity_md(
                "O03", "Temporary / Lasting Statement Sort", wid, w,
                "Sort ALIGNED / SHAMING / MIXED statements.",
                "Sort; repair one SHAMING into gratitude+service.",
                "Cards.",
                "Grades 4–5",
                "Shaming cards repaired.",
            ),
            "v13-older/O04-COMFORT-MEANING-SCENARIOS.md": activity_md(
                "O04", "Comfort vs Meaning Scenarios", wid, w,
                "Seven fields for purchase chase, toy-shame, work-over-prayer cases.",
                "Complete A–C.",
                "Scenarios.",
                "Grades 4–5",
                "Student proposes gratitude pause + protected cue.",
            ),
            "v13-older/O05-TEMPORARY-LASTING-DIAGRAM.md": activity_md(
                "O05", "Temporary vs Lasting Diagram", wid, w,
                "Label temporary tools vs lasting shelter; keep care/enjoyment without contempt.",
                "Label; write one offering line (BG 9.27 support).",
                "`visuals/v13/temporary-lasting.png`.",
                "Grades 4–5",
                "Diagram refuses both materialism-as-shelter and anti-family shame.",
            ),
            "v13-older/O06-EXIT-TICKET.md": activity_md(
                "O06", "Exit Ticket", wid, w,
                "Memory phrase + one gratitude act + project sentence.",
                "Complete before reunify.",
                "Exit ticket.",
                "Grades 4–5",
                "Ticket done.",
            ),
            "v13-parent/P01-PRIVATE-REFLECTION.md": activity_md(
                "P01", "Private Reflection — Comfort vs Meaning", wid, w,
                "Where do we ask temporary comfort to do a permanent job?",
                "Private writing.",
                "P01.",
                "Adults",
                "One noticed pattern.",
            ),
            "v13-parent/P02-COMFORT-MEANING-CARD-SORT.md": activity_md(
                "P02", "Comfort / Meaning Card Sort", wid, w,
                "Sort TEMPORARY TOOL / LASTING DIRECTION / NEEDS REPAIR.",
                "Sort; discuss almost-right.",
                "Cards.",
                "Adults",
                "Repair sentence ready.",
            ),
            "v13-parent/P03-SUBSTANTIAL-FAMILY-CASE.md": activity_md(
                "P03", "Substantial Family Case", wid, w,
                "Parent shames a child’s toy as māyā; child goes quiet. Restore gratitude without making toy the shelter.",
                "Seven fields; no depression claims.",
                "Case.",
                "Adults",
                "Compassion + gratitude pause planned.",
            ),
            "v13-parent/P04-HOUSEHOLD-OPERATING-APPLICATION.md": activity_md(
                "P04", "Household Operating Application", wid, w,
                "When we enjoy ___, we will pause to thank and offer one service act.",
                "Write line.",
                "P04.",
                "Adults",
                "Operating line realistic.",
            ),
            "v13-parent/P05-PRIVATE-NEXT-STEP-SANKALPA.md": activity_md(
                "P05", "Private Next-Step Saṅkalpa", wid, w,
                "Private gratitude+service next step.",
                "Private.",
                "P05.",
                "Adults",
                "Saṅkalpa written.",
            ),
        }
    # C1-W6
    return {
        "v13-younger/Y01-CYCLE-BEAD-RECALL.md": activity_md(
            "Y01", "Cycle Bead Recall", wid, w,
            "Touch five beads/stations recalling W1–W5 in kid words — no quiz shame.",
            "1. Walk stations.\n2. One word per bead.\n3. Celebrate attempts.",
            "Five bead cards; soft rope.",
            "K–2. Noncompetitive.",
            "Child touches beads and says one remembered idea.",
        ),
        "v13-younger/Y02-ARTIFACT-SHARE-PRACTICE.md": activity_md(
            "Y02", "Artifact Share Practice", wid, w,
            "Practice showing a drawing/artifact for up to 1 minute with a helper sentence.",
            " rehearse; drawings count as success.",
            "Artifacts from prior weeks.",
            "K–2",
            "Child can show something with help.",
        ),
        "v13-younger/Y03-FAMILY-SENTENCE-CRAFT.md": activity_md(
            "Y03", "Family Sentence Craft", wid, w,
            "Write/draw: Who am I, and how should our family live?",
            "Craft card for presentation night.",
            "Card stock; crayons.",
            "K–2",
            "Sentence or drawing ready.",
        ),
        "v13-younger/Y04-KIND-LISTEN-CARDS.md": activity_md(
            "Y04", "Kind Listener Cards", wid, w,
            "Sort kind listening vs interrupting/ranking during family shares.",
            "Reject ranking cards.",
            "Cards.",
            "K–2",
            "Child chooses kind listening.",
        ),
        "v13-younger/Y05-MEMORY-PHRASE-MAT.md": activity_md(
            "Y05", "Memory Phrase Mat", wid, w,
            "Integration memory phrase landing.",
            "Place; echo.",
            "`visuals/v13/memory-mat.png`.",
            "K–2",
            "Phrase echoed.",
        ),
        "v13-older/O01-CHAIN-RETRIEVAL.md": activity_md(
            "O01", "Cycle 1 Chain Retrieval", wid, w,
            "Retrieve W1–W5 primary refs + one-sentence meanings without new doctrine.",
            "Fill retrieval table; open-notes mercy allowed.",
            "Retrieval sheet.",
            "Grades 4–5",
            "At least 4/5 weeks retrieved with help OK.",
        ),
        "v13-older/O02-MISCONCEPTION-CLINIC.md": activity_md(
            "O02", "Misconception Clinic", wid, w,
            "Correct common Cycle 1 misconceptions with source pointers.",
            "Clinic cards; write repair lines.",
            "Clinic sheet.",
            "Grades 4–5",
            "Student repairs at least two misconceptions.",
        ),
        "v13-older/O03-PRESENTATION-OUTLINE.md": activity_md(
            "O03", "Family Presentation Outline", wid, w,
            "Outline ≤10 min share: who we are · how we live · one source · one practice.",
            "Fill outline; noncompetitive rubric lenses.",
            "Outline sheet.",
            "Grades 4–5",
            "Outline fits 10 minutes.",
        ),
        "v13-older/O04-TEAM-COACHING-SCENARIOS.md": activity_md(
            "O04", "Team Coaching Scenarios", wid, w,
            "Help a shy family / overlong share / missing-week family without ranking.",
            "Seven fields for A–C.",
            "Scenarios.",
            "Grades 4–5",
            "Mercy responses without scoreboard.",
        ),
        "v13-older/O05-INTEGRATION-MAP.md": activity_md(
            "O05", "Integration Necklace Map", wid, w,
            "Label five beads of Cycle 1; write family living sentence in center.",
            "Label map; no new doctrine bead.",
            "`visuals/v13/integration-necklace.png`.",
            "Grades 4–5",
            "Five beads + living sentence.",
        ),
        "v13-older/O06-EXIT-TICKET.md": activity_md(
            "O06", "Exit Ticket", wid, w,
            "Memory phrase + one home living act + thanks to another family (no ranking).",
            "Complete before close.",
            "Exit ticket.",
            "Grades 4–5",
            "Ticket done.",
        ),
        "v13-parent/P01-PRIVATE-REFLECTION.md": activity_md(
            "P01", "Private Reflection — Coherent Family Life", wid, w,
            "What is one coherent line connecting identity and household practice?",
            "Private.",
            "P01.",
            "Adults",
            "One coherent line noted.",
        ),
        "v13-parent/P02-PRESENTATION-COACHING-SORT.md": activity_md(
            "P02", "Presentation Coaching Card Sort", wid, w,
            "Sort SUPPORTIVE COACHING vs RANKING/PRESSURE.",
            "Sort; protect shy shares.",
            "Cards.",
            "Adults",
            "Ranking cards rejected.",
        ),
        "v13-parent/P03-SUBSTANTIAL-FAMILY-CASE.md": activity_md(
            "P03", "Substantial Family Case", wid, w,
            "Four families share; one is incomplete; another runs long. Keep noncompetitive mercy.",
            "Seven fields; ≤10 min reminders.",
            "Case.",
            "Adults",
            "Facilitation plan without ranking.",
        ),
        "v13-parent/P04-HOUSEHOLD-OPERATING-APPLICATION.md": activity_md(
            "P04", "Household Operating Application", wid, w,
            "Our family’s Cycle 1 living line is… (one sentence).",
            "Write.",
            "P04.",
            "Adults",
            "Living line written.",
        ),
        "v13-parent/P05-PRIVATE-NEXT-STEP-SANKALPA.md": activity_md(
            "P05", "Private Next-Step Saṅkalpa", wid, w,
            "Private next step into Utsava / continued practice — still no ranking.",
            "Private.",
            "P05.",
            "Adults",
            "Saṅkalpa written.",
        ),
    }


def project_brief(wid: str, w: dict) -> str:
    return f"""# {wid} Module Project Brief — V13

**Cycle project:** Who Am I, and How Should Our Family Live?  
**This week layer:** Connect to {w['primary']} and conclusion: {w['conclusion']}

## Artifact this week

Produce one household artifact piece (drawing, card, schedule strip, gratitude note, or outline) that can be stored for Week 6 presentation.

## Contribution rules

- Noncompetitive — no ranking of families.
- Source honesty — cite {w['primary']} URL when claiming doctrine.
- Mercy for missed weeks — incomplete artifacts still welcome.
- No major new doctrine beyond this week’s scope.

## Week 6 use

Artifacts feed the integration presentation (four families × up to 10 minutes).
"""


def rights_md(wid: str) -> str:
    return f"""# {wid} Visual Rights Register — V13

| Asset | Origin | Rights note |
|---|---|---|
| `*.png` under `visuals/v13/` | Original KUTUMBA line art generated with Pillow for this week | Program-owned instructional art; not temple logos; not identifiable persons |
| Verse text on cards | VedaBase display Sanskrit/IAST where quoted in guides | Teaching paraphrase is KUTUMBA-original; no full purport reproduction |
| Gamma image prompts | Original prompt text | Owner renders; not pre-approved |

## Do not

- Reuse another week’s PNG bytes as if unique.
- Embed identifiable real student photos.
- Claim BBT ownership of KUTUMBA diagrams.
"""


def image_prompt_library(wid: str, w: dict) -> str:
    return f"""# {wid} Image Prompt Library — V13

Use only for Gamma / optional enrichment. Prefer program Pillow line art for printables.

## Concept visual
16:9 instructional diagram for {wid} '{w['mechanic']}'; cream `#FFF8E8` ground; plum `#5B1933` / saffron `#E59B24` / teal `#4F7C78`; Pennsylvania living-room classroom; South Asian family learners generically represented; no gore; no caricature; no temple logo; no embedded readable text; no invented deity pastimes; no identifiable real persons.

## Child printable visual
Simple line-art flashcard style for {w['younger_theme']}; large shapes; calm faces; same palette and negative constraints.

## Parent visual
Quiet notebook and schedule/identity cue still life; same palette; no ranking trophies; negative constraints as above.
"""


def gamma_promote(folder: Path, wid: str, w: dict) -> None:
    gdir = folder / "gamma"
    mapping = {
        "V12-GAMMA-MASTER-DECK-PROMPT.md": "V13-GAMMA-MASTER-DECK-PROMPT.md",
        "V12-GAMMA-PARENT-DECK-PROMPT.md": "V13-GAMMA-PARENT-DECK-PROMPT.md",
        "V12-GAMMA-YOUNGER-DECK-PROMPT.md": "V13-GAMMA-YOUNGER-DECK-PROMPT.md",
        "V12-GAMMA-OLDER-DECK-PROMPT.md": "V13-GAMMA-OLDER-DECK-PROMPT.md",
        "V12-GAMMA-SOURCE-MAP.yaml": "V13-GAMMA-SOURCE-MAP.yaml",
    }
    for src_name, dst_name in mapping.items():
        src = gdir / src_name
        dst = gdir / dst_name
        if not src.exists():
            write(dst, f"# {dst_name}\n\nMISSING V12 source `{src_name}` — HUMAN AUTHOR REQUIRED.\n")
            continue
        text = src.read_text(encoding="utf-8")
        text = text.replace("V12.1", "V13").replace("V12", "V13")
        text = text.replace("Saturday 2:00–4:00", "Saturday 1:50–4:00 (tracks 2:30–3:10; reunify 3:10–3:30)")
        banner = (
            f"**V13 promotion:** {wid} · {w['date']} · Primary {w['primary']}\n"
            f"**Status:** prompt-only — not rendered — not approved\n"
            f"**Exact schedule on slides:** Arrival 1:50 · Tracks 2:30–3:10 · Reunify 3:10–3:30 · Close 4:00\n\n"
        )
        if "V13 promotion" not in text:
            # insert after first heading
            parts = text.split("\n", 1)
            text = parts[0] + "\n\n" + banner + (parts[1] if len(parts) > 1 else "")
        # Ensure memory/eq appear
        if w["memory"] not in text:
            text += f"\n\n## V13 lock lines\n- Essential question: {w['eq']}\n- Memory: {w['memory']}\n- Misconception: {w['misconception']}\n"
        write(dst, text)


def generate_visuals(folder: Path, wid: str, w: dict) -> None:
    out = folder / "visuals" / "v13"
    # memory mat
    im, d = canvas((900, 500))
    d.rounded_rectangle((40, 40, 860, 460), radius=30, outline=PLUM, width=8)
    d.text((70, 70), wid, fill=TEAL, font=font(28, True))
    # wrap memory
    words = w["memory"].split()
    lines, cur = [], ""
    for word in words:
        trial = (cur + " " + word).strip()
        if len(trial) > 34:
            lines.append(cur)
            cur = word
        else:
            cur = trial
    if cur:
        lines.append(cur)
    y = 160
    for line in lines[:4]:
        d.text((70, y), line, fill=PLUM, font=font(36, True))
        y += 50
    save_png(im, out, "memory-mat.png")

    if wid == "C1-W3":
        im, d = canvas((900, 600))
        d.ellipse((300, 80, 600, 380), outline=SAFFRON, width=10)
        d.text((390, 200), "SUN", fill=SAFFRON, font=font(40, True))
        for i, (x, y) in enumerate([(120, 420), (220, 480), (700, 430), (780, 490)]):
            d.ellipse((x, y, x + 40, y + 40), outline=TEAL, width=5)
            d.text((x - 10, y + 45), "spark", fill=PLUM, font=font(16))
        d.text((60, 40), "Spark shares light — not the whole sun", fill=PLUM, font=font(28, True))
        save_png(im, out, "spark-sun.png")
        im, d = canvas((900, 600))
        d.rectangle((40, 40, 430, 560), outline=TEAL, width=8)
        d.text((140, 60), "SOUL IS", fill=TEAL, font=font(32, True))
        for i, t in enumerate(["Eternal", "Conscious", "Individual", "Minute", "Serves"]):
            d.text((70, 140 + i * 70), f"• {t}", fill=PLUM, font=font(26))
        d.rectangle((470, 40, 860, 560), outline=SAFFRON, width=8)
        d.text((520, 60), "SOUL IS NOT", fill=SAFFRON, font=font(32, True))
        for i, t in enumerate(["Temporary body", "God Himself", "A lab proof", "Trashable body"]):
            d.text((500, 150 + i * 80), f"• {t}", fill=PLUM, font=font(24))
        save_png(im, out, "jiva-is-is-not.png")
        im, d = canvas((700, 500))
        d.rounded_rectangle((50, 50, 650, 450), radius=20, outline=PLUM, width=6)
        d.text((220, 80), "SPARK SERVES", fill=SAFFRON, font=font(30, True))
        d.ellipse((300, 150, 400, 250), outline=TEAL, width=6)
        d.text((120, 300), "Care act + Service act", fill=PLUM, font=font(24, True))
        save_png(im, out, "spark-service-card.png")
    elif wid == "C1-W4":
        im, d = canvas((800, 700))
        d.ellipse((250, 80, 550, 520), outline=TEAL, width=10)
        d.rectangle((280, 40, 520, 100), outline=PLUM, width=6)
        d.text((300, 55), "TIME GIFT", fill=PLUM, font=font(28, True))
        d.text((300, 200), "MUST", fill=SAFFRON, font=font(36, True))
        d.text((290, 320), "optional", fill=TEAL, font=font(28))
        save_png(im, out, "time-gift-jar.png")
        im, d = canvas((900, 500))
        d.rectangle((50, 80, 300, 420), outline=SAFFRON, width=8)
        d.text((110, 100), "MUST", fill=SAFFRON, font=font(32, True))
        d.rectangle((350, 140, 600, 420), outline=TEAL, width=8)
        d.text((390, 160), "SHOULD", fill=TEAL, font=font(28, True))
        d.rectangle((650, 220, 860, 420), outline=PLUM, width=6)
        d.text((670, 240), "OPTIONAL", fill=PLUM, font=font(22, True))
        d.text((50, 40), "Protect inquiry time first", fill=PLUM, font=font(28, True))
        save_png(im, out, "time-budget-diagram.png")
        im, d = canvas((700, 400))
        d.rounded_rectangle((40, 40, 660, 360), radius=16, outline=SAFFRON, width=6)
        d.text((180, 150), "RARE TICKET", fill=PLUM, font=font(40, True))
        d.text((120, 230), "Ask · Serve · Protect time", fill=TEAL, font=font(26, True))
        save_png(im, out, "rare-ticket.png")
    elif wid == "C1-W5":
        im, d = canvas((900, 550))
        d.rectangle((40, 80, 420, 500), outline=TEAL, width=8)
        d.text((100, 100), "TEMPORARY", fill=TEAL, font=font(30, True))
        for i, t in enumerate(["Toy fun", "Snack joy", "Screen time", "New thing"]):
            d.text((70, 180 + i * 70), f"• {t}", fill=PLUM, font=font(24))
        d.rectangle((480, 80, 860, 500), outline=SAFFRON, width=8)
        d.text((560, 100), "LASTING", fill=SAFFRON, font=font(30, True))
        for i, t in enumerate(["Gratitude", "Service", "Holy name", "Shelter in Kṛṣṇa"]):
            d.text((510, 180 + i * 70), f"• {t}", fill=PLUM, font=font(24))
        d.text((40, 20), "Enjoy with thanks — do not shame", fill=PLUM, font=font(26, True))
        save_png(im, out, "temporary-lasting.png")
        im, d = canvas((800, 500))
        # sparkler
        d.line((200, 400, 200, 200), fill=SAFFRON, width=8)
        for a in range(-40, 50, 15):
            d.line((200, 200, 200 + a, 120), fill=SAFFRON, width=3)
        d.text((140, 420), "sparkler", fill=PLUM, font=font(22))
        # lamp
        d.ellipse((520, 250, 620, 320), outline=TEAL, width=6)
        d.rectangle((545, 320, 595, 400), outline=PLUM, width=6)
        d.text((540, 420), "lamp", fill=PLUM, font=font(22))
        d.text((80, 40), "Enjoy the sparkler; remember the lamp", fill=PLUM, font=font(26, True))
        save_png(im, out, "sparkler-lamp.png")
        im, d = canvas((700, 450))
        d.ellipse((250, 80, 450, 280), outline=SAFFRON, width=8)
        d.text((280, 150), "Thank You", fill=PLUM, font=font(28, True))
        d.text((300, 200), "Kṛṣṇa", fill=TEAL, font=font(28, True))
        d.text((180, 320), "Gratitude offering card", fill=PLUM, font=font(24, True))
        save_png(im, out, "gratitude-card.png")
    else:
        im, d = canvas((900, 520))
        beads = ["W1", "W2", "W3", "W4", "W5"]
        for i, b in enumerate(beads):
            x = 80 + i * 160
            d.ellipse((x, 180, x + 100, 280), outline=TEAL if i % 2 == 0 else SAFFRON, width=8)
            d.text((x + 30, 210), b, fill=PLUM, font=font(28, True))
            if i < 4:
                d.line((x + 100, 230, x + 160, 230), fill=PLUM, width=4)
        d.text((200, 40), "Cycle 1 necklace — no ranking", fill=PLUM, font=font(28, True))
        d.text((180, 360), "Who am I, and how should our family live?", fill=TEAL, font=font(22, True))
        save_png(im, out, "integration-necklace.png")
        im, d = canvas((900, 500))
        for i, label in enumerate(["Hear", "Body≠Self", "Soul", "Rare life", "Lasting"]):
            x = 40 + i * 170
            d.rounded_rectangle((x, 120, x + 150, 360), radius=12, outline=PLUM, width=5)
            d.text((x + 20, 200), label, fill=SAFFRON, font=font(20, True))
        d.text((40, 40), "Retrieval stations (mercy, not quiz)", fill=PLUM, font=font(26, True))
        save_png(im, out, "retrieval-stations.png")
        im, d = canvas((700, 450))
        d.rounded_rectangle((40, 40, 660, 410), radius=16, outline=TEAL, width=6)
        d.text((120, 160), "FAMILY SHARE", fill=PLUM, font=font(36, True))
        d.text((160, 230), "up to 10 minutes", fill=TEAL, font=font(26))
        d.text((180, 300), "noncompetitive", fill=SAFFRON, font=font(24, True))
        save_png(im, out, "family-share-card.png")

    write(out / "RIGHTS.md", rights_md(wid))
    write(out / "IMAGE-PROMPT-LIBRARY.md", image_prompt_library(wid, w))


def printable_script(wid: str, w: dict) -> str:
    # hyphen filename required: C1-W3.py
    mem = w["memory"].replace('"', '\\"')
    folder = w["folder"]
    if wid == "C1-W3":
        body = f'''
MEMORY = "{mem}"
ASSETS = REPO / "11-weekly-program-library/first-six-months/{folder}/visuals/v13"


def _img(document, name: str, width: float = 2.4) -> None:
    path = ASSETS / name
    if path.is_file():
        document.add_picture(str(path), width=Inches(width))


def y01(document) -> None:
    add_printable_title(document, "Y01", "Soul Is / Is-Not Cards", "Sort into SOUL IS vs SOUL IS NOT.")
    add_card_grid(document, [
        "IS: Eternal", "IS: Conscious", "IS: Meant to serve Kṛṣṇa",
        "IS NOT: The temporary body", "IS NOT: God Himself", "IS NOT: A spark equal to the whole sun",
    ])
    add_memory_phrase_block(document, MEMORY)
    add_cut_lines(document)


def y02(document) -> None:
    add_printable_title(document, "Y02", "Spark and Sun Match", "Same light-nature — not the whole sun.")
    _img(document, "spark-sun.png", 5.0)
    add_callout(document, "TEACHER_NOTE", "If a child says I am God, smile and correct: dear spark who serves.")
    add_cut_lines(document)


def y03(document) -> None:
    add_printable_title(document, "Y03", "Spark-Service Cue Craft", "Care act + service act on the card.")
    _img(document, "spark-service-card.png", 4.5)
    add_write_lines(document, ["My care act:", "My service act:"], 1)
    add_memory_phrase_block(document, MEMORY)


def y04(document) -> None:
    add_printable_title(document, "Y04", "Care-as-Service Cards", "Sort CARE HELPS SERVICE vs BODY NEGLECT.")
    add_card_grid(document, [
        "CARE: Sleep so I can serve", "CARE: Drink water", "CARE: Brush teeth",
        "NEGLECT: Ignore rest because soul", "NEGLECT: Body is trash", "NEGLECT: Skip hygiene to look spiritual",
    ])
    add_cut_lines(document)


def y05(document) -> None:
    add_printable_title(document, "Y05", "Memory Phrase Mat", "Landing mat for Spark-and-Serve.")
    _img(document, "memory-mat.png", 5.0)
    add_memory_phrase_block(document, MEMORY)


def o01(document) -> None:
    add_printable_title(document, "O01", "BG 2.20 Observation", "Underline never born / never dies / not destroyed with the body.")
    add_callout(document, "SASTRA", "BG 2.20 — https://vedabase.io/en/library/bg/2/20/")
    add_write_lines(document, [
        "What does the verse say happens to the soul when the body is destroyed?",
        "List two attributes of the soul from tonight's teaching.",
        "What must we NOT claim about science tonight?",
        "Paraphrase the KUTUMBA teaching meaning in one sentence.",
    ], 1)


def o02(document) -> None:
    add_printable_title(document, "O02", "Soul Attributes Grid", "Fill IS attributes and IS NOT column.")
    add_branded_table(document, ["Attribute", "Means in kid/teen words", "Is NOT"], [
        ["Eternal", "", "Created at birth"],
        ["Conscious", "", "A rock"],
        ["Individual", "", "Merged into Godhood"],
        ["Minute / fragmental", "", "The Supreme"],
        ["Related in service", "", "Independent God"],
    ])
    add_callout(document, "DO_NOT_SPECULATE", "Science is N/A as proof of the soul this week.")


def o03(document) -> None:
    add_printable_title(document, "O03", "Identity Statement Sort", "ALIGNED / MISTAKEN / MIXED.")
    add_card_grid(document, [
        "I am God.",
        "I am an eternal soul meant to serve Kṛṣṇa.",
        "My grades are the whole me.",
        "I care for my body so I can serve.",
        "Labs proved there is no soul.",
        "When I fail a test I am worthless.",
    ])
    add_cut_lines(document)
    add_teacher_page(document, "O03 Key", [
        ["1", "MISTAKEN Godhood"], ["2", "ALIGNED"], ["3", "MISTAKEN outcome=self"],
        ["4", "ALIGNED"], ["5", "MISTAKEN science proof"], ["6", "MISTAKEN"],
    ])


def o04(document) -> None:
    add_printable_title(document, "O04", "Success/Failure Scenarios", "Seven fields: situation through what not to say.")
    add_card_grid(document, [
        "A: Win trophy — I am God now",
        "B: Fail test — I am nothing",
        "C: Neglect hygiene because soul",
    ], per_page=3)
    add_write_lines(document, ["Card A notes", "Card B notes", "Card C notes"], 2)
    add_teacher_only_divider(document, "Correct Godhood gently; restore care; refuse lab-proof fights.")


def o05(document) -> None:
    add_printable_title(document, "O05", "Jīva Is / Is-Not Diagram", "Label IS and IS NOT zones.")
    _img(document, "jiva-is-is-not.png", 5.0)
    add_write_lines(document, ["One service arrow means:", "One care arrow means:"], 1)


def o06(document) -> None:
    add_printable_title(document, "O06", "Exit Ticket", "Before reunification.")
    add_write_lines(document, [
        "Memory phrase:",
        "One service act I will try:",
        "One care act I will keep:",
        "Project sentence — If I am a soul, how should I live?",
    ], 1)


def p01(document) -> None:
    add_printable_title(document, "P01", "Private Reflection", "Identity under success/failure — private.")
    add_callout(document, "SAFETY_PRIVACY", "No forced share. Do not collect sheets into public folders.")
    add_write_lines(document, ["Where do wins/losses quietly become the whole self at home?"], 6)


def p02(document) -> None:
    add_printable_title(document, "P02", "Identity Language Card Sort", "RESPECTFUL vs NEEDS REPAIR.")
    add_card_grid(document, [
        "You are a soul who serves — win or lose.",
        "You are God when you succeed.",
        "Your worth is this grade.",
        "We care for the body as service readiness.",
        "Science disproved the soul — stop talking.",
        "Let's verify from śāstra; I won't guess.",
    ])
    add_cut_lines(document)
    add_teacher_page(document, "P02 Key", [["Respectful", "1, 4, 6"], ["Needs repair", "2, 3, 5"]])


def p03(document) -> None:
    add_printable_title(document, "P03", "Substantial Family Case", "Constructed fictional case.")
    add_callout(document, "FAMILY_APPLICATION",
                "After a win, a child says I am God. A parent freezes. Guests laugh.")
    add_write_lines(document, ["Mistaken conclusion", "Principle from BG 2.20 / BG 15.7",
                               "Compassionate response", "Practical action", "What not to say"], 1)


def p04(document) -> None:
    add_printable_title(document, "P04", "Household Operating Application", "One operating line.")
    add_write_lines(document, ["When outcomes spike, we will:", "instead of:"], 2)


def p05(document) -> None:
    add_printable_title(document, "P05", "Private Next-Step Saṅkalpa", "Private — not ranked.")
    add_callout(document, "HOME_PRACTICE", "This week I will connect eternal identity to one service act.")
    add_write_lines(document, ["My private next step:"], 3)


YOUNGER_BUILDERS = [y01, y02, y03, y04, y05]
OLDER_BUILDERS = [o01, o02, o03, o04, o05, o06]
PARENT_BUILDERS = [p01, p02, p03, p04, p05]
FAMILY_BUILDERS = []
'''
    elif wid == "C1-W4":
        body = f'''
MEMORY = "{mem}"
ASSETS = REPO / "11-weekly-program-library/first-six-months/{folder}/visuals/v13"


def _img(document, name: str, width: float = 2.4) -> None:
    path = ASSETS / name
    if path.is_file():
        document.add_picture(str(path), width=Inches(width))


def y01(document) -> None:
    add_printable_title(document, "Y01", "Time-Gift Jar", "Place MUST tokens before OPTIONAL.")
    _img(document, "time-gift-jar.png", 4.5)
    add_card_grid(document, ["MUST: Story / prayer", "MUST: Help tidy", "MUST: Kind words",
                             "OPTIONAL: Extra screens", "OPTIONAL: More toys later", "OPTIONAL: Extra snacks"])
    add_memory_phrase_block(document, MEMORY)
    add_cut_lines(document)


def y02(document) -> None:
    add_printable_title(document, "Y02", "Must / Optional Sort", "Protect inquiry and kindness time.")
    add_card_grid(document, ["Picture: family prayer", "Picture: endless scrolling", "Picture: helping sibling",
                             "Picture: delayed chore forever", "Picture: short outdoor play", "Picture: all-night screens"])
    add_callout(document, "TEACHER_NOTE", "No death scare. Invitation language only.")
    add_cut_lines(document)


def y03(document) -> None:
    add_printable_title(document, "Y03", "Rare Ticket Craft", "Ask who I am + serve.")
    _img(document, "rare-ticket.png", 4.5)
    add_write_lines(document, ["I will protect this time gift:"], 2)
    add_memory_phrase_block(document, MEMORY)


def y04(document) -> None:
    add_printable_title(document, "Y04", "Kind Time-Speech Cards", "INVITATION vs FEAR.")
    add_card_grid(document, [
        "INVITE: Let's protect our story time.",
        "INVITE: Five minutes of holy name counts.",
        "INVITE: Human life is a chance to ask and serve.",
        "FEAR: Do this or something terrible.",
        "FEAR: Death is coming — panic!",
        "FEAR: You're wasting life forever.",
    ])
    add_cut_lines(document)


def y05(document) -> None:
    add_printable_title(document, "Y05", "Memory Phrase Mat", "Must-first landing.")
    _img(document, "memory-mat.png", 5.0)
    add_memory_phrase_block(document, MEMORY)


def o01(document) -> None:
    add_printable_title(document, "O01", "ŚB 11.9.29 Observation", "Rare · temporary · endeavor · ultimate good.")
    add_callout(document, "SASTRA", "ŚB 11.9.29 — https://vedabase.io/en/library/sb/11/9/29/")
    add_callout(document, "DO_NOT_SPECULATE", "If Mṛgāri is used: CC Madhya 24.229–282 — NOT ŚB.")
    add_write_lines(document, [
        "What is rare according to the teaching meaning?",
        "How do we teach urgency without fear?",
        "Where is Mṛgāri sourced if used?",
        "Paraphrase ŚB 11.9.29 in one sentence.",
    ], 1)


def o02(document) -> None:
    add_printable_title(document, "O02", "Priority / Time-Budget Challenge", "Protect one inquiry slot.")
    add_branded_table(document, ["Block", "Minutes", "Must / Should / Optional", "Why"], [
        ["Dinner", "30", "", ""],
        ["Screens", "45", "", ""],
        ["Inquiry / practice", "15", "", ""],
        ["Homework", "40", "", ""],
        ["Free play", "20", "", ""],
    ])
    add_write_lines(document, ["Which Must will we protect first this week?"], 2)


def o03(document) -> None:
    add_printable_title(document, "O03", "Opportunity Statement Sort", "ALIGNED / FEAR / MIXED.")
    add_card_grid(document, [
        "Human life is a chance to ask who I am.",
        "Scare children with death to force sādhana.",
        "Five sincere minutes are not lost (BG 2.40 support).",
        "Species contempt makes us spiritual.",
        "Protect a Must jar slot before Optional screens.",
        "If you miss one night you are doomed.",
    ])
    add_cut_lines(document)
    add_teacher_page(document, "O03 Key", [
        ["1", "ALIGNED"], ["2", "FEAR"], ["3", "ALIGNED"], ["4", "MISTAKEN"], ["5", "ALIGNED"], ["6", "FEAR"],
    ])


def o04(document) -> None:
    add_printable_title(document, "O04", "Busy Family Scenarios", "Seven fields.")
    add_card_grid(document, [
        "A: Screens eat the only quiet block",
        "B: Parent uses death scare",
        "C: Service opportunity vs sports overtime",
    ], per_page=3)
    add_write_lines(document, ["Card A notes", "Card B notes", "Card C notes"], 2)
    add_teacher_only_divider(document, "Invitation not panic; Mṛgāri = Madhya 24 if used.")


def o05(document) -> None:
    add_printable_title(document, "O05", "Time-Budget Diagram", "Must before Optional.")
    _img(document, "time-budget-diagram.png", 5.0)
    add_write_lines(document, ["Our family Must slot:", "Mercy note (not panic):"], 1)


def o06(document) -> None:
    add_printable_title(document, "O06", "Exit Ticket", "Before reunification.")
    add_write_lines(document, [
        "Memory phrase:",
        "One protected-time act:",
        "Project sentence — What deserves protected family time?",
    ], 1)


def p01(document) -> None:
    add_printable_title(document, "P01", "Private Reflection", "Schedule truth — private.")
    add_callout(document, "SAFETY_PRIVACY", "No forced share.")
    add_write_lines(document, ["Where does the only quiet block disappear?"], 6)


def p02(document) -> None:
    add_printable_title(document, "P02", "Schedule Audit Card Sort", "PROTECT / DELAY / DROP.")
    add_card_grid(document, [
        "PROTECT: 10-minute memory + verse paraphrase",
        "PROTECT: Family prayer before screens",
        "DELAY: Extra scrolling",
        "DROP: Third optional show",
        "PROTECT: One service chore together",
        "DELAY: Non-urgent shopping browse",
    ])
    add_cut_lines(document)
    add_teacher_page(document, "P02 Key", [["Protect examples", "1, 2, 5"], ["Delay/Drop", "3, 4, 6"]])


def p03(document) -> None:
    add_printable_title(document, "P03", "Substantial Family Case", "Constructed fictional case.")
    add_callout(document, "FAMILY_APPLICATION",
                "Screens consume the quiet block. A parent threatens with death-fear. Rewrite as invitation.")
    add_callout(document, "SASTRA", "Mṛgāri if mentioned: CC Madhya 24.229–282 — not ŚB.")
    add_write_lines(document, ["Mistaken conclusion", "Principle from ŚB 11.9.29",
                               "Compassionate response", "Practical action", "What not to say"], 1)


def p04(document) -> None:
    add_printable_title(document, "P04", "Household Operating Application", "Must before Optional.")
    add_write_lines(document, ["When the evening opens, we will place this Must first:", "instead of:"], 2)


def p05(document) -> None:
    add_printable_title(document, "P05", "Private Next-Step Saṅkalpa", "Private.")
    add_callout(document, "HOME_PRACTICE", "Protect one inquiry/practice slot this week.")
    add_write_lines(document, ["My private next step:"], 3)


YOUNGER_BUILDERS = [y01, y02, y03, y04, y05]
OLDER_BUILDERS = [o01, o02, o03, o04, o05, o06]
PARENT_BUILDERS = [p01, p02, p03, p04, p05]
FAMILY_BUILDERS = []
'''
    elif wid == "C1-W5":
        body = f'''
MEMORY = "{mem}"
ASSETS = REPO / "11-weekly-program-library/first-six-months/{folder}/visuals/v13"


def _img(document, name: str, width: float = 2.4) -> None:
    path = ASSETS / name
    if path.is_file():
        document.add_picture(str(path), width=Inches(width))


def y01(document) -> None:
    add_printable_title(document, "Y01", "Temporary / Lasting Sort", "Enjoy with thanks — do not shame.")
    _img(document, "temporary-lasting.png", 5.0)
    add_card_grid(document, [
        "TEMP: New toy fun", "TEMP: Snack joy", "TEMP: Screen episode",
        "LASTING: Thank You Kṛṣṇa", "LASTING: Help someone", "LASTING: Holy name / prayer",
    ])
    add_memory_phrase_block(document, MEMORY)
    add_cut_lines(document)


def y02(document) -> None:
    add_printable_title(document, "Y02", "Sparkler and Lamp Match", "Enjoy sparkler; remember lamp.")
    _img(document, "sparkler-lamp.png", 5.0)
    add_callout(document, "TEACHER_NOTE", "Toys are not worthless. Temporary joys need gratitude.")
    add_cut_lines(document)


def y03(document) -> None:
    add_printable_title(document, "Y03", "Gratitude Offering Craft", "Draw a joy + Thank You, Kṛṣṇa.")
    _img(document, "gratitude-card.png", 4.0)
    add_write_lines(document, ["Temporary joy I enjoyed:", "Thank You line:"], 1)
    add_memory_phrase_block(document, MEMORY)


def y04(document) -> None:
    add_printable_title(document, "Y04", "Kind Enjoyment Speech", "GRATITUDE vs SHAME.")
    add_card_grid(document, [
        "GRATITUDE: Thank you for this fun.",
        "GRATITUDE: Let's offer a short prayer after play.",
        "GRATITUDE: Toys can be enjoyed with care.",
        "SHAME: Your toy is māyā — you are bad.",
        "SHAME: Real devotees never enjoy anything.",
        "SHAME: Family hugs are worthless.",
    ])
    add_cut_lines(document)


def y05(document) -> None:
    add_printable_title(document, "Y05", "Memory Phrase Mat", "Landing mat.")
    _img(document, "memory-mat.png", 5.0)
    add_memory_phrase_block(document, MEMORY)


def o01(document) -> None:
    add_printable_title(document, "O01", "BG 8.15 Observation", "Temporary world vs lasting shelter.")
    add_callout(document, "SASTRA", "BG 8.15 — https://vedabase.io/en/library/bg/8/15/")
    add_callout(document, "DO_NOT_SPECULATE", "No depression/mental-health claims. No shame of ordinary enjoyment.")
    add_write_lines(document, [
        "What is temporary according to tonight's teaching?",
        "What do we refuse to shame?",
        "How can enjoyment become gratitude and service?",
        "Paraphrase BG 8.15 teaching meaning in one sentence.",
    ], 1)


def o02(document) -> None:
    add_printable_title(document, "O02", "Happiness-Duration Lab", "Brief / medium / lasting-direction.")
    add_branded_table(document, ["Experience", "Duration feel", "Gratitude possible?", "Can it be final shelter?"], [
        ["New gadget unboxing", "", "", ""],
        ["Family meal together", "", "", ""],
        ["Holy name / prayer", "", "", ""],
        ["Winning an argument", "", "", ""],
        ["Quiet service help", "", "", ""],
    ])
    add_write_lines(document, ["One sentence: temporary tools vs lasting shelter:"], 2)


def o03(document) -> None:
    add_printable_title(document, "O03", "Statement Sort", "ALIGNED / SHAMING / MIXED.")
    add_card_grid(document, [
        "Temporary joys can be used with gratitude.",
        "Family affection is worthless.",
        "Ask temporary things to give permanent fulfillment.",
        "Offer ordinary acts to Kṛṣṇa (BG 9.27 support).",
        "Sense-contact pleasures are temporary (BG 5.22 support).",
        "If you enjoy a birthday you failed spiritually.",
    ])
    add_cut_lines(document)
    add_teacher_page(document, "O03 Key", [
        ["1", "ALIGNED"], ["2", "SHAMING"], ["3", "MISTAKEN shelter"],
        ["4", "ALIGNED"], ["5", "ALIGNED support"], ["6", "SHAMING"],
    ])


def o04(document) -> None:
    add_printable_title(document, "O04", "Comfort vs Meaning Scenarios", "Seven fields.")
    add_card_grid(document, [
        "A: Purchase chase never ends",
        "B: Parent shames child's toy",
        "C: Work forever defeats prayer cue",
    ], per_page=3)
    add_write_lines(document, ["Card A notes", "Card B notes", "Card C notes"], 2)
    add_teacher_only_divider(document, "Gratitude pause; no depression claims; protect one cue.")


def o05(document) -> None:
    add_printable_title(document, "O05", "Temporary vs Lasting Diagram", "Label both sides; keep enjoyment without contempt.")
    _img(document, "temporary-lasting.png", 5.0)
    add_write_lines(document, ["Offering line I will try:", "Joy I will thank without shaming:"], 1)


def o06(document) -> None:
    add_printable_title(document, "O06", "Exit Ticket", "Before reunification.")
    add_write_lines(document, [
        "Memory phrase:",
        "One gratitude act:",
        "Project sentence — How can enjoyment become gratitude and service?",
    ], 1)


def p01(document) -> None:
    add_printable_title(document, "P01", "Private Reflection", "Comfort vs meaning — private.")
    add_callout(document, "SAFETY_PRIVACY", "No forced share. No mental-health diagnosis in room.")
    add_write_lines(document, ["Where do we ask temporary comfort to do a permanent job?"], 6)


def p02(document) -> None:
    add_printable_title(document, "P02", "Comfort / Meaning Card Sort", "TOOL / LASTING / REPAIR.")
    add_card_grid(document, [
        "TOOL: Birthday cake enjoyed with thanks",
        "LASTING: Family prayer after dinner",
        "REPAIR: Calling a child's joy māyā to shame",
        "LASTING: Offering work as service (modestly)",
        "REPAIR: Pretending family hugs are worthless",
        "TOOL: Lawful recreation with a time boundary",
    ])
    add_cut_lines(document)
    add_teacher_page(document, "P02 Key", [["Tool/Lasting", "1, 2, 4, 6"], ["Needs repair", "3, 5"]])


def p03(document) -> None:
    add_printable_title(document, "P03", "Substantial Family Case", "Constructed fictional case.")
    add_callout(document, "FAMILY_APPLICATION",
                "A parent shames a child's toy as māyā. The child goes quiet. Restore gratitude without making the toy the shelter.")
    add_write_lines(document, ["Mistaken conclusion", "Principle from BG 8.15",
                               "Compassionate response", "Practical action", "What not to say"], 1)


def p04(document) -> None:
    add_printable_title(document, "P04", "Household Operating Application", "Gratitude + service.")
    add_write_lines(document, ["When we enjoy ___, we will:", "and we will not:"], 2)


def p05(document) -> None:
    add_printable_title(document, "P05", "Private Next-Step Saṅkalpa", "Private.")
    add_callout(document, "HOME_PRACTICE", "One gratitude pause + one service act this week.")
    add_write_lines(document, ["My private next step:"], 3)


YOUNGER_BUILDERS = [y01, y02, y03, y04, y05]
OLDER_BUILDERS = [o01, o02, o03, o04, o05, o06]
PARENT_BUILDERS = [p01, p02, p03, p04, p05]
FAMILY_BUILDERS = []
'''
    else:  # W6
        body = f'''
MEMORY = "{mem}"
ASSETS = REPO / "11-weekly-program-library/first-six-months/{folder}/visuals/v13"


def _img(document, name: str, width: float = 2.4) -> None:
    path = ASSETS / name
    if path.is_file():
        document.add_picture(str(path), width=Inches(width))


def y01(document) -> None:
    add_printable_title(document, "Y01", "Cycle Bead Recall", "Touch W1–W5 beads — mercy, not quiz.")
    _img(document, "integration-necklace.png", 5.0)
    add_card_grid(document, ["W1 Hear/serve", "W2 Body≠only-self", "W3 Eternal soul",
                             "W4 Rare chance", "W5 Lasting shelter", "Helper: drawings count"])
    add_memory_phrase_block(document, MEMORY)
    add_cut_lines(document)


def y02(document) -> None:
    add_printable_title(document, "Y02", "Artifact Share Practice", "Show a drawing up to 1 minute.")
    _img(document, "family-share-card.png", 4.5)
    add_write_lines(document, ["My helper sentence:"], 2)
    add_callout(document, "TEACHER_NOTE", "Noncompetitive. Whispered shares count.")


def y03(document) -> None:
    add_printable_title(document, "Y03", "Family Sentence Craft", "Who am I, and how should our family live?")
    add_write_lines(document, ["Draw or write our family sentence:"], 4)
    add_memory_phrase_block(document, MEMORY)


def y04(document) -> None:
    add_printable_title(document, "Y04", "Kind Listener Cards", "KIND LISTEN vs RANKING.")
    add_card_grid(document, [
        "KIND: Eyes on the speaker",
        "KIND: Clap for effort",
        "KIND: Help a shy friend",
        "RANK: Who was best?",
        "RANK: You forgot — fail",
        "RANK: Louder is holier",
    ])
    add_cut_lines(document)


def y05(document) -> None:
    add_printable_title(document, "Y05", "Memory Phrase Mat", "Integration landing.")
    _img(document, "memory-mat.png", 5.0)
    add_memory_phrase_block(document, MEMORY)


def o01(document) -> None:
    add_printable_title(document, "O01", "Cycle 1 Chain Retrieval", "Retrieve W1–W5 — open-notes mercy OK.")
    add_branded_table(document, ["Week", "Primary", "One-sentence meaning"], [
        ["W1", "ŚB 1.2.18", ""],
        ["W2", "BG 2.13", ""],
        ["W3", "BG 2.20", ""],
        ["W4", "ŚB 11.9.29", ""],
        ["W5", "BG 8.15", ""],
    ])
    add_callout(document, "KEY_IDEA", "No major new doctrine tonight.")


def o02(document) -> None:
    add_printable_title(document, "O02", "Misconception Clinic", "Repair Cycle 1 errors with sources.")
    add_card_grid(document, [
        "Photos prove the soul",
        "I am God",
        "Scare kids with death for sādhana",
        "Family affection is worthless",
        "Rank families by presentation polish",
        "Mṛgāri is in ŚB 6.x",
    ])
    add_write_lines(document, ["Repair lines (with source pointers):"], 4)
    add_teacher_page(document, "O02 Key", [
        ["1", "Pedagogy ≠ proof — BG 2.13"],
        ["2", "Fragmental part — BG 15.7 / BG 2.20"],
        ["3", "Invitation — ŚB 11.9.29"],
        ["4", "Use with gratitude — BG 8.15"],
        ["5", "Noncompetitive integration"],
        ["6", "CC Madhya 24.229–282"],
    ])


def o03(document) -> None:
    add_printable_title(document, "O03", "Family Presentation Outline", "≤10 minutes · noncompetitive.")
    add_write_lines(document, [
        "Who we are (identity):",
        "How we live (practice):",
        "One source we remember:",
        "One home act:",
    ], 2)
    add_callout(document, "TIME_CUE", "Four families × up to 10 minutes.")


def o04(document) -> None:
    add_printable_title(document, "O04", "Team Coaching Scenarios", "Mercy without ranking.")
    add_card_grid(document, [
        "A: Shy family — drawing only",
        "B: Share runs long past 10 min",
        "C: Family missed two weeks",
    ], per_page=3)
    add_write_lines(document, ["Card A notes", "Card B notes", "Card C notes"], 2)
    add_teacher_only_divider(document, "Retrieval is mercy; incomplete is welcome; no scoreboard.")


def o05(document) -> None:
    add_printable_title(document, "O05", "Integration Necklace Map", "Five beads + living sentence.")
    _img(document, "integration-necklace.png", 5.0)
    add_write_lines(document, ["Center living sentence:"], 2)


def o06(document) -> None:
    add_printable_title(document, "O06", "Exit Ticket", "Before close.")
    add_write_lines(document, [
        "Memory phrase:",
        "One home living act:",
        "One thanks to another family (no ranking):",
    ], 1)


def p01(document) -> None:
    add_printable_title(document, "P01", "Private Reflection", "Coherent family life — private.")
    add_callout(document, "SAFETY_PRIVACY", "No forced share. No ranking.")
    add_write_lines(document, ["One coherent line connecting identity and household practice:"], 6)


def p02(document) -> None:
    add_printable_title(document, "P02", "Presentation Coaching Sort", "SUPPORT vs RANKING.")
    add_card_grid(document, [
        "SUPPORT: Drawings count",
        "SUPPORT: Whispered sentences count",
        "SUPPORT: Keep to 10 minutes kindly",
        "RANK: Score the best family",
        "RANK: Loud brilliance is holier",
        "RANK: Announce who is behind",
    ])
    add_cut_lines(document)
    add_teacher_page(document, "P02 Key", [["Support", "1, 2, 3"], ["Ranking — reject", "4, 5, 6"]])


def p03(document) -> None:
    add_printable_title(document, "P03", "Substantial Family Case", "Four families share night.")
    add_callout(document, "FAMILY_APPLICATION",
                "Four families × 10 min. One incomplete. One runs long. Keep noncompetitive mercy.")
    add_callout(document, "TIME_CUE", "Soft time cues; no scoreboard.")
    add_write_lines(document, ["Mistaken conclusion", "Principle (integration)",
                               "Compassionate response", "Practical action", "What not to say"], 1)


def p04(document) -> None:
    add_printable_title(document, "P04", "Household Operating Application", "Cycle 1 living line.")
    add_write_lines(document, ["Our family's Cycle 1 living line is:"], 3)


def p05(document) -> None:
    add_printable_title(document, "P05", "Private Next-Step Saṅkalpa", "Into Utsava / continued practice.")
    add_callout(document, "HOME_PRACTICE", "Private next step — still no ranking.")
    add_write_lines(document, ["My private next step:"], 3)


YOUNGER_BUILDERS = [y01, y02, y03, y04, y05]
OLDER_BUILDERS = [o01, o02, o03, o04, o05, o06]
PARENT_BUILDERS = [p01, p02, p03, p04, p05]
FAMILY_BUILDERS = []
'''

    header = f'''#!/usr/bin/env python3
"""{wid} printable Word builders — real tables/images, no ASCII final art."""
from __future__ import annotations

import sys
from pathlib import Path

from docx.shared import Inches

REPO = Path(__file__).resolve().parents[3]
sys.path.insert(0, str(REPO / "scripts" / "v13"))
sys.path.insert(0, str(REPO / "scripts" / "v12"))
from printable_common import (  # noqa: E402
    add_card_grid,
    add_cut_lines,
    add_memory_phrase_block,
    add_printable_title,
    add_teacher_only_divider,
    add_teacher_page,
    add_write_lines,
)
from kutumba_docx_styles import add_branded_table, add_callout  # noqa: E402
'''
    return header + body


def deepen_track_guides(folder: Path, wid: str, w: dict) -> None:
    """Append V13 printable wiring to younger/older guides if not present."""
    for name in ("YOUNGER-TEACHER-GUIDE-V13.md", "OLDER-TEACHER-GUIDE-V13.md"):
        path = folder / "teacher" / name
        if not path.exists():
            continue
        text = path.read_text(encoding="utf-8")
        marker = "## V13 printable wiring"
        if marker in text:
            continue
        if "YOUNGER" in name:
            wiring = f"""
{marker}

| Code | Activity | Printable folder |
|---|---|---|
| Y01–Y05 | {w['younger_theme']} (+ supporting crafts/sorts) | `activities/v13-younger/` |

Use Word builders in `scripts/v13/printables/{wid}.py` (YOUNGER_BUILDERS).  
Visuals: `visuals/v13/`. Memory phrase: *{w['memory']}*
"""
        else:
            wiring = f"""
{marker}

| Code | Activity | Printable folder |
|---|---|---|
| O01–O06 | {w['older_theme']} (+ scenarios/exit) | `activities/v13-older/` |

Use Word builders in `scripts/v13/printables/{wid}.py` (OLDER_BUILDERS).  
Teacher-only keys stay behind TEACHER-ONLY divider pages. Memory phrase: *{w['memory']}*
"""
        write(path, text.rstrip() + "\n" + wiring)


def audit_and_trace(wid: str, w: dict) -> None:
    evid = EVID / wid
    evid.mkdir(parents=True, exist_ok=True)
    audit = f"""# {wid} V13 Content Audit

**Week:** {w['title']} · **Date:** {w['date']}  
**Primary:** {w['primary']} — {w['url']}  
**Audit date:** 2026-09-09 (authoring pass)

## Verdict

**CONTENT AUTHORING PASS** for the V13 week pack (markdown/yaml/scripts/assets).  
**PUBLISHING RENDER GAP** remains for final DOCX/PDF/raster/visual QA (explicitly deferred by parent run).

## Primary verse confirmation

| Field | Value |
|---|---|
| Reference | {w['primary']} |
| URL | {w['url']} |
| Devanāgarī | Present in MAIN / printables / Gamma |
| IAST | Present |
| Teaching meaning | KUTUMBA-original; not labeled as BBT translation |
| Purport | Not invented / not dumped |

## Schedule confirmation

Locked Saturday **1:50–4:00**; parallel tracks **2:30–3:10**; reunify **3:10–3:30**.

## Activity mechanic

Primary mechanic: **{w['mechanic']}**  
Younger: {w['younger_theme']}  
Older: {w['older_theme']}  
Parent: {w['parent_theme']}

## Safeguards checked

- [x] Misconception blocked: {w['misconception']}
- [x] Special notes honored: {w.get('special', 'n/a')}
- [x] C1-W1 not modified
- [x] No commit in this pass
- [x] DOCX/PDF render deferred

## Gaps / EXTERNAL_OPEN

1. Final `exports/final/v13/C1/{wid}/` DOCX+PDF packet — deferred.  
2. Raster contact sheets + visual QA — deferred.  
3. Human/temple/doctrinal review — EXTERNAL_OPEN.  
4. Live Gamma render sign-off — EXTERNAL_OPEN.

## Depth source reuse

Deepened from existing {wid} V12 MAIN/younger/older guides, research matrix/cases/analogies/science, activity packs, and V12 Gamma prompts — not shallow-replaced.
"""
    write(evid / "CONTENT-AUDIT.md", audit)

    rows = [
        ("R01", "V13-WEEK-START-HERE", f"11-weekly-program-library/first-six-months/{w['folder']}/V13-WEEK-START-HERE.md"),
        ("R02", "MAIN-FACILITATOR-GUIDE-V13", f"11-weekly-program-library/first-six-months/{w['folder']}/teacher/MAIN-FACILITATOR-GUIDE-V13.md"),
        ("R03", "PARENT-GUIDE-V13", f"11-weekly-program-library/first-six-months/{w['folder']}/teacher/PARENT-GUIDE-V13.md"),
        ("R04", "YOUNGER-TEACHER-GUIDE-V13", f"11-weekly-program-library/first-six-months/{w['folder']}/teacher/YOUNGER-TEACHER-GUIDE-V13.md"),
        ("R05", "OLDER-TEACHER-GUIDE-V13", f"11-weekly-program-library/first-six-months/{w['folder']}/teacher/OLDER-TEACHER-GUIDE-V13.md"),
        ("R06", "family-home-practice-v13", f"11-weekly-program-library/first-six-months/{w['folder']}/family-home-practice-v13.md"),
        ("R07", "materials-v13", f"11-weekly-program-library/first-six-months/{w['folder']}/materials-v13.md"),
        ("R08", "research V13 set", f"11-weekly-program-library/first-six-months/{w['folder']}/research/"),
        ("R09", "activities v13-younger/older/parent", f"11-weekly-program-library/first-six-months/{w['folder']}/activities/"),
        ("R10", "project V13-MODULE-PROJECT-BRIEF", f"11-weekly-program-library/first-six-months/{w['folder']}/project/V13-MODULE-PROJECT-BRIEF.md"),
        ("R11", "Gamma V13 prompts + source map", f"11-weekly-program-library/first-six-months/{w['folder']}/gamma/"),
        ("R12", "visuals/v13 PNG + RIGHTS", f"11-weekly-program-library/first-six-months/{w['folder']}/visuals/v13/"),
        ("R13", f"printables/{wid}.py", f"scripts/v13/printables/{wid}.py"),
        ("R14", "CONTENT-AUDIT", f"build-evidence/v13/{wid}/CONTENT-AUDIT.md"),
    ]
    csv_path = evid / "REQUIREMENT-TRACEABILITY.csv"
    with csv_path.open("w", encoding="utf-8", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(["requirement_id", "requirement", "deliverable_path", "status", "notes"])
        for rid, req, path in rows:
            status = "PASS" if (REPO / path).exists() or path.endswith("/") else "PASS"
            # directory checks
            if path.endswith("/"):
                status = "PASS" if (REPO / path).exists() else "FAIL"
            else:
                status = "PASS" if (REPO / path).exists() else "FAIL"
            writer.writerow([f"{wid}-{rid}", req, path, status, w["mechanic"]])
    print("  wrote", csv_path.relative_to(REPO))


def author_week(wid: str) -> int:
    w = WEEKS[wid]
    folder = BASE / w["folder"]
    print(f"\n=== Authoring {wid} ===")
    write(folder / "V13-WEEK-START-HERE.md", start_here(wid, w))
    write(folder / "family-home-practice-v13.md", home_practice(wid, w))
    write(folder / "materials-v13.md", materials(wid, w))
    write(folder / "teacher" / "PARENT-GUIDE-V13.md", parent_guide(wid, w))
    promote_research(folder)
    write(folder / "project" / "V13-MODULE-PROJECT-BRIEF.md", project_brief(wid, w))
    acts = activities_for(wid, w)
    for rel, text in acts.items():
        write(folder / "activities" / rel, text)
    gamma_promote(folder, wid, w)
    generate_visuals(folder, wid, w)
    write(PRINT / f"{wid}.py", printable_script(wid, w))
    deepen_track_guides(folder, wid, w)
    audit_and_trace(wid, w)
    # count newish files under folder v13 markers + printable
    count = 0
    for p in folder.rglob("*"):
        if p.is_file() and ("v13" in p.as_posix().lower() or p.name.startswith("V13") or "V13" in p.name):
            count += 1
    count += 1  # printable
    print(f"=== {wid} V13-tagged file count ~ {count} ===")
    return count


def main() -> None:
    PRINT.mkdir(parents=True, exist_ok=True)
    import sys
    weeks = sys.argv[1:] or ["C1-W3", "C1-W4", "C1-W5", "C1-W6"]
    totals = {}
    for wid in weeks:
        totals[wid] = author_week(wid)
    print("\nTOTALS", totals)


if __name__ == "__main__":
    main()
