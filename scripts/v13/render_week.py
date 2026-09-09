#!/usr/bin/env python3
"""Render one V13 week into seven DOCX+PDF packets under exports/final/v13/."""
from __future__ import annotations

import argparse
import hashlib
import importlib.util
import json
import os
import re
import sys
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Callable

import pymupdf as fitz
import yaml
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.shared import Inches

REPO = Path(__file__).resolve().parents[2]
V13 = Path(__file__).resolve().parent
for item in (REPO, REPO / "scripts" / "v12", V13):
    sys.path.insert(0, str(item))

from kutumba_docx_styles import (  # noqa: E402
    add_branded_table, add_callout, add_cover, add_header_footer, apply_kutumba_styles,
)
from printable_common import add_teacher_only_divider  # noqa: E402
from render_markdown_safe import render_markdown_safe  # noqa: E402
from render_publication_docs import convert_with_word  # noqa: E402

REGISTRY = V13 / "week_registry.yaml"
FIRST_SIX = REPO / "11-weekly-program-library" / "first-six-months"
STATUS = "Internal founding-cohort teaching material — human/temple review EXTERNAL_OPEN"
SCHEDULE = [
    ("1:50–2:00", "Arrival"),
    ("2:00–2:10", "Opening mantras + welcome"),
    ("2:10–2:30", "All-family launch/orientation"),
    ("2:30–3:10", "Parallel tracks"),
    ("3:10–3:30", "Reunification + bhakti laboratory"),
    ("3:30–3:40", "Snack + water"),
    ("3:40–3:55", "Saṅkalpa + project + home practice"),
    ("3:55–4:00", "Close"),
]
PACKET_SUFFIXES = [
    ("START-HERE", "Start Here", "Owner-facing operational entry point", "Program owner"),
    ("MAIN-FACILITATOR", "Main Facilitator", "Speaking map and session leadership", "Main facilitator"),
    ("PARENT-TRACK", "Parent Track", "Facilitator guide and parent printables", "Parent facilitator"),
    ("YOUNGER-TEACHER-PACK", "Younger Teacher Pack", "K–2 run sheet and activities", "Younger teacher"),
    ("OLDER-TEACHER-PACK", "Older Teacher Pack", "Grades 4–5 run sheet and activities", "Older teacher"),
    ("FAMILY-HANDOUT", "Family Handout", "Family take-home and practice sheet", "Families"),
    ("SATURDAY-PRINT-PACKET", "Saturday Print Packet", "Complete ordered participant and teacher packet", "Families and facilitators"),
]


@dataclass(frozen=True)
class WeekEntry:
    id: str
    title: str
    cycle: str
    folder: str
    primary_verse: str
    primary_url: str
    schedule_date: str
    kind: str
    export_subdir: str

    @property
    def module_dir(self) -> Path:
        return FIRST_SIX / self.folder

    @property
    def export_dir(self) -> Path:
        return REPO / "exports" / "final" / "v13" / Path(self.export_subdir)


def load_registry() -> list[WeekEntry]:
    data = yaml.safe_load(REGISTRY.read_text(encoding="utf-8"))
    return [WeekEntry(**row) for row in data["weeks"]]


def get_week(week_id: str) -> WeekEntry:
    for entry in load_registry():
        if entry.id == week_id:
            return entry
    known = ", ".join(e.id for e in load_registry())
    raise SystemExit(f"Unknown week id {week_id!r}. Known: {known}")


def resolve_folder(entry: WeekEntry) -> Path:
    """Resolve registry folder against filesystem (diacritics may vary)."""
    exact = FIRST_SIX / entry.folder
    if exact.is_dir():
        return exact
    needle = entry.folder.casefold()
    for child in FIRST_SIX.iterdir():
        if child.is_dir() and child.name.casefold() == needle:
            return child
    # Fuzzy: match id prefix like c1-w2-
    slug = entry.id.casefold()
    for child in FIRST_SIX.iterdir():
        if not child.is_dir():
            continue
        name = child.name.casefold()
        if name.startswith(slug + "-"):
            return child
    if entry.kind in {"utsava", "mela"}:
        raise SystemExit(
            f"Module folder not found for {entry.id}: expected "
            f"{exact.relative_to(REPO).as_posix()} (create when authoring this pack)."
        )
    raise SystemExit(f"Module folder not found for {entry.id}: {entry.folder}")


def load_printables(week_id: str) -> Any:
    module_path = V13 / "printables" / f"{week_id}.py"
    if not module_path.is_file():
        raise SystemExit(
            f"Missing week printable builder: {module_path.relative_to(REPO).as_posix()}\n"
            f"Create scripts/v13/printables/{week_id}.py exporting "
            f"YOUNGER_BUILDERS, OLDER_BUILDERS, PARENT_BUILDERS "
            f"(and optional FAMILY_BUILDERS)."
        )
    spec = importlib.util.spec_from_file_location(f"v13_printables_{week_id.replace('-', '_')}", module_path)
    if spec is None or spec.loader is None:
        raise SystemExit(f"Cannot import printable module: {module_path}")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    for required in ("YOUNGER_BUILDERS", "OLDER_BUILDERS", "PARENT_BUILDERS"):
        if not hasattr(module, required):
            raise SystemExit(f"{module_path.name} must export {required}")
    return module


def _relative(path: Path) -> str:
    return path.resolve().relative_to(REPO.resolve()).as_posix()


def _text(path: Path) -> str:
    if not path.is_file():
        raise FileNotFoundError(path)
    text = path.read_text(encoding="utf-8-sig")
    return "\n".join(line for line in text.splitlines() if "2:35" not in line and "3:25" not in line)


def _first_existing(base: Path, candidates: list[str]) -> Path | None:
    for name in candidates:
        path = base / name
        if path.is_file():
            return path
    return None


def require_source(base: Path, candidates: list[str], label: str) -> Path:
    found = _first_existing(base, candidates)
    if found is None:
        raise SystemExit(f"Missing {label} under {base.relative_to(REPO).as_posix()}: tried {candidates}")
    return found


def _source(document, path: Path, heading: str, with_hash: bool = True) -> None:
    document.add_heading(heading, level=1)
    if with_hash:
        digest = hashlib.sha256(path.read_bytes()).hexdigest()
        document.add_paragraph(
            f"Canonical source: {_relative(path)}  •  SHA-256: {digest}",
            style="KUTUMBA Rights",
        )
    title_match = re.search(r"^#\s+(.+)$", _text(path), re.M)
    render_markdown_safe(document, _text(path), title_match.group(1).strip() if title_match else heading)


def _page_source(document, path: Path, heading: str, with_hash: bool = True) -> None:
    document.add_page_break()
    _source(document, path, heading, with_hash=with_hash)


def _page_builder(document, builder: Callable) -> None:
    document.add_page_break()
    builder(document)


def _schedule(document, entry: WeekEntry, heading: str = "Locked Saturday Schedule") -> None:
    document.add_heading(heading, level=1)
    if entry.kind == "topic":
        add_branded_table(document, ["Time", "Block"], SCHEDULE)
        add_callout(document, "TIME_CUE", "Track handoff is at 2:30; all groups reunite at 3:10.")
    else:
        add_branded_table(document, ["Time", "Block"], SCHEDULE)
        add_callout(
            document,
            "TIME_CUE",
            f"{entry.kind.title()} planning date {entry.schedule_date}. "
            "Use 2:30 / 3:10 cues when parallel tracks run; local tithi remains EXTERNAL_OPEN.",
        )
    document.add_paragraph(f"Primary verse: {entry.primary_verse}")
    document.add_paragraph(f"Source URL: {entry.primary_url}")


def _new(entry: WeekEntry, title: str, subtitle: str, audience: str):
    doc = Document()
    apply_kutumba_styles(doc)
    add_cover(
        doc,
        title,
        subtitle,
        f"{entry.id} • {entry.title}",
        version="V13",
        status=STATUS,
    )
    add_header_footer(doc, f"{entry.id} • V13", audience)
    for section in doc.sections:
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
    return doc


def build_start(document, entry: WeekEntry, base: Path) -> None:
    _schedule(document, entry)
    path = require_source(base, ["V13-WEEK-START-HERE.md", "complete-week.md"], "Start Here")
    _page_source(document, path, "Operational Entry Point")


def build_main(document, entry: WeekEntry, base: Path) -> None:
    _schedule(document, entry)
    path = require_source(
        base / "teacher",
        ["MAIN-FACILITATOR-GUIDE-V13.md", "MAIN-FACILITATOR-GUIDE-V12.md"],
        "Main Facilitator guide",
    )
    _page_source(document, path, "Main Facilitator Guide")


def build_track(document, entry: WeekEntry, guide: Path, guide_title: str, builders, rules: bool = False) -> None:
    _schedule(document, entry)
    _page_source(document, guide, guide_title)
    if rules:
        rules_path = REPO / "launch" / "CHILD-HOUSE-RULES.md"
        if rules_path.is_file():
            _page_source(document, rules_path, "Child House Rules")
    for builder in builders:
        _page_builder(document, builder)


def build_family(document, entry: WeekEntry, base: Path, builders) -> None:
    _schedule(document, entry)
    handout = _first_existing(
        base,
        ["family-home-practice-v13.md", "family-home-practice.md", "home-practice.md"],
    )
    if handout:
        _page_source(document, handout, "Family Home Practice", with_hash=False)
    else:
        document.add_heading("Family Home Practice", level=1)
        document.add_paragraph("Complete the tiny home practice at least once before next Saturday.")
        document.add_paragraph(f"Remember this week's primary verse: {entry.primary_verse}")
    for builder in builders:
        _page_builder(document, builder)


def build_saturday(document, entry: WeekEntry, base: Path, printables) -> None:
    _schedule(document, entry, "Quick Schedule")
    mantras = REPO / "launch" / "OPENING-MANTRAS-HANDOUT.md"
    if mantras.is_file():
        _page_source(document, mantras, "Opening Mantras")
    parent_guide = require_source(
        base / "teacher",
        ["PARENT-GUIDE-V13.md", "PARENT-GUIDE.md"],
        "Parent guide",
    )
    _page_source(document, parent_guide, "Parent Track Run Sheet")
    for builder in printables.PARENT_BUILDERS:
        _page_builder(document, builder)
    younger_guide = require_source(
        base / "teacher",
        ["YOUNGER-TEACHER-GUIDE-V13.md", "YOUNGER-TEACHER-GUIDE.md"],
        "Younger guide",
    )
    older_guide = require_source(
        base / "teacher",
        ["OLDER-TEACHER-GUIDE-V13.md", "OLDER-TEACHER-GUIDE.md"],
        "Older guide",
    )
    rules = REPO / "launch" / "CHILD-HOUSE-RULES.md"
    if rules.is_file():
        _page_source(document, rules, "Younger Child Rules")
    for builder in printables.YOUNGER_BUILDERS:
        _page_builder(document, builder)
    if rules.is_file():
        _page_source(document, rules, "Older Child Rules")
    for builder in printables.OLDER_BUILDERS:
        _page_builder(document, builder)
    project = _first_existing(base / "project", ["V13-MODULE-PROJECT-BRIEF.md", "MODULE-PROJECT-BRIEF.md"])
    if project:
        _page_source(document, project, "Module Project")
    home = _first_existing(base, ["family-home-practice-v13.md", "family-home-practice.md", "home-practice.md"])
    if home:
        _page_source(document, home, "Week Home Practice", with_hash=False)
    add_teacher_only_divider(document)
    _page_source(document, younger_guide, "Younger Teacher Run Sheet")
    _page_source(document, older_guide, "Older Teacher Run Sheet")
    main = require_source(
        base / "teacher",
        ["MAIN-FACILITATOR-GUIDE-V13.md", "MAIN-FACILITATOR-GUIDE-V12.md"],
        "Main Facilitator guide",
    )
    _page_source(document, main, "Main Facilitator Speaking Map")
    document.add_page_break()
    document.add_heading("Room / Material / Snack / Cleanup Checklist", level=1)
    add_branded_table(
        document,
        ["Done", "Area", "Check"],
        [
            ("☐", "Room", "Parent circle, younger floor zone, older table zone, clear aisles"),
            ("☐", "Materials", "Labeled track folders, pencils, crayons, tape, timer"),
            ("☐", "Privacy", "Completed forms remain private and never enter Git"),
            ("☐", "Snack", "Snack and water staged for 3:30; allergy plan confirmed"),
            ("☐", "Safety", "Calm corner visible; private rooms closed; exits clear"),
            ("☐", "Cleanup", "Tables wiped and host home reset assigned"),
        ],
    )


def strip_trailing_empty(pdf_path: Path) -> int:
    doc = fitz.open(pdf_path)
    removed = 0
    while doc.page_count > 1:
        blocks = doc[-1].get_text("blocks")
        body = " ".join(block[4] for block in blocks if block[1] > 60 and block[3] < 740).strip()
        if body:
            break
        doc.delete_page(-1)
        removed += 1
    if removed:
        temp = pdf_path.with_suffix(".tmp.pdf")
        doc.save(temp, deflate=True)
        doc.close()
        os.replace(temp, pdf_path)
    else:
        doc.close()
    return removed


def _forbidden_check(pdf_path: Path, entry: WeekEntry) -> int:
    forbidden = [
        "<div", "</div>", "page-break-after:", "```",
        "[ ear /", "[ beads /", "[ helping hands ]", "[ heart / home ]",
    ] + list("┌┐└┘─│├┤┬┴┼")
    with fitz.open(pdf_path) as pdf:
        text = "\n".join(page.get_text() for page in pdf)
        pages = pdf.page_count
    found = [token for token in forbidden if token.lower() in text.lower()]
    if found:
        raise RuntimeError(f"{pdf_path.name} contains forbidden output: {found}")
    for token in ("2:30", "3:10"):
        if token not in text:
            raise RuntimeError(f"{pdf_path.name} missing {token}")
    for token in ("2:35", "3:25"):
        if token in text:
            raise RuntimeError(f"{pdf_path.name} contains obsolete {token}")
    verse_token = entry.primary_verse.split("(")[0].strip()
    if verse_token and verse_token not in text and entry.primary_verse not in text:
        # Allow short tokens like "BG 2.13"
        short = verse_token[:32]
        if short not in text:
            raise RuntimeError(f"{pdf_path.name} missing primary verse marker {entry.primary_verse!r}")
    return pages


def packet_stem(entry: WeekEntry, suffix: str) -> str:
    return f"KUTUMBA-{entry.id}-{suffix}-V13"


def render_week(entry: WeekEntry) -> dict:
    base = resolve_folder(entry)
    printables = load_printables(entry.id)
    out = entry.export_dir
    out.mkdir(parents=True, exist_ok=True)

    parent_guide = require_source(base / "teacher", ["PARENT-GUIDE-V13.md", "PARENT-GUIDE.md"], "Parent guide")
    younger_guide = require_source(
        base / "teacher", ["YOUNGER-TEACHER-GUIDE-V13.md", "YOUNGER-TEACHER-GUIDE.md"], "Younger guide"
    )
    older_guide = require_source(
        base / "teacher", ["OLDER-TEACHER-GUIDE-V13.md", "OLDER-TEACHER-GUIDE.md"], "Older guide"
    )
    family_builders = getattr(printables, "FAMILY_BUILDERS", [])

    builders_map: dict[str, Callable] = {
        "START-HERE": lambda d: build_start(d, entry, base),
        "MAIN-FACILITATOR": lambda d: build_main(d, entry, base),
        "PARENT-TRACK": lambda d: build_track(
            d, entry, parent_guide, "Parent Facilitator Run Sheet", printables.PARENT_BUILDERS
        ),
        "YOUNGER-TEACHER-PACK": lambda d: build_track(
            d, entry, younger_guide, "Younger Teacher Run Sheet", printables.YOUNGER_BUILDERS, True
        ),
        "OLDER-TEACHER-PACK": lambda d: build_track(
            d, entry, older_guide, "Older Teacher Run Sheet", printables.OLDER_BUILDERS
        ),
        "FAMILY-HANDOUT": lambda d: build_family(d, entry, base, family_builders),
        "SATURDAY-PRINT-PACKET": lambda d: build_saturday(d, entry, base, printables),
    }

    results = []
    for suffix, title, subtitle, audience in PACKET_SUFFIXES:
        stem = packet_stem(entry, suffix)
        docx = out / f"{stem}.docx"
        pdf = out / f"{stem}.pdf"
        document = _new(entry, f"{entry.id} — {title}", subtitle, audience)
        builders_map[suffix](document)
        add_callout(
            document,
            "RIGHTS_NOTE",
            "KUTUMBA-original framing; cited third-party material retains its own rights. "
            "Human/temple review remains EXTERNAL_OPEN.",
        )
        document.save(docx)
        if pdf.exists():
            pdf.unlink()
        reported = convert_with_word(docx, pdf)
        strip_trailing_empty(pdf)
        pages = _forbidden_check(pdf, entry)
        results.append({
            "name": stem,
            "docx": _relative(docx),
            "docx_bytes": docx.stat().st_size,
            "pdf": _relative(pdf),
            "pdf_bytes": pdf.stat().st_size,
            "pages": pages,
            "word_reported_pages": reported,
        })
        print(f"{stem}: {pages} pages")

    manifest = out / f"{entry.id}-render-manifest.json"
    payload = {
        "version": "V13",
        "week": entry.id,
        "title": entry.title,
        "kind": entry.kind,
        "module": _relative(base),
        "export_subdir": entry.export_subdir,
        "packets": results,
    }
    manifest.write_text(json.dumps(payload, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print(f"Wrote {_relative(manifest)}")
    return payload


def main() -> int:
    parser = argparse.ArgumentParser(description="Render one V13 week packet set.")
    parser.add_argument("--week", required=True, help="Week id from week_registry.yaml, e.g. C1-W2")
    args = parser.parse_args()
    entry = get_week(args.week)
    render_week(entry)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
