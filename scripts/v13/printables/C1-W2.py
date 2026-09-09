#!/usr/bin/env python3
"""C1-W2 V13 printable Word builders for render_week.py (week id C1-W2).

Real Word tables/images. DOCX/PDF emission is owned by scripts/v13/render_week.py.
"""
from __future__ import annotations

import sys
from pathlib import Path

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

REPO = Path(__file__).resolve().parents[3]
sys.path.insert(0, str(REPO / "scripts" / "v13"))
sys.path.insert(0, str(REPO / "scripts" / "v12"))
from printable_common import (  # noqa: E402
    PLUM,
    SAFFRON,
    TEAL,
    add_border,
    add_card_grid,
    add_cut_lines,
    add_memory_phrase_block,
    add_page,
    add_printable_title,
    add_teacher_only_divider,
    add_teacher_page,
    add_write_lines,
)
from kutumba_docx_styles import add_branded_table, add_callout, add_verse_card  # noqa: E402

ASSETS = (
    REPO
    / "11-weekly-program-library"
    / "first-six-months"
    / "c1-w2-i-am-not-this-body"
    / "visuals"
    / "v13"
)
MEMORY = "My body changes; I continue as the conscious self."
VERSE_URL = "https://vedabase.io/en/library/bg/2/13/"
DEV = "देहिनोऽस्मिन् यथा देहे कौमारं यौवनं जरा । तथा देहान्तरप्राप्तिर्धीरस्तत्र न मुह्यति ॥ १३ ॥"
IAST = "dehino ’smin yathā dehe kaumāraṁ yauvanaṁ jarā / tathā dehāntara-prāptir dhīras tatra na muhyati"
MEANING = (
    "Just as the embodied self passes through childhood, youth, and old age in one body, "
    "the sober person is not bewildered when the self passes to another body."
)


def _img(document, name: str, width: float = 2.4) -> None:
    path = ASSETS / name
    if path.is_file():
        p = document.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(path), width=Inches(width))


def y01(document) -> None:
    add_printable_title(
        document,
        "Y01",
        "Life-Stage Sequence Cards",
        "Cut cards. Place in order: child → youth → elder. Bodies change; the person continues.",
    )
    _img(document, "life-stage-sequence.png", 6.0)
    for name, label, line in (
        ("life-stage-child.png", "CHILD", "A smaller body. Growing."),
        ("life-stage-youth.png", "YOUTH", "A growing body. Changing."),
        ("life-stage-elder.png", "ELDER", "An older body. Still a person."),
    ):
        document.add_heading(label, level=2)
        _img(document, name, 2.2)
        p = document.add_paragraph(line)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_memory_phrase_block(document, MEMORY)
    add_cut_lines(document)
    add_callout(document, "TEACHER_NOTE", "Drawings are pedagogy — not proof of the soul. No death scare.")


def y02(document) -> None:
    add_printable_title(
        document,
        "Y02",
        "Same-Person / Changing-Body Matching",
        "Cut and match pairs. Body changed; person continues. Kind words only.",
    )
    add_card_grid(
        document,
        [
            "Baby shoes",
            "Bigger shoes",
            "Short sleeve last year",
            "Longer arms this year",
            "First tooth smile",
            "Missing tooth smile",
            "Small bicycle",
            "Bigger bicycle seat",
        ],
        per_page=8,
    )
    add_cut_lines(document)
    add_teacher_page(
        document,
        "Y02 Answer Key",
        [
            ["1–2", "Same child; feet grew"],
            ["3–4", "Same child; body grew"],
            ["5–6", "Same child; body changed"],
            ["7–8", "Same rider; body grew"],
        ],
    )


def y03(document) -> None:
    add_printable_title(
        document,
        "Y03",
        "Care / Self Cue Craft",
        "Color Care (body) and Self (I). Practice Care → Self with the memory phrase.",
    )
    _img(document, "care-self-craft.png", 6.0)
    if not (ASSETS / "care-self-craft.png").is_file():
        _img(document, "care-self-card.png", 4.5)
    add_write_lines(document, ["Draw one care act or kind speech bubble:"], 2)
    add_memory_phrase_block(document, MEMORY)
    add_callout(document, "TEACHER_NOTE", "Care remains good. Do not teach body neglect or ‘body is trash.’")


def y04(document) -> None:
    add_printable_title(
        document,
        "Y04",
        "Kind Body-Speech Cards",
        "Cut cards. After a redirect, children point to a KIND card.",
    )
    _img(document, "kind-speech-icons.png", 5.5)
    add_card_grid(
        document,
        [
            "You’re a friend.",
            "Bodies change. We stay kind.",
            "I care for my body with sleep and water.",
            "I’m sorry. That was not kind.",
            "Someone teases about hair. → Choose a kind repair.",
            "Someone laughs about height. → Choose a kind repair.",
        ],
    )
    add_cut_lines(document)
    add_teacher_page(
        document,
        "Y04 Teacher Demo (not ‘funny’ handouts)",
        [
            ["Mistaken demo", "You’re weird looking. → Stop; choose kind card"],
            ["Mistaken demo", "It’s just a joke. → Stop; choose repair card"],
            ["Guidance", "Any KIND card OK; prefer repair after hurt"],
        ],
    )


def y05(document) -> None:
    add_printable_title(
        document,
        "Y05",
        "Memory Phrase Mat",
        "Place on floor/table. Tap each line while echoing during Grow-and-Freeze.",
    )
    if (ASSETS / "memory-phrase-mat.png").is_file():
        _img(document, "memory-phrase-mat.png", 6.2)
    else:
        _img(document, "memory-mat.png", 5.0)
    add_memory_phrase_block(document, MEMORY)
    document.add_paragraph("Care for the body. Don’t tease bodies. · BG 2.13 family week")


def o01(document) -> None:
    add_printable_title(document, "O01", "BG 2.13 Observation Worksheet")
    add_verse_card(
        document,
        "BG 2.13",
        DEV,
        IAST,
        MEANING,
        VERSE_URL,
        "Scripture display/source: VedaBase; teaching meaning: KUTUMBA-original (not labeled as BBT translation).",
    )
    add_write_lines(
        document,
        [
            "Body stages named in the verse:",
            "Who is not bewildered? What does sober family speech look like?",
            "What does this verse NOT authorize? (mockery / body contempt as ‘spiritual’ / lab proof of ātman)",
            "One sentence paraphrase in my own words:",
            "Why don’t photos prove the soul?",
            "One question for the facilitator:",
        ],
        1,
    )
    document.add_paragraph("Quick checks: Photos prove the soul? Yes / No · Ignore medicine if ‘spiritual’? Yes / No")


def o02(document) -> None:
    add_printable_title(
        document,
        "O02",
        "Life-Stage Evidence Timeline",
        "Timeline is pedagogy about change/continuity — not soul-proof.",
    )
    _img(document, "life-stage-sequence.png", 5.5)
    add_branded_table(
        document,
        ["Stage", "One body change (no shame language)", "What continues?"],
        [
            ["Childhood", "________________________", "________________________"],
            ["Youth", "________________________", "________________________"],
            ["Elder", "________________________", "________________________"],
        ],
    )
    add_write_lines(
        document,
        [
            "Why is this timeline helpful for teaching BG 2.13?",
            "Why is this timeline not proof of the soul?",
        ],
        2,
    )


def o03(document) -> None:
    add_printable_title(
        document,
        "O03",
        "Identity Statement Sort",
        "Sort onto SELF / CONTINUES · BODY / CHANGES · CARE DUTY. Discuss mistaken cards.",
    )
    mats = document.add_table(rows=1, cols=3)
    add_border(mats, PLUM, 10)
    for cell, label in zip(mats.rows[0].cells, ("SELF / CONTINUES", "BODY / CHANGES", "CARE DUTY")):
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(label)
        r.bold = True
        r.font.size = Pt(14)
        r.font.color.rgb = RGBColor.from_string(PLUM)
    add_page(document)
    add_card_grid(
        document,
        [
            "I can choose kind words.",
            "My height changes over years.",
            "I sleep so I can serve with energy.",
            "I can love and remember Kṛṣṇa.",
            "My shoe size changes.",
            "I drink water and eat to care for this body.",
            "I am the same “I” across childhood and youth.",
            "My hair length changes.",
            "I take medicine when needed.",
            "My whole worth is my appearance. (discuss)",
            "Bodies are trash; skip care. (discuss)",
            "Photos prove the soul. (discuss)",
        ],
        per_page=12,
    )
    add_cut_lines(document)
    add_teacher_page(
        document,
        "O03 Answer Key",
        [
            ["SELF", "kind words; love/remember Kṛṣṇa; same I across stages"],
            ["BODY", "height; shoe size; hair length"],
            ["CARE DUTY", "sleep; water/food; medicine"],
            ["MISTAKEN discuss", "appearance = worth; bodies trash; photos prove soul"],
        ],
    )


def o04(document) -> None:
    add_printable_title(
        document,
        "O04",
        "Respectful Language Scenarios",
        "For each scenario: mistaken conclusion · principle · response · action · what not to say.",
    )
    scenarios = [
        "A: Sibling appearance tease at snack — “Just a joke.”",
        "B: Teen identity = fitness metrics — “The body project is the self.”",
        "C: Relative asks weight publicly — “Public body talk is fine.”",
    ]
    for i, scenario in enumerate(scenarios, 1):
        if i > 1:
            add_page(document)
        document.add_heading(f"Scenario {i}", level=2)
        document.add_paragraph(scenario)
        add_write_lines(
            document,
            [
                "Mistaken conclusion:",
                "Principle from BG 2.13:",
                "Compassionate response:",
                "Better family action:",
                "What not to say:",
            ],
            2,
        )
    add_teacher_page(
        document,
        "O04 Answer Key",
        [
            ["A", "Stop/affirm/repair; speech pledge; not ‘toughen up’ / ‘just a joke’"],
            ["B", "Affirm health; separate care from identity; keep routine + memory/service; not ‘fitness is māyā’"],
            ["C", "Warm redirect; protect child; agree redirect line; not announce stats / angry lecture in front of child"],
        ],
    )


def o05(document) -> None:
    add_printable_title(
        document,
        "O05",
        "Changing-Body / Enduring-Self Diagram",
        "Stages → self continues → kind speech + responsible care.",
    )
    if (ASSETS / "care-self-diagram.png").is_file():
        _img(document, "care-self-diagram.png", 6.0)
    else:
        _img(document, "body-self-diagram.png", 4.2)
    add_write_lines(
        document,
        [
            "Write the short echo inside Kind speech:",
            "One low-risk example of body change:",
            "One care duty that remains good:",
            "Why doesn’t this diagram prove the soul?",
        ],
        2,
    )
    document.add_paragraph("Short echo: Care for the body; don’t confuse it with the self.")


def o06(document) -> None:
    add_printable_title(document, "O06", "Exit Ticket", "Name: __________________________  Date: ______________")
    add_write_lines(
        document,
        [
            "BG 2.13 in one sentence (my words):",
            "One analogy and its limit:",
            "One respectful body-speech rule for home:",
            "True/False — Photos prove the soul:",
            "True/False — Ignore medicine if ‘spiritual’:",
            "One question I still have:",
            "Project line — This week our family will speak about bodies by:",
        ],
        2,
    )


def p01(document) -> None:
    add_printable_title(document, "P01", "Private Reflection (Labels / Roles / Body Changes)")
    _img(document, "parent-icon.png", 3.2)
    add_callout(
        document,
        "SAFETY_PRIVACY",
        "Write privately. Sharing is optional. Do not collect or place completed sheets in Git.",
    )
    _img(document, "parent-care-continuity.png", 3.5)
    p = document.add_paragraph(
        "Which labels, roles, or body-change comments do we use at home that quietly treat the body as the whole self?"
    )
    p.runs[0].bold = True
    add_write_lines(document, ["Private notes:", "Habit to watch this week:"], 4)


def p02(document) -> None:
    add_printable_title(document, "P02", "Respectful Family Language Card Sort", "Sort onto RESPECTFUL vs MISTAKEN.")
    mats = document.add_table(rows=1, cols=2)
    add_border(mats, PLUM, 10)
    for cell, label in zip(mats.rows[0].cells, ("RESPECTFUL", "MISTAKEN")):
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(label)
        r.bold = True
        r.font.size = Pt(22)
        r.font.color.rgb = RGBColor.from_string(TEAL if label == "RESPECTFUL" else SAFFRON)
    add_page(document)
    add_card_grid(
        document,
        [
            "Bodies change. We stay kind.",
            "It’s just a joke — toughen up.",
            "I’m sorry — that joke about looks was not okay.",
            "Your whole worth is how you look.",
            "We care for health; we are not only metrics.",
            "Bodies are trash; real devotees ignore the body.",
            "Let’s change the subject — body stats stay private.",
            "Tell everyone their weight — honesty!",
            "You are a person who can love and serve — not a punchline.",
            "Photos prove we have a soul.",
            "Sleep and medicine are care, not a failure of spirituality.",
            "If you were spiritual, you wouldn’t need rest.",
        ],
        prefix="L",
        per_page=12,
    )
    add_write_lines(
        document,
        [
            "One ‘almost right’ sentence our home sometimes uses:",
            "Better rewrite with a limit:",
        ],
        3,
    )
    add_teacher_page(
        document,
        "P02 Facilitator Match Key",
        [
            ["RESPECTFUL", "L1, L3, L5, L7, L9, L11"],
            ["MISTAKEN", "L2, L4, L6, L8, L10, L12"],
        ],
    )


def p03(document) -> None:
    add_printable_title(document, "P03", "Substantial Family Case")
    add_callout(document, "SAFETY_PRIVACY", "Fictional case. No compelled confession of real events.")
    document.add_paragraph(
        "A middle-school child comes home upset after a relative joked about their height. "
        "A sibling repeats the joke at dinner. One parent says “toughen up.” Another freezes."
    )
    add_write_lines(
        document,
        [
            "Tempting mistaken conclusion:",
            "Principle from BG 2.13:",
            "Compassionate response:",
            "One seven-day household action:",
            "What we should not say:",
        ],
        3,
    )
    add_teacher_page(
        document,
        "P03 Facilitator Answer Direction",
        [
            ["Mistaken", "Jokes harmless / silence enough / toughness fixes hurt"],
            ["Principle", "Body changes; person continues; speech honors dignity (BG 2.13)"],
            ["Response", "Stop repeat; affirm child; privately coach sibling; prepare relative redirect"],
            ["Action", "Speech pledge + nightly memory-line cue"],
            ["Do not say", "Toughen up; spiritual people ignore feelings; announce stats"],
        ],
    )


def p04(document) -> None:
    add_printable_title(document, "P04", "Household Operating Application")
    add_callout(document, "SAFETY_PRIVACY", "No ranking and no public reading required.")
    add_write_lines(
        document,
        [
            "One body-care duty we keep as service readiness:",
            "One speech rule we keep for seven days:",
            "One redirect line when teasing starts:",
            "Who starts the nightly cue if others are tired:",
            "Minimum version on a hard day:",
        ],
        3,
    )
    document.add_paragraph("Family initials (optional): ____________________  Date: ______________")
    add_memory_phrase_block(document, MEMORY)


def p05(document) -> None:
    add_printable_title(document, "P05", "Private Next-Step Saṅkalpa")
    add_callout(document, "KEY_IDEA", "specific action + frequency + trigger + minimum version")
    add_write_lines(
        document,
        [
            "Specific action:",
            "Frequency:",
            "Trigger (when and where):",
            "Minimum version for a hard day:",
            "Where we will place the reminder:",
        ],
        2,
    )
    _img(document, "kind-speech-icons.png", 5.0)
    document.add_paragraph("No photo proof. No ranking. The minimum version counts as success.")


YOUNGER_BUILDERS = [y01, y02, y03, y04, y05]
OLDER_BUILDERS = [o01, o02, o03, o04, o05, o06]
PARENT_BUILDERS = [p01, p02, p03, p04, p05]
FAMILY_BUILDERS = []
