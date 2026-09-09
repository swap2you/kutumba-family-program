#!/usr/bin/env python3
"""C2-W4 V13 printable Word builders for render_week.py (week id C2-W4).

Real Word tables/images. DOCX/PDF emission is owned by scripts/v13/render_week.py.
"""
from __future__ import annotations

import sys
from pathlib import Path

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

REPO = Path(__file__).resolve().parents[3]
sys.path.insert(0, str(REPO / "scripts" / "v13"))
sys.path.insert(0, str(REPO / "scripts" / "v12"))
from printable_common import (  # noqa: E402
    PLUM,
    SAFFRON,
    TEAL,
    add_border,
    add_card_grid,
    add_cut_lines,
    add_memory_phrase_block,
    add_page,
    add_printable_title,
    add_teacher_page,
    add_write_lines,
)
from kutumba_docx_styles import add_branded_table, add_callout, add_verse_card  # noqa: E402

ASSETS = REPO / "11-weekly-program-library/first-six-months/c2-w4-the-three-modes-of-material-nature" / "visuals" / "v13"
MEMORY = 'Calm helps me hear Kṛṣṇa.'
VERSE_REF = 'BG 14.5'
VERSE_URL = 'https://vedabase.io/en/library/bg/14/5/'
DEV = 'सत्त्वं रजस्तम इति गुणा: प्रकृतिसम्भवा: । निबध्नन्ति महाबाहो देहे देहिनमव्ययम् ॥ ५ ॥'
IAST = 'sattvaṁ rajas tama iti guṇāḥ prakṛti-sambhavāḥ / nibadhnanti mahā-bāho dehe dehinam avyayam'
MEANING = 'Goodness, passion, and ignorance — born of material nature — bind the eternal embodied self within the body.'


def _img(document, name: str, width: float = 2.4) -> None:
    path = ASSETS / name
    if path.is_file():
        p = document.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(path), width=Inches(width))


def y01(document) -> None:
    add_printable_title(document, "Y01", "Mode weather corners", "Sunny/windy/foggy corners for calm/busy/sleepy moments — not people.")
    _img(document, "mode-weather.png", 5.5)
    add_card_grid(document, ['Sunny calm', 'Windy rush', 'Foggy sleepy', 'Reset corner', 'Quiet voice', 'Call someone a mode'], per_page=6)
    add_memory_phrase_block(document, MEMORY)
    add_cut_lines(document)
    add_callout(document, "TEACHER_NOTE", 'Never label persons. Avoid food policing. Environment supports practice; devotion is the goal.')

def y02(document) -> None:
    add_printable_title(document, "Y02", "Moment matching", "Match scenes to weather icons.")
    _img(document, "scenario-board.png", 5.5)
    add_card_grid(document, ['Sunny calm', 'Windy rush', 'Foggy sleepy', 'Reset corner', 'Quiet voice', 'Call someone a mode'], per_page=6)
    add_memory_phrase_block(document, MEMORY)
    add_cut_lines(document)
    add_callout(document, "TEACHER_NOTE", 'Never label persons. Avoid food policing. Environment supports practice; devotion is the goal.')

def y03(document) -> None:
    add_printable_title(document, "Y03", "Reset-corner craft", "Make a calm-corner card for home.")
    _img(document, "clarity-reset.png", 5.5)
    add_card_grid(document, ['Sunny calm', 'Windy rush', 'Foggy sleepy', 'Reset corner', 'Quiet voice', 'Call someone a mode'], per_page=6)
    add_memory_phrase_block(document, MEMORY)
    add_cut_lines(document)
    add_callout(document, "TEACHER_NOTE", 'Never label persons. Avoid food policing. Environment supports practice; devotion is the goal.')

def y04(document) -> None:
    add_printable_title(document, "Y04", "Kind reset cards", "Words for resetting a chaotic room.")
    _img(document, "memory-mat.png", 5.5)
    add_card_grid(document, ['Sunny calm', 'Windy rush', 'Foggy sleepy', 'Reset corner', 'Quiet voice', 'Call someone a mode'], per_page=6)
    add_memory_phrase_block(document, MEMORY)
    add_cut_lines(document)
    add_callout(document, "TEACHER_NOTE", 'Never label persons. Avoid food policing. Environment supports practice; devotion is the goal.')

def y05(document) -> None:
    add_printable_title(document, "Y05", "Memory phrase mat", "Echo: Calm helps me hear Kṛṣṇa.")
    _img(document, "memory-mat.png", 6.0)
    add_memory_phrase_block(document, MEMORY)
    document.add_paragraph('BG 14.5' + " · family week")


def o01(document) -> None:
    add_printable_title(document, "O01", "BG 14.5 observation", "Name three modes; what they do; what they are not.")
    add_verse_card(document, VERSE_REF, DEV, IAST, MEANING, VERSE_URL,
        "Scripture display/source: VedaBase; teaching meaning: KUTUMBA-original (not labeled as BBT translation).")
    add_write_lines(document, [
        "What does the verse teach in my words?",
        "What misconception does this week block?",
        "One question for the facilitator:",
    ], 2)

def o02(document) -> None:
    add_printable_title(document, "O02", "Mode scenario board", "Sort scenarios into moment influences.")
    _img(document, "scenario-board.png", 5.0)
    add_branded_table(document, ["Prompt", "My notes"], [["Mode scenario board", "________________"], ["Limit / care", "________________"]])
    add_write_lines(document, ["Application:", "What not to say:"], 2)
    add_callout(document, "KEY_IDEA", 'Sattva, rajas, and tamas condition the embodied self; notice the moment; design supports; never label persons.')

def o03(document) -> None:
    add_printable_title(document, "O03", "Support vs label sort", "Helpful environment support vs person-labeling.")
    _img(document, "clarity-reset.png", 5.0)
    add_branded_table(document, ["Prompt", "My notes"], [["Support vs label sort", "________________"], ["Limit / care", "________________"]])
    add_write_lines(document, ["Application:", "What not to say:"], 2)
    add_callout(document, "KEY_IDEA", 'Sattva, rajas, and tamas condition the embodied self; notice the moment; design supports; never label persons.')

def o04(document) -> None:
    add_printable_title(document, "O04", "Habit/environment scenarios", "Three household chaos/clarity cases.")
    _img(document, "memory-mat.png", 5.0)
    add_branded_table(document, ["Prompt", "My notes"], [["Habit/environment scenarios", "________________"], ["Limit / care", "________________"]])
    add_write_lines(document, ["Application:", "What not to say:"], 2)
    add_callout(document, "KEY_IDEA", 'Sattva, rajas, and tamas condition the embodied self; notice the moment; design supports; never label persons.')

def o05(document) -> None:
    add_printable_title(document, "O05", "Clarity-before-chanting diagram", "Reset → practice → review.")
    _img(document, "parent-environment.png", 5.0)
    add_branded_table(document, ["Prompt", "My notes"], [["Clarity-before-chanting diagram", "________________"], ["Limit / care", "________________"]])
    add_write_lines(document, ["Application:", "What not to say:"], 2)
    add_callout(document, "KEY_IDEA", 'Sattva, rajas, and tamas condition the embodied self; notice the moment; design supports; never label persons.')

def o06(document) -> None:
    add_printable_title(document, "O06", "Exit ticket", "One moment notice + one support + no person label.")
    add_write_lines(document, [
        "Primary in one sentence:",
        "One analogy or mechanic and its limit:",
        "One home action:",
        "One question I still have:",
    ], 2)


def p01(document) -> None:
    add_printable_title(document, "P01", "Private reflection — habit environment", "Which room cues push restless or dull moments?")
    _img(document, "parent-icon.png", 3.2)
    add_callout(document, "SAFETY_PRIVACY", "Write privately. Sharing is optional. Do not collect or place completed sheets in Git.")
    _img(document, "parent-environment.png", 3.5)
    p = document.add_paragraph('Which room cues push restless or dull moments?')
    p.runs[0].bold = True
    add_write_lines(document, ["Private notes:", "Habit to watch this week:"], 4)

def p02(document) -> None:
    add_printable_title(document, "P02", "Support vs label card sort", "Sort SUPPORT / LABELING / FOOD POLICING.")
    mats = document.add_table(rows=1, cols=2)
    add_border(mats, PLUM, 10)
    for cell, label in zip(mats.rows[0].cells, ("HELPFUL", "MISTAKEN")):
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(label)
        r.bold = True
        r.font.size = Pt(22)
        r.font.color.rgb = RGBColor.from_string(TEAL if label == "HELPFUL" else SAFFRON)
    add_page(document)
    add_card_grid(document, ['Clear the table before reading', 'You are tamasic', 'Dim lights for bedtime', 'Our family is rajasic', 'Phone basket at door', 'Food police lecture'], prefix="L", per_page=12)
    add_callout(document, "TEACHER_NOTE", 'Labeling people/families as a mode; food policing as spiritual superiority.')
    _img(document, "parent-environment.png", 3.0)

def p03(document) -> None:
    add_printable_title(document, "P03", "Substantial family case", "Parent calls a child “tamasic” after late screens.")
    add_callout(document, "SAFETY_PRIVACY", "Fictional case. No compelled confession of real events.")
    document.add_paragraph('After late screens, a parent announces “you’re being tamasic again” in front of siblings. The child shuts down; another parent starts food policing as spiritual ranking.')
    add_write_lines(document, [
        "Tempting mistaken conclusion:",
        "Principle from primary:",
        "Compassionate response:",
        "One seven-day household action:",
        "What we should not say:",
    ], 3)
    _img(document, "mode-weather.png", 4.0)

def p04(document) -> None:
    add_printable_title(document, "P04", "Household operating application", "One five-minute reset ritual for seven days.")
    add_callout(document, "SAFETY_PRIVACY", "No ranking and no public reading required.")
    add_write_lines(document, [
        "One household action for seven days:",
        "Trigger (when/where):",
        "Who starts if others are tired:",
        "Minimum version on a hard day:",
    ], 3)
    add_memory_phrase_block(document, MEMORY)
    _img(document, "parent-environment.png", 3.2)

def p05(document) -> None:
    add_printable_title(document, "P05", "Private next-step saṅkalpa", "Environment support + minimum version.")
    add_callout(document, "KEY_IDEA", "specific action + frequency + trigger + minimum version")
    add_write_lines(document, [
        "Specific action:",
        "Frequency:",
        "Trigger (when and where):",
        "Minimum version for a hard day:",
        "Where we will place the reminder:",
    ], 2)
    _img(document, "parent-environment.png", 4.5)
    document.add_paragraph("No ranking. The minimum version counts as success.")


YOUNGER_BUILDERS = [y01, y02, y03, y04, y05]
OLDER_BUILDERS = [o01, o02, o03, o04, o05, o06]
PARENT_BUILDERS = [p01, p02, p03, p04, p05]
FAMILY_BUILDERS = []
