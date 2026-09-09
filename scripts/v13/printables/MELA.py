#!/usr/bin/env python3
"""MELA V13 printable Word builders for render_week.py.

Real Word tables/images. DOCX/PDF emission owned by scripts/v13/render_week.py.
This pass authors builders only — do not render here.
"""
from __future__ import annotations

import sys
from pathlib import Path

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches

REPO = Path(__file__).resolve().parents[3]
sys.path.insert(0, str(REPO / "scripts" / "v13"))
sys.path.insert(0, str(REPO / "scripts" / "v12"))
from printable_common import (  # noqa: E402
    add_card_grid,
    add_cut_lines,
    add_memory_phrase_block,
    add_printable_title,
    add_write_lines,
)
from kutumba_docx_styles import add_callout, add_verse_card  # noqa: E402

ASSETS = (
    REPO
    / "11-weekly-program-library"
    / "first-six-months"
    / "mela-six-month-reflection-family-mela"
    / "visuals"
    / "v13"
)
MEMORY = "Continue · Review · Strengthen — without comparison."
VERSE_URL = "https://vedabase.io/en/library/sb/1/2/18/"


def _img(document, name: str, width: float = 2.4) -> None:
    path = ASSETS / name
    if path.is_file():
        p = document.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(path), width=Inches(width))



def y01(document) -> None:
    add_printable_title(document, "Y01", "Exhibition Walk", "See · say one kind word")
    _img(document, "exhibition-feedback-loop.png", 6.0)
    add_memory_phrase_block(document, MEMORY)


def y02(document) -> None:
    add_printable_title(document, "Y02", "Appreciation Cards", "No ranking winners")
    add_card_grid(document, ["I noticed...", "I learned...", "Thank you for...", "This helped me..."], per_page=4)
    add_callout(document, "CHILD_ACTIVITY", "Appreciation only — no scoreboard.")


def y03(document) -> None:
    add_printable_title(document, "Y03", "Favorite Memory Craft", "")
    add_write_lines(document, ["My favorite learning memory:", "One person who helped me:"], 3)


def y04(document) -> None:
    add_printable_title(document, "Y04", "Feedback Smile Cards", "Kind · specific · hopeful")
    add_write_lines(document, ["Kind:", "Specific:", "Hopeful:"], 2)


def y05(document) -> None:
    add_printable_title(document, "Y05", "Memory Phrase Mat", MEMORY)
    add_memory_phrase_block(document, MEMORY)


def o01(document) -> None:
    add_printable_title(document, "O01", "Gentle Verse-Chain Quiz", "Not a tournament")
    add_write_lines(document, ["One verse I remember:", "One week theme I remember:", "One family practice I remember:"], 2)
    add_callout(document, "TEACHER_NOTE", "Celebrate recall; do not rank scores publicly.")


def o02(document) -> None:
    add_printable_title(document, "O02", "Exhibition Host Sheet", "")
    add_write_lines(document, ["Welcome line:", "One-sentence explain:", "Thank-you line:"], 2)


def o03(document) -> None:
    add_printable_title(document, "O03", "Feedback That Helps", "")
    add_write_lines(document, ["Comparative line to avoid:", "Helpful rewrite:"], 3)


def o04(document) -> None:
    add_printable_title(document, "O04", "Continue / Review / Strengthen", "")
    _img(document, "continue-review-strengthen.png", 5.5)
    add_write_lines(document, ["CONTINUE:", "REVIEW:", "STRENGTHEN:"], 2)


def o05(document) -> None:
    add_printable_title(document, "O05", "Next-Phase Preview", "High-level only")
    add_write_lines(document, ["One curiosity for next phase:", "One support we need:"], 2)


def o06(document) -> None:
    add_printable_title(document, "O06", "Exit Ticket", "")
    add_write_lines(document, ["Continue:", "Review:", "Strengthen:"], 2)


def p01(document) -> None:
    add_printable_title(document, "P01", "Private Six-Month Reflection", "")
    _img(document, "parent-icon.png", 3.2)
    add_write_lines(document, ["Gift of these months:", "Hard stretch:", "Grace noticed:"], 3)
    add_callout(document, "SAFETY_PRIVACY", "Private. Optional share. No ranking.")


def p02(document) -> None:
    add_printable_title(document, "P02", "Review-Chain Observation", "ŚB 1.2.18 center")
    add_verse_card(document, "ŚB 1.2.18", "", "", "Hearing and serving culture nourishes bhakti — review center for tonight.", VERSE_URL, 'Sanskrit/Bengali display + KUTUMBA teaching meaning; no full purport')
    add_write_lines(document, ["What we heard well:", "What we want to hear better:"], 3)


def p03(document) -> None:
    add_printable_title(document, "P03", "Testimony Prep (optional)", "")
    _img(document, "family-testimony-frame.png", 5.0)
    add_write_lines(document, ["One sentence without comparison:", "Opt-out is OK because:"], 2)
    add_callout(document, "SAFETY_PRIVACY", "Invitation only.")


def p04(document) -> None:
    add_printable_title(document, "P04", "Continue / Review / Strengthen", "")
    add_write_lines(document, ["CONTINUE:", "REVIEW:", "STRENGTHEN:", "Who supports this at home:"], 2)
    add_callout(document, "FAMILY_APPLICATION", "Decision without competition.")


def p05(document) -> None:
    add_printable_title(document, "P05", "Private Next-Phase Saṅkalpa", "")
    add_callout(document, "KEY_IDEA", "specific action + frequency + trigger + minimum version")
    add_write_lines(document, ["Specific action:", "Frequency:", "Trigger:", "Minimum version:"], 2)


YOUNGER_BUILDERS = [y01, y02, y03, y04, y05]
OLDER_BUILDERS = [o01, o02, o03, o04, o05, o06]
PARENT_BUILDERS = [p01, p02, p03, p04, p05]


FAMILY_BUILDERS = []
