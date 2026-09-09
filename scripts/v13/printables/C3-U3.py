#!/usr/bin/env python3
"""C3-U3 V13 printable Word builders for render_week.py.

Real Word tables/images. DOCX/PDF emission owned by scripts/v13/render_week.py.
This pass authors builders only — do not render here.
"""
from __future__ import annotations

import sys
from pathlib import Path

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches

REPO = Path(__file__).resolve().parents[3]
sys.path.insert(0, str(REPO / "scripts" / "v13"))
sys.path.insert(0, str(REPO / "scripts" / "v12"))
from printable_common import (  # noqa: E402
    add_card_grid,
    add_cut_lines,
    add_memory_phrase_block,
    add_printable_title,
    add_write_lines,
)
from kutumba_docx_styles import add_callout, add_verse_card  # noqa: E402

ASSETS = (
    REPO
    / "11-weekly-program-library"
    / "first-six-months"
    / "c3-u3-nityananda-trayodasi-family-utsava"
    / "visuals"
    / "v13"
)
MEMORY = "Nityānanda is an incarnation of mercy."
VERSE_URL = "https://vedabase.io/en/library/cc/adi/5/208/"


def _img(document, name: str, width: float = 2.4) -> None:
    path = ASSETS / name
    if path.is_file():
        p = document.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(path), width=Inches(width))



def y01(document) -> None:
    add_printable_title(document, "Y01", "Mercy Chain", "Hear → mercy → serve → invite")
    _img(document, "mercy-chain.png", 6.0)
    add_memory_phrase_block(document, MEMORY)
    add_callout(document, "CHILD_ACTIVITY", "No scary or violent pictures.")


def y02(document) -> None:
    add_printable_title(document, "Y02", "Service Challenge Cards", "")
    _img(document, "service-challenge-cards.png", 5.5)
    add_cut_lines(document)


def y03(document) -> None:
    add_printable_title(document, "Y03", "Help-Hands Craft", "Trace and write one help.")
    add_write_lines(document, ["My helping hand will:", "For whom:"], 2)


def y04(document) -> None:
    add_printable_title(document, "Y04", "Kind Invitation Cards", "")
    add_card_grid(document, ["Please sing with us", "Listening is welcome", "We can help together", "Come share kindness"], per_page=4)


def y05(document) -> None:
    add_printable_title(document, "Y05", "Memory Phrase Mat", MEMORY)
    add_memory_phrase_block(document, MEMORY)


def o01(document) -> None:
    add_printable_title(document, "O01", "CC Ādi 5.208 Observation", "")
    add_verse_card(document, "CC Ādi 5.208", "", "", "Nityānanda is an incarnation of mercy and does not distinguish good and bad.", VERSE_URL, 'Sanskrit/Bengali display + KUTUMBA teaching meaning; no full purport')
    add_write_lines(document, ["What is mercy here?", "What must we not dramatize?"], 3)


def o02(document) -> None:
    add_printable_title(document, "O02", "Mercy vs Judgment Sort", "")
    add_card_grid(document, ["Invite kindly", "Mock the fallen", "Serve quietly", "Public shame", "Listen in kīrtana", "Scoreboard seva"], per_page=6)


def o03(document) -> None:
    add_printable_title(document, "O03", "Age-Sensitive Boundaries", "Rewrite without graphic violence.")
    add_write_lines(document, ["Harsh line (do not perform):", "Mercy-focused rewrite:", "Why the rewrite is better:"], 3)
    add_callout(document, "TEACHER_NOTE", "No graphic Jagāi-Mādhāi violence in materials.")


def o04(document) -> None:
    add_printable_title(document, "O04", "Service Plan", "Who can we serve?")
    add_write_lines(document, ["Person/place:", "Help act:", "When:", "Minimum version:"], 2)


def o05(document) -> None:
    add_printable_title(document, "O05", "Saṅkīrtana Invitation Map", "")
    _img(document, "sankirtana-invitation.png", 4.5)
    add_write_lines(document, ["Invitation line:", "Listening counts because:"], 2)


def o06(document) -> None:
    add_printable_title(document, "O06", "Exit Ticket", "")
    add_write_lines(document, ["Memory line:", "One mercy act:", "One person to serve:"], 2)


def p01(document) -> None:
    add_printable_title(document, "P01", "Private Reflection", "")
    _img(document, "parent-icon.png", 3.2)
    add_write_lines(document, ["Where judgment blocks mercy:", "One softer response:"], 3)
    add_callout(document, "SAFETY_PRIVACY", "Private optional share.")


def p02(document) -> None:
    add_printable_title(document, "P02", "Source Observation", "")
    add_verse_card(document, "CC Ādi 5.208", "", "", "Incarnation of mercy.", VERSE_URL, 'Sanskrit/Bengali display + KUTUMBA teaching meaning; no full purport')
    add_write_lines(document, ["Takeaway:", "Calendar honesty:"], 3)


def p03(document) -> None:
    add_printable_title(document, "P03", "Age-Sensitive Case", "Violent dramatization request")
    add_write_lines(document, ["Mistaken goal:", "Principle:", "Redirect:"], 3)
    add_callout(document, "DO_NOT_SPECULATE", "No graphic violence; no miracle embellishment.")


def p04(document) -> None:
    add_printable_title(document, "P04", "Who Can We Serve?", "")
    add_write_lines(document, ["Household:", "Community:", "Temple/service space:", "Minimum version:"], 2)
    add_callout(document, "FAMILY_APPLICATION", "One real act beats many slogans.")


def p05(document) -> None:
    add_printable_title(document, "P05", "Private Saṅkalpa", "")
    add_callout(document, "KEY_IDEA", "specific action + frequency + trigger + minimum version")
    add_write_lines(document, ["Specific action:", "Frequency:", "Trigger:", "Minimum version:"], 2)


YOUNGER_BUILDERS = [y01, y02, y03, y04, y05]
OLDER_BUILDERS = [o01, o02, o03, o04, o05, o06]
PARENT_BUILDERS = [p01, p02, p03, p04, p05]


FAMILY_BUILDERS = []
