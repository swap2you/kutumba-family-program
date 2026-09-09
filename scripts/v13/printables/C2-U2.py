#!/usr/bin/env python3
"""C2-U2 V13 printable Word builders for render_week.py.

Real Word tables/images. DOCX/PDF emission owned by scripts/v13/render_week.py.
This pass authors builders only — do not render here.
"""
from __future__ import annotations

import sys
from pathlib import Path

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches

REPO = Path(__file__).resolve().parents[3]
sys.path.insert(0, str(REPO / "scripts" / "v13"))
sys.path.insert(0, str(REPO / "scripts" / "v12"))
from printable_common import (  # noqa: E402
    add_card_grid,
    add_cut_lines,
    add_memory_phrase_block,
    add_printable_title,
    add_write_lines,
)
from kutumba_docx_styles import add_callout, add_verse_card  # noqa: E402

ASSETS = (
    REPO
    / "11-weekly-program-library"
    / "first-six-months"
    / "c2-u2-gita-jayanti-family-utsava"
    / "visuals"
    / "v13"
)
MEMORY = "Think of Kṛṣṇa; become His devotee."
VERSE_URL = "https://vedabase.io/en/library/bg/18/65/"


def _img(document, name: str, width: float = 2.4) -> None:
    path = ASSETS / name
    if path.is_file():
        p = document.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(path), width=Inches(width))



def y01(document) -> None:
    add_printable_title(document, "Y01", "18-Chapter Map", "Point and wonder. No forced memorization.")
    _img(document, "gita-18-chapter-map.png", 6.2)
    add_memory_phrase_block(document, MEMORY)


def y02(document) -> None:
    add_printable_title(document, "Y02", "Find the Gītā Theme", "Match theme cards.")
    add_card_grid(document, ["Remember Kṛṣṇa", "Be a devotee", "Worship", "Offer homage", "Kind share", "Not a contest"], per_page=6)
    add_cut_lines(document)


def y03(document) -> None:
    add_printable_title(document, "Y03", "Devotee Craft", "I can think of Kṛṣṇa when…")
    add_write_lines(document, ["When I wake:", "When I help:", "When I sing:"], 2)


def y04(document) -> None:
    add_printable_title(document, "Y04", "Book Gratitude Card", "Thank you for the Gītā.")
    add_write_lines(document, ["I am grateful for:", "One kind thing I can do:"], 2)
    add_callout(document, "TEACHER_NOTE", "No child book-distribution quotas.")


def y05(document) -> None:
    add_printable_title(document, "Y05", "Memory Phrase Mat", MEMORY)
    add_memory_phrase_block(document, MEMORY)


def o01(document) -> None:
    add_printable_title(document, "O01", "BG 18.65 Observation", "")
    add_verse_card(document, "BG 18.65", "", "", "Think of Me; become My devotee; worship Me; offer homage.", VERSE_URL, 'Sanskrit/Bengali display + KUTUMBA teaching meaning; no full purport')
    add_write_lines(document, ["Four invitations I see:", "One promise I notice:", "One family application:"], 3)


def o02(document) -> None:
    add_printable_title(document, "O02", "18-Chapter Map Worksheet", "")
    _img(document, "gita-18-chapter-map.png", 6.0)
    add_write_lines(document, ["A chapter number I notice:", "Why a map helps:", "Map limit:"], 2)


def o03(document) -> None:
    add_printable_title(document, "O03", "Theme Sort", "")
    add_card_grid(document, ["man-manā", "mad-bhakta", "mad-yājī", "namaskuru", "ranking contest", "quota pressure"], per_page=6)


def o04(document) -> None:
    add_printable_title(document, "O04", "Arjuna Decision Case", "Paraphrase only — no invented dialogue.")
    _img(document, "arjuna-decision-frame.png", 5.5)
    add_write_lines(document, ["Duty tension:", "Fear tension:", "Kṛṣṇa’s instruction direction:"], 3)


def o05(document) -> None:
    add_printable_title(document, "O05", "Family Reading Plan", "")
    _img(document, "family-gita-reading-plan.png", 5.0)
    add_write_lines(document, ["When:", "How long (minutes):", "Minimum version:"], 2)


def o06(document) -> None:
    add_printable_title(document, "O06", "Exit Ticket", "")
    add_write_lines(document, ["Memory line:", "One BG 18.65 word:", "Home reading cue:"], 2)


def p01(document) -> None:
    add_printable_title(document, "P01", "Private Reflection", "")
    _img(document, "parent-icon.png", 3.2)
    add_write_lines(document, ["Where Gītā already touches our home:", "Where pressure sneaks in:"], 3)
    add_callout(document, "SAFETY_PRIVACY", "Private writing. Optional share.")


def p02(document) -> None:
    add_printable_title(document, "P02", "Source Observation", "BG 18.65")
    add_verse_card(document, "BG 18.65", "", "", "Always think of Me; become My devotee.", VERSE_URL, 'Sanskrit/Bengali display + KUTUMBA teaching meaning; no full purport')
    add_write_lines(document, ["What stands out?", "Calendar honesty note:"], 3)


def p03(document) -> None:
    add_printable_title(document, "P03", "Family Case", "Child distribution quota pressure")
    add_write_lines(document, ["Mistaken conclusion:", "Principle:", "Action:"], 3)
    add_callout(document, "DO_NOT_SPECULATE", "No child quotas. Local tithi EXTERNAL_OPEN.")


def p04(document) -> None:
    add_printable_title(document, "P04", "Household Reading Plan", "")
    add_write_lines(document, ["Cue:", "Duration:", "Who starts:", "Minimum version:"], 2)
    add_callout(document, "FAMILY_APPLICATION", "Tiny and regular.")


def p05(document) -> None:
    add_printable_title(document, "P05", "Private Saṅkalpa", "")
    add_callout(document, "KEY_IDEA", "specific action + frequency + trigger + minimum version")
    add_write_lines(document, ["Specific action:", "Frequency:", "Trigger:", "Minimum version:"], 2)


YOUNGER_BUILDERS = [y01, y02, y03, y04, y05]
OLDER_BUILDERS = [o01, o02, o03, o04, o05, o06]
PARENT_BUILDERS = [p01, p02, p03, p04, p05]


FAMILY_BUILDERS = []
