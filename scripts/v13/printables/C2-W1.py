#!/usr/bin/env python3
"""C2-W1 V13 printable Word builders for render_week.py (week id C2-W1).

Real Word tables/images. DOCX/PDF emission is owned by scripts/v13/render_week.py.
"""
from __future__ import annotations

import sys
from pathlib import Path

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

REPO = Path(__file__).resolve().parents[3]
sys.path.insert(0, str(REPO / "scripts" / "v13"))
sys.path.insert(0, str(REPO / "scripts" / "v12"))
from printable_common import (  # noqa: E402
    PLUM,
    SAFFRON,
    TEAL,
    add_border,
    add_card_grid,
    add_cut_lines,
    add_memory_phrase_block,
    add_page,
    add_printable_title,
    add_teacher_page,
    add_write_lines,
)
from kutumba_docx_styles import add_branded_table, add_callout, add_verse_card  # noqa: E402

ASSETS = REPO / "11-weekly-program-library/first-six-months/c2-w1-action-and-reaction-how-karma-binds" / "visuals" / "v13"
MEMORY = 'Choices are seeds. I want to plant service seeds.'
VERSE_REF = 'BG 3.9'
VERSE_URL = 'https://vedabase.io/en/library/bg/3/9/'
DEV = 'यज्ञार्थात्कर्मणोऽन्यत्र लोकोऽयं कर्मबन्धनः । तदर्थं कर्म कौन्तेय मुक्तसङ्गः समाचर ॥ ९ ॥'
IAST = 'yajñārthāt karmaṇo ’nyatra loko ’yaṁ karma-bandhanaḥ / tad-arthaṁ karma kaunteya mukta-saṅgaḥ samācara'
MEANING = 'Work done for the satisfaction of the Lord frees us from bondage; otherwise action binds. Therefore act with detachment for His purpose.'


def _img(document, name: str, width: float = 2.4) -> None:
    path = ASSETS / name
    if path.is_file():
        p = document.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(path), width=Inches(width))


def y01(document) -> None:
    add_printable_title(document, "Y01", "Domino / cause-effect path (limits)", "Place domino cards: action → next → result. Say the limit aloud.")
    _img(document, "cause-effect-path.png", 5.5)
    add_card_grid(document, ['Plant a help seed', 'Send a sharp text seed', 'Share a toy seed', 'Hide a mess seed', 'Say sorry seed', 'Blame loudly seed'], per_page=6)
    add_memory_phrase_block(document, MEMORY)
    add_cut_lines(document)
    add_callout(document, "TEACHER_NOTE", 'No victim-blaming. No graphic hell imagery with children.')

def y02(document) -> None:
    add_printable_title(document, "Y02", "Seed choice matching", "Match planting pictures to kind/service vs sharp/selfish seeds.")
    _img(document, "seed-cards.png", 5.5)
    add_card_grid(document, ['Plant a help seed', 'Send a sharp text seed', 'Share a toy seed', 'Hide a mess seed', 'Say sorry seed', 'Blame loudly seed'], per_page=6)
    add_memory_phrase_block(document, MEMORY)
    add_cut_lines(document)
    add_callout(document, "TEACHER_NOTE", 'No victim-blaming. No graphic hell imagery with children.')

def y03(document) -> None:
    add_printable_title(document, "Y03", "Before–During–After craft", "Fold a three-panel card for one small help at home.")
    _img(document, "bda-panels.png", 5.5)
    add_card_grid(document, ['Plant a help seed', 'Send a sharp text seed', 'Share a toy seed', 'Hide a mess seed', 'Say sorry seed', 'Blame loudly seed'], per_page=6)
    add_memory_phrase_block(document, MEMORY)
    add_cut_lines(document)
    add_callout(document, "TEACHER_NOTE", 'No victim-blaming. No graphic hell imagery with children.')

def y04(document) -> None:
    add_printable_title(document, "Y04", "Soft-words redirect cards", "After hot words, point to a soft repair card.")
    _img(document, "memory-mat.png", 5.5)
    add_card_grid(document, ['Plant a help seed', 'Send a sharp text seed', 'Share a toy seed', 'Hide a mess seed', 'Say sorry seed', 'Blame loudly seed'], per_page=6)
    add_memory_phrase_block(document, MEMORY)
    add_cut_lines(document)
    add_callout(document, "TEACHER_NOTE", 'No victim-blaming. No graphic hell imagery with children.')

def y05(document) -> None:
    add_printable_title(document, "Y05", "Memory phrase mat", "Tap each line while echoing the memory phrase.")
    _img(document, "memory-mat.png", 6.0)
    add_memory_phrase_block(document, MEMORY)
    document.add_paragraph('BG 3.9' + " · family week")


def o01(document) -> None:
    add_printable_title(document, "O01", "BG 3.9 observation", "Observe verse layers; paraphrase; name what the verse does not authorize.")
    add_verse_card(document, VERSE_REF, DEV, IAST, MEANING, VERSE_URL,
        "Scripture display/source: VedaBase; teaching meaning: KUTUMBA-original (not labeled as BBT translation).")
    add_write_lines(document, [
        "What does the verse teach in my words?",
        "What misconception does this week block?",
        "One question for the facilitator:",
    ], 2)

def o02(document) -> None:
    add_printable_title(document, "O02", "Action → intention → consequence map", "Fill three boxes for a school/home scenario.")
    _img(document, "seed-cards.png", 5.0)
    add_branded_table(document, ["Prompt", "My notes"], [["Action → intention → consequence map", "________________"], ["Limit / care", "________________"]])
    add_write_lines(document, ["Application:", "What not to say:"], 2)
    add_callout(document, "KEY_IDEA", 'Action with selfish attachment binds; action for the Lord’s purpose can be purified.')

def o03(document) -> None:
    add_printable_title(document, "O03", "Service vs selfish-attachment sort", "Sort cards; discuss almost-right traps.")
    _img(document, "bda-panels.png", 5.0)
    add_branded_table(document, ["Prompt", "My notes"], [["Service vs selfish-attachment sort", "________________"], ["Limit / care", "________________"]])
    add_write_lines(document, ["Application:", "What not to say:"], 2)
    add_callout(document, "KEY_IDEA", 'Action with selfish attachment binds; action for the Lord’s purpose can be purified.')

def o04(document) -> None:
    add_printable_title(document, "O04", "Choice–consequence scenarios", "Work three cases with principle + better action.")
    _img(document, "memory-mat.png", 5.0)
    add_branded_table(document, ["Prompt", "My notes"], [["Choice–consequence scenarios", "________________"], ["Limit / care", "________________"]])
    add_write_lines(document, ["Application:", "What not to say:"], 2)
    add_callout(document, "KEY_IDEA", 'Action with selfish attachment binds; action for the Lord’s purpose can be purified.')

def o05(document) -> None:
    add_printable_title(document, "O05", "Before–During–After diagram", "Map one duty with offered intention.")
    _img(document, "parent-consequence.png", 5.0)
    add_branded_table(document, ["Prompt", "My notes"], [["Before–During–After diagram", "________________"], ["Limit / care", "________________"]])
    add_write_lines(document, ["Application:", "What not to say:"], 2)
    add_callout(document, "KEY_IDEA", 'Action with selfish attachment binds; action for the Lord’s purpose can be purified.')

def o06(document) -> None:
    add_printable_title(document, "O06", "Exit ticket", "One-sentence paraphrase + one home action.")
    add_write_lines(document, [
        "Primary in one sentence:",
        "One analogy or mechanic and its limit:",
        "One home action:",
        "One question I still have:",
    ], 2)


def p01(document) -> None:
    add_printable_title(document, "P01", "Private reflection — binding speech/habits", "Where does our home act for face/control rather than service?")
    _img(document, "parent-icon.png", 3.2)
    add_callout(document, "SAFETY_PRIVACY", "Write privately. Sharing is optional. Do not collect or place completed sheets in Git.")
    _img(document, "parent-consequence.png", 3.5)
    p = document.add_paragraph('Where does our home act for face/control rather than service?')
    p.runs[0].bold = True
    add_write_lines(document, ["Private notes:", "Habit to watch this week:"], 4)

def p02(document) -> None:
    add_printable_title(document, "P02", "Consequence case card sort", "Sort RESPONSIBLE REPAIR vs MISTAKEN DIAGNOSIS.")
    mats = document.add_table(rows=1, cols=2)
    add_border(mats, PLUM, 10)
    for cell, label in zip(mats.rows[0].cells, ("HELPFUL", "MISTAKEN")):
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(label)
        r.bold = True
        r.font.size = Pt(22)
        r.font.color.rgb = RGBColor.from_string(TEAL if label == "HELPFUL" else SAFFRON)
    add_page(document)
    add_card_grid(document, ['Quiet kitchen help', 'Gift for applause', 'Repair private apology', 'Diagnose cousin’s illness as karma', 'Offer food first', 'Win the argument at all costs'], prefix="L", per_page=12)
    add_callout(document, "TEACHER_NOTE", 'Fatalism removes responsibility; do not diagnose others’ hardship as specific past karma.')
    _img(document, "parent-consequence.png", 3.0)

def p03(document) -> None:
    add_printable_title(document, "P03", "Substantial family case", "Public blame text after a missed pickup.")
    add_callout(document, "SAFETY_PRIVACY", "Fictional case. No compelled confession of real events.")
    document.add_paragraph('After a missed pickup, a parent sends a sharp group-chat blame message copying relatives. The logistics need was real; the public heat was extra.')
    add_write_lines(document, [
        "Tempting mistaken conclusion:",
        "Principle from primary:",
        "Compassionate response:",
        "One seven-day household action:",
        "What we should not say:",
    ], 3)
    _img(document, "cause-effect-path.png", 4.0)

def p04(document) -> None:
    add_printable_title(document, "P04", "Household operating application", "One Before–During–After duty for seven days.")
    add_callout(document, "SAFETY_PRIVACY", "No ranking and no public reading required.")
    add_write_lines(document, [
        "One household action for seven days:",
        "Trigger (when/where):",
        "Who starts if others are tired:",
        "Minimum version on a hard day:",
    ], 3)
    add_memory_phrase_block(document, MEMORY)
    _img(document, "parent-consequence.png", 3.2)

def p05(document) -> None:
    add_printable_title(document, "P05", "Private next-step saṅkalpa", "Specific action + frequency + trigger + minimum version.")
    add_callout(document, "KEY_IDEA", "specific action + frequency + trigger + minimum version")
    add_write_lines(document, [
        "Specific action:",
        "Frequency:",
        "Trigger (when and where):",
        "Minimum version for a hard day:",
        "Where we will place the reminder:",
    ], 2)
    _img(document, "parent-consequence.png", 4.5)
    document.add_paragraph("No ranking. The minimum version counts as success.")


YOUNGER_BUILDERS = [y01, y02, y03, y04, y05]
OLDER_BUILDERS = [o01, o02, o03, o04, o05, o06]
PARENT_BUILDERS = [p01, p02, p03, p04, p05]
FAMILY_BUILDERS = []
