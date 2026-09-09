#!/usr/bin/env python3
"""C3-W1 V13 printable Word builders for render_week.py (week id C3-W1).

Real Word tables/images. DOCX/PDF emission is owned by scripts/v13/render_week.py.
"""
from __future__ import annotations

import sys
from pathlib import Path

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

REPO = Path(__file__).resolve().parents[3]
sys.path.insert(0, str(REPO / "scripts" / "v13"))
sys.path.insert(0, str(REPO / "scripts" / "v12"))
from printable_common import (  # noqa: E402
    PLUM,
    SAFFRON,
    TEAL,
    add_border,
    add_card_grid,
    add_cut_lines,
    add_memory_phrase_block,
    add_page,
    add_printable_title,
    add_teacher_page,
    add_write_lines,
)
from kutumba_docx_styles import add_branded_table, add_callout, add_verse_card  # noqa: E402

ASSETS = REPO / "11-weekly-program-library/first-six-months/c3-w1-who-is-god-the-supreme-enjoyer-proprietor-and-friend" / "visuals" / "v13"
MEMORY = 'Kṛṣṇa owns, enjoys, befriends.'
VERSE_REF = 'BG 5.29'
VERSE_URL = 'https://vedabase.io/en/library/bg/5/29/'
DEV = 'भोक्तारं यज्ञतपसां सर्वलोकमहेश्वरम् । सुहृदं सर्वभूतानां ज्ञात्वा मां शान्तिमृच्छति ॥ २९ ॥'
IAST = 'bhoktāraṁ yajña-tapasāṁ sarva-loka-maheśvaram / suhṛdaṁ sarva-bhūtānāṁ jñātvā māṁ śāntim ṛcchati'
MEANING = 'Knowing Kṛṣṇa as the supreme enjoyer of sacrifice and austerity, the Lord of all worlds, and the friend of all beings, one attains peace.'


def _img(document, name: str, width: float = 2.4) -> None:
    path = ASSETS / name
    if path.is_file():
        p = document.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(path), width=Inches(width))


def y01(document) -> None:
    add_printable_title(document, "Y01", "Enjoyer/owner/friend icons", "Place three icon cards; say each lens.")
    _img(document, "three-lens-map.png", 5.5)
    add_card_grid(document, ['Enjoyer', 'Owner', 'Friend', 'Share kindly', 'Grab and shout mine', 'Care for a plant'], per_page=6)
    add_memory_phrase_block(document, MEMORY)
    add_cut_lines(document)
    add_callout(document, "TEACHER_NOTE", 'Age-appropriate ownership/friendship/service; no speculative iconography.')

def y02(document) -> None:
    add_printable_title(document, "Y02", "Lens matching", "Match home scenes to enjoyer/owner/friend.")
    _img(document, "lens-icons.png", 5.5)
    add_card_grid(document, ['Enjoyer', 'Owner', 'Friend', 'Share kindly', 'Grab and shout mine', 'Care for a plant'], per_page=6)
    add_memory_phrase_block(document, MEMORY)
    add_cut_lines(document)
    add_callout(document, "TEACHER_NOTE", 'Age-appropriate ownership/friendship/service; no speculative iconography.')

def y03(document) -> None:
    add_printable_title(document, "Y03", "Friend-of-all craft", "Craft a friend-heart for kind service.")
    _img(document, "friend-heart.png", 5.5)
    add_card_grid(document, ['Enjoyer', 'Owner', 'Friend', 'Share kindly', 'Grab and shout mine', 'Care for a plant'], per_page=6)
    add_memory_phrase_block(document, MEMORY)
    add_cut_lines(document)
    add_callout(document, "TEACHER_NOTE", 'Age-appropriate ownership/friendship/service; no speculative iconography.')

def y04(document) -> None:
    add_printable_title(document, "Y04", "Steward cards", "We care for things because Kṛṣṇa owns.")
    _img(document, "memory-mat.png", 5.5)
    add_card_grid(document, ['Enjoyer', 'Owner', 'Friend', 'Share kindly', 'Grab and shout mine', 'Care for a plant'], per_page=6)
    add_memory_phrase_block(document, MEMORY)
    add_cut_lines(document)
    add_callout(document, "TEACHER_NOTE", 'Age-appropriate ownership/friendship/service; no speculative iconography.')

def y05(document) -> None:
    add_printable_title(document, "Y05", "Memory phrase mat", "Echo: Kṛṣṇa owns, enjoys, befriends.")
    _img(document, "memory-mat.png", 6.0)
    add_memory_phrase_block(document, MEMORY)
    document.add_paragraph('BG 5.29' + " · family week")


def o01(document) -> None:
    add_printable_title(document, "O01", "BG 5.29 observation", "Three titles → peace.")
    add_verse_card(document, VERSE_REF, DEV, IAST, MEANING, VERSE_URL,
        "Scripture display/source: VedaBase; teaching meaning: KUTUMBA-original (not labeled as BBT translation).")
    add_write_lines(document, [
        "What does the verse teach in my words?",
        "What misconception does this week block?",
        "One question for the facilitator:",
    ], 2)

def o02(document) -> None:
    add_printable_title(document, "O02", "Three-lens worksheet", "Map enjoyer / proprietor / friend applications.")
    _img(document, "lens-icons.png", 5.0)
    add_branded_table(document, ["Prompt", "My notes"], [["Three-lens worksheet", "________________"], ["Limit / care", "________________"]])
    add_write_lines(document, ["Application:", "What not to say:"], 2)
    add_callout(document, "KEY_IDEA", 'Kṛṣṇa is supreme enjoyer, proprietor of all worlds, and friend of all beings — knowing this brings peace.')

def o03(document) -> None:
    add_printable_title(document, "O03", "Ownership/service sort", "Sort domination vs stewardship.")
    _img(document, "friend-heart.png", 5.0)
    add_branded_table(document, ["Prompt", "My notes"], [["Ownership/service sort", "________________"], ["Limit / care", "________________"]])
    add_write_lines(document, ["Application:", "What not to say:"], 2)
    add_callout(document, "KEY_IDEA", 'Kṛṣṇa is supreme enjoyer, proprietor of all worlds, and friend of all beings — knowing this brings peace.')

def o04(document) -> None:
    add_printable_title(document, "O04", "Peace application scenarios", "Three ownership conflict cases.")
    _img(document, "memory-mat.png", 5.0)
    add_branded_table(document, ["Prompt", "My notes"], [["Peace application scenarios", "________________"], ["Limit / care", "________________"]])
    add_write_lines(document, ["Application:", "What not to say:"], 2)
    add_callout(document, "KEY_IDEA", 'Kṛṣṇa is supreme enjoyer, proprietor of all worlds, and friend of all beings — knowing this brings peace.')

def o05(document) -> None:
    add_printable_title(document, "O05", "Triple-lens diagram", "One Lord; three lenses; peace.")
    _img(document, "parent-steward.png", 5.0)
    add_branded_table(document, ["Prompt", "My notes"], [["Triple-lens diagram", "________________"], ["Limit / care", "________________"]])
    add_write_lines(document, ["Application:", "What not to say:"], 2)
    add_callout(document, "KEY_IDEA", 'Kṛṣṇa is supreme enjoyer, proprietor of all worlds, and friend of all beings — knowing this brings peace.')

def o06(document) -> None:
    add_printable_title(document, "O06", "Exit ticket", "One lens to practice at home.")
    add_write_lines(document, [
        "Primary in one sentence:",
        "One analogy or mechanic and its limit:",
        "One home action:",
        "One question I still have:",
    ], 2)


def p01(document) -> None:
    add_printable_title(document, "P01", "Private reflection — ownership tone", "Where does our home clutch rather than steward?")
    _img(document, "parent-icon.png", 3.2)
    add_callout(document, "SAFETY_PRIVACY", "Write privately. Sharing is optional. Do not collect or place completed sheets in Git.")
    _img(document, "parent-steward.png", 3.5)
    p = document.add_paragraph('Where does our home clutch rather than steward?')
    p.runs[0].bold = True
    add_write_lines(document, ["Private notes:", "Habit to watch this week:"], 4)

def p02(document) -> None:
    add_printable_title(document, "P02", "Ownership/service card sort", "Sort STEWARD / CLUTCH / FRIEND.")
    mats = document.add_table(rows=1, cols=2)
    add_border(mats, PLUM, 10)
    for cell, label in zip(mats.rows[0].cells, ("HELPFUL", "MISTAKEN")):
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(label)
        r.bold = True
        r.font.size = Pt(22)
        r.font.color.rgb = RGBColor.from_string(TEAL if label == "HELPFUL" else SAFFRON)
    add_page(document)
    add_card_grid(document, ['Offer first', 'Dominate siblings', 'Steward shared space', 'Kṛṣṇa is only a force', 'Befriend with service', 'Clutch the remote'], prefix="L", per_page=12)
    add_callout(document, "TEACHER_NOTE", 'God as remote force only; ownership means domination without friendship/service.')
    _img(document, "parent-steward.png", 3.0)

def p03(document) -> None:
    add_printable_title(document, "P03", "Substantial family case", "Sibling fight over “my room / my phone.”")
    add_callout(document, "SAFETY_PRIVACY", "Fictional case. No compelled confession of real events.")
    document.add_paragraph('Siblings fight over “my room” and “my phone.” Parents escalate ownership lectures without naming Kṛṣṇa as proprietor or modeling friendship.')
    add_write_lines(document, [
        "Tempting mistaken conclusion:",
        "Principle from primary:",
        "Compassionate response:",
        "One seven-day household action:",
        "What we should not say:",
    ], 3)
    _img(document, "three-lens-map.png", 4.0)

def p04(document) -> None:
    add_printable_title(document, "P04", "Household operating application", "One stewardship sentence + one friend act.")
    add_callout(document, "SAFETY_PRIVACY", "No ranking and no public reading required.")
    add_write_lines(document, [
        "One household action for seven days:",
        "Trigger (when/where):",
        "Who starts if others are tired:",
        "Minimum version on a hard day:",
    ], 3)
    add_memory_phrase_block(document, MEMORY)
    _img(document, "parent-steward.png", 3.2)

def p05(document) -> None:
    add_printable_title(document, "P05", "Private next-step saṅkalpa", "Practice one lens daily.")
    add_callout(document, "KEY_IDEA", "specific action + frequency + trigger + minimum version")
    add_write_lines(document, [
        "Specific action:",
        "Frequency:",
        "Trigger (when and where):",
        "Minimum version for a hard day:",
        "Where we will place the reminder:",
    ], 2)
    _img(document, "parent-steward.png", 4.5)
    document.add_paragraph("No ranking. The minimum version counts as success.")


YOUNGER_BUILDERS = [y01, y02, y03, y04, y05]
OLDER_BUILDERS = [o01, o02, o03, o04, o05, o06]
PARENT_BUILDERS = [p01, p02, p03, p04, p05]
FAMILY_BUILDERS = []
