#!/usr/bin/env python3
"""C1-U1 V13 printable Word builders for render_week.py.

Real Word tables/images. DOCX/PDF emission owned by scripts/v13/render_week.py.
This pass authors builders only — do not render here.
"""
from __future__ import annotations

import sys
from pathlib import Path

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches

REPO = Path(__file__).resolve().parents[3]
sys.path.insert(0, str(REPO / "scripts" / "v13"))
sys.path.insert(0, str(REPO / "scripts" / "v12"))
from printable_common import (  # noqa: E402
    add_card_grid,
    add_cut_lines,
    add_memory_phrase_block,
    add_printable_title,
    add_write_lines,
)
from kutumba_docx_styles import add_callout, add_verse_card  # noqa: E402

ASSETS = (
    REPO
    / "11-weekly-program-library"
    / "first-six-months"
    / "c1-u1-kartika-damodara-family-utsava"
    / "visuals"
    / "v13"
)
MEMORY = "Kṛṣṇa is bound by love."
VERSE_URL = "https://vedabase.io/en/library/sb/10/9/18/"


def _img(document, name: str, width: float = 2.4) -> None:
    path = ASSETS / name
    if path.is_file():
        p = document.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(path), width=Inches(width))



def y01(document) -> None:
    add_printable_title(document, "Y01", "Dāmodara Story Sequence", "Order the cards. Tell the story kindly.")
    _img(document, "damodara-story-sequence.png", 6.0)
    add_card_grid(document, ["Butter pot", "Search", "Rope try", "Bound by love", "Gratitude"], per_page=5)
    add_memory_phrase_block(document, MEMORY)
    add_cut_lines(document)
    add_callout(document, "CHILD_ACTIVITY", "Kind voices only. No mocking mother Yaśodā.")


def y02(document) -> None:
    add_printable_title(document, "Y02", "Rope-of-Love Craft", "Write love/service words on paper strips. Link them.")
    add_write_lines(document, ["Love word 1:", "Love word 2:", "Service word:"], 2)
    add_callout(document, "SAFETY_PRIVACY", "Paper craft only. No real rope play-bondage.")


def y03(document) -> None:
    add_printable_title(document, "Y03", "LED Gratitude Lamp Card", "Color the card. Write one thank-you.")
    _img(document, "led-gratitude-lamp-card.png", 4.5)
    add_write_lines(document, ["I thank Kṛṣṇa for:", "I can help at home by:"], 2)
    add_callout(document, "SAFETY_PRIVACY", "LED default. No child open flame. No child fasting.")


def y04(document) -> None:
    add_printable_title(document, "Y04", "Kind Līlā Speech", "Choose kind words about the butter story.")
    add_card_grid(document, ["Kind: Mother loved Kṛṣṇa", "Unkind: mocking laugh", "Kind: Kṛṣṇa is merciful", "Unkind: teasing fasting"], per_page=4)
    add_cut_lines(document)


def y05(document) -> None:
    add_printable_title(document, "Y05", "Memory Phrase Mat", MEMORY)
    add_memory_phrase_block(document, MEMORY)
    add_callout(document, "KEY_IDEA", MEMORY)


def o01(document) -> None:
    add_printable_title(document, "O01", "ŚB 10.9.18 Observation", "Observe the verse layer. No purport dump.")
    add_verse_card(document, "ŚB 10.9.18", "", "", "Seeing mother Yaśodā tired, Kṛṣṇa mercifully agreed to be bound.", VERSE_URL, 'Sanskrit/Bengali display + KUTUMBA teaching meaning; no full purport')
    add_write_lines(document, ["Who is tired?", "What does Kṛṣṇa do?", "What does mercy look like?"], 3)


def o02(document) -> None:
    add_printable_title(document, "O02", "Dāmodara Narrative Map", "Map SB 10.9 paraphrase stations. No invented dialogue.")
    _img(document, "damodara-rope-of-love.png", 5.5)
    add_write_lines(document, ["Station 1:", "Station 2:", "Station 3:", "Mercy moment:"], 2)


def o03(document) -> None:
    add_printable_title(document, "O03", "Love vs Force Sort", "Sort into Bound-by-love / Bound-by-force / Not sure.")
    add_card_grid(document, ["Patient motherly effort", "Harsh control", "Gratitude service", "Public ranking of devotion", "Private saṅkalpa", "Child fasting pressure"], per_page=6)
    add_cut_lines(document)


def o04(document) -> None:
    add_printable_title(document, "O04", "Kārtika Scenarios", "Choose the KUTUMBA-aligned response.")
    add_write_lines(document, ["Scenario: child fasting pressure →", "Scenario: open flame request →", "Scenario: ranking lamps →"], 3)
    add_callout(document, "TEACHER_NOTE", "Correct redirects: no child fasting; LED default; no ranking.")


def o05(document) -> None:
    add_printable_title(document, "O05", "Gratitude / Service Diagram", "Remembrance → Gratitude → One service act.")
    add_write_lines(document, ["Remembrance cue:", "Gratitude sentence:", "Service act this week:"], 2)


def o06(document) -> None:
    add_printable_title(document, "O06", "Exit Ticket", "Three lines only.")
    add_write_lines(document, ["Memory line:", "One thing Yaśodā shows:", "One safe home act:"], 2)


def p01(document) -> None:
    add_printable_title(document, "P01", "Private Reflection", "Sharing optional.")
    _img(document, "parent-icon.png", 3.2)
    add_write_lines(document, ["Where do we force outcomes?", "Where could love/patience lead instead?"], 3)
    add_callout(document, "SAFETY_PRIVACY", "Write privately. Do not collect sheets into Git.")


def p02(document) -> None:
    add_printable_title(document, "P02", "Source Observation", "ŚB 10.9.18")
    add_verse_card(document, "ŚB 10.9.18", "", "", "Kṛṣṇa agrees to be bound when He sees mother’s fatigue.", VERSE_URL, 'Sanskrit/Bengali display + KUTUMBA teaching meaning; no full purport')
    add_write_lines(document, ["What stands out?", "What must we not claim about Oct 31?"], 3)


def p03(document) -> None:
    add_printable_title(document, "P03", "Substantial Family Case", "Child-fasting pressure during Kārtika.")
    add_write_lines(document, ["Mistaken conclusion:", "Principle:", "Family action:"], 3)
    add_callout(document, "DO_NOT_SPECULATE", "No child fasting instructions. Local tithi EXTERNAL_OPEN.")


def p04(document) -> None:
    add_printable_title(document, "P04", "Household Operating Application", "Kārtika remembrance without ranking.")
    add_write_lines(document, ["Remembrance cue:", "Service act:", "LED/no-flame plan:", "Minimum version:"], 2)
    add_callout(document, "FAMILY_APPLICATION", "Tiny and regular beats dramatic and ranked.")


def p05(document) -> None:
    add_printable_title(document, "P05", "Private Next-Step Saṅkalpa", "")
    add_callout(document, "KEY_IDEA", "specific action + frequency + trigger + minimum version")
    add_write_lines(document, ["Specific action:", "Frequency:", "Trigger:", "Minimum version:"], 2)


YOUNGER_BUILDERS = [y01, y02, y03, y04, y05]
OLDER_BUILDERS = [o01, o02, o03, o04, o05, o06]
PARENT_BUILDERS = [p01, p02, p03, p04, p05]


FAMILY_BUILDERS = []
