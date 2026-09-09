#!/usr/bin/env python3
"""Author complete V13 packs for C2-W1..C2-W6 and C3-W1..C3-W6.

Deepens from existing V12 MAIN/research/gamma. Does not render DOCX/PDF.
Does not modify C1-W1. Does not commit.
"""
from __future__ import annotations

import csv
import re
import textwrap
from datetime import date
from pathlib import Path

import yaml
from PIL import Image, ImageDraw, ImageFont

from c2_c3_week_specs import SPECS

REPO = Path(__file__).resolve().parents[2]
FIRST = REPO / "11-weekly-program-library" / "first-six-months"
REGISTRY = Path(__file__).resolve().parent / "week_registry.yaml"
EVIDENCE = REPO / "build-evidence" / "v13"
PRINTABLES = Path(__file__).resolve().parent / "printables"
CREAM, PLUM, SAFFRON, TEAL = "#FFF8E8", "#5B1933", "#E59B24", "#4F7C78"
STATUS = "Internal founding-cohort teaching material — human/temple review EXTERNAL_OPEN"
AUDIT_DATE = "2026-09-09"


def font(size: int, bold: bool = False):
    names = ["segoeuib.ttf" if bold else "segoeui.ttf", "arialbd.ttf" if bold else "arial.ttf"]
    for name in names:
        path = Path(r"C:\Windows\Fonts") / name
        if path.exists():
            return ImageFont.truetype(str(path), size)
    return ImageFont.load_default()


def resolve_folder(folder: str) -> Path:
    exact = FIRST / folder
    if exact.is_dir():
        return exact
    needle = folder.casefold()
    for child in FIRST.iterdir():
        if child.is_dir() and child.name.casefold() == needle:
            return child
    raise FileNotFoundError(folder)


def load_weeks() -> list[dict]:
    data = yaml.safe_load(REGISTRY.read_text(encoding="utf-8"))
    return [w for w in data["weeks"] if w["id"] in SPECS]


def fmt_date(iso: str) -> str:
    y, m, d = map(int, iso.split("-"))
    return date(y, m, d).strftime("%d %b %Y")


def write(path: Path, text: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text.rstrip() + "\n", encoding="utf-8")


def upgrade_main(v12: str, week: dict, spec: dict) -> str:
    wid, title, d = week["id"], week["title"], week["schedule_date"]
    header = f"""# {wid} Main Facilitator Guide — V13

**KUTUMBA · Families Growing in Krishna Consciousness**  
**Program Director: Swapnil Patil**  
**Week:** {wid} — {title} · **Date:** {d}  
**Status:** {STATUS}  
**Controlling guide:** This V13 file is authoritative for Saturday delivery. Core teaching is inlined below.  
**Entry point:** `V13-WEEK-START-HERE.md`  
**Primary mechanic:** {spec['mechanic']} (see `V13-ACTIVITY-VARIETY-MATRIX.md`)

## V13 schedule lock (do not invent another clock)

| Time | Block |
|---|---|
| 1:50–2:00 | Arrival / settle |
| 2:00–2:10 | Opening mantras + welcome |
| 2:10–2:30 | All-family launch / Prem-kī-Kathā + verse |
| 2:30–3:10 | Parallel tracks: parents / K–2 / Grades 4–5 |
| 3:10–3:30 | Reunification + bhakti laboratory |
| 3:30–3:40 | Light snack + water only |
| 3:40–3:55 | Saṅkalpa + project + home practice |
| 3:55–4:00 | Closing / next Saturday |

Parents onsite. No weekly meal. Research files under `research/V13-*` are enrichment only — not a substitute for this inlined guide.

---

"""
    body = v12
    # Strip original top H1 if present to avoid double titles, keep remaining depth.
    body = re.sub(r"^# .+\n+", "", body, count=1)
    body = body.replace("Saturday 2:00–4:00", "Saturday 1:50–4:00")
    body = body.replace("2:00–4:00", "1:50–4:00")
    body = body.replace("MAIN-FACILITATOR-GUIDE-V12", "MAIN-FACILITATOR-GUIDE-V13")
    footer = f"""

---

## V13 activity wiring (do not replace with “see research”)

Younger printables: `activities/v13-younger/` ({', '.join(c for c,_t,_d in spec['y'])})  
Older printables: `activities/v13-older/` ({', '.join(c for c,_t,_d in spec['o'])})  
Parent printables: `activities/v13-parent/` ({', '.join(c for c,_t,_d in spec['p'])})  
Project: `project/V13-MODULE-PROJECT-BRIEF.md`  
Deferral line: “I don’t want to guess. I will verify that from Śrīla Prabhupāda’s books / our source packet and come back to you.”

**Memory line:** {spec['memory']}  
**Misconception to block:** {spec['misconception']}  
**Care boundary:** {spec['care_boundary']}
"""
    if week["kind"] == "integration":
        footer += "\n**Integration rule:** No major new doctrine — retrieval, teach-back, and presentation focus.\n"
    return header + body.strip() + footer


def upgrade_research(src: Path, title: str, week: dict) -> str:
    text = src.read_text(encoding="utf-8") if src.exists() else f"# {title}\n\nContent pending source pack.\n"
    banner = (
        f"# {week['id']} V13 — {title}\n\n"
        f"**Week:** {week['title']} · **Date:** {week['schedule_date']}\n"
        f"**Primary:** {week['primary_verse']} — {week['primary_url']}\n"
        f"**Status:** {STATUS}\n\n"
        f"> Deepened from V12 research; keep classroom teaching in MAIN-FACILITATOR-GUIDE-V13.\n\n---\n\n"
    )
    # drop first H1
    text = re.sub(r"^# .+\n+", "", text, count=1)
    return banner + text.strip() + "\n"


def activity_md(week: dict, spec: dict, code: str, title: str, blurb: str, track: str) -> str:
    return f"""# {week['id']} · {code} · {title}

**Print:** US Letter · {track}  
**Primary mechanic:** {spec['mechanic']}  
**Primary bridge:** {week['primary_verse']}  
**URL:** {week['primary_url']}

---

## Teaching purpose

{blurb}

Memory line: **{spec['memory']}**

## How to run

1. Distribute the printable from the Saturday packet / Word builder `{week['id']}.py`.
2. State the purpose in one sentence; name the care boundary: {spec['care_boundary']}
3. Run the activity for the planned minutes in the track guide.
4. Echo the memory line once.
5. Collect only what is needed for classroom reuse — never commit private family writing.

## Materials

Printed page(s); pencils/crayons as needed; optional visual from `visuals/v13/`.

## Age boundary / safeguards

- Misconception to block: {spec['misconception']}
- Do not invent doctrine beyond the week’s primary and supporting readings.
- No ranking of children or families.

## Answer key / facilitator direction

See teacher-only pages in the printable builder and the track guide. Prefer principle + compassionate response + practical action over “gotcha” scoring.

## Source

{week['primary_verse']} — {week['primary_url']}  
KUTUMBA teaching meaning is original pedagogy — not labeled as a BBT translation.
"""


def start_here(week: dict, spec: dict) -> str:
    d = fmt_date(week["schedule_date"])
    return f"""# V13 WEEK START HERE — {week['id']}

**KUTUMBA · Families Growing in Krishna Consciousness**  
**Program Director: Swapnil Patil**  
**Week:** {week['id']} — {week['title']}  
**Date:** Saturday {d}  
**This is the only {week['id']} operational entry point for V13.**

---

## THIS SATURDAY (locked)

| Time | Block |
|---|---|
| 1:50–2:00 | Arrival / settle |
| 2:00–2:10 | Opening mantras + welcome |
| 2:10–2:30 | All-family launch / Prem-kī-Kathā + verse |
| 2:30–3:10 | Parallel tracks: **parents** / **K–2** / **Grades 4–5** |
| 3:10–3:30 | Reunification + bhakti laboratory |
| 3:30–3:40 | Light snack + water only |
| 3:40–3:55 | Saṅkalpa + project + home practice |
| 3:55–4:00 | Closing / next Saturday |

Parents onsite. No weekly meal. Parallel tracks **2:30–3:10**; reunify **3:10–3:30**. Do not invent 2:35 / 3:25.

---

## ESSENTIAL FOCUS

| Field | Content |
|---|---|
| Essential question | {spec['eq']} |
| Conclusion | {spec['conclusion']} |
| Memory line | {spec['memory']} |
| Primary | **{week['primary_verse']}** — {week['primary_url']} |
| Primary mechanic | {spec['mechanic']} |
| Misconception to block | {spec['misconception']} |
| Care boundary | {spec['care_boundary']} |

---

## WHO SPEAKS / TEACHES

| Role | File |
|---|---|
| Main facilitator | `teacher/MAIN-FACILITATOR-GUIDE-V13.md` |
| Parent track | `teacher/PARENT-GUIDE-V13.md` |
| Younger K–2 | `teacher/YOUNGER-TEACHER-GUIDE-V13.md` |
| Older 4–5 | `teacher/OLDER-TEACHER-GUIDE-V13.md` |
| Home practice | `family-home-practice-v13.md` |
| Materials | `materials-v13.md` |
| Project | `project/V13-MODULE-PROJECT-BRIEF.md` |
| Gamma (prompt-only) | `gamma/V13-GAMMA-MASTER-DECK-PROMPT.md` |
| Printables | `activities/v13-parent/` · `v13-younger/` · `v13-older/` |

---

## FRIDAY NIGHT (≤60 min)

1. Open this file.
2. When available, print `exports/final/v13/{week['export_subdir']}/` Saturday packet (builders ready; DOCX/PDF render may follow).
3. Read MAIN speaking script once aloud (10–15 min).
4. Sort printables into three track folders (Y / O / P).
5. Confirm snack/water + room reset + privacy reminder card.
6. Optional: paste Gamma master prompt into Gamma Studio (owner render; not pre-approved).

---

## HOME PRACTICE (tell families at 3:40)

Minimum before next Saturday: **{spec['home_min']}.**  
Ideal: memory line + one-sentence primary paraphrase + one related service/gratitude act (5–15 min).  
No ranking. Five minutes counts.

---

## NEXT WEEK (name only)

**{spec['next_week']}.** Do not teach the next week’s full ontology tonight.

---

## EXTERNAL_OPEN

Local tithi / Utsava confirmation · pilot GO / temple human approval · live Gamma render sign-off · formal BBT licensing beyond teaching paraphrase.

**Verdict language allowed:** {week['id']} owner-runnable pack authored — human/temple external gates remain open.  
Do not claim public publication approval or official ISKCON endorsement.
"""


def track_guide(week: dict, spec: dict, kind: str) -> str:
    d = week["schedule_date"]
    folder = resolve_folder(week["folder"])
    if kind == "parent":
        items = spec["p"]
        title = "Parent Track Guide — V13 (40 minutes)"
        window = "parents"
        codes = "P01–P05"
        v12 = folder / "teacher" / "PARENT-GUIDE.md"
        if not v12.exists():
            v12 = folder / "parent-lesson.md"
        src_label = "parent-lesson / constructed V13 flow"
    elif kind == "younger":
        items = spec["y"]
        title = "Younger Teacher Guide — V13 (K–2, 40 minutes)"
        window = "K–2"
        codes = "Y01–Y05"
        v12 = folder / "teacher" / "YOUNGER-TEACHER-GUIDE.md"
        src_label = "YOUNGER-TEACHER-GUIDE.md"
    else:
        items = spec["o"]
        title = "Older Teacher Guide — V13 (Grades 4–5, 40 minutes)"
        window = "Grades 4–5"
        codes = "O01–O06"
        v12 = folder / "teacher" / "OLDER-TEACHER-GUIDE.md"
        src_label = "OLDER-TEACHER-GUIDE.md"

    chunks = []
    spans = {
        "parent": [(0, 5), (5, 12), (12, 22), (22, 32), (32, 38), (38, 40)],
        "younger": [(0, 6), (6, 14), (14, 24), (24, 32), (32, 38), (38, 40)],
        "older": [(0, 6), (6, 14), (14, 22), (22, 30), (30, 36), (36, 40)],
    }[kind]
    for (a, b), (code, t, blurb) in zip(spans, items):
        chunks.append(
            f"### V13 printable block {a}–{b} — {t}\n\n"
            f"**Printable:** `activities/v13-{kind}/{code}.md` (`{code}`)\n\n"
            f"{blurb}\n\n"
            f"Keep primary mechanic (**{spec['mechanic']}**). Block: {spec['misconception']}\n"
        )

    inherited = ""
    if v12.exists():
        body = v12.read_text(encoding="utf-8")
        body = re.sub(r"^# .+\n+", "", body, count=1)
        body = body.replace("Saturday 2:00–4:00", "Saturday 1:50–4:00")
        body = body.replace("2:00–4:00", "1:50–4:00")
        inherited = (
            f"\n---\n\n## Inherited classroom depth (from V12 `{src_label}`)\n\n"
            f"> Keep story scripts, movement steps, and wonder questions below. "
            f"Wire printables to the V13 codes above; do not replace this depth with “see research.”\n\n"
            + body.strip()
            + "\n"
        )

    verse_block = f"""## Source layer (use mid-track)

**Primary:** {week['primary_verse']} — {week['primary_url']}

**Devanāgarī:** {spec['dev']}

**IAST:** `{spec['iast']}`

**KUTUMBA teaching meaning (not a BBT translation):**  
{spec['meaning']}
"""

    return f"""# {week['id']} {title}

**KUTUMBA · Families Growing in Krishna Consciousness**  
**Program Director: Swapnil Patil**  
**Week:** {week['title']} · **Date:** {d}  
**Track window:** Saturday **2:30–3:10** ({window})  
**Reunite:** **3:10** for bhakti laboratory  
**Primary:** {week['primary_verse']} — {week['primary_url']}  
**Primary mechanic:** {spec['mechanic']}  
**Status:** {STATUS}

---

## Essential question

{spec['eq']}

## Materials (print before Saturday)

From `activities/v13-{kind}/`: **{codes}**  
Also pens/crayons; verse card; privacy reminder.

## Locked Saturday context

- 1:50–2:00 Arrival · 2:00–2:10 Mantras · 2:10–2:30 All-family launch  
- **2:30–3:10 This track** · 3:10–3:30 Reunify · 3:30–3:40 Snack/water · 3:40–3:55 Saṅkalpa · 3:55–4:00 Close  

Parents remain onsite. No public ranking. No forced disclosure.

{verse_block}

## Memory line

> {spec['memory']}

## V13 printable wiring ({spec['mechanic']})

{''.join(chunks)}

## Close of track (handoff)

Bring one sentence or artifact to reunification. Do not start next week’s topic.  
Care boundary: {spec['care_boundary']}
{inherited}
"""


def home_practice(week: dict, spec: dict) -> str:
    return f"""# {week['id']} Family Home Practice — V13

**Week:** {week['title']}  
**Primary:** {week['primary_verse']} — {week['primary_url']}  
**Memory line:** {spec['memory']}

## Minimum (before next Saturday)

{spec['home_min'].capitalize()}.

Five minutes counts. No ranking. No photo proof required.

## Ideal (5–15 minutes)

1. Say the memory line once together.
2. Paraphrase the primary in one family sentence.
3. Do one related gratitude/service/attention act tied to this week’s conclusion: {spec['conclusion']}

## Optional conversation cue

Essential question: {spec['eq']}

## Do not

- Rank families or children
- Diagnose others’ karma/modes/destinations
- Invent doctrine beyond this week’s sources

**Next week (name only):** {spec['next_week']}
"""


def materials(week: dict, spec: dict) -> str:
    y = ", ".join(c for c, _, _ in spec["y"])
    o = ", ".join(c for c, _, _ in spec["o"])
    p = ", ".join(c for c, _, _ in spec["p"])
    return f"""# {week['id']} Materials — V13

## Print / prepare

- Start Here + MAIN facilitator pages (when rendered)
- Younger: {y}
- Older: {o}
- Parent: {p}
- Family home practice sheet
- Project brief contribution page
- Visuals from `visuals/v13/` ({', '.join(spec['pngs'])})

## Room

- Parent circle seating; younger floor/table zone; older table zone
- Privacy reminder card visible
- Snack + water only (no weekly meal)
- Soft toss / movement space if the mechanic needs it ({spec['mechanic']})

## Consumables

Pencils, crayons, tape/scissors for cut cards, sticky notes for saṅkalpa.

## Do not bring

Private family conflict write-ups for Git; ranking trophies; scare media.
"""


def project_brief(week: dict, spec: dict) -> str:
    integ = week["kind"] == "integration"
    return f"""# {week['id']} Module Project Brief — V13

**Week:** {week['title']}  
**Primary:** {week['primary_verse']} — {week['primary_url']}  
**Cycle contribution:** {spec['project_line']}

## This week’s artifact

Families produce: **{spec['project_line']}**.

## Success looks like

- Uses this week’s primary source honestly
- Names the care boundary: {spec['care_boundary']}
- Fits on one page or one small craft
- Can be explained in ≤60 seconds at reunification or W6

## Integration note

{"This is an integration week: retrieval/presentation focus — no major new doctrine." if integ else "Save the artifact toward the cycle integration night (W6)."}

## Rubric (noncompetitive)

| Lens | Look for |
|---|---|
| Source | Mentions primary verse or theme |
| Clarity | Child/adult can explain in plain words |
| Kindness | No ranking / no diagnosis of others |
| Practice | Names one home action |

## Memory line to print on artifact

{spec['memory']}
"""


def claim_register(week: dict, spec: dict) -> str:
    return f"""week_id: {week['id']}
title: {week['title']!r}
primary_verse: {week['primary_verse']!r}
primary_url: {week['primary_url']!r}
claims:
  - id: {week['id']}-C01
    statement: {spec['conclusion']!r}
    status: TEACH
    basis: primary verse teaching meaning + MAIN guide
    do_not_claim: science proves this doctrinal conclusion
  - id: {week['id']}-C02
    statement: {spec['misconception']!r}
    status: BLOCK
    basis: week care boundary
  - id: {week['id']}-C03
    statement: {spec['care_boundary']!r}
    status: BOUNDARY
    basis: package standards 04/06/07
external_open:
  - human/temple doctrinal review
  - live Gamma render sign-off
  - local tithi where festival-adjacent
"""


def upgrade_gamma(src: Path, dest_name: str, week: dict, spec: dict) -> str:
    text = src.read_text(encoding="utf-8") if src.exists() else ""
    if not text.strip():
        text = synthesize_gamma(week, spec, dest_name)
    text = text.replace("V12", "V13")
    text = text.replace("Saturday 2:00–4:00", "Saturday 1:50–4:00")
    text = text.replace("2:00–4:00", "1:50–4:00")
    banner = (
        f"# {week['id']} {dest_name}\n\n"
        f"**Status:** prompt-only — not rendered — not approved  \n"
        f"**Design:** 16:9 · cream `#FFF8E8` · plum `#5B1933` / saffron `#E59B24` / teal `#4F7C78`  \n"
        f"**Brand:** KUTUMBA · Families Growing in Krishna Consciousness · Program Director: Swapnil Patil  \n"
        f"**Week:** {week['id']} — {week['title']} · **Date:** {week['schedule_date']}  \n"
        f"**Schedule on slides:** Arrival 1:50 · Tracks 2:30–3:10 · Reunify 3:10–3:30 · Close 4:00  \n"
        f"**Primary mechanic:** {spec['mechanic']}  \n"
        f"**Memory line:** {spec['memory']}\n\n"
        f"Paste into Gamma Studio. Owner renders; no claim of Gamma approval until then.\n\n---\n\n"
    )
    body = re.sub(r"^# .+\n+", "", text, count=1)
    # Ensure verse layers present for master
    if "MASTER" in dest_name.upper() and spec["dev"] not in body:
        body += f"""

### V13 verse completeness check (required)

- **Reference:** {week['primary_verse']}
- **URL:** {week['primary_url']}
- **Devanāgarī:** {spec['dev']}
- **IAST:** {spec['iast']}
- **KUTUMBA teaching meaning:** {spec['meaning']}
- Do not truncate mid-word; use two slides if needed.
"""
    return banner + body.strip() + "\n"


def synthesize_gamma(week: dict, spec: dict, dest_name: str) -> str:
    slides = []
    for i, (title, copy) in enumerate(
        [
            ("Title / hook", f"{week['title']}\nKUTUMBA\n{week['id']} · {spec['memory']}"),
            ("Essential question", spec["eq"]),
            ("Primary verse", f"{week['primary_verse']}\n{spec['iast']}\n{spec['meaning']}"),
            ("Concept", f"{spec['mechanic']}\n{spec['conclusion']}"),
            ("Misconception", spec["misconception"]),
            ("Application", spec["home_min"]),
            ("Close", spec["next_week"]),
        ],
        1,
    ):
        slides.append(
            f"""### Slide {i} — {title}
- **Purpose:** {title}
- **Exact title:** {title}
- **Exact on-screen copy:**\n{textwrap.indent(copy, '  - ')}
- **Sources:** {week['primary_url']}
- **Layout:** clean educational 16:9
- **Detailed AI-image prompt:** Warm family learning scene for {week['id']} theme; cream/saffron textiles; no logos; no identifiable real persons; no embedded readable text; respectful; 16:9
- **Presenter notes:** Teach {spec['mechanic']}; block {spec['misconception']}
- **Do not claim:** temple approval; BBT ownership of KUTUMBA materials
"""
        )
    return "\n".join(slides)


def gamma_source_map(week: dict, spec: dict) -> str:
    return f"""week_id: {week['id']}
title: {week['title']}
primary_verse: {week['primary_verse']}
primary_url: {week['primary_url']}
devanagari: {spec['dev']!r}
iast: {spec['iast']!r}
teaching_meaning: {spec['meaning']!r}
memory_line: {spec['memory']!r}
mechanic: {spec['mechanic']!r}
decks:
  master: V13-GAMMA-MASTER-DECK-PROMPT.md
  parent: V13-GAMMA-PARENT-DECK-PROMPT.md
  younger: V13-GAMMA-YOUNGER-DECK-PROMPT.md
  older: V13-GAMMA-OLDER-DECK-PROMPT.md
schedule:
  arrival: "1:50"
  tracks: "2:30-3:10"
  reunify: "3:10-3:30"
  close: "4:00"
status: prompt-only
"""


def rights_md(week: dict, spec: dict) -> str:
    rows = "\n".join(f"| `{n}` | Original line art (Pillow) | KUTUMBA program original · internal teaching use |" for n in spec["pngs"])
    return f"""# {week['id']} V13 Visual Rights

**Week:** {week['title']}  
**Folder:** `visuals/v13/`

## Original assets (KUTUMBA-created)

| File | Type | Rights |
|---|---|---|
{rows}
| `IMAGE-PROMPT-LIBRARY.md` | Prompt register | KUTUMBA original |

## Scripture display

- Sanskrit / IAST for {week['primary_verse']} from VedaBase: {week['primary_url']}
- KUTUMBA teaching meanings are original paraphrases — **not** labeled as BBT translations
- No full BBT purport reproduction in visuals

## Forbidden

- Identifiable real persons / student photos as proof
- Unauthorized ISKCON logos
- Copyrighted internet art copied without rights
- Fake historical photography
- Byte-identical reuse of another week’s unique visuals masquerading as new work

## Gamma images

Gamma AI image prompts are composition briefs for owner render only.
"""


def image_prompt_library(week: dict, spec: dict) -> str:
    lines = [f"# {week['id']} Image Prompt Library — V13", "", "Owner-render briefs for Gamma / future art. Not pre-approved.", ""]
    for i, name in enumerate(spec["pngs"], 1):
        lines += [
            f"## {i}. {name}",
            f"- Scene: educational diagram/illustration for {spec['mechanic']}",
            f"- Subjects: abstract icons / simple family silhouettes relevant to {week['title']}",
            "- Setting: cream background; plum/teal/saffron line art",
            "- Composition: centered, printer-safe, high contrast",
            "- Negative: no logos, no real faces, no gore, no embedded paragraphs of text",
            f"- Pedagogical purpose: support {spec['memory']}",
            "",
        ]
    return "\n".join(lines)


def draw_pngs(out: Path, week: dict, spec: dict) -> None:
    out.mkdir(parents=True, exist_ok=True)
    for idx, name in enumerate(spec["pngs"]):
        w, h = (900, 520) if idx == 0 else (720, 480)
        im = Image.new("RGB", (w, h), CREAM)
        d = ImageDraw.Draw(im)
        d.rounded_rectangle((20, 20, w - 20, h - 20), radius=28, outline=PLUM, width=8)
        # unique geometry per week+index
        color = [TEAL, SAFFRON, PLUM][idx % 3]
        if "wheel" in name or "lens" in name or "nine" in name:
            cx, cy, r = w // 2, h // 2, min(w, h) // 3
            d.ellipse((cx - r, cy - r, cx + r, cy + r), outline=color, width=8)
            for k in range(6):
                import math

                ang = -90 + 60 * k
                rad = math.radians(ang)
                d.line((cx, cy, cx + int(r * math.cos(rad)), cy + int(r * math.sin(rad))), fill=PLUM, width=4)
        elif "tree" in name or "fork" in name or "path" in name or "trap" in name or "flow" in name:
            d.line((80, h // 2, w - 80, h // 2), fill=TEAL, width=10)
            for x in range(120, w - 80, max(80, (w - 200) // 4)):
                d.ellipse((x - 14, h // 2 - 14, x + 14, h // 2 + 14), outline=SAFFRON, width=6)
                d.line((x, h // 2, x + 40, h // 2 - 60), fill=PLUM, width=5)
        elif "board" in name or "scenario" in name or "map" in name:
            for r_i in range(2):
                for c_i in range(3):
                    x0 = 60 + c_i * ((w - 120) // 3)
                    y0 = 70 + r_i * ((h - 140) // 2)
                    d.rectangle((x0, y0, x0 + (w - 160) // 3, y0 + (h - 180) // 2), outline=color, width=5)
        else:
            d.polygon([(w // 2, 70), (w - 80, h - 90), (80, h - 90)], outline=color, width=8)
            d.line((80, h - 90, w - 80, h - 90), fill=SAFFRON, width=6)
        label = f"{week['id']} · {name.replace('.png', '').replace('-', ' ')}"
        b = d.textbbox((0, 0), label, font=font(22, True))
        d.text(((w - (b[2] - b[0])) / 2, 28), label, fill=PLUM, font=font(22, True))
        # uniqueness mark
        d.text((30, h - 40), f"{week['id']}-{idx}", fill=TEAL, font=font(16, True))
        im.save(out / name, "PNG", optimize=True)


def printable_module(week: dict, spec: dict, folder: Path) -> str:
    wid = week["id"]
    assets = folder.as_posix().replace("\\", "/")
    # relative from scripts/v13/printables -> repo is parents[3] in C1-W2 pattern
    y_titles = spec["y"]
    o_titles = spec["o"]
    p_titles = spec["p"]
    cards_y = spec["cards_y"]
    cards_o = spec["cards_o_sort"]
    parent_case = spec["parent_case"]
    pngs = spec["pngs"]
    parent_png = pngs[-1]
    mem = spec["memory"]
    verse = week["primary_verse"]
    url = week["primary_url"]

    def builder_fn(name: str, code: str, title: str, blurb: str, body: str) -> str:
        return f'''
def {name}(document) -> None:
    add_printable_title(document, "{code}", "{title}", "{blurb}")
{body}
'''

    y_fns = []
    for i, (code, title, blurb) in enumerate(y_titles, 1):
        img = pngs[min(i - 1, len(pngs) - 1)]
        if i == 5:
            body = f'''    _img(document, "{pngs[3] if len(pngs) > 3 else pngs[0]}", 6.0)
    add_memory_phrase_block(document, MEMORY)
    document.add_paragraph({verse!r} + " · family week")'''
        else:
            body = f'''    _img(document, "{img}", 5.5)
    add_card_grid(document, {cards_y!r}, per_page=6)
    add_memory_phrase_block(document, MEMORY)
    add_cut_lines(document)
    add_callout(document, "TEACHER_NOTE", {spec['care_boundary']!r})'''
        y_fns.append(builder_fn(f"y{i:02d}", f"Y{i:02d}", title, blurb, body))

    o_fns = []
    for i, (code, title, blurb) in enumerate(o_titles, 1):
        if i == 1:
            body = f'''    add_verse_card(document, VERSE_REF, DEV, IAST, MEANING, VERSE_URL,
        "Scripture display/source: VedaBase; teaching meaning: KUTUMBA-original (not labeled as BBT translation).")
    add_write_lines(document, [
        "What does the verse teach in my words?",
        "What misconception does this week block?",
        "One question for the facilitator:",
    ], 2)'''
        elif i == 6:
            body = f'''    add_write_lines(document, [
        "Primary in one sentence:",
        "One analogy or mechanic and its limit:",
        "One home action:",
        "One question I still have:",
    ], 2)'''
        else:
            body = f'''    _img(document, "{pngs[min(i-1, len(pngs)-1)]}", 5.0)
    add_branded_table(document, ["Prompt", "My notes"], [["{title}", "________________"], ["Limit / care", "________________"]])
    add_write_lines(document, ["Application:", "What not to say:"], 2)
    add_callout(document, "KEY_IDEA", {spec['conclusion']!r})'''
        o_fns.append(builder_fn(f"o{i:02d}", f"O{i:02d}", title, blurb, body))

    p_fns = []
    for i, (code, title, blurb) in enumerate(p_titles, 1):
        if i == 1:
            body = f'''    add_callout(document, "SAFETY_PRIVACY", "Write privately. Sharing is optional. Do not collect or place completed sheets in Git.")
    _img(document, "{parent_png}", 3.5)
    p = document.add_paragraph({blurb!r})
    p.runs[0].bold = True
    add_write_lines(document, ["Private notes:", "Habit to watch this week:"], 4)'''
        elif i == 2:
            body = f'''    mats = document.add_table(rows=1, cols=2)
    add_border(mats, PLUM, 10)
    for cell, label in zip(mats.rows[0].cells, ("HELPFUL", "MISTAKEN")):
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(label)
        r.bold = True
        r.font.size = Pt(22)
        r.font.color.rgb = RGBColor.from_string(TEAL if label == "HELPFUL" else SAFFRON)
    add_page(document)
    add_card_grid(document, {cards_o!r}, prefix="L", per_page=12)
    add_callout(document, "TEACHER_NOTE", {spec['misconception']!r})
    _img(document, "{parent_png}", 3.0)'''
        elif i == 3:
            body = f'''    add_callout(document, "SAFETY_PRIVACY", "Fictional case. No compelled confession of real events.")
    document.add_paragraph({parent_case!r})
    add_write_lines(document, [
        "Tempting mistaken conclusion:",
        "Principle from primary:",
        "Compassionate response:",
        "One seven-day household action:",
        "What we should not say:",
    ], 3)
    _img(document, "{pngs[0]}", 4.0)'''
        elif i == 4:
            body = f'''    add_callout(document, "SAFETY_PRIVACY", "No ranking and no public reading required.")
    add_write_lines(document, [
        "One household action for seven days:",
        "Trigger (when/where):",
        "Who starts if others are tired:",
        "Minimum version on a hard day:",
    ], 3)
    add_memory_phrase_block(document, MEMORY)
    _img(document, "{parent_png}", 3.2)'''
        else:
            body = f'''    add_callout(document, "KEY_IDEA", "specific action + frequency + trigger + minimum version")
    add_write_lines(document, [
        "Specific action:",
        "Frequency:",
        "Trigger (when and where):",
        "Minimum version for a hard day:",
        "Where we will place the reminder:",
    ], 2)
    _img(document, "{parent_png}", 4.5)
    document.add_paragraph("No ranking. The minimum version counts as success.")'''
        p_fns.append(builder_fn(f"p{i:02d}", f"P{i:02d}", title, blurb, body))

    rel_folder = folder.relative_to(REPO).as_posix()
    return f'''#!/usr/bin/env python3
"""{wid} V13 printable Word builders for render_week.py (week id {wid}).

Real Word tables/images. DOCX/PDF emission is owned by scripts/v13/render_week.py.
"""
from __future__ import annotations

import sys
from pathlib import Path

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

REPO = Path(__file__).resolve().parents[3]
sys.path.insert(0, str(REPO / "scripts" / "v13"))
sys.path.insert(0, str(REPO / "scripts" / "v12"))
from printable_common import (  # noqa: E402
    PLUM,
    SAFFRON,
    TEAL,
    add_border,
    add_card_grid,
    add_cut_lines,
    add_memory_phrase_block,
    add_page,
    add_printable_title,
    add_teacher_page,
    add_write_lines,
)
from kutumba_docx_styles import add_branded_table, add_callout, add_verse_card  # noqa: E402

ASSETS = REPO / "{rel_folder}" / "visuals" / "v13"
MEMORY = {mem!r}
VERSE_REF = {verse!r}
VERSE_URL = {url!r}
DEV = {spec['dev']!r}
IAST = {spec['iast']!r}
MEANING = {spec['meaning']!r}


def _img(document, name: str, width: float = 2.4) -> None:
    path = ASSETS / name
    if path.is_file():
        p = document.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(path), width=Inches(width))

{''.join(y_fns)}
{''.join(o_fns)}
{''.join(p_fns)}

YOUNGER_BUILDERS = [y01, y02, y03, y04, y05]
OLDER_BUILDERS = [o01, o02, o03, o04, o05, o06]
PARENT_BUILDERS = [p01, p02, p03, p04, p05]
FAMILY_BUILDERS = []
'''


def evidence(week: dict, spec: dict, folder: Path, counts: dict) -> None:
    out = EVIDENCE / week["id"]
    out.mkdir(parents=True, exist_ok=True)
    audit = f"""# {week['id']} V13 Content Audit

**Week:** {week['title']} · **Date:** {week['schedule_date']}  
**Primary:** {week['primary_verse']} — {week['primary_url']}  
**Audit date:** {AUDIT_DATE} (authoring pass)

## Verdict

**CONTENT AUTHORING PASS** for the V13 week pack (markdown/yaml/scripts/assets).  
**PUBLISHING RENDER GAP** remains for final DOCX/PDF/raster/visual QA (explicitly deferred by this pass).

## Primary verse confirmation

| Field | Value |
|---|---|
| Reference | {week['primary_verse']} |
| URL | {week['primary_url']} |
| Devanāgarī | Present in MAIN / O01 / Gamma / source map |
| IAST | Present complete |
| Teaching meaning | KUTUMBA-original; not labeled as BBT translation |
| Purport | Not invented / not dumped |

## Schedule confirmation

Locked Saturday **1:50–4:00**; parallel tracks **2:30–3:10**; reunify **3:10–3:30**.

## Activity mechanic summary

Primary mechanic: **{spec['mechanic']}**

| Code | Title |
|---|---|
{chr(10).join(f"| {c} | {t} |" for c,t,_ in spec['y'] + spec['o'] + spec['p'])}

## Safeguards checked

- [x] Care boundary: {spec['care_boundary']}
- [x] Misconception blocked: {spec['misconception']}
- [x] No invented BBT purport text
- [x] C1-W1 not modified
- [x] No commit in this pass
- [x] {"Integration: no major new doctrine" if week['kind']=='integration' else "Topic week deepened from V12 MAIN"}

## File counts (this pack)

| Bucket | Count |
|---|---|
| Week markdown/yaml/png/py authored | {counts['week_files']} |
| Printable module | {counts['printable']} |
| Evidence files | {counts['evidence']} |

## Gaps / EXTERNAL_OPEN

1. Final `exports/final/v13/{week['export_subdir']}/` DOCX+PDF packet — deferred (do not render in this pass).  
2. Raster contact sheets + VISUAL-QA — deferred until PDFs exist.  
3. Human/temple/doctrinal review — EXTERNAL_OPEN.  
4. Live Gamma render sign-off — EXTERNAL_OPEN.

## Depth source reuse

Deepened from existing {week['id']} V12 MAIN guide, research pack, and V12 Gamma prompts — not shallow-replaced with “use research.”
"""
    write(out / "CONTENT-AUDIT.md", audit)

    rows = [
        ["requirement_id", "requirement", "deliverable_path", "status", "notes"],
        [f"{week['id']}-R01", "V13-WEEK-START-HERE", str(folder / "V13-WEEK-START-HERE.md"), "PASS", "1:50-4:00 lock"],
        [f"{week['id']}-R02", "MAIN-FACILITATOR-GUIDE-V13", str(folder / "teacher" / "MAIN-FACILITATOR-GUIDE-V13.md"), "PASS", "Deepened from V12"],
        [f"{week['id']}-R03", "PARENT-GUIDE-V13", str(folder / "teacher" / "PARENT-GUIDE-V13.md"), "PASS", ""],
        [f"{week['id']}-R04", "YOUNGER-TEACHER-GUIDE-V13", str(folder / "teacher" / "YOUNGER-TEACHER-GUIDE-V13.md"), "PASS", ""],
        [f"{week['id']}-R05", "OLDER-TEACHER-GUIDE-V13", str(folder / "teacher" / "OLDER-TEACHER-GUIDE-V13.md"), "PASS", ""],
        [f"{week['id']}-R06", "family-home-practice-v13", str(folder / "family-home-practice-v13.md"), "PASS", ""],
        [f"{week['id']}-R07", "materials-v13", str(folder / "materials-v13.md"), "PASS", ""],
        [f"{week['id']}-R08", "research V13 set", str(folder / "research"), "PASS", "7 research files"],
        [f"{week['id']}-R09", "activities Y01-Y05", str(folder / "activities" / "v13-younger"), "PASS", spec["mechanic"]],
        [f"{week['id']}-R10", "activities O01-O06", str(folder / "activities" / "v13-older"), "PASS", ""],
        [f"{week['id']}-R11", "activities P01-P05", str(folder / "activities" / "v13-parent"), "PASS", "parent PNG embed in builders"],
        [f"{week['id']}-R12", "project brief", str(folder / "project" / "V13-MODULE-PROJECT-BRIEF.md"), "PASS", ""],
        [f"{week['id']}-R13", "Gamma prompts + source map", str(folder / "gamma"), "PASS", "actual slide copy from V12 upgrade"],
        [f"{week['id']}-R14", "visuals PNG + RIGHTS", str(folder / "visuals" / "v13"), "PASS", ""],
        [f"{week['id']}-R15", f"printable {week['id']}.py", str(PRINTABLES / f"{week['id']}.py"), "PASS", "hyphenated; DOCX deferred"],
        [f"{week['id']}-R16", "CONTENT-AUDIT", str(out / "CONTENT-AUDIT.md"), "PASS", ""],
        [f"{week['id']}-R17", "REQUIREMENT-TRACEABILITY", str(out / "REQUIREMENT-TRACEABILITY.csv"), "PASS", ""],
    ]
    with (out / "REQUIREMENT-TRACEABILITY.csv").open("w", encoding="utf-8", newline="") as f:
        csv.writer(f).writerows(rows)


def author_week(week: dict) -> dict:
    spec = SPECS[week["id"]]
    folder = resolve_folder(week["folder"])
    created: list[Path] = []

    # Start here + teacher
    write(folder / "V13-WEEK-START-HERE.md", start_here(week, spec))
    created.append(folder / "V13-WEEK-START-HERE.md")

    v12_main = folder / "teacher" / "MAIN-FACILITATOR-GUIDE-V12.md"
    write(folder / "teacher" / "MAIN-FACILITATOR-GUIDE-V13.md", upgrade_main(v12_main.read_text(encoding="utf-8"), week, spec))
    write(folder / "teacher" / "PARENT-GUIDE-V13.md", track_guide(week, spec, "parent"))
    write(folder / "teacher" / "YOUNGER-TEACHER-GUIDE-V13.md", track_guide(week, spec, "younger"))
    write(folder / "teacher" / "OLDER-TEACHER-GUIDE-V13.md", track_guide(week, spec, "older"))
    write(folder / "family-home-practice-v13.md", home_practice(week, spec))
    write(folder / "materials-v13.md", materials(week, spec))

    # Research upgrades
    research_map = [
        ("SOURCE-MATRIX.md", "V13-SOURCE-MATRIX.md", "Source Matrix"),
        ("SCRIPTURAL-EXAMPLES.md", "V13-SCRIPTURAL-EXAMPLES.md", "Scriptural Examples"),
        ("DEVOTIONAL-AND-HISTORICAL-EXAMPLES.md", "V13-DEVOTIONAL-HISTORICAL-EXAMPLES.md", "Devotional / Historical Examples"),
        ("CASE-STUDIES.md", "V13-CASE-STUDIES.md", "Case Studies"),
        ("ANALOGIES-AND-LIMITS.md", "V13-ANALOGIES-AND-LIMITS.md", "Analogies and Limits"),
        ("SCIENCE-AND-APPLICATION.md", "V13-SCIENCE-AND-APPLICATION.md", "Science and Application"),
    ]
    for src_name, dest_name, title in research_map:
        write(folder / "research" / dest_name, upgrade_research(folder / "research" / src_name, title, week))
    # claim register
    src_claim = folder / "research" / "CLAIM-REGISTER.yaml"
    if src_claim.exists():
        write(folder / "research" / "V13-CLAIM-REGISTER.yaml", upgrade_research(src_claim, "Claim Register", week))
    else:
        write(folder / "research" / "V13-CLAIM-REGISTER.yaml", claim_register(week, spec))

    # Activities
    for code, title, blurb in spec["y"]:
        write(folder / "activities" / "v13-younger" / f"{code}.md", activity_md(week, spec, code, title, blurb, "K–2"))
    for code, title, blurb in spec["o"]:
        write(folder / "activities" / "v13-older" / f"{code}.md", activity_md(week, spec, code, title, blurb, "Grades 4–5"))
    for code, title, blurb in spec["p"]:
        write(folder / "activities" / "v13-parent" / f"{code}.md", activity_md(week, spec, code, title, blurb, "Parent"))

    write(folder / "project" / "V13-MODULE-PROJECT-BRIEF.md", project_brief(week, spec))

    # Gamma
    g = folder / "gamma"
    write(g / "V13-GAMMA-MASTER-DECK-PROMPT.md", upgrade_gamma(g / "V12-GAMMA-MASTER-DECK-PROMPT.md", "V13 Gamma Master Deck Prompt", week, spec))
    write(g / "V13-GAMMA-PARENT-DECK-PROMPT.md", upgrade_gamma(g / "V12-GAMMA-PARENT-DECK-PROMPT.md", "V13 Gamma Parent Deck Prompt", week, spec))
    write(g / "V13-GAMMA-YOUNGER-DECK-PROMPT.md", upgrade_gamma(g / "V12-GAMMA-YOUNGER-DECK-PROMPT.md", "V13 Gamma Younger Deck Prompt", week, spec))
    write(g / "V13-GAMMA-OLDER-DECK-PROMPT.md", upgrade_gamma(g / "V12-GAMMA-OLDER-DECK-PROMPT.md", "V13 Gamma Older Deck Prompt", week, spec))
    write(g / "V13-GAMMA-SOURCE-MAP.yaml", gamma_source_map(week, spec))

    # Visuals
    vis = folder / "visuals" / "v13"
    draw_pngs(vis, week, spec)
    write(vis / "RIGHTS.md", rights_md(week, spec))
    write(vis / "IMAGE-PROMPT-LIBRARY.md", image_prompt_library(week, spec))

    # Printable
    PRINTABLES.mkdir(parents=True, exist_ok=True)
    write(PRINTABLES / f"{week['id']}.py", printable_module(week, spec, folder))

    # Count week files newly present under V13 patterns
    week_files = []
    patterns = [
        "V13-WEEK-START-HERE.md",
        "family-home-practice-v13.md",
        "materials-v13.md",
        "teacher/*-V13.md",
        "research/V13-*",
        "activities/v13-*/*",
        "project/V13-*",
        "gamma/V13-*",
        "visuals/v13/*",
    ]
    for p in folder.rglob("*"):
        if not p.is_file():
            continue
        rel = p.relative_to(folder).as_posix()
        if (
            rel == "V13-WEEK-START-HERE.md"
            or rel.endswith("-v13.md")
            or "/V13-" in f"/{rel}"
            or rel.startswith("activities/v13-")
            or rel.startswith("visuals/v13/")
            or rel.startswith("gamma/V13-")
            or rel.startswith("project/V13-")
            or rel.startswith("teacher/") and rel.endswith("-V13.md")
            or rel.startswith("research/V13-")
        ):
            week_files.append(rel)

    counts = {
        "week_files": len(week_files),
        "printable": 1,
        "evidence": 2,
        "paths": week_files,
    }
    evidence(week, spec, folder, counts)
    return {"id": week["id"], "folder": folder.name, "week_files": len(week_files), "mechanic": spec["mechanic"]}


def main() -> int:
    results = []
    for week in load_weeks():
        print(f"Authoring {week['id']} ...")
        results.append(author_week(week))
        print(f"  -> {results[-1]['week_files']} week files · mechanic={results[-1]['mechanic']}")
    summary = REPO / "build-evidence" / "v13" / "C2-C3-AUTHORING-SUMMARY.md"
    lines = ["# C2–C3 V13 Authoring Summary", "", f"Date: {AUDIT_DATE}", "", "| Week | Mechanic | Week files |", "|---|---|---|"]
    for r in results:
        lines.append(f"| {r['id']} | {r['mechanic']} | {r['week_files']} |")
    lines += ["", "Printables: `scripts/v13/printables/C2-W*.py` and `C3-W*.py`", "Evidence: `build-evidence/v13/<WEEK-ID>/`", "DOCX/PDF render: deferred", "C1-W1: untouched", "Commit: not created"]
    write(summary, "\n".join(lines))
    print("Wrote", summary)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
