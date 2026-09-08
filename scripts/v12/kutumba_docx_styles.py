#!/usr/bin/env python3
"""KUTUMBA V12 publication styles and reusable DOCX components."""
from __future__ import annotations

import os
from pathlib import Path
from typing import Iterable, Sequence

from docx import Document
from docx.document import Document as DocumentObject
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

BRAND = "KUTUMBA"
TAGLINE = "Families Growing in Krishna Consciousness"
DIRECTOR = "Program Director: Swapnil Patil"

# Hex values are intentionally centralized so renderers can use the same palette.
CREAM = "FFF8E8"
PLUM = "5B1933"
MAROON = "7A263A"
SAFFRON = "E59B24"
GOLD = "C98916"
TEAL = "4F7C78"
CHARCOAL = "2E3033"
WHITE = "FFFFFF"
PALE_GOLD = "F8E8BF"
PALE_TEAL = "E4F0ED"
PALE_ROSE = "F4E4E8"
PALE_GRAY = "EFF0F1"

CALLOUTS = {
    "SASTRA": ("ŚĀSTRA", PLUM, PALE_GOLD),
    "KEY_IDEA": ("KEY IDEA", SAFFRON, CREAM),
    "TEACHER_NOTE": ("TEACHER NOTE", TEAL, PALE_TEAL),
    "DO_NOT_SPECULATE": ("DO NOT SPECULATE", MAROON, PALE_ROSE),
    "FAMILY_APPLICATION": ("FAMILY APPLICATION", TEAL, PALE_TEAL),
    "CHILD_ACTIVITY": ("CHILD ACTIVITY", SAFFRON, CREAM),
    "MATERIALS": ("MATERIALS", GOLD, PALE_GOLD),
    "HOME_PRACTICE": ("HOME PRACTICE", PLUM, PALE_GOLD),
    "SAFETY_PRIVACY": ("SAFETY & PRIVACY", MAROON, PALE_ROSE),
    "TIME_CUE": ("TIME CUE", TEAL, PALE_GRAY),
    "RIGHTS_NOTE": ("RIGHTS NOTE", CHARCOAL, PALE_GRAY),
}


def _rgb(hex_value: str) -> RGBColor:
    return RGBColor.from_string(hex_value)


def _installed_font(preferred: Sequence[str], fallback: str) -> str:
    """Return the first commonly installed Windows font, without extra deps."""
    windows = Path(os.environ.get("WINDIR", r"C:\Windows")) / "Fonts"
    known_files = {
        "Nirmala UI": ("Nirmala.ttf", "NirmalaB.ttf"),
        "Mangal": ("mangal.ttf", "mangalb.ttf"),
        "Segoe UI": ("segoeui.ttf",),
        "Cambria": ("cambria.ttc", "cambria.ttf"),
        "Calibri": ("calibri.ttf", "calibri.ttf"),
    }
    for name in preferred:
        if any((windows / candidate).exists() for candidate in known_files.get(name, ())):
            return name
    return fallback


DEVANAGARI_FONT = _installed_font(("Nirmala UI", "Mangal"), "Nirmala UI")


def _set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def _set_cell_margins(cell, top: int = 100, start: int = 120, bottom: int = 100, end: int = 120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_table_borders(table, color: str = "D8C8B2", size: int = 6) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), str(size))
        tag.set(qn("w:color"), color)


def _set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)


def _set_keep_with_next(paragraph, value: bool = True) -> None:
    paragraph.paragraph_format.keep_with_next = value


def _set_run_font(run, name: str, size: float | None = None) -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run._element.rPr.rFonts.set(qn("w:cs"), name)
    if size is not None:
        run.font.size = Pt(size)


def _add_field(run, instruction: str) -> None:
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    fallback = OxmlElement("w:t")
    fallback.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instr, separate, fallback, end))


def _set_page_background(document: DocumentObject, color: str = CREAM) -> None:
    background = document._element.find(qn("w:background"))
    if background is None:
        background = OxmlElement("w:background")
        document._element.insert(0, background)
    background.set(qn("w:color"), color)
    settings = document.settings._element
    display = settings.find(qn("w:displayBackgroundShape"))
    if display is None:
        settings.append(OxmlElement("w:displayBackgroundShape"))


def _style_font(style, name: str, size: float, color: str = CHARCOAL, bold: bool = False) -> None:
    style.font.name = name
    style.font.size = Pt(size)
    style.font.color.rgb = _rgb(color)
    style.font.bold = bold
    style._element.rPr.rFonts.set(qn("w:eastAsia"), name)


def _get_or_add_style(document: DocumentObject, name: str, style_type=WD_STYLE_TYPE.PARAGRAPH):
    try:
        return document.styles[name]
    except KeyError:
        return document.styles.add_style(name, style_type)


def apply_kutumba_styles(document: DocumentObject) -> DocumentObject:
    """Apply the complete V12 page, paragraph, character, and table style set."""
    _set_page_background(document)
    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)
    section.header_distance = Inches(0.28)
    section.footer_distance = Inches(0.28)

    normal = document.styles["Normal"]
    _style_font(normal, "Calibri", 10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08
    normal.paragraph_format.widow_control = True

    title = document.styles["Title"]
    _style_font(title, "Cambria", 32, PLUM, True)
    title.paragraph_format.space_after = Pt(10)

    subtitle = document.styles["Subtitle"]
    _style_font(subtitle, "Segoe UI", 14, TEAL)
    subtitle.paragraph_format.space_after = Pt(8)

    heading_specs = {
        "Heading 1": (20, PLUM, 16, 7),
        "Heading 2": (15, MAROON, 12, 5),
        "Heading 3": (12, TEAL, 9, 4),
        "Heading 4": (10.5, PLUM, 7, 3),
    }
    for name, (size, color, before, after) in heading_specs.items():
        style = document.styles[name]
        _style_font(style, "Cambria", size, color, True)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Bullet 2", "List Number", "List Number 2"):
        style = document.styles[name]
        _style_font(style, "Calibri", 10.5)
        style.paragraph_format.space_after = Pt(3)

    quote = document.styles["Quote"]
    _style_font(quote, "Cambria", 10.5, PLUM)
    quote.font.italic = True
    quote.paragraph_format.left_indent = Inches(0.25)
    quote.paragraph_format.right_indent = Inches(0.15)
    quote.paragraph_format.space_before = Pt(5)
    quote.paragraph_format.space_after = Pt(7)

    caption = document.styles["Caption"]
    _style_font(caption, "Segoe UI", 8.5, TEAL, True)

    for name, size, color, bold in (
        ("KUTUMBA Eyebrow", 9, SAFFRON, True),
        ("KUTUMBA Meta", 9, TEAL, False),
        ("KUTUMBA Callout Label", 8.5, WHITE, True),
        ("KUTUMBA Callout Body", 10.25, CHARCOAL, False),
        ("KUTUMBA Verse Devanagari", 15, PLUM, False),
        ("KUTUMBA Verse IAST", 10.5, TEAL, False),
        ("KUTUMBA Rights", 8, CHARCOAL, False),
    ):
        style = _get_or_add_style(document, name)
        font = DEVANAGARI_FONT if "Devanagari" in name else ("Segoe UI" if size <= 9 else "Calibri")
        _style_font(style, font, size, color, bold)
        style.paragraph_format.space_after = Pt(4)
        if "Verse" in name:
            style.paragraph_format.keep_together = True

    table_style = _get_or_add_style(document, "KUTUMBA Table", WD_STYLE_TYPE.TABLE)
    table_style.font.name = "Calibri"
    table_style.font.size = Pt(9)
    table_style.font.color.rgb = _rgb(CHARCOAL)
    return document


def add_cover(
    document: DocumentObject,
    title: str,
    subtitle: str,
    scope: str,
    version: str = "V12",
    status: str = "Internal founding-cohort teaching material — human/temple review EXTERNAL_OPEN",
) -> None:
    """Add a branded cover and begin body content on a fresh page."""
    apply_kutumba_styles(document)
    p = document.add_paragraph()
    p.paragraph_format.space_before = Inches(0.65)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(BRAND)
    _set_run_font(run, "Segoe UI", 18)
    run.bold = True
    run.font.color.rgb = _rgb(SAFFRON)

    rule = document.add_table(rows=1, cols=1)
    rule.alignment = WD_TABLE_ALIGNMENT.CENTER
    rule.autofit = False
    rule.columns[0].width = Inches(2.0)
    cell = rule.cell(0, 0)
    _set_cell_shading(cell, PLUM)
    cell.height = Pt(3)
    cell.text = ""

    title_p = document.add_paragraph(style="Title")
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_before = Inches(0.55)
    title_p.add_run(title)

    if subtitle:
        sub_p = document.add_paragraph(style="Subtitle")
        sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        sub_p.add_run(subtitle)

    scope_p = document.add_paragraph()
    scope_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    scope_p.paragraph_format.space_before = Pt(18)
    scope_run = scope_p.add_run(scope)
    _set_run_font(scope_run, "Cambria", 12)
    scope_run.bold = True
    scope_run.font.color.rgb = _rgb(MAROON)

    meta = document.add_table(rows=2, cols=1)
    meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta.autofit = False
    meta.columns[0].width = Inches(5.8)
    for cell in meta.column_cells(0):
        _set_cell_margins(cell, 130, 180, 130, 180)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    _set_cell_shading(meta.cell(0, 0), PALE_GOLD)
    _set_cell_shading(meta.cell(1, 0), PALE_TEAL)
    meta.cell(0, 0).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.cell(0, 0).paragraphs[0].add_run(
        f"{version}  •  {TAGLINE}\n{DIRECTOR}"
    ).bold = True
    meta.cell(1, 0).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    status_run = meta.cell(1, 0).paragraphs[0].add_run(status)
    status_run.italic = True
    status_run.font.color.rgb = _rgb(PLUM)
    _set_table_borders(meta, "E3CA91", 5)

    brand_p = document.add_paragraph()
    brand_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    brand_p.paragraph_format.space_before = Inches(1.0)
    run = brand_p.add_run(f"{BRAND}  |  {TAGLINE}")
    _set_run_font(run, "Segoe UI", 9)
    run.font.color.rgb = _rgb(TEAL)

    document.add_page_break()


def add_header_footer(document: DocumentObject, week_or_scope: str, audience: str) -> None:
    """Add consistent headers and rights-aware PAGE/NUMPAGES footers."""
    for section in document.sections:
        header = section.header
        hp = header.paragraphs[0]
        hp.clear()
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        left = hp.add_run(f"{BRAND}  •  ")
        _set_run_font(left, "Segoe UI", 8.5)
        left.bold = True
        left.font.color.rgb = _rgb(PLUM)
        detail = hp.add_run(f"{week_or_scope}  |  {audience}")
        _set_run_font(detail, "Segoe UI", 8.5)
        detail.font.color.rgb = _rgb(TEAL)
        p_pr = hp._p.get_or_add_pPr()
        borders = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "8")
        bottom.set(qn("w:space"), "4")
        bottom.set(qn("w:color"), SAFFRON)
        borders.append(bottom)
        p_pr.append(borders)

        footer = section.footer
        fp = footer.paragraphs[0]
        fp.clear()
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        brand = fp.add_run(f"{BRAND} • {TAGLINE}  |  ")
        _set_run_font(brand, "Segoe UI", 7.5)
        brand.bold = True
        brand.font.color.rgb = _rgb(PLUM)
        rights = fp.add_run("KUTUMBA-original framing; cited third-party material retains its own rights.  |  Page ")
        _set_run_font(rights, "Segoe UI", 7.5)
        rights.font.color.rgb = _rgb(CHARCOAL)
        page_run = fp.add_run()
        _set_run_font(page_run, "Segoe UI", 7.5)
        _add_field(page_run, "PAGE")
        of_run = fp.add_run(" of ")
        _set_run_font(of_run, "Segoe UI", 7.5)
        pages_run = fp.add_run()
        _set_run_font(pages_run, "Segoe UI", 7.5)
        _add_field(pages_run, "NUMPAGES")


def add_callout(document: DocumentObject, kind: str, text: str):
    """Add a two-cell branded callout and return the created table."""
    normalized = kind.strip().upper().replace(" ", "_").replace("-", "_")
    if normalized not in CALLOUTS:
        allowed = ", ".join(CALLOUTS)
        raise ValueError(f"Unknown callout kind {kind!r}; expected one of: {allowed}")
    label, accent, fill = CALLOUTS[normalized]
    table = document.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(1.35)
    table.columns[1].width = Inches(5.65)
    label_cell, body_cell = table.rows[0].cells
    _set_cell_shading(label_cell, accent)
    _set_cell_shading(body_cell, fill)
    for cell in (label_cell, body_cell):
        _set_cell_margins(cell, 130, 140, 130, 140)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    label_p = label_cell.paragraphs[0]
    label_p.style = "KUTUMBA Callout Label"
    label_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    label_p.add_run(label)
    body_p = body_cell.paragraphs[0]
    body_p.style = "KUTUMBA Callout Body"
    body_p.add_run(text)
    _set_table_borders(table, accent, 8)
    document.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_verse_card(
    document: DocumentObject,
    reference: str,
    devanagari: str,
    iast: str,
    teaching_meaning: str,
    url: str,
    rights_status: str,
):
    """Add a source-explicit verse card without implying translation ownership."""
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(6.9)
    cell = table.cell(0, 0)
    _set_cell_shading(cell, CREAM)
    _set_cell_margins(cell, 180, 220, 180, 220)
    _set_table_borders(table, GOLD, 14)

    ref_p = cell.paragraphs[0]
    ref_p.style = "Heading 3"
    ref_p.add_run(f"VERSE CARD  •  {reference}")

    dev_p = cell.add_paragraph(style="KUTUMBA Verse Devanagari")
    dev_run = dev_p.add_run(devanagari)
    _set_run_font(dev_run, DEVANAGARI_FONT, 15)

    iast_p = cell.add_paragraph(style="KUTUMBA Verse IAST")
    iast_run = iast_p.add_run(iast)
    iast_run.italic = True

    meaning_p = cell.add_paragraph()
    label = meaning_p.add_run("KUTUMBA teaching meaning: ")
    label.bold = True
    label.font.color.rgb = _rgb(PLUM)
    meaning_p.add_run(teaching_meaning)

    source_p = cell.add_paragraph(style="KUTUMBA Rights")
    source_p.add_run(f"Source: {url}\nRights/status: {rights_status}")
    document.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_branded_table(
    document: DocumentObject,
    headers: Sequence[object],
    rows: Iterable[Sequence[object]],
):
    """Add a banded, repeating-header table."""
    headers = [str(value) for value in headers]
    if not headers:
        raise ValueError("headers must contain at least one column")
    normalized_rows = [list(row) for row in rows]
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "KUTUMBA Table"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    _set_repeat_table_header(table.rows[0])
    for index, value in enumerate(headers):
        cell = table.rows[0].cells[index]
        _set_cell_shading(cell, PLUM)
        _set_cell_margins(cell)
        run = cell.paragraphs[0].add_run(value)
        run.bold = True
        run.font.color.rgb = _rgb(WHITE)
    for row_index, values in enumerate(normalized_rows):
        cells = table.add_row().cells
        for column_index in range(len(headers)):
            cell = cells[column_index]
            _set_cell_margins(cell)
            if row_index % 2:
                _set_cell_shading(cell, CREAM)
            value = values[column_index] if column_index < len(values) else ""
            cell.text = str(value)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
    _set_table_borders(table)
    document.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_toc_placeholder(document: DocumentObject) -> None:
    """Insert an updateable Word TOC field for Heading 1–3."""
    heading = document.add_heading("Contents", level=1)
    _set_keep_with_next(heading)
    paragraph = document.add_paragraph()
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin.set(qn("w:dirty"), "true")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = r'TOC \o "1-3" \h \z \u'
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    display = OxmlElement("w:t")
    display.text = "Right-click and choose Update Field to build the table of contents."
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instruction, separate, display, end))
    document.add_page_break()


__all__ = [
    "apply_kutumba_styles",
    "add_cover",
    "add_header_footer",
    "add_callout",
    "add_verse_card",
    "add_branded_table",
    "add_toc_placeholder",
]
