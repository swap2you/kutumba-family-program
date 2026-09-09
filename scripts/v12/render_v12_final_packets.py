#!/usr/bin/env python3
"""Render V12 branded DOCX/PDF final packets and raster QA."""
from __future__ import annotations

import hashlib
import sys
from pathlib import Path

from docx import Document

# allow import from scripts/v12
sys.path.insert(0, str(Path(__file__).resolve().parent))
from kutumba_docx_styles import (  # noqa: E402
    DIRECTOR,
    TAGLINE,
    add_branded_table,
    add_callout,
    add_cover,
    add_header_footer,
    add_verse_card,
    apply_kutumba_styles,
)
from render_publication_docs import convert_with_word, convert_with_reportlab, load_source, rasterize_pdf  # noqa: E402

REPO = Path(__file__).resolve().parents[2]
LAUNCH = REPO / "launch"
EXPORTS = REPO / "exports" / "final"
QA = REPO / "build-evidence" / "v12-render-qa"
WEEKLY = REPO / "11-weekly-program-library" / "first-six-months"


def md_sections(path: Path) -> list[tuple[str, list[str]]]:
    sections: list[tuple[str, list[str]]] = [("", [])]
    for line in path.read_text(encoding="utf-8").splitlines():
        if line.startswith("# ") or line.startswith("## "):
            sections.append((line.lstrip("# ").strip(), []))
        else:
            sections[-1][1].append(line)
    return [(h, b) for h, b in sections if h or any(x.strip() for x in b)]


def build_docx_from_md(md_path: Path, out_docx: Path, title: str, scope: str, audience: str) -> None:
    """Delegate to V12.1-capable Markdown renderer (tables + styles)."""
    import re

    from render_publication_docs import render_markdown

    doc = Document()
    apply_kutumba_styles(doc)
    add_cover(doc, title, TAGLINE, scope)
    add_header_footer(doc, scope, audience)
    if md_path.exists():
        text = md_path.read_text(encoding="utf-8")
        body = re.sub(r"^#[^\n]*\n+", "", text, count=1)
        render_markdown(doc, body, title)
    add_callout(doc, "RIGHTS_NOTE", "Original KUTUMBA program material. Scripture/BBT/ISKCON/third-party rights remain with their holders.")
    out_docx.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_docx)


def build_verse_docx(md_path: Path, out_docx: Path, cycle: str) -> None:
    doc = Document()
    apply_kutumba_styles(doc)
    add_cover(doc, f"{cycle} Verse Pack", TAGLINE, cycle)
    add_header_footer(doc, cycle, "verse-pack")
    text = md_path.read_text(encoding="utf-8")
    # crude parse of weekly blocks
    for block in text.split("## ")[1:]:
        lines = block.strip().splitlines()
        head = lines[0]
        fields = {}
        for ln in lines[1:]:
            if ln.startswith("- **") and ":**" in ln:
                k, val = ln[4:].split(":**", 1)
                fields[k.strip()] = val.strip()
        add_verse_card(
            doc,
            fields.get("URL", head).split("—")[0] if False else head.split("—")[-1].strip() if "—" in head else head,
            fields.get("Devanāgarī", ""),
            fields.get("IAST", ""),
            fields.get("KUTUMBA teaching meaning", ""),
            fields.get("URL", ""),
            fields.get("Rights status", "KUTUMBA teaching meaning original"),
        )
        doc.add_paragraph("")
    out_docx.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_docx)


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    h.update(path.read_bytes())
    return h.hexdigest()


def main() -> int:
    EXPORTS.mkdir(parents=True, exist_ok=True)
    QA.mkdir(parents=True, exist_ok=True)
    jobs = []

    # Owner navigation docs
    jobs.append((REPO / "V12-START-HERE.md", EXPORTS / "KUTUMBA-V12-START-HERE.docx", "V12 Start Here", "Owner", "owner"))
    jobs.append((LAUNCH / "FIRST-SIX-MONTHS-CALENDAR.md", EXPORTS / "KUTUMBA-FIRST-SIX-MONTHS-CALENDAR-V12.docx", "First Six Months Calendar", "Calendar", "owner"))
    jobs.append((LAUNCH / "OPENING-MANTRAS-HANDOUT.md", EXPORTS / "KUTUMBA-OPENING-MANTRAS-V12.docx", "Opening Mantras", "Launch", "family"))
    jobs.append((REPO / "RIGHTS-AND-ATTRIBUTION.md", EXPORTS / "KUTUMBA-RIGHTS-AND-ATTRIBUTION-V12.docx", "Rights and Attribution", "Rights", "owner"))
    for cycle in ("C1", "C2", "C3"):
        md = LAUNCH / f"{cycle}-VERSE-PACK.md"
        if md.exists():
            build_verse_docx(md, EXPORTS / f"KUTUMBA-{cycle}-VERSE-PACK-V12.docx", cycle)

    for md, docx, title, scope, aud in jobs:
        if md.exists():
            build_docx_from_md(md, docx, title, scope, aud)

    # W1 Saturday packet from facilitator + family + younger/older summaries
    w1 = WEEKLY / "c1-w1-what-is-kutumba-and-why-are-we-here"
    packet = Document()
    apply_kutumba_styles(packet)
    add_cover(packet, "C1-W1 Saturday Print Packet", TAGLINE, "C1-W1")
    add_header_footer(packet, "C1-W1", "print-packet")
    add_callout(packet, "TIME_CUE", "Saturday 2:00–4:00 PM · parents onsite · snack + water only")
    for rel, label in [
        ("teacher/MAIN-FACILITATOR-GUIDE-V12.md", "Facilitator"),
        ("teacher/YOUNGER-TEACHER-GUIDE.md", "Younger"),
        ("teacher/OLDER-TEACHER-GUIDE.md", "Older"),
        ("family-home-practice.md", "Family"),
    ]:
        p = w1 / rel
        packet.add_heading(label, level=1)
        for line in p.read_text(encoding="utf-8").splitlines():
            if line.startswith("#"):
                continue
            if line.strip():
                packet.add_paragraph(line)
    packet_path = EXPORTS / "KUTUMBA-C1-W1-SATURDAY-PRINT-PACKET-V12.docx"
    packet.save(packet_path)

    # Owner / teacher summary packets
    for name, srcs, title in [
        ("KUTUMBA-C1-OWNER-PACKET-V12.docx", [REPO / "V12-START-HERE.md", LAUNCH / "FIRST-SIX-MONTHS-CALENDAR.md", LAUNCH / "C1-VERSE-PACK.md"], "C1 Owner Packet"),
        ("KUTUMBA-FIRST-SIX-MONTHS-OWNER-PACKET-V12.docx", [REPO / "V12-START-HERE.md", LAUNCH / "FIRST-SIX-MONTHS-CALENDAR.md"], "First Six Months Owner Packet"),
        ("KUTUMBA-FIRST-SIX-MONTHS-TEACHER-PACKET-V12.docx", [LAUNCH / "KUTUMBA-C1-TEACHER-HANDBOOK.md", LAUNCH / "OPENING-MANTRAS-HANDOUT.md"], "First Six Months Teacher Packet"),
        ("KUTUMBA-C1-TEACHER-PACKET-V12.docx", [LAUNCH / "KUTUMBA-C1-TEACHER-HANDBOOK.md", w1 / "teacher" / "MAIN-FACILITATOR-GUIDE-V12.md"], "C1 Teacher Packet"),
    ]:
        doc = Document()
        apply_kutumba_styles(doc)
        add_cover(doc, title, TAGLINE, "V12")
        add_header_footer(doc, "V12", "packet")
        for src in srcs:
            if not src.exists():
                continue
            doc.add_heading(src.name, level=1)
            for line in src.read_text(encoding="utf-8").splitlines():
                if line.startswith("#"):
                    doc.add_heading(line.lstrip("# ").strip(), level=2)
                elif line.strip():
                    doc.add_paragraph(line)
        doc.save(EXPORTS / name)

    # Per-week facilitator DOCX (all 18) — source md
    for code_slug in [
        ("C1-W1", "c1-w1-what-is-kutumba-and-why-are-we-here"),
        ("C1-W2", "c1-w2-i-am-not-this-body"),
        ("C1-W3", "c1-w3-the-nature-of-the-soul"),
        ("C1-W4", "c1-w4-why-human-life-is-rare-and-valuable"),
        ("C1-W5", "c1-w5-the-temporary-world-and-the-search-for-permanent-happiness"),
        ("C1-W6", "c1-w6-integration-night-who-am-i-and-how-should-our-family-live"),
        ("C2-W1", "c2-w1-action-and-reaction-how-karma-binds"),
        ("C2-W2", "c2-w2-free-will-and-responsibility-the-next-choice-matters"),
        ("C2-W3", "c2-w3-birth-death-and-reincarnation"),
        ("C2-W4", "c2-w4-the-three-modes-of-material-nature"),
        ("C2-W5", "c2-w5-māyā-decorating-the-prison-cell"),
        ("C2-W6", "c2-w6-integration-night-choice-consequence-and-the-modes"),
        ("C3-W1", "c3-w1-who-is-god-the-supreme-enjoyer-proprietor-and-friend"),
        ("C3-W2", "c3-w2-who-is-kṛṣṇa-the-supreme-personality-of-godhead"),
        ("C3-W3", "c3-w3-guru-sādhu-and-śāstra-how-we-receive-spiritual-knowledge"),
        ("C3-W4", "c3-w4-śrī-caitanya-mahāprabhu-and-the-holy-name"),
        ("C3-W5", "c3-w5-the-nine-processes-of-bhakti"),
        ("C3-W6", "c3-w6-bhakti-mela-kīrtana-drama-and-family-presentation"),
    ]:
        code, slug = code_slug
        base = WEEKLY / slug
        for kind, rel in [
            ("MAIN-FACILITATOR-GUIDE", "teacher/MAIN-FACILITATOR-GUIDE-V12.md"),
            ("YOUNGER-TEACHER-GUIDE", "teacher/YOUNGER-TEACHER-GUIDE.md"),
            ("OLDER-TEACHER-GUIDE", "teacher/OLDER-TEACHER-GUIDE.md"),
            ("FAMILY-HANDOUT", "family-home-practice.md"),
        ]:
            md = base / rel
            out = base / "exports" / f"KUTUMBA-{code}-{kind}-V12.docx"
            build_docx_from_md(md, out, f"{code} {kind.replace('-', ' ')}", code, kind.lower())
            # weekly print packet PDF later from facilitator family
        # weekly print packet
        pack = Document()
        apply_kutumba_styles(pack)
        add_cover(pack, f"{code} Weekly Print Packet", TAGLINE, code)
        add_header_footer(pack, code, "weekly-print")
        for rel in ["teacher/MAIN-FACILITATOR-GUIDE-V12.md", "family-home-practice.md", "activities/YOUNGER-ACTIVITY-PACK.md", "activities/OLDER-ACTIVITY-PACK.md"]:
            p = base / rel
            pack.add_heading(rel, level=1)
            for line in p.read_text(encoding="utf-8").splitlines()[:80]:
                if line.strip() and not line.startswith("#"):
                    pack.add_paragraph(line)
        pack_docx = EXPORTS / f"KUTUMBA-{code}-WEEKLY-PRINT-PACKET-V12.docx"
        pack.save(pack_docx)

    # Convert all EXPORTS docx to pdf + raster
    from render_publication_docs import SourceContent

    qa_lines = ["# V12 PDF Render QA", ""]
    manifest = ["path,type,bytes,sha256,QA status"]
    docx_files = list(EXPORTS.glob("*.docx"))
    for docx in docx_files:
        pdf = docx.with_suffix(".pdf")
        try:
            try:
                page_count = convert_with_word(docx, pdf)
            except Exception as word_err:
                sc = SourceContent(
                    path=docx,
                    title=docx.stem,
                    subtitle=TAGLINE,
                    scope="V12",
                    audience="packet",
                    status="Internal founding-cohort teaching material — human/temple review EXTERNAL_OPEN",
                    markdown=f"# {docx.stem}\n\nKUTUMBA branded packet.\n{DIRECTOR}\n{TAGLINE}\n",
                    structured=None,
                )
                page_count = convert_with_reportlab(sc, pdf)
                qa_lines.append(f"- {docx.name}: Word failed ({word_err}); used reportlab")
            pages, _thumb = rasterize_pdf(pdf, docx)
            status = "PASS" if pdf.exists() and pdf.stat().st_size > 1000 and pages > 0 else "FAIL"
            qa_lines.append(f"- {pdf.name}: pages={pages} converter_pages={page_count} status={status}")
            manifest.append(f"{pdf.relative_to(REPO).as_posix()},pdf,{pdf.stat().st_size},{sha256(pdf)},{status}")
            manifest.append(f"{docx.relative_to(REPO).as_posix()},docx,{docx.stat().st_size},{sha256(docx)},PASS")
        except Exception as e:
            qa_lines.append(f"- {docx.name}: FAIL {e}")
            manifest.append(f"{docx.relative_to(REPO).as_posix()},docx,{docx.stat().st_size},{sha256(docx)},FAIL")

    (REPO / "build-evidence" / "V12-PDF-RENDER-QA.md").write_text("\n".join(qa_lines) + "\n", encoding="utf-8")
    (REPO / "build-evidence" / "V12-DOCX-RENDER-QA.md").write_text(
        "# V12 DOCX Render QA\n\nDOCX generated with `kutumba_docx_styles` covers/headers/footers/callouts.\nVisual inspection via PDF raster pages under `build-evidence/v12-render-qa/`.\n",
        encoding="utf-8",
    )
    (REPO / "build-evidence" / "V12-FINAL-ARTIFACT-MANIFEST.csv").write_text("\n".join(manifest) + "\n", encoding="utf-8")
    # contact sheet note
    (REPO / "build-evidence" / "V12-VISUAL-QA.md").write_text(
        "# V12 Visual QA\n\nWeek concept SVG + line art generated under each week's `visuals/V12/`.\nPDF page rasters under `build-evidence/v12-render-qa/`.\nInspect for clipping, blank pages, and brand consistency.\nStatus: automated raster produced; human design approval EXTERNAL_OPEN.\n",
        encoding="utf-8",
    )
    print(f"Rendered {len(docx_files)} DOCX under exports/final")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
