#!/usr/bin/env python3
"""Render KUTUMBA V12 publication sources to branded DOCX, PDF, and QA PNGs.

Examples:
    python scripts/v12/render_publication_docs.py --sample
    python scripts/v12/render_publication_docs.py --all
    python scripts/v12/render_publication_docs.py --input path/to/file.md
"""
from __future__ import annotations

import argparse
import html
import json
import re
import shutil
import sys
import traceback
import zipfile
from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable, Sequence

import pymupdf as fitz
import yaml
from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

try:
    from kutumba_docx_styles import (
        CHARCOAL,
        CREAM,
        MAROON,
        PLUM,
        SAFFRON,
        TAGLINE,
        TEAL,
        add_branded_table,
        add_callout,
        add_cover,
        add_header_footer,
        add_toc_placeholder,
        add_verse_card,
        apply_kutumba_styles,
    )
except ImportError:  # Supports module execution from repository root.
    from scripts.v12.kutumba_docx_styles import (
        CHARCOAL,
        CREAM,
        MAROON,
        PLUM,
        SAFFRON,
        TAGLINE,
        TEAL,
        add_branded_table,
        add_callout,
        add_cover,
        add_header_footer,
        add_toc_placeholder,
        add_verse_card,
        apply_kutumba_styles,
    )

REPO = Path(__file__).resolve().parents[2]
WEEKLY = REPO / "11-weekly-program-library" / "first-six-months"
OUTPUT_ROOT = REPO / "build-evidence" / "v12-publication-docs"
QA_ROOT = REPO / "build-evidence" / "v12-render-qa"
DOCX_QA = REPO / "build-evidence" / "V12-DOCX-RENDER-QA.md"
PDF_QA = REPO / "build-evidence" / "V12-PDF-RENDER-QA.md"
WORD_EXE = Path(r"C:\Program Files\Microsoft Office\root\Office16\WINWORD.EXE")
DEFAULT_STATUS = (
    "Internal founding-cohort teaching material — "
    "human/temple review EXTERNAL_OPEN"
)

PUBLICATION_NAMES = {
    "overview.md",
    "facilitator-guide.md",
    "parent-lesson.md",
    "opening-hook.md",
    "worksheet.md",
    "bhakti-lab.md",
    "family-application.md",
    "lesson.md",
    "kisora-kisori-lesson.md",
    "lala-lali-lesson.md",
    "shared-family-transition.md",
    "MAIN-FACILITATOR-GUIDE-V11.md",
    "OLDER-ACTIVITY-PACK.md",
    "YOUNGER-ACTIVITY-PACK.md",
}

CALLOUT_ALIASES = {
    "SASTRA": "SASTRA",
    "ŚĀSTRA": "SASTRA",
    "KEY IDEA": "KEY_IDEA",
    "TEACHER NOTE": "TEACHER_NOTE",
    "DO NOT SPECULATE": "DO_NOT_SPECULATE",
    "FAMILY APPLICATION": "FAMILY_APPLICATION",
    "CHILD ACTIVITY": "CHILD_ACTIVITY",
    "MATERIALS": "MATERIALS",
    "HOME PRACTICE": "HOME_PRACTICE",
    "SAFETY": "SAFETY_PRIVACY",
    "SAFETY & PRIVACY": "SAFETY_PRIVACY",
    "SAFETY AND PRIVACY": "SAFETY_PRIVACY",
    "PRIVACY": "SAFETY_PRIVACY",
    "TIME CUE": "TIME_CUE",
    "RIGHTS NOTE": "RIGHTS_NOTE",
    "RIGHTS": "RIGHTS_NOTE",
}


@dataclass
class SourceContent:
    path: Path
    title: str
    subtitle: str
    scope: str
    audience: str
    status: str = DEFAULT_STATUS
    metadata: dict[str, Any] = field(default_factory=dict)
    markdown: str | None = None
    structured: Any = None


@dataclass
class RenderResult:
    source: Path
    docx: Path
    pdf: Path
    png: Path | None = None
    docx_pages: int = 0
    pdf_pages: int = 0
    conversion: str = ""
    docx_open: bool = False
    pdf_open: bool = False
    error: str = ""

    @property
    def docx_pass(self) -> bool:
        return (
            self.docx_open
            and self.docx_pages > 0
            and self.docx.stat().st_size > 0
            and self.png is not None
            and self.png.exists()
            and self.png.stat().st_size > 0
        )

    @property
    def pdf_pass(self) -> bool:
        return (
            self.pdf_open
            and self.pdf_pages > 0
            and self.pdf.stat().st_size > 0
            and self.png is not None
            and self.png.exists()
            and self.png.stat().st_size > 0
        )


def _humanize(value: str) -> str:
    value = re.sub(r"^c\d+-w\d+-", "", value, flags=re.IGNORECASE)
    value = value.replace("-", " ").replace("_", " ")
    return " ".join(word if word.isupper() else word.capitalize() for word in value.split())


def _relative(path: Path) -> str:
    try:
        return path.resolve().relative_to(REPO.resolve()).as_posix()
    except ValueError:
        return path.resolve().as_posix()


def _week_scope(path: Path, metadata: dict[str, Any]) -> str:
    for key in ("week_code", "week", "scope", "module"):
        if metadata.get(key):
            return str(metadata[key])
    match = re.search(r"\b(c\d+-w\d+)", path.as_posix(), re.IGNORECASE)
    return match.group(1).upper() if match else _humanize(path.parent.name)


def _audience(path: Path, metadata: dict[str, Any]) -> str:
    if metadata.get("audience"):
        return str(metadata["audience"])
    parts = {part.lower() for part in path.parts}
    name = path.name.lower()
    if "children" in parts:
        if "kisora" in name:
            return "Older children / youth"
        if "lala" in name:
            return "Younger children"
        return "Children"
    if "teacher" in parts or "facilitator" in name:
        return "Facilitators"
    if "parent" in name:
        return "Parents / caregivers"
    if "worksheet" in name or "activity" in name:
        return "Family activity"
    return "Families and facilitators"


def _split_front_matter(text: str) -> tuple[dict[str, Any], str]:
    if not text.startswith("---"):
        return {}, text
    match = re.match(r"^---\s*\n(.*?)\n---\s*(?:\n|$)", text, re.DOTALL)
    if not match:
        return {}, text
    try:
        metadata = yaml.safe_load(match.group(1)) or {}
        if not isinstance(metadata, dict):
            metadata = {}
    except yaml.YAMLError:
        metadata = {}
    return metadata, text[match.end() :]


def load_source(path: Path) -> SourceContent:
    """Load Markdown, YAML, or JSON into normalized publication content."""
    suffix = path.suffix.lower()
    text = path.read_text(encoding="utf-8-sig")
    if suffix in {".md", ".markdown", ".txt"}:
        metadata, body = _split_front_matter(text)
        heading = re.search(r"^#\s+(.+?)\s*$", body, re.MULTILINE)
        title = str(metadata.get("title") or metadata.get("week_title") or (heading.group(1) if heading else _humanize(path.stem)))
        subtitle = str(metadata.get("subtitle") or TAGLINE)
        return SourceContent(
            path=path,
            title=title,
            subtitle=subtitle,
            scope=_week_scope(path, metadata),
            audience=_audience(path, metadata),
            status=str(metadata.get("status") or DEFAULT_STATUS),
            metadata=metadata,
            markdown=body,
        )
    if suffix == ".json":
        data = json.loads(text)
    elif suffix in {".yaml", ".yml"}:
        data = yaml.safe_load(text)
    else:
        raise ValueError(f"Unsupported source type: {path.suffix}")
    data = data or {}
    metadata = data if isinstance(data, dict) else {}
    title = str(metadata.get("title") or metadata.get("week_title") or _humanize(path.stem))
    return SourceContent(
        path=path,
        title=title,
        subtitle=str(metadata.get("subtitle") or TAGLINE),
        scope=_week_scope(path, metadata),
        audience=_audience(path, metadata),
        status=str(metadata.get("status") or DEFAULT_STATUS),
        metadata=metadata,
        structured=data,
    )


def _add_hyperlink(paragraph, label: str, url: str) -> None:
    relationship = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship)
    run = OxmlElement("w:r")
    run_properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), TEAL)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    run_properties.extend((color, underline))
    run.append(run_properties)
    text = OxmlElement("w:t")
    text.text = label
    run.append(text)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


INLINE_RE = re.compile(
    r"(\[([^\]]+)\]\((https?://[^)]+)\)|\*\*([^*]+)\*\*|(?<!\*)\*([^*\n]+)\*(?!\*)|`([^`]+)`)"
)


def _add_inline(paragraph, text: str) -> None:
    """Add basic Markdown emphasis and links as actual Word runs."""
    position = 0
    for match in INLINE_RE.finditer(text):
        if match.start() > position:
            paragraph.add_run(text[position : match.start()])
        if match.group(2) is not None:
            _add_hyperlink(paragraph, match.group(2), match.group(3))
        elif match.group(4) is not None:
            paragraph.add_run(match.group(4)).bold = True
        elif match.group(5) is not None:
            paragraph.add_run(match.group(5)).italic = True
        elif match.group(6) is not None:
            run = paragraph.add_run(match.group(6))
            run.font.name = "Consolas"
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor.from_string(MAROON)
        position = match.end()
    if position < len(text):
        paragraph.add_run(text[position:])


def _table_separator(line: str) -> bool:
    cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells)


def _table_cells(line: str) -> list[str]:
    # Escaped pipes are uncommon in curriculum tables; preserve them via sentinel.
    sentinel = "\u0000"
    line = line.replace(r"\|", sentinel)
    return [cell.strip().replace(sentinel, "|") for cell in line.strip().strip("|").split("|")]


def _callout_from_text(text: str) -> tuple[str, str] | None:
    cleaned = text.strip().lstrip(">").strip()
    alert = re.match(r"^\[!([A-Za-z_ &ŚĀ]+)\]\s*(.*)$", cleaned)
    if alert:
        key = alert.group(1).replace("_", " ").upper()
        return CALLOUT_ALIASES.get(key, key.replace(" ", "_")), alert.group(2)
    labeled = re.match(r"^\*{0,2}([^:*]{2,35})\*{0,2}:\s+(.+)$", cleaned)
    if labeled:
        key = labeled.group(1).strip().upper()
        if key in CALLOUT_ALIASES:
            return CALLOUT_ALIASES[key], labeled.group(2).strip()
    return None


def _parse_verse_directive(lines: Sequence[str], start: int) -> tuple[dict[str, str], int] | None:
    if lines[start].strip().upper() != ":::VERSE":
        return None
    values: dict[str, str] = {}
    index = start + 1
    while index < len(lines) and lines[index].strip() != ":::":
        if ":" in lines[index]:
            key, value = lines[index].split(":", 1)
            values[key.strip().lower()] = value.strip().strip('"')
        index += 1
    if index >= len(lines):
        return None
    return values, index + 1


def render_markdown(document, body: str, title: str) -> None:
    """Render Markdown semantically: headings, lists, callouts, tables, and verse cards."""
    lines = body.splitlines()
    index = 0
    skipped_title = False
    while index < len(lines):
        raw = lines[index]
        line = raw.rstrip()
        if not line.strip():
            index += 1
            continue

        directive = _parse_verse_directive(lines, index)
        if directive:
            values, index = directive
            add_verse_card(
                document,
                values.get("reference", "Reference pending"),
                values.get("devanagari", ""),
                values.get("iast", ""),
                values.get("teaching_meaning", values.get("meaning", "")),
                values.get("url", ""),
                values.get("rights_status", "Rights review pending"),
            )
            continue

        heading = re.match(r"^(#{1,6})\s+(.+)$", line)
        if heading:
            level = len(heading.group(1))
            text = heading.group(2).strip().rstrip("#").strip()
            if level == 1 and not skipped_title and text.casefold() == title.casefold():
                skipped_title = True
            else:
                paragraph = document.add_heading(level=min(level, 4))
                _add_inline(paragraph, text)
            index += 1
            continue

        if line.lstrip().startswith("|") and index + 1 < len(lines) and _table_separator(lines[index + 1]):
            headers = _table_cells(line)
            index += 2
            rows: list[list[str]] = []
            while index < len(lines) and lines[index].lstrip().startswith("|"):
                rows.append(_table_cells(lines[index]))
                index += 1
            add_branded_table(document, headers, rows)
            continue

        if line.lstrip().startswith(">"):
            quote_lines: list[str] = []
            while index < len(lines) and lines[index].lstrip().startswith(">"):
                quote_lines.append(lines[index].lstrip()[1:].strip())
                index += 1
            quote_text = " ".join(part for part in quote_lines if part)
            callout = _callout_from_text(quote_text)
            if callout and callout[0] in {
                "SASTRA", "KEY_IDEA", "TEACHER_NOTE", "DO_NOT_SPECULATE",
                "FAMILY_APPLICATION", "CHILD_ACTIVITY", "MATERIALS",
                "HOME_PRACTICE", "SAFETY_PRIVACY", "TIME_CUE", "RIGHTS_NOTE",
            }:
                add_callout(document, *callout)
            else:
                paragraph = document.add_paragraph(style="Quote")
                _add_inline(paragraph, quote_text)
            continue

        callout = _callout_from_text(line)
        if callout:
            add_callout(document, *callout)
            index += 1
            continue

        bullet = re.match(r"^(\s*)[-+*]\s+(.+)$", line)
        number = re.match(r"^(\s*)\d+[.)]\s+(.+)$", line)
        if bullet or number:
            match = bullet or number
            assert match is not None
            depth = min(2, len(match.group(1).expandtabs(4)) // 2)
            base = "List Bullet" if bullet else "List Number"
            style = base if depth == 0 else f"{base} 2"
            paragraph = document.add_paragraph(style=style)
            _add_inline(paragraph, match.group(2))
            index += 1
            continue

        if re.fullmatch(r"[-*_]{3,}", line.strip()):
            document.add_paragraph("◆  ◆  ◆", style="KUTUMBA Meta").alignment = 1
            index += 1
            continue

        paragraph_lines = [line.strip()]
        index += 1
        while index < len(lines):
            candidate = lines[index]
            if (
                not candidate.strip()
                or re.match(r"^(#{1,6})\s+", candidate)
                or re.match(r"^\s*([-+*]|\d+[.)])\s+", candidate)
                or candidate.lstrip().startswith((">", "|", ":::VERSE"))
            ):
                break
            paragraph_lines.append(candidate.strip())
            index += 1
        paragraph = document.add_paragraph()
        _add_inline(paragraph, " ".join(paragraph_lines))


def _scalar(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "Yes" if value else "No"
    return str(value)


def render_structured(document, value: Any, heading: str | None = None, level: int = 1) -> None:
    """Render nested JSON/YAML as sections, lists, metadata tables, and verse cards."""
    if heading:
        document.add_heading(_humanize(heading), level=min(level, 4))
    if isinstance(value, dict):
        verse_keys = {"reference", "devanagari", "iast", "url"}
        if verse_keys.issubset(value):
            add_verse_card(
                document,
                _scalar(value["reference"]),
                _scalar(value["devanagari"]),
                _scalar(value["iast"]),
                _scalar(value.get("teaching_meaning") or value.get("kutumba_teaching_meaning")),
                _scalar(value["url"]),
                _scalar(value.get("rights_status") or "Rights review pending"),
            )
            extras = {key: item for key, item in value.items() if key not in verse_keys | {"teaching_meaning", "kutumba_teaching_meaning", "rights_status"}}
            if extras:
                add_branded_table(document, ["Field", "Value"], [(key, _scalar(item)) for key, item in extras.items()])
            return
        scalar_rows = [(key, _scalar(item)) for key, item in value.items() if not isinstance(item, (dict, list))]
        if scalar_rows:
            add_branded_table(document, ["Field", "Value"], scalar_rows)
        for key, item in value.items():
            if isinstance(item, (dict, list)):
                render_structured(document, item, str(key), level + 1)
    elif isinstance(value, list):
        if value and all(isinstance(item, dict) for item in value):
            keys: list[str] = []
            for item in value:
                for key, nested in item.items():
                    if not isinstance(nested, (dict, list)) and key not in keys:
                        keys.append(key)
            if keys:
                add_branded_table(document, [_humanize(key) for key in keys], [[_scalar(item.get(key)) for key in keys] for item in value])
                for row_index, item in enumerate(value, 1):
                    nested = {key: nested_value for key, nested_value in item.items() if isinstance(nested_value, (dict, list))}
                    if nested:
                        render_structured(document, nested, f"Item {row_index} details", level + 1)
            else:
                for row_index, item in enumerate(value, 1):
                    render_structured(document, item, f"Item {row_index}", level + 1)
        else:
            for item in value:
                paragraph = document.add_paragraph(style="List Bullet")
                paragraph.add_run(_scalar(item))
    else:
        document.add_paragraph(_scalar(value))


def _output_base(source: Path) -> Path:
    try:
        relative = source.resolve().relative_to(REPO.resolve())
    except ValueError:
        relative = Path("external-inputs") / source.name
    return OUTPUT_ROOT / relative.parent / source.stem


def render_docx(source: SourceContent) -> Path:
    document = Document()
    apply_kutumba_styles(document)
    add_cover(document, source.title, source.subtitle, source.scope, status=source.status)
    add_header_footer(document, source.scope, source.audience)
    if (source.markdown and source.markdown.count("\n") > 120) or (
        source.structured is not None and len(json.dumps(source.structured, ensure_ascii=False, default=str)) > 12000
    ):
        add_toc_placeholder(document)
    if source.metadata:
        visible_meta = []
        for key in ("week_code", "week_title", "component", "source_id", "source_hash", "derived_from"):
            if key in source.metadata:
                value = source.metadata[key]
                visible_meta.append(
                    (
                        _humanize(key),
                        json.dumps(value, ensure_ascii=False, default=str)
                        if isinstance(value, dict)
                        else value,
                    )
                )
        if visible_meta:
            add_branded_table(document, ["Document metadata", "Value"], visible_meta)
    if source.markdown is not None:
        render_markdown(document, source.markdown, source.title)
    else:
        render_structured(document, source.structured)
    add_callout(
        document,
        "RIGHTS_NOTE",
        "KUTUMBA-original framing and teaching aids are identified as such. "
        "Scriptural and third-party sources remain subject to their respective rights; "
        "this document does not claim BBT or ISKCON ownership.",
    )
    target = _output_base(source.path).with_suffix(".docx")
    target.parent.mkdir(parents=True, exist_ok=True)
    document.save(target)
    return target


def _docx_opens(path: Path) -> bool:
    if not path.exists() or path.stat().st_size == 0 or not zipfile.is_zipfile(path):
        return False
    try:
        checked = Document(path)
        return len(checked.sections) > 0
    except Exception:
        return False


def convert_with_word(docx_path: Path, pdf_path: Path) -> int:
    """Convert with installed Microsoft Word and return Word's page count."""
    if not WORD_EXE.exists():
        raise FileNotFoundError(f"Microsoft Word executable not found at {WORD_EXE}")
    import comtypes.client

    word = None
    opened = None
    try:
        word = comtypes.client.CreateObject("Word.Application")
        word.Visible = False
        word.DisplayAlerts = 0
        opened = word.Documents.Open(str(docx_path.resolve()), ReadOnly=True, AddToRecentFiles=False)
        for section in opened.Sections:
            for field in section.Headers(1).Range.Fields:
                field.Update()
            for field in section.Footers(1).Range.Fields:
                field.Update()
        for field in opened.Fields:
            field.Update()
        opened.Repaginate()
        # wdStatisticPages=2; wdExportFormatPDF=17.
        page_count = int(opened.ComputeStatistics(2))
        pdf_path.parent.mkdir(parents=True, exist_ok=True)
        opened.ExportAsFixedFormat(str(pdf_path.resolve()), 17)
        return max(1, page_count)
    finally:
        if opened is not None:
            opened.Close(False)
        if word is not None:
            word.Quit()


def _reportlab_font() -> str:
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    candidates = [
        ("Nirmala", Path(r"C:\Windows\Fonts\Nirmala.ttf")),
        ("Mangal", Path(r"C:\Windows\Fonts\mangal.ttf")),
        ("SegoeUI", Path(r"C:\Windows\Fonts\segoeui.ttf")),
    ]
    for name, path in candidates:
        if path.exists():
            try:
                pdfmetrics.registerFont(TTFont(name, str(path)))
                return name
            except Exception:
                continue
    return "Helvetica"


def _plain_source_text(source: SourceContent) -> str:
    if source.markdown is not None:
        text = re.sub(r"^---.*?---\s*", "", source.markdown, flags=re.DOTALL)
        text = re.sub(r"!\[([^\]]*)\]\([^)]+\)", r"\1", text)
        text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
        text = re.sub(r"[#>*_`|]", " ", text)
        return re.sub(r"\s+", " ", text).strip()
    return json.dumps(source.structured, ensure_ascii=False, indent=2, default=str)


def convert_with_reportlab(source: SourceContent, pdf_path: Path) -> int:
    """Create a branded multipage content summary when Word automation fails."""
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import LETTER
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer

    font = _reportlab_font()
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "KutumbaTitle", parent=styles["Title"], fontName=font, fontSize=28,
        leading=33, textColor=colors.HexColor(f"#{PLUM}"), alignment=TA_CENTER,
    )
    subtitle_style = ParagraphStyle(
        "KutumbaSubtitle", parent=styles["Normal"], fontName=font, fontSize=13,
        leading=18, textColor=colors.HexColor(f"#{TEAL}"), alignment=TA_CENTER,
    )
    body_style = ParagraphStyle(
        "KutumbaBody", parent=styles["BodyText"], fontName=font, fontSize=10,
        leading=14, textColor=colors.HexColor(f"#{CHARCOAL}"), spaceAfter=7,
    )
    pdf_path.parent.mkdir(parents=True, exist_ok=True)

    def decorate(canvas, doc) -> None:
        canvas.saveState()
        width, height = LETTER
        canvas.setFillColor(colors.HexColor(f"#{CREAM}"))
        canvas.rect(0, 0, width, height, stroke=0, fill=1)
        canvas.setStrokeColor(colors.HexColor(f"#{SAFFRON}"))
        canvas.line(0.7 * inch, height - 0.48 * inch, width - 0.7 * inch, height - 0.48 * inch)
        canvas.setFillColor(colors.HexColor(f"#{PLUM}"))
        canvas.setFont(font, 7)
        canvas.drawString(0.7 * inch, 0.36 * inch, f"KUTUMBA • {TAGLINE}")
        canvas.drawRightString(width - 0.7 * inch, 0.36 * inch, f"Page {doc.page}")
        canvas.restoreState()

    story: list[Any] = [
        Spacer(1, 1.2 * inch),
        Paragraph(html.escape(source.title), title_style),
        Spacer(1, 0.2 * inch),
        Paragraph(html.escape(source.subtitle), subtitle_style),
        Spacer(1, 0.25 * inch),
        Paragraph(html.escape(f"{source.scope} • {source.audience}"), subtitle_style),
        Spacer(1, 0.35 * inch),
        Paragraph(html.escape(source.status), subtitle_style),
        PageBreak(),
        Paragraph("Content summary", title_style),
        Spacer(1, 0.2 * inch),
    ]
    plain = _plain_source_text(source)
    chunks = re.split(r"(?<=[.!?])\s+|\n{2,}", plain)
    buffer = ""
    for chunk in chunks:
        if len(buffer) + len(chunk) < 900:
            buffer = f"{buffer} {chunk}".strip()
        else:
            story.append(Paragraph(html.escape(buffer), body_style))
            buffer = chunk
    if buffer:
        story.append(Paragraph(html.escape(buffer), body_style))
    story.extend(
        [
            Spacer(1, 0.2 * inch),
            Paragraph(
                html.escape(
                    "Rights note: KUTUMBA-original framing; cited scriptural and "
                    "third-party material retains its own rights. No BBT or ISKCON "
                    "ownership is claimed."
                ),
                body_style,
            ),
        ]
    )
    doc = SimpleDocTemplate(
        str(pdf_path),
        pagesize=LETTER,
        leftMargin=0.75 * inch,
        rightMargin=0.75 * inch,
        topMargin=0.7 * inch,
        bottomMargin=0.65 * inch,
        title=source.title,
        author="KUTUMBA — Program Director: Swapnil Patil",
    )
    doc.build(story, onFirstPage=decorate, onLaterPages=decorate)
    with fitz.open(pdf_path) as opened:
        return opened.page_count


def rasterize_pdf(pdf_path: Path, source: Path) -> tuple[int, Path]:
    """Open PDF and rasterize every page under the required QA directory."""
    relative_name = re.sub(r"[^A-Za-z0-9._-]+", "-", _relative(source)).strip("-")
    output_dir = QA_ROOT / relative_name
    if output_dir.exists():
        shutil.rmtree(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    with fitz.open(pdf_path) as opened:
        count = opened.page_count
        if count <= 0:
            raise ValueError("PDF contains no pages")
        matrix = fitz.Matrix(1.5, 1.5)
        for number, page in enumerate(opened, 1):
            pixmap = page.get_pixmap(matrix=matrix, alpha=False)
            pixmap.save(output_dir / f"page-{number:03d}.png")
    return count, output_dir / "page-001.png"


def render_one(path: Path) -> RenderResult:
    source = load_source(path)
    docx = render_docx(source)
    pdf = docx.with_suffix(".pdf")
    result = RenderResult(source=path, docx=docx, pdf=pdf)
    result.docx_open = _docx_opens(docx)
    if not result.docx_open:
        result.error = "Generated DOCX did not reopen successfully."
        return result
    try:
        result.docx_pages = convert_with_word(docx, pdf)
        result.conversion = "Word COM"
    except Exception as exc:
        result.conversion = "reportlab fallback"
        result.error = f"Word COM unavailable/failed: {type(exc).__name__}: {exc}"
        result.docx_pages = convert_with_reportlab(source, pdf)
    try:
        result.pdf_pages, result.png = rasterize_pdf(pdf, path)
        result.pdf_open = True
        if result.docx_pages <= 0:
            result.docx_pages = result.pdf_pages
    except Exception as exc:
        result.error = f"{result.error}; PDF QA failed: {type(exc).__name__}: {exc}".strip("; ")
    return result


def discover_all() -> list[Path]:
    sources = [
        path for path in WEEKLY.rglob("*.md")
        if path.name in PUBLICATION_NAMES and "myDropbox" not in path.parts
    ]
    return sorted(sources, key=lambda item: item.as_posix().casefold())


def discover_sample() -> list[Path]:
    preferred = [
        WEEKLY / "c1-w1-what-is-kutumba-and-why-are-we-here" / "facilitator-guide.md",
        WEEKLY / "c1-w1-what-is-kutumba-and-why-are-we-here" / "parent-lesson.md",
        WEEKLY / "c1-w1-what-is-kutumba-and-why-are-we-here" / "children" / "lesson.md",
    ]
    existing = [path for path in preferred if path.exists()]
    if existing:
        return existing
    return discover_all()[:3]


def _qa_row(result: RenderResult, kind: str) -> str:
    if kind == "docx":
        target, pages, opened, passed = result.docx, result.docx_pages, result.docx_open, result.docx_pass
        sample = _relative(result.png) if result.png else "missing"
    else:
        target, pages, opened, passed = result.pdf, result.pdf_pages, result.pdf_open, result.pdf_pass
        sample = _relative(result.png) if result.png else "missing"
    size = target.stat().st_size if target.exists() else 0
    note = result.conversion
    if result.error:
        note = f"{note}; {result.error}".strip("; ")
    safe_note = note.replace("|", "/").replace("\n", " ")
    return (
        f"| `{_relative(result.source)}` | `{_relative(target)}` | {pages} | "
        f"{size} | {'yes' if opened else 'no'} | `{sample}` | "
        f"**{'PASS' if passed else 'FAIL'}** | {safe_note} |"
    )


def write_qa_reports(results: Sequence[RenderResult], command: str) -> None:
    generated = datetime.now(timezone.utc).isoformat(timespec="seconds")
    docx_pass = all(item.docx_pass for item in results) and bool(results)
    pdf_pass = all(item.pdf_pass for item in results) and bool(results)
    common = (
        f"- Generated: `{generated}`\n"
        f"- Command: `{command}`\n"
        f"- Files assessed: **{len(results)}**\n"
        "- Automated rendering QA only; this is not doctrinal, rights, safeguarding, "
        "pedagogy, temple, or publication approval.\n\n"
    )
    DOCX_QA.write_text(
        "# V12 DOCX Render QA\n\n"
        + f"## Overall: {'PASS' if docx_pass else 'FAIL'}\n\n"
        + common
        + "| Source | DOCX | Pages | Bytes | Opens | Sample raster | Result | Notes |\n"
        + "|---|---|---:|---:|---|---|---|---|\n"
        + "\n".join(_qa_row(item, "docx") for item in results)
        + "\n",
        encoding="utf-8",
    )
    PDF_QA.write_text(
        "# V12 PDF Render QA\n\n"
        + f"## Overall: {'PASS' if pdf_pass else 'FAIL'}\n\n"
        + common
        + "| Source | PDF | Pages | Bytes | Opens | Sample raster | Result | Notes |\n"
        + "|---|---|---:|---:|---|---|---|---|\n"
        + "\n".join(_qa_row(item, "pdf") for item in results)
        + "\n",
        encoding="utf-8",
    )


def parse_args(argv: Sequence[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    mode = parser.add_mutually_exclusive_group(required=True)
    mode.add_argument("--all", action="store_true", help="Render all curated first-six-month publication sources.")
    mode.add_argument("--sample", action="store_true", help="Render a small representative source set.")
    mode.add_argument(
        "--input",
        action="append",
        type=Path,
        metavar="PATH",
        help="Render one Markdown, YAML, or JSON file; repeat for multiple files.",
    )
    return parser.parse_args(argv)


def main(argv: Sequence[str] | None = None) -> int:
    args = parse_args(argv)
    if args.all:
        paths = discover_all()
        command = "python scripts/v12/render_publication_docs.py --all"
    elif args.sample:
        paths = discover_sample()
        command = "python scripts/v12/render_publication_docs.py --sample"
    else:
        paths = [(path if path.is_absolute() else REPO / path).resolve() for path in args.input]
        command = "python scripts/v12/render_publication_docs.py --input …"
    paths = [path for path in paths if path.exists() and "myDropbox" not in path.parts]
    if not paths:
        print("No eligible source files found.", file=sys.stderr)
        return 2

    OUTPUT_ROOT.mkdir(parents=True, exist_ok=True)
    QA_ROOT.mkdir(parents=True, exist_ok=True)
    results: list[RenderResult] = []
    for number, path in enumerate(paths, 1):
        print(f"[{number}/{len(paths)}] Rendering {_relative(path)}")
        try:
            result = render_one(path)
        except Exception as exc:
            base = _output_base(path)
            result = RenderResult(
                source=path,
                docx=base.with_suffix(".docx"),
                pdf=base.with_suffix(".pdf"),
                error=f"{type(exc).__name__}: {exc}",
            )
            traceback.print_exc()
        results.append(result)
        print(
            f"  DOCX={'PASS' if result.docx_pass else 'FAIL'} "
            f"PDF={'PASS' if result.pdf_pass else 'FAIL'} "
            f"pages={result.pdf_pages} via {result.conversion or 'none'}"
        )
    write_qa_reports(results, command)
    passed = all(item.docx_pass and item.pdf_pass for item in results)
    print(f"Wrote {_relative(DOCX_QA)} and {_relative(PDF_QA)}")
    return 0 if passed else 1


if __name__ == "__main__":
    raise SystemExit(main())
