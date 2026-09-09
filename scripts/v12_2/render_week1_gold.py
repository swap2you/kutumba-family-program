#!/usr/bin/env python3
"""Build the seven KUTUMBA C1-W1 V12.2 gold-standard DOCX/PDF packets."""
from __future__ import annotations

import hashlib
import json
import re
import sys
import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import Callable, Iterable

import pymupdf as fitz
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

REPO = Path(__file__).resolve().parents[2]
V12_SCRIPTS = REPO / "scripts" / "v12"
for import_path in (REPO, V12_SCRIPTS):
    if str(import_path) not in sys.path:
        sys.path.insert(0, str(import_path))

from kutumba_docx_styles import (  # noqa: E402
    BRAND,
    CHARCOAL,
    CREAM,
    GOLD,
    MAROON,
    PLUM,
    SAFFRON,
    TEAL,
    add_branded_table,
    add_callout,
    add_cover,
    add_header_footer,
    add_verse_card,
    apply_kutumba_styles,
)
from render_publication_docs import convert_with_word, render_markdown  # noqa: E402

BASE = (
    REPO
    / "11-weekly-program-library"
    / "first-six-months"
    / "c1-w1-what-is-kutumba-and-why-are-we-here"
)
LAUNCH = REPO / "launch"
OUTPUT = REPO / "exports" / "final" / "week1"

SCHEDULE = [
    ("1:50–2:00", "Arrival"),
    ("2:00–2:10", "Opening mantras + welcome"),
    ("2:10–2:30", "All-family launch/orientation"),
    ("2:30–3:10", "Parallel tracks"),
    ("3:10–3:30", "Reunification + bhakti laboratory"),
    ("3:30–3:40", "Snack + water"),
    ("3:40–3:55", "Saṅkalpa + Cycle 1 project + home practice"),
    ("3:55–4:00", "Close"),
]
FORBIDDEN_TIMES = ("2:" + "35", "3:" + "25")
STATUS = (
    "Internal founding-cohort teaching material — "
    "human/temple review EXTERNAL_OPEN"
)


@dataclass(frozen=True)
class Packet:
    filename: str
    title: str
    subtitle: str
    audience: str
    build_body: Callable[[Document], None]


def _relative(path: Path) -> str:
    return path.resolve().relative_to(REPO.resolve()).as_posix()


def _require(path: Path) -> Path:
    if not path.is_file():
        raise FileNotFoundError(f"Required Week-1 source is missing: {_relative(path)}")
    return path


def _source_text(path: Path) -> str:
    text = _require(path).read_text(encoding="utf-8-sig")
    # Legacy negative-control prose is not useful in a participant packet and
    # must never leak obsolete clock values into a rendered deliverable.
    for token in FORBIDDEN_TIMES:
        text = "\n".join(line for line in text.splitlines() if token not in line)
    return text


def _markdown_title(text: str, fallback: str) -> str:
    match = re.search(r"^#\s+(.+?)\s*$", text, re.MULTILINE)
    return match.group(1).strip() if match else fallback


def _source_trace(document: Document, path: Path) -> None:
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    paragraph = document.add_paragraph(style="KUTUMBA Rights")
    paragraph.add_run(
        f"Canonical source: {_relative(path)}  •  SHA-256: {digest}"
    )


def add_source(document: Document, path: Path, section_title: str | None = None) -> None:
    text = _source_text(path)
    source_title = _markdown_title(text, path.stem.replace("-", " ").title())
    document.add_heading(section_title or source_title, level=1)
    _source_trace(document, path)
    render_markdown(document, text, source_title)


def page_then_source(
    document: Document, path: Path, section_title: str | None = None
) -> None:
    document.add_page_break()
    add_source(document, path, section_title)


def add_schedule(document: Document, heading: str = "Locked Saturday Schedule") -> None:
    document.add_heading(heading, level=1)
    add_branded_table(document, ["Time", "Block"], SCHEDULE)
    add_callout(
        document,
        "TIME_CUE",
        "Track handoff is at 2:30; all groups reunite at 3:10.",
    )


def add_generated_section(document: Document, title: str, body: str) -> None:
    document.add_heading(title, level=1)
    paragraph = document.add_paragraph(style="KUTUMBA Rights")
    paragraph.add_run("Source ID: KUTUMBA-V12.2-W1-GENERATED")
    render_markdown(document, body.strip(), title)


def _configure_letter(document: Document) -> None:
    for section in document.sections:
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)


def _new_document(title: str, subtitle: str, audience: str) -> Document:
    document = Document()
    apply_kutumba_styles(document)
    add_cover(
        document,
        title,
        subtitle,
        "C1-W1 • What Is KUTUMBA, and Why Are We Here?",
        version="V12.2",
        status=STATUS,
    )
    add_header_footer(document, "C1-W1 • V12.2", audience)
    _configure_letter(document)
    return document


def _glob_required(directory: Path, pattern: str) -> list[Path]:
    paths = sorted(directory.glob(pattern), key=lambda item: item.name.casefold())
    if not paths:
        raise FileNotFoundError(
            f"No sources match {_relative(directory)}/{pattern}"
        )
    return paths


def build_start_here(document: Document) -> None:
    add_schedule(document)
    page_then_source(document, REPO / "V12_2-WEEK1-START-HERE.md", "Operational Entry Point")


def build_owner_runbook(document: Document) -> None:
    add_schedule(document)
    page_then_source(document, BASE / "OWNER-RUNBOOK-V12.2.md", "Owner Runbook")


def build_orientation(document: Document) -> None:
    add_schedule(document)
    sources = [
        (LAUNCH / "KUTUMBA-C1-FAMILY-ORIENTATION.md", "Family Orientation"),
        (LAUNCH / "FAMILY-COVENANT.md", "Family Covenant"),
        (
            LAUNCH / "FAMILY-COVENANT-ACKNOWLEDGEMENT-TEMPLATE.md",
            "Blank Covenant Acknowledgement",
        ),
        (LAUNCH / "CHILD-HOUSE-RULES.md", "Child House Rules"),
    ]
    for path, title in sources:
        page_then_source(document, path, title)


def build_track(
    document: Document, guide: Path, activities: Iterable[Path], extra: Path | None = None
) -> None:
    add_schedule(document)
    page_then_source(document, guide, "Teacher / Facilitator Run Sheet")
    if extra:
        page_then_source(document, extra, "Child House Rules")
    for path in activities:
        page_then_source(document, path)


def _teacher_divider(document: Document) -> None:
    document.add_page_break()
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Inches(2.4)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("TEACHER-ONLY")
    run.bold = True
    run.font.size = Pt(32)
    run.font.color.rgb = __import__("docx").shared.RGBColor.from_string(PLUM)
    note = document.add_paragraph(
        "Remove this section before distributing participant copies."
    )
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER


def _add_checklist(document: Document) -> None:
    document.add_heading("Room / Material / Snack / Cleanup Checklist", level=1)
    rows = [
        ("☐", "Room", "Parent circle, younger floor zone, older table zone, clear aisles"),
        ("☐", "Materials", "Three labeled track folders; pencils; crayons; tape; timer"),
        ("☐", "Forms", "Blank acknowledgements only; private completed forms never enter Git"),
        ("☐", "Snack", "Light snack and water staged for 3:30; allergy plan confirmed"),
        ("☐", "Safety", "Calm corner visible; private rooms closed; exits unobstructed"),
        ("☐", "Cleanup", "Recycle/trash bins ready; tables wiped; host home reset assigned"),
    ]
    add_branded_table(document, ["Done", "Area", "Check"], rows)


def build_saturday(document: Document) -> None:
    # 1. Cover is created by _new_document; quick schedule follows it.
    add_schedule(document, "Quick Schedule")

    ordered_sources: list[tuple[Path, str]] = [
        (LAUNCH / "OPENING-MANTRAS-HANDOUT.md", "2. Opening Mantras"),
        (LAUNCH / "KUTUMBA-C1-FAMILY-ORIENTATION.md", "3. Family Orientation"),
        (LAUNCH / "FIRST-SIX-MONTHS-CALENDAR.md", "4. First-Six-Month Roadmap"),
        (LAUNCH / "FAMILY-COVENANT.md", "5. Family Covenant Summary"),
        (
            LAUNCH / "FAMILY-COVENANT-ACKNOWLEDGEMENT-TEMPLATE.md",
            "6. Blank Acknowledgement",
        ),
        (BASE / "teacher" / "PARENT-GUIDE.md", "7. Parent Track Run Sheet"),
    ]
    for path, title in ordered_sources:
        page_then_source(document, path, title)
    for path in _glob_required(BASE / "activities" / "v12_2-parent", "P0*.md"):
        page_then_source(document, path)

    page_then_source(document, LAUNCH / "CHILD-HOUSE-RULES.md", "8. Younger Child Rules")
    younger = BASE / "activities" / "v12_2-younger"
    for code, title in [
        ("Y01-*.md", "9. Younger Four Corners Signs"),
        ("Y02-*.md", "9. Younger Four Corners Cards"),
        ("Y03-*.md", "10. Younger Bhakti Garden"),
        ("Y04-*.md", "11. Younger Memory Card"),
        ("Y05-*.md", "11b. Younger House-Rule Sort"),
    ]:
        page_then_source(document, _glob_required(younger, code)[0], title)

    page_then_source(document, LAUNCH / "CHILD-HOUSE-RULES.md", "12. Older Child Rules")
    older = BASE / "activities" / "v12_2-older"
    for number, code, title in [
        (13, "O01-*.md", "Older Verse Observation"),
        (14, "O02-*.md", "Older Is / Is Not"),
        (15, "O03-*.md", "Older Six-Purpose Challenge"),
        (16, "O04-*.md", "Older Scenario Challenge"),
        (17, "O05-*.md", "Older Family Compass"),
        (18, "O06-*.md", "Older Exit Ticket"),
    ]:
        page_then_source(
            document, _glob_required(older, code)[0], f"{number}. {title}"
        )

    page_then_source(
        document,
        BASE / "activities" / "v12_2-parent" / "P05-FAMILY-SANKALPA-BUILDER.md",
        "19. Family Saṅkalpa Builder",
    )
    page_then_source(
        document,
        BASE / "project" / "MODULE-PROJECT-BRIEF.md",
        "20. Cycle 1 Project Introduction",
    )
    document.add_page_break()
    add_generated_section(
        document,
        "21. Week-1 Home Practice",
        """
Minimum: **5 minutes**, at least once before next Saturday.

1. Offer one prayer.
2. Chant three mahā-mantras.
3. Each person offers one appreciation.

Standard rhythm: twice during the week. No photo proof. No ranking.
""",
    )

    _teacher_divider(document)  # 22
    page_then_source(
        document,
        BASE / "teacher" / "YOUNGER-TEACHER-GUIDE.md",
        "23. Younger Teacher Run Sheet",
    )
    page_then_source(
        document,
        BASE / "teacher" / "OLDER-TEACHER-GUIDE.md",
        "24. Older Teacher Run Sheet",
    )
    document.add_page_break()
    document.add_heading("25. Older Answer Keys", level=1)
    for code in ("O02-*.md", "O03-*.md", "O04-*.md"):
        add_source(document, _glob_required(older, code)[0])
    answer_key = BASE / "activities" / "OLDER-ANSWER-KEY.md"
    if answer_key.exists():
        add_source(document, answer_key, "Additional Older Answer Key")

    document.add_page_break()
    add_generated_section(
        document,
        "26. Parent Facilitator Answers / Prompts",
        """
- Protect private writing; never force sharing.
- In the two-family case, affirm what is healthy before naming what is missing.
- Balanced repair: protect Saturday association and a tiny home practice without shame.
- Saṅkalpa formula: specific action + frequency + trigger + minimum version.
- Keep completed family agreements private; do not collect them into Git.
""",
    )
    page_then_source(
        document,
        BASE / "teacher" / "MAIN-FACILITATOR-GUIDE-V12.md",
        "27. Main Facilitator Speaking Map",
    )
    document.add_page_break()
    _add_checklist(document)  # 28
    document.add_page_break()
    add_generated_section(
        document,
        "29. Rights / Privacy Note",
        """
Status: **EXTERNAL_OPEN**. Human and temple review gates remain open.

Do not apply false approval labels or imply official ISKCON, temple, or BBT endorsement.
KUTUMBA-original framing is identified as such; cited third-party material retains its
own rights. Keep completed acknowledgements, reflections, and family commitments
private. Do not store private family data, photographs, or proof-of-practice in Git.
""",
    )


def _packet_specs() -> list[Packet]:
    parent_activities = _glob_required(BASE / "activities" / "v12_2-parent", "P0*.md")
    younger_activities = _glob_required(BASE / "activities" / "v12_2-younger", "Y0*.md")
    older_activities = _glob_required(BASE / "activities" / "v12_2-older", "O0*.md")
    return [
        Packet(
            "KUTUMBA-C1-W1-START-HERE-V12.2",
            "Week 1 — Start Here",
            "Owner-facing operational entry point",
            "Program owner",
            build_start_here,
        ),
        Packet(
            "KUTUMBA-C1-W1-OWNER-RUNBOOK-V12.2",
            "Week 1 — Owner Runbook",
            "Friday preparation and Saturday execution",
            "Program owner",
            build_owner_runbook,
        ),
        Packet(
            "KUTUMBA-C1-W1-FAMILY-ORIENTATION-AND-COVENANT-V12.2",
            "Family Orientation and Covenant",
            "Welcome, commitments, acknowledgement, and house rules",
            "Families",
            build_orientation,
        ),
        Packet(
            "KUTUMBA-C1-W1-PARENT-TRACK-V12.2",
            "Parent Track",
            "Facilitator guide and complete parent printables",
            "Parent facilitator",
            lambda doc: build_track(
                doc, BASE / "teacher" / "PARENT-GUIDE.md", parent_activities
            ),
        ),
        Packet(
            "KUTUMBA-C1-W1-YOUNGER-TEACHER-PACK-V12.2",
            "Younger Teacher Pack",
            "K–2 run sheet, child rules, and complete activities",
            "Younger teacher",
            lambda doc: build_track(
                doc,
                BASE / "teacher" / "YOUNGER-TEACHER-GUIDE.md",
                younger_activities,
                LAUNCH / "CHILD-HOUSE-RULES.md",
            ),
        ),
        Packet(
            "KUTUMBA-C1-W1-OLDER-TEACHER-PACK-V12.2",
            "Older Teacher Pack",
            "Grades 4–5 run sheet and complete activities",
            "Older teacher",
            lambda doc: build_track(
                doc, BASE / "teacher" / "OLDER-TEACHER-GUIDE.md", older_activities
            ),
        ),
        Packet(
            "KUTUMBA-C1-W1-SATURDAY-PRINT-PACKET-V12.2",
            "Saturday Print Packet",
            "Complete ordered participant and teacher packet",
            "Families and facilitators",
            build_saturday,
        ),
    ]


def _docx_text(path: Path) -> str:
    with zipfile.ZipFile(path) as archive:
        xml = archive.read("word/document.xml")
    from lxml import etree

    root = etree.fromstring(xml)
    return "\n".join(root.xpath("//w:t/text()", namespaces={"w": qn("w:t").split("}")[0][1:]}))


def _strip_trailing_empty_pdf_pages(pdf_path: Path) -> int:
    """Remove trailing pages that contain only header/footer chrome."""
    import os

    doc = fitz.open(pdf_path)
    removed = 0
    while doc.page_count > 1:
        page = doc[-1]
        blocks = page.get_text("blocks")
        body = " ".join(
            block[4] for block in blocks if block[1] > 60 and block[3] < 740
        ).strip()
        if body:
            break
        doc.delete_page(-1)
        removed += 1
    if removed:
        tmp = pdf_path.with_suffix(".tmp.pdf")
        doc.save(tmp, deflate=True)
        doc.close()
        os.replace(tmp, pdf_path)
    else:
        doc.close()
    return removed


def _validate_output(docx_path: Path, pdf_path: Path) -> int:
    _strip_trailing_empty_pdf_pages(pdf_path)
    docx_text = _docx_text(docx_path)
    with fitz.open(pdf_path) as pdf:
        if pdf.page_count < 1:
            raise RuntimeError(f"PDF has no pages: {pdf_path}")
        pdf_text = "\n".join(page.get_text() for page in pdf)
        page_count = pdf.page_count
    for required in ("2:30", "3:10"):
        if required not in docx_text or required not in pdf_text:
            raise RuntimeError(f"Locked time {required} missing from {pdf_path.name}")
    for forbidden in FORBIDDEN_TIMES:
        if forbidden in docx_text or forbidden in pdf_text:
            raise RuntimeError(
                f"Obsolete schedule time detected in rendered output {pdf_path.name}"
            )
    if "<w:tbl" not in zipfile.ZipFile(docx_path).read("word/document.xml").decode(
        "utf-8"
    ):
        raise RuntimeError(f"No real Word table found in {docx_path.name}")
    return page_count


def main() -> int:
    OUTPUT.mkdir(parents=True, exist_ok=True)
    results = []
    for packet in _packet_specs():
        docx_path = OUTPUT / f"{packet.filename}.docx"
        pdf_path = OUTPUT / f"{packet.filename}.pdf"
        document = _new_document(packet.title, packet.subtitle, packet.audience)
        packet.build_body(document)
        add_callout(
            document,
            "RIGHTS_NOTE",
            "KUTUMBA-original framing; cited third-party material retains its own "
            "rights. Human/temple review remains EXTERNAL_OPEN.",
        )
        _configure_letter(document)
        document.save(docx_path)
        if pdf_path.exists():
            pdf_path.unlink()
        word_pages = convert_with_word(docx_path, pdf_path)
        if not pdf_path.is_file() or pdf_path.stat().st_size == 0:
            raise RuntimeError(f"Word did not create a usable PDF: {pdf_path}")
        pdf_pages = _validate_output(docx_path, pdf_path)
        result = {
            "name": packet.filename,
            "docx": _relative(docx_path),
            "docx_bytes": docx_path.stat().st_size,
            "pdf": _relative(pdf_path),
            "pdf_bytes": pdf_path.stat().st_size,
            "pages": pdf_pages,
            "word_reported_pages": word_pages,
        }
        results.append(result)
        print(
            f"{packet.filename}: {pdf_pages} pages; "
            f"DOCX {result['docx_bytes']:,} bytes; PDF {result['pdf_bytes']:,} bytes"
        )
    manifest = OUTPUT / "week1-render-manifest.json"
    manifest.write_text(
        json.dumps({"version": "V12.2", "packets": results}, indent=2, ensure_ascii=False)
        + "\n",
        encoding="utf-8",
    )
    print(f"Wrote {_relative(manifest)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
