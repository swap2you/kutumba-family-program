#!/usr/bin/env python3
"""V11 Cycle 1 production generator — launch packs, research, projects, guides, gamma, visuals, DOCX."""
from __future__ import annotations

import textwrap
from pathlib import Path

from docx import Document
from docx.shared import Pt

REPO = Path(__file__).resolve().parents[2]
WEEKLY = REPO / "11-weekly-program-library" / "first-six-months"
LAUNCH = REPO / "launch"

WEEKS = [
    {
        "slug": "c1-w1-what-is-kutumba-and-why-are-we-here",
        "code": "C1-W1",
        "title": "What Is KUTUMBA, and Why Are We Here?",
        "conclusion": "Regular family hearing and practice creates a protected path for growth.",
        "verse": "Śrīmad-Bhāgavatam 1.2.18",
        "verse_url": "https://vedabase.io/en/library/sb/1/2/18/",
        "memory": "By regular Bhāgavata hearing and service, the heart becomes steadied in devotion.",
        "question": "Why are we committing as a family?",
        "project": "Family purpose / first saṅkalpa card",
        "science_theme": "family routines, habit cues, repeated practice, parent modeling",
        "misconception": "Attendance without home practice equals growth",
    },
    {
        "slug": "c1-w2-i-am-not-this-body",
        "code": "C1-W2",
        "title": "I Am Not This Body",
        "conclusion": "The body changes; the conscious self continues.",
        "verse": "Bhagavad-gītā 2.13",
        "verse_url": "https://vedabase.io/en/library/bg/2/13/",
        "memory": "As the embodied soul continuously passes through childhood, youth and old age, the soul similarly passes into another body.",
        "question": "How does identity change how we speak about bodies?",
        "project": "Changing-body / enduring-self artifact",
        "science_theme": "developmental continuity and respectful body language — never proof of ātman",
        "misconception": "Psychology proves the soul",
    },
    {
        "slug": "c1-w3-the-nature-of-the-soul",
        "code": "C1-W3",
        "title": "The Nature of the Soul",
        "conclusion": "The jīva is eternal, conscious, individual, minute, and related to Kṛṣṇa in service.",
        "verse": "Bhagavad-gītā 2.20",
        "verse_url": "https://vedabase.io/en/library/bg/2/20/",
        "memory": "The soul is never born and never dies; it is eternal and unchanging.",
        "question": "If I am a soul, how should I live?",
        "project": "Positive nature-of-soul artifact",
        "science_theme": "limits of third-person measurement only with careful caveats — no proof science",
        "misconception": "All souls are God / we are the Supreme",
    },
    {
        "slug": "c1-w4-why-human-life-is-rare-and-valuable",
        "code": "C1-W4",
        "title": "Why Human Life Is Rare and Valuable",
        "conclusion": "Human life gives a rare opportunity for deliberate self-realization.",
        "verse": "Śrīmad-Bhāgavatam 11.9.29",
        "verse_url": "https://vedabase.io/en/library/sb/11/9/29/",
        "memory": "After many births one attains the rare human form — use it for self-realization.",
        "question": "What deserves protected family time?",
        "project": "Priority / time-use artifact",
        "science_theme": "time allocation, goal salience, values-based action",
        "misconception": "Use fear or death-pressure to motivate children",
    },
    {
        "slug": "c1-w5-the-temporary-world-and-the-search-for-permanent-happiness",
        "code": "C1-W5",
        "title": "The Temporary World and the Search for Permanent Happiness",
        "conclusion": "Temporary things can be used well but cannot provide permanent fulfillment.",
        "verse": "Bhagavad-gītā 5.22",
        "verse_url": "https://vedabase.io/en/library/bg/5/22/",
        "memory": "Pleasures born of contact have a beginning and end; the wise do not seek fulfillment in them alone.",
        "question": "How can enjoyment become gratitude and service?",
        "project": "Temporary vs lasting happiness application artifact",
        "science_theme": "hedonic adaptation and gratitude/prosocial research — carefully framed",
        "misconception": "Material things are worthless / family affection is meaningless",
    },
    {
        "slug": "c1-w6-integration-night-who-am-i-and-how-should-our-family-live",
        "code": "C1-W6",
        "title": "Integration Night: Who Am I, and How Should Our Family Live?",
        "conclusion": "Identity, purpose, and practice must form one coherent family life.",
        "verse": "Cycle 1 review (SB 1.2.18 · BG 2.13 · BG 2.20 · SB 11.9.29 · BG 5.22)",
        "verse_url": "https://vedabase.io/en/library/sb/1/2/18/",
        "memory": "We remember who we are and how our family chooses to live.",
        "question": "Can our family explain and apply what we learned?",
        "project": "10-minute family presentation",
        "science_theme": "retrieval practice, explanation-to-learn, formative assessment",
        "misconception": "Competition or ranking of families",
    },
]


def write(path: Path, content: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(content.strip() + "\n", encoding="utf-8")


def md_to_docx(md_path: Path, docx_path: Path, title: str) -> None:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    doc.add_heading(title, level=0)
    for line in md_path.read_text(encoding="utf-8").splitlines():
        if line.startswith("# "):
            continue
        if line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=1)
        elif line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=2)
        elif line.startswith("|") or line.startswith("- ") or line.startswith("> ") or line.startswith("1.") or line.startswith("**"):
            doc.add_paragraph(line)
        elif line.strip() == "":
            continue
        else:
            doc.add_paragraph(line)
    docx_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(docx_path)


def project_files(w: dict) -> None:
    base = WEEKLY / w["slug"] / "project"
    write(
        base / "MODULE-PROJECT-BRIEF.md",
        f"""# {w['code']} Module Project Brief

## Cycle 1 project title

**Who Am I, and How Should Our Family Live?**

## Module contribution

{w['code']} contributes: **{w['project']}**.

## Objective link

- Primary conclusion: {w['conclusion']}
- Family question: {w['question']}
- Key scripture: {w['verse']}

## Optional outputs (choose one)

| Output | Description | Band |
|---|---|---|
| Drawing / poster | Age-appropriate visual of this week's idea | Younger / Older |
| Short speech | 1–2 minutes shared at reunification or W6 | Older / Parents |
| Action artifact | Saṅkalpa, timeline, priority map, or gratitude card | All |
| Family interview | Two questions answered together at home | All |

## Burden

Normal weeks: ≤20 minutes optional project work. W6 uses a low-burden presentation option.

## Privacy

No required public disclosure of private struggles. No real names in shared gallery labels unless families consent privately outside Git.

## Sources

{w['verse']} — {w['verse_url']}
""",
    )
    write(
        base / "CYCLE-CONTRIBUTION.md",
        f"""# {w['code']} Cycle Contribution

## Cumulative project

**Who Am I, and How Should Our Family Live?**

## This week's layer

{w['project']}

## How it connects

| Prior | This week | Next |
|---|---|---|
| Previous C1 layers (if any) | {w['code']}: {w['project']} | Continues toward W6 presentation |

## Family action

1. Complete one artifact linked to: {w['question']}
2. Keep it simple and age-appropriate.
3. Bring or photograph (with consent) for the Cycle 1 gallery / W6 share.

## Facilitation note

Celebrate effort and understanding — never rank families or children.
""",
    )
    write(
        base / "PRESENTATION-RUBRIC.md",
        f"""# {w['code']} Presentation Rubric (Non-Competitive)

Use for formative encouragement. **No ranking. No public failure labels.**

| Dimension | Emerging | Developing | Secure |
|---|---|---|---|
| Understanding | Names topic vaguely | States week's conclusion in own words | Connects conclusion to {w['verse']} paraphrase |
| Application | No home link | One concrete family action | Clear saṅkalpa / habit cue |
| Teamwork | One person only | Partial family participation | Parents + children share roles |
| Source accuracy | Mixes weeks / invents | Stays in week scope | Distinguishes scripture vs analogy |

## W6 note

If understanding is weak across dimensions, recommend **C1 review/extension** before Cycle 2 — privately, with dignity.
""",
    )


def research_files(w: dict) -> None:
    base = WEEKLY / w["slug"] / "research"
    write(
        base / "SCRIPTURAL-EXAMPLES.md",
        f"""# {w['code']} Scriptural Examples

## Primary anchor

| Reference | URL | Use | Limitation |
|---|---|---|---|
| {w['verse']} | {w['verse_url']} | Core teaching paraphrase | No full purport dump in Git |

## Supporting (tier 1)

| Reference | URL | Use | Limitation |
|---|---|---|---|
| Related Prabhupāda book context | VedaBase library root https://vedabase.io/en/library/ | Clarify scope | Quote only short verified lines; prefer paraphrase |

## Teaching note

Present as KUTUMBA paraphrase unless an exact short quotation is verified from VedaBase. Never invent verse wording.
""",
    )
    write(
        base / "DEVOTIONAL-AND-HISTORICAL-EXAMPLES.md",
        f"""# {w['code']} Devotional and Historical Examples

## Allowed use

Authorized ISKCON / Prabhupāda lecture or conversation examples with date/place/URL when available.

## Example slots

| Example | Provenance | Use | Status |
|---|---|---|---|
| Regular hearing / practice anecdote tied to week's theme | Must have VedaBase or authorized lecture link | Illustration | TODO if link not verified |
| Temple-family cooperation example | Local authorized story only with permission | Illustration | Optional / pending |

## Omitted pending provenance

- Bhakta-mālā episodes: **omit** unless exact edition/provenance is available.
- Rāmāyaṇa dialogue: **omit invented dialogue**; use only traceable paraphrase with edition named.

## Labeling

Mark supplementary/hagiographic material clearly. Never use as sole doctrinal proof.
""",
    )
    write(
        base / "CASE-STUDIES.md",
        f"""# {w['code']} Case Studies

## Categories

1. Scriptural case (from śāstra narrative already in curriculum)
2. Historical/devotional example (provenance required)
3. Constructed contemporary family case (**fictional / anonymized**)

## Constructed teaching case A

**Label:** Constructed teaching case — not a real family.

- **Situation:** A family pattern related to: {w['question']}
- **Tempting mistaken conclusion:** {w['misconception']}
- **Relevant principle:** {w['conclusion']}
- **Compassionate response:** Acknowledge effort; avoid shame; return to one small practice.
- **Family action:** One 5–15 minute home practice linked to {w['verse']}.
- **What not to say:** Comparisons, spiritual threats, or public exposure of private struggles.

## Constructed teaching case B

**Label:** Constructed teaching case — not a real family.

- **Situation:** Screens / overload / comparison pressure intersects this week's theme.
- **Tempting mistaken conclusion:** More information without practice will fix everything.
- **Relevant principle:** Protected rhythm + humble application.
- **Compassionate response:** Simplify; one cue; one minimum version.
- **Family action:** Choose one habit cue for the week.
- **What not to say:** "Other families are more serious than you."
""",
    )
    write(
        base / "SCIENCE-AND-APPLICATION.md",
        f"""# {w['code']} Science and Application

## Scope

Theme: **{w['science_theme']}**

Science may support pedagogy and family application. **Science never proves** ātman, karma, rebirth, or Kṛṣṇa.

## Suggested empirical supports (cite before claiming)

| Topic | Example literature class | Use | Do not claim |
|---|---|---|---|
| Habit / cueing | Behavioral science reviews on implementation intentions | Home practice design | Proves devotion |
| Family routines | Developmental psychology on family rituals | Saturday rhythm | Proves doctrine |
| Retrieval practice | Cognitive science on testing effect | W6 review stations | Proves śāstra |
| Hedonic adaptation | Well-being research on material consumption | W5 framing only | Proves temporary world ontology |

## Citation rule

If a specific study is named in teaching, include DOI/URL and year. If not verified, keep the claim general ("research on habit cues suggests…") or omit.

## Boundary

Metaphysical conclusions remain grounded in śāstra / Prabhupāda teachings, not lab results.
""",
    )
    write(
        base / "ANALOGIES-AND-LIMITS.md",
        f"""# {w['code']} Analogies and Limits

## Week-scoped analogies

| Analogy | Teaches | Does not teach |
|---|---|---|
| Week-appropriate classroom analogy | {w['conclusion']} | Full ontology of other weeks |
| Garden / path / map metaphors as relevant | Protected practice rhythm | Guaranteed spiritual advancement |

## Hard limits

- Do not import other weeks' core analogies as if they were this week's doctrine.
- Label analogy as pedagogy, not as śāstra quotation.
- For W1: use hearing/practice/community analogies — **not** body/soul garment analogies (those belong to W2).
""",
    )


def teacher_guides(w: dict) -> None:
    base = WEEKLY / w["slug"]
    write(
        base / "teacher" / "YOUNGER-TEACHER-GUIDE.md",
        f"""# {w['code']} Younger Teacher Guide (K–2 / Lāla–Lālī adapt)

## Objective

Children taste: **{w['conclusion']}**

## Memory phrase

{w['memory']}

## 40-minute core

| Min | Block |
|---|---|
| 0–5 | Settle + memory phrase echo |
| 5–15 | Story / Prem-kī-Kathā listening |
| 15–25 | Movement + hands-on object |
| 25–33 | Coloring / matching / craft |
| 33–38 | Memory card + take-home cue |
| 38–40 | Cleanup + parent handoff |

## Materials

See `materials.md` + `activities/YOUNGER-ACTIVITY-PACK.md`.

## Boundaries

- No forced chanting volume.
- No public shame.
- Defer hard doctrine: use approved deferral line from `launch/TEACHER-READINESS-STANDARD.md`.

## Key source

{w['verse']} — {w['verse_url']} (paraphrase for K–2)
""",
    )
    write(
        base / "teacher" / "OLDER-TEACHER-GUIDE.md",
        f"""# {w['code']} Older Teacher Guide (Grades 4–5)

## Objective

Students can explain: **{w['conclusion']}**

## Essential question

{w['question']}

## 40-minute core

| Min | Block |
|---|---|
| 0–5 | Hook + essential question |
| 5–15 | Short primary text observation ({w['verse']}) |
| 15–25 | Diagram + case discussion |
| 25–33 | Worksheet / game |
| 33–38 | Application + project link |
| 38–40 | Cleanup + parent handoff |

## Answer keys

See `activities/OLDER-ANSWER-KEY.md`.

## Boundaries

- Do not teach full other-week ontology.
- Misconception to block: {w['misconception']}
- Science may illustrate habits only — never prove metaphysics.

## Key source

{w['verse']} — {w['verse_url']}
""",
    )
    write(
        base / "teacher" / "PRE-WEEK-CHECKLIST.md",
        f"""# {w['code']} Teacher Pre-Week Checklist

- [ ] Objective: {w['conclusion']}
- [ ] Verse open in browser: {w['verse_url']}
- [ ] Misconception boundary: {w['misconception']}
- [ ] Younger materials packed
- [ ] Older worksheet + answer key packed
- [ ] Low-prep backup game ready
- [ ] Reunification time known (≈3:10)
- [ ] No private family data in shared notes
""",
    )
    write(
        base / "teacher" / "MAIN-FACILITATOR-GUIDE-V11.md",
        f"""# {w['code']} Main Facilitator Guide — Saturday 2:00–4:00

## Two-minute summary

{w['conclusion']} Key scripture: {w['verse']}.

## Essential question

{w['question']}

## Locked flow

| Time | Block |
|---|---|
| 1:50–2:00 | Arrival / settle |
| 2:00–2:10 | Opening mantras / short kīrtana |
| 2:10–2:30 | Shared family opening / Prem-kī-Kathā |
| 2:30–3:10 | Parallel parent + younger + older |
| 3:10–3:30 | Family reunification / bhakti lab |
| 3:30–3:40 | Child snack + water (no weekly meal) |
| 3:40–3:55 | Saṅkalpa / project / questions / next week |
| 3:55–4:00 | Closing |
| 4:00 | End on time |

## Tracks

- Parents: `parent-lesson.md`
- Younger: `teacher/YOUNGER-TEACHER-GUIDE.md`
- Older: `teacher/OLDER-TEACHER-GUIDE.md`

## Home practice

5–15 minutes — see `family-home-practice.md` / W1 launch-pack home practice.

## Do not claim

Human approval, temple approval, or publication readiness.
""",
    )


def activities(w: dict) -> None:
    base = WEEKLY / w["slug"] / "activities"
    write(
        base / "YOUNGER-ACTIVITY-PACK.md",
        f"""# {w['code']} Younger Activity Pack (K–2)

## 1. Coloring / line-art

Theme: {w['title']} — family hearing / practice scene (no deity caricature).  
File: `../visuals/V11/line-art-younger.svg`

## 2. Movement / matching

Simple sort: "Hearing / Practice / Respect / Cleanup" cards matching today's objective.

## 3. Craft / foldable

Fold a card with memory phrase on the inside:  
**{w['memory'][:80]}...**

## 4. Memory card

Front: picture cue. Back: short phrase from this week.

## 5. Take-home family cue

Ask at home: **{w['question']}** (one sentence each)

## Low-prep backup

Circle time echo of memory phrase + stretch + sit.

## Extension

Help set snack cups / cleanup leadership with adult.
""",
    )
    write(
        base / "OLDER-ACTIVITY-PACK.md",
        f"""# {w['code']} Older Activity Pack (Grades 4–5)

## 1. Concept worksheet

Explain in your words: {w['conclusion']}

## 2. Puzzle / game

Word search or matching: key terms for this week (answer key required).

## 3. Scenario cards

Use constructed case from `research/CASE-STUDIES.md`.

## 4. Diagram

Complete `../visuals/V11/concept-diagram.mmd` labels.

## 5. Project / poster component

Advance: {w['project']}

## 6. Optional reflection

Where did I confuse scripture with analogy this week?

## Low-prep backup

Pair-share essential question for 3 minutes + one written sentence.
""",
    )
    write(
        base / "OLDER-ANSWER-KEY.md",
        f"""# {w['code']} Older Answer Key

## Worksheet — sample secure answers

1. Primary conclusion: {w['conclusion']}
2. Key scripture: {w['verse']}
3. Misconception to avoid: {w['misconception']}
4. Family question: {w['question']}

## Matching

| Term | Match |
|---|---|
| Scripture | Primary doctrinal control |
| Analogy | Pedagogy with limits |
| Constructed case | Fictional teaching story |
| Science note | Application only — not siddhānta proof |

## Puzzle note

If a printable puzzle is generated later, record answers here before session.
""",
    )


def gamma_v11(w: dict) -> None:
    g = WEEKLY / w["slug"] / "gamma"
    for audience, name in [
        ("master", "V11-GAMMA-MASTER-DECK-PROMPT.md"),
        ("parent", "V11-GAMMA-PARENT-DECK-PROMPT.md"),
        ("younger", "V11-GAMMA-YOUNGER-DECK-PROMPT.md"),
        ("older", "V11-GAMMA-OLDER-DECK-PROMPT.md"),
    ]:
        slides = []
        for i in range(1, 9):
            slides.append(
                f"""### Slide {i}

- **Title:** {w['code']} / {audience} / slide {i}
- **Audience:** {audience}
- **Purpose:** Teach week's conclusion without overclaiming
- **On-slide content:** Short bullets on {w['conclusion'] if i > 1 else w['title']}
- **Presenter notes:** Stay in week scope; Saturday 2–4; no ranking
- **Primary source:** {w['verse']} ({w['verse_url']})
- **Suggested visual:** V11 diagram / line-art
- **Image/diagram prompt:** Simple instructional graphic; no deity caricature
- **Interaction:** One question or echo phrase
- **Do-not-claim:** Not human-approved; not publication-ready; science does not prove doctrine
- **Accessibility note:** High contrast; large type; read aloud key lines
"""
            )
        extra = ""
        if w["code"] == "C1-W1" and audience == "master":
            extra = """
## W1 master must cover

Program identity; is/isn't; six-month roadmap; C1/C2/C3 preview; Saturday flow; cycle rhythm; parent commitment; attendance; child rules; correction ladder; teacher standard; privacy/no comparison/no gossip; host rules; mantra sequence; Cycle 1 project; Week 1 philosophy/key verse; home practice; feedback; next week.
"""
        write(
            g / name,
            f"""# {w['code']} V11 Gamma {audience.title()} Deck Prompt

**Status:** prompt-only — **not rendered** — **not approved**

## Deck identity

- Module: {w['code']} — {w['title']}
- Format: 16:9
- Session: Saturday 2:00–4:00 PM

{extra}

## Slides

{''.join(slides)}
""",
        )
    write(
        g / "V11-GAMMA-SOURCE-MAP.yaml",
        f"""module: {w['code']}
render_status: prompt-only-not-rendered
session: Saturday-14:00-16:00
primary_verse: {w['verse']}
primary_url: {w['verse_url']}
do_not_claim:
  - human-approved
  - publication-ready
  - gamma-rendered
slides_require:
  - primary_source
  - do_not_claim
  - accessibility_note
""",
    )


def visuals_v11(w: dict) -> None:
    v = WEEKLY / w["slug"] / "visuals" / "V11"
    write(
        v / "concept-diagram.mmd",
        f"""flowchart LR
    A["Hear"] --> B["Discuss"]
    B --> C["Practice"]
    C --> D["Reflect"]
    D --> E["Serve"]
    E --> A
    F["{w['code']} focus"] --> C
""",
    )
    # Unique-ish SVG per week via code hash in title position
    y_off = 40 + (int(w["code"][-1]) * 3 if w["code"][-1].isdigit() else 0)
    write(
        v / "concept-diagram.svg",
        f"""<?xml version="1.0" encoding="UTF-8"?>
<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 800 420" role="img">
  <title>{w['code']} concept diagram</title>
  <rect width="800" height="420" fill="#faf8f5"/>
  <text x="400" y="{y_off}" text-anchor="middle" font-size="18" font-family="Segoe UI,sans-serif">{w['code']} — Growth loop</text>
  <rect x="40" y="100" width="120" height="50" rx="8" fill="#e8f0fe" stroke="#4a6fa5"/>
  <text x="100" y="130" text-anchor="middle" font-size="12">Hear</text>
  <rect x="180" y="100" width="120" height="50" rx="8" fill="#e8f0fe" stroke="#4a6fa5"/>
  <text x="240" y="130" text-anchor="middle" font-size="12">Discuss</text>
  <rect x="320" y="100" width="120" height="50" rx="8" fill="#e8f0fe" stroke="#4a6fa5"/>
  <text x="380" y="130" text-anchor="middle" font-size="12">Practice</text>
  <rect x="460" y="100" width="120" height="50" rx="8" fill="#e8f0fe" stroke="#4a6fa5"/>
  <text x="520" y="130" text-anchor="middle" font-size="12">Reflect</text>
  <rect x="600" y="100" width="120" height="50" rx="8" fill="#e8f0fe" stroke="#4a6fa5"/>
  <text x="660" y="130" text-anchor="middle" font-size="12">Serve</text>
  <text x="400" y="220" text-anchor="middle" font-size="14">{w['conclusion'][:70]}</text>
  <text x="400" y="260" text-anchor="middle" font-size="12">{w['verse']}</text>
  <text x="400" y="390" text-anchor="middle" font-size="10" fill="#666">KUTUMBA-original · {w['code']} · human review required</text>
</svg>
""",
    )
    write(
        v / "line-art-younger.svg",
        f"""<?xml version="1.0" encoding="UTF-8"?>
<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 612 792" role="img">
  <title>{w['code']} younger line art</title>
  <rect width="612" height="792" fill="#ffffff"/>
  <text x="306" y="60" text-anchor="middle" font-size="16">{w['code']} coloring sheet</text>
  <rect x="80" y="120" width="450" height="280" fill="none" stroke="#000" stroke-width="2"/>
  <circle cx="200" cy="280" r="40" fill="none" stroke="#000" stroke-width="2"/>
  <circle cx="400" cy="280" r="40" fill="none" stroke="#000" stroke-width="2"/>
  <rect x="150" y="450" width="300" height="80" fill="none" stroke="#000" stroke-width="2"/>
  <text x="306" y="580" text-anchor="middle" font-size="12">Family hearing and practice — color the scene</text>
  <text x="306" y="740" text-anchor="middle" font-size="10">Original KUTUMBA line art · no copyrighted tracing</text>
</svg>
""",
    )
    write(
        v / "IMAGE-GENERATION-PROMPTS.md",
        f"""# {w['code']} Image Generation Prompts

Status: prompts only — do not pretend raster art was generated if unavailable.

## Family deck hero

- Objective: warm family learning scene for {w['title']}
- Subjects: parents + children silhouette-friendly; books; simple home shrine suggestion without deity caricature
- Constraints: no frightening imagery; no invented sacred iconography
- Aspect: 16:9

## Younger coloring scene

- Simple outlines; large shapes; {w['code']} theme
- Print US Letter

## Older concept illustration

- Diagrammatic; labels left for teacher overlay
- Theme: {w['conclusion']}

## Project/poster visual

- Support: {w['project']}

## Scriptural scene (optional)

- Only if source scene is known; mark what must not be invented
""",
    )
    write(
        v / "VISUAL-RIGHTS-REGISTER.yaml",
        f"""module: {w['code']}
assets:
  - id: {w['code'].lower()}-v11-concept
    title: Concept diagram SVG
    source_url: null
    creator: KUTUMBA
    license_rights_status: kutumba-original
    use_allowed: true
    modification_allowed: true
    intended_use: teaching-diagram
    reviewer_note: human-review-required
  - id: {w['code'].lower()}-v11-line-art
    title: Younger line art SVG
    source_url: null
    creator: KUTUMBA
    license_rights_status: kutumba-original
    use_allowed: true
    modification_allowed: true
    intended_use: printable-coloring
    reviewer_note: human-review-required
""",
    )


def w1_launch_pack() -> None:
    lp = WEEKLY / WEEKS[0]["slug"] / "launch-pack"
    files = {
        "SATURDAY-RUN-OF-SHOW.md": """# Saturday Run of Show — C1-W1 (2:00–4:00)

| Time | Block | Owner |
|---|---|---|
| 1:50–2:00 | Arrival / settle | Host + greeter |
| 2:00–2:10 | Opening mantras | Main facilitator |
| 2:10–2:30 | Orientation + Prem-kī-Kathā | Main facilitator |
| 2:30–3:10 | Parallel tracks | Parent / younger / older teachers |
| 3:10–3:30 | Reunification + bhakti lab | All |
| 3:30–3:40 | Snack + water | Parents |
| 3:40–3:55 | Saṅkalpa + project intro | Main facilitator |
| 3:55–4:00 | Closing | Main facilitator |
| 4:00 | End | — |

No weekly meal. End on time.
""",
        "MAIN-FACILITATOR-SCRIPT.md": """# Main Facilitator Script — C1-W1

## Opening

Welcome families. State: KUTUMBA is a family sādhana community — not a drop-off program, substitute temple, or ranking contest.

## Orientation points (brief)

- Saturday 2–4 locked
- Six weeks + off week + Utsava candidate
- Parents onsite
- Covenant + child rules
- Home practice 5–15 minutes
- Cumulative project introduction

## Philosophy core

SB 1.2.18 paraphrase: regular Bhāgavata hearing and service steadies the heart.  
Do not teach W2 body/soul as this week's core.

## Close

Thank host. Remind next Saturday. Private feedback welcome.
""",
        "PARENT-ORIENTATION-HANDOUT.md": """# Parent Orientation Handout — Cycle 1

- What KUTUMBA is / is not
- Saturday 2:00–4:00 PM
- Whole-family attendance; parents onsite
- Minimum home practice
- No comparison / no public sādhana scoring
- Respectful pause/exit allowed
- Snack/water only — no weekly meal
- See `launch/FAMILY-COVENANT.md`
""",
        "SIX-MONTH-ROADMAP-HANDOUT.md": """# Six-Month Roadmap Handout

| Cycle | Focus |
|---|---|
| C1 | Identity — body/soul foundations + family practice |
| C2 | Karma, modes, responsibility |
| C3 | Bhakti — Kṛṣṇa, guru-sādhu-śāstra, holy name |

Reasonable destination: stronger vocabulary, clearer foundations, small sustainable home rhythm, age-appropriate explanation by children, respectful speech and temple connection.

Do not promise certification, initiation eligibility, or guaranteed advancement.
""",
        "CHILD-RULES-YOUNGER.md": """# Child Rules — Younger (K–2)

1. Safe body  
2. Kind words  
3. Follow freeze cue  
4. Feet on floor  
5. Ask before private rooms  
6. Devices away  
7. Snack only in snack area  
8. Help clean up  

Correction: reminder → redirect → quiet reset → parent help. Never shame.
""",
        "CHILD-RULES-OLDER.md": """# Child Rules — Older (Grades 4–5)

1. Safe body and respectful speech  
2. Follow facilitator safety cue  
3. No furniture jumping / indoor throwing  
4. Ask before private rooms / belongings  
5. Devices away unless assigned  
6. Model kindness for younger children  
7. Snack only in allowed area  
8. Cleanup leadership  

Correction ladder applies. Never public shame or comparison.
""",
        "TEACHER-BRIEFING.md": """# Teacher Briefing — C1-W1

Read `launch/TEACHER-READINESS-STANDARD.md`.  
Know objective, SB 1.2.18 link, boundaries, deferral line, and reunification time.
""",
        "PRINT-CHECKLIST.md": """# Print Checklist — C1-W1

- [ ] Parent orientation handout
- [ ] Six-month roadmap
- [ ] Child rules younger/older
- [ ] Opening mantras handout
- [ ] Family saṅkalpa card
- [ ] Covenant acknowledgement template (blank)
- [ ] Younger coloring sheet
- [ ] Older worksheet
""",
        "ROOM-SETUP.md": """# Room Setup — C1-W1

- Opening circle space
- Three parallel areas: parent / younger / older
- Visible quiet-reset spot
- Snack table separate from teaching rugs
- Host private rooms closed/signed
""",
        "SNACK-AND-ALLERGY-CHECKLIST-TEMPLATE.md": """# Snack and Allergy Checklist Template

**Do not commit completed forms.**

| Item | Notes |
|---|---|
| Water available | |
| One simple child snack | |
| Allergy info collected privately | Offline only |
| No shared utensils pressure | |

No weekly meal.
""",
        "FAMILY-SANKALPA-CARD.md": """# Family Saṅkalpa Card — W1

Formula: **specific action + frequency + trigger + minimum version**

Example: "After dinner on Sunday, we read one verse paraphrase for 5 minutes; minimum = one sentence of gratitude."

Family alias (optional): __________
""",
        "CYCLE-1-PROJECT-INTRO.md": """# Cycle 1 Project Intro

**Title:** Who Am I, and How Should Our Family Live?

W1 layer: family purpose + first saṅkalpa.  
Builds through W6 10-minute family presentation (non-competitive).
""",
        "WEEK-1-HOME-PRACTICE.md": """# Week 1 Home Practice

5–15 minutes: choose one —  
1) Read SB 1.2.18 paraphrase together  
2) One family gratitude sentence  
3) Soft mahā-mantra for 2 minutes (no force)

Saturday next: C1-W2.
""",
        "OWNER-NIGHT-BEFORE-CHECKLIST.md": """# Owner Night-Before Checklist

- [ ] Confirm host + rooms
- [ ] Prints stacked
- [ ] Teachers briefed
- [ ] Mantra handout ready
- [ ] Snack plan (simple)
- [ ] No private data in shared drive folders
""",
        "OWNER-60-MINUTE-PREP-PLAN.md": """# Owner 60-Minute Prep Plan

| Min | Task |
|---|---|
| 0–10 | Open run-of-show + print checklist |
| 10–25 | Set rooms / signs |
| 25–40 | Lay materials per track |
| 40–50 | Teacher 5-minute sync |
| 50–60 | Water/snack table + start music/silence plan |
""",
    }
    for name, content in files.items():
        write(lp / name, content)


def orientation_and_handbook() -> None:
    write(
        LAUNCH / "KUTUMBA-C1-FAMILY-ORIENTATION.md",
        """# KUTUMBA Cycle 1 Family Orientation

## What KUTUMBA is / is not

See Family Covenant. KUTUMBA is a disciplined family sādhana community with parents onsite.

## Why families commit

To protect weekly hearing, build shared vocabulary, and practice small sustainable devotion at home.

## Saturday 2:00–4:00 PM

Arrival 1:50. End 4:00. No weekly meal — water + one simple child snack.

## First six months

C1 identity → C2 karma/modes → C3 bhakti. Six-week rhythm, protected off week, Utsava candidate.

## Home practice

Usually 5–15 minutes. Minimum versions allowed.

## Culture

No comparison, no public private-sādhana reporting, respectful pause/exit, host-home respect, child rules, teacher readiness, privacy.

## Reasonable six-month destination

Stronger vocabulary; clearer body/soul foundations; beginning karma/modes understanding; clearer Kṛṣṇa / guru-sādhu-śāstra / holy name orientation; small home rhythm; age-appropriate child explanation; respectful speech and service orientation.

Do not promise certification, initiation eligibility, or guaranteed advancement.
""",
    )
    write(
        LAUNCH / "KUTUMBA-C1-TEACHER-HANDBOOK.md",
        """# KUTUMBA Cycle 1 Teacher Handbook

Combines:

- Teacher readiness standard
- Pre-week checklist
- Child house rules + correction ladder
- Opening mantras usage notes
- Saturday flow
- Deferral line
- Privacy boundaries

Teachers teach from week packets. Do not invent doctrine. Do not claim approvals.
""",
    )


def fix_w1_contamination() -> None:
    write(
        WEEKLY / WEEKS[0]["slug"] / "analogy-and-application.md",
        """---
week_code: C1-W1
week_title: What Is KUTUMBA, and Why Are We Here?
---

## Analogy and Practical Example

W1 analogies support **family hearing, protected practice, and community boundaries** — not body/soul identity (see C1-W2).

| Analogy | Source status | Teaching value | Limitation | Family example |
| --- | --- | --- | --- | --- |
| Protected garden plot | kutumba-summary | Regular watering grows roots; skipped care weakens plants | Not a guarantee of realization | Family protects Saturday 2–4 like garden time |
| Weekly team practice | kutumba-summary | Showing up + home drills both matter | Sports glory is not the goal | Attend session and do 5–15 min home practice |
| Path with fence | kutumba-summary | Boundaries (is/isn't) keep the path clear | Boundaries are not hostility | Review KUTUMBA is/isn't cards |
| Lamp needing oil | śāstra-adjacent pedagogy | Hearing is fuel for steady light | Do not over-literalize | Read one verse paraphrase at home |

_See `research/ANALOGIES-AND-LIMITS.md` and `research/MISCONCEPTIONS-AND-BOUNDARIES.md`._
""",
    )
    # Friday → Saturday in parent lesson
    pl = WEEKLY / WEEKS[0]["slug"] / "parent-lesson.md"
    text = pl.read_text(encoding="utf-8")
    text = text.replace("attends every Friday", "attends every Saturday")
    pl.write_text(text, encoding="utf-8")


def c1_index() -> None:
    write(
        WEEKLY / "C1-V11-OWNER-INDEX.md",
        """# Cycle 1 V11 Owner Index

Saturday founding cohort package — **not** human/temple approved for public distribution.

## Start here

1. [launch/C1-SATURDAY-CALENDAR.md](../../launch/C1-SATURDAY-CALENDAR.md)
2. [launch/FAMILY-COVENANT.md](../../launch/FAMILY-COVENANT.md)
3. [launch/OPENING-MANTRAS-HANDOUT.md](../../launch/OPENING-MANTRAS-HANDOUT.md)
4. [c1-w1.../launch-pack/SATURDAY-RUN-OF-SHOW.md](c1-w1-what-is-kutumba-and-why-are-we-here/launch-pack/SATURDAY-RUN-OF-SHOW.md)
5. [c1-w1.../launch-pack/MAIN-FACILITATOR-SCRIPT.md](c1-w1-what-is-kutumba-and-why-are-we-here/launch-pack/MAIN-FACILITATOR-SCRIPT.md)
6. [c1-w1.../launch-pack/PRINT-CHECKLIST.md](c1-w1-what-is-kutumba-and-why-are-we-here/launch-pack/PRINT-CHECKLIST.md)

## Per-week guides

| Week | Main facilitator | Younger | Older | Family handout / home practice |
|---|---|---|---|---|
| W1 | [guide](c1-w1-what-is-kutumba-and-why-are-we-here/teacher/MAIN-FACILITATOR-GUIDE-V11.md) | [younger](c1-w1-what-is-kutumba-and-why-are-we-here/teacher/YOUNGER-TEACHER-GUIDE.md) | [older](c1-w1-what-is-kutumba-and-why-are-we-here/teacher/OLDER-TEACHER-GUIDE.md) | [home](c1-w1-what-is-kutumba-and-why-are-we-here/family-home-practice.md) |
| W2 | [guide](c1-w2-i-am-not-this-body/teacher/MAIN-FACILITATOR-GUIDE-V11.md) | [younger](c1-w2-i-am-not-this-body/teacher/YOUNGER-TEACHER-GUIDE.md) | [older](c1-w2-i-am-not-this-body/teacher/OLDER-TEACHER-GUIDE.md) | [home](c1-w2-i-am-not-this-body/family-home-practice.md) |
| W3 | [guide](c1-w3-the-nature-of-the-soul/teacher/MAIN-FACILITATOR-GUIDE-V11.md) | [younger](c1-w3-the-nature-of-the-soul/teacher/YOUNGER-TEACHER-GUIDE.md) | [older](c1-w3-the-nature-of-the-soul/teacher/OLDER-TEACHER-GUIDE.md) | [home](c1-w3-the-nature-of-the-soul/family-home-practice.md) |
| W4 | [guide](c1-w4-why-human-life-is-rare-and-valuable/teacher/MAIN-FACILITATOR-GUIDE-V11.md) | [younger](c1-w4-why-human-life-is-rare-and-valuable/teacher/YOUNGER-TEACHER-GUIDE.md) | [older](c1-w4-why-human-life-is-rare-and-valuable/teacher/OLDER-TEACHER-GUIDE.md) | [home](c1-w4-why-human-life-is-rare-and-valuable/family-home-practice.md) |
| W5 | [guide](c1-w5-the-temporary-world-and-the-search-for-permanent-happiness/teacher/MAIN-FACILITATOR-GUIDE-V11.md) | [younger](c1-w5-the-temporary-world-and-the-search-for-permanent-happiness/teacher/YOUNGER-TEACHER-GUIDE.md) | [older](c1-w5-the-temporary-world-and-the-search-for-permanent-happiness/teacher/OLDER-TEACHER-GUIDE.md) | [home](c1-w5-the-temporary-world-and-the-search-for-permanent-happiness/family-home-practice.md) |
| W6 | [guide](c1-w6-integration-night-who-am-i-and-how-should-our-family-live/teacher/MAIN-FACILITATOR-GUIDE-V11.md) | [younger](c1-w6-integration-night-who-am-i-and-how-should-our-family-live/teacher/YOUNGER-TEACHER-GUIDE.md) | [older](c1-w6-integration-night-who-am-i-and-how-should-our-family-live/teacher/OLDER-TEACHER-GUIDE.md) | [rubric](c1-w6-integration-night-who-am-i-and-how-should-our-family-live/project/PRESENTATION-RUBRIC.md) |
""",
    )


def export_docx() -> None:
    # Launch docs
    pairs = [
        (LAUNCH / "KUTUMBA-C1-FAMILY-ORIENTATION.md", LAUNCH / "KUTUMBA-C1-FAMILY-ORIENTATION.docx", "KUTUMBA C1 Family Orientation"),
        (LAUNCH / "KUTUMBA-C1-TEACHER-HANDBOOK.md", LAUNCH / "KUTUMBA-C1-TEACHER-HANDBOOK.docx", "KUTUMBA C1 Teacher Handbook"),
        (LAUNCH / "OPENING-MANTRAS-HANDOUT.md", LAUNCH / "OPENING-MANTRAS-HANDOUT.docx", "Opening Mantras Handout"),
    ]
    for md, docx, title in pairs:
        md_to_docx(md, docx, title)
    for w in WEEKS:
        base = WEEKLY / w["slug"]
        exp = base / "exports"
        md_to_docx(base / "teacher" / "MAIN-FACILITATOR-GUIDE-V11.md", exp / f"{w['code']}-MAIN-FACILITATOR-GUIDE.docx", f"{w['code']} Main Facilitator Guide")
        md_to_docx(base / "teacher" / "YOUNGER-TEACHER-GUIDE.md", exp / f"{w['code']}-YOUNGER-TEACHER-GUIDE.docx", f"{w['code']} Younger Teacher Guide")
        md_to_docx(base / "teacher" / "OLDER-TEACHER-GUIDE.md", exp / f"{w['code']}-OLDER-TEACHER-GUIDE.docx", f"{w['code']} Older Teacher Guide")
        # family handout from home practice if exists else orientation snippet
        home = base / "family-home-practice.md"
        if not home.exists():
            write(home, f"# {w['code']} Family Home Practice\n\n5–15 minutes linked to {w['verse']}.\n")
        md_to_docx(home, exp / f"{w['code']}-FAMILY-HANDOUT.docx", f"{w['code']} Family Handout")


def main() -> int:
    orientation_and_handbook()
    fix_w1_contamination()
    w1_launch_pack()
    c1_index()
    for w in WEEKS:
        project_files(w)
        research_files(w)
        teacher_guides(w)
        activities(w)
        gamma_v11(w)
        visuals_v11(w)
    export_docx()
    print("V11 C1 generation complete")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
