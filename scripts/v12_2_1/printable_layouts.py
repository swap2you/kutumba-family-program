#!/usr/bin/env python3
"""Real Word-table and image layouts for C1-W1 participant printables."""
from __future__ import annotations

import sys
from pathlib import Path

from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

REPO = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(REPO / "scripts" / "v12"))
from kutumba_docx_styles import CREAM, PLUM, SAFFRON, TEAL, add_branded_table, add_callout, add_verse_card  # noqa: E402

ASSETS = REPO / "11-weekly-program-library" / "first-six-months" / "c1-w1-what-is-kutumba-and-why-are-we-here" / "visuals" / "v12_2_1"
MEMORY = "Our family helps one another remember Kṛṣṇa."
PURPOSES = [
    "Home chanting / hearing", "Temple service / association",
    "Respectful family correction", "Festival participation",
    "Study readiness / source discipline", "Family cooperation / service",
]
EXAMPLES = [
    "After dinner we chant three mahā-mantras together.",
    "We help set up chairs for a temple program.",
    "A parent corrects a child privately with a calm voice.",
    "Our family joins a festival with service.",
    "Before teaching, we open the verse link and check the source.",
    "Siblings share clean-up jobs without ranking.",
    "A parent reads one short Bhāgavata passage at home.",
    "We greet devotees and stay for association.",
    "When speech becomes harsh, we pause and restart with respect.",
    "We help with prasādam distribution or festival decorating.",
    "We write a question and check śāstra instead of guessing.",
    "Everyone carries one bag and thanks the host.",
]
EXAMPLE_KEYS = [1, 2, 3, 4, 5, 6, 1, 2, 3, 4, 5, 6]


def _border(table, color: str = PLUM, size: int = 10) -> None:
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), str(size))
        element.set(qn("w:color"), color)
        borders.append(element)
    table._tbl.tblPr.append(borders)


def _shade(cell, color: str) -> None:
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color)
    cell._tc.get_or_add_tcPr().append(shd)


def _title(document, code: str, title: str, instruction: str = "") -> None:
    document.add_heading(f"{code} — {title}", level=1)
    if instruction:
        p = document.add_paragraph(instruction)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER


def _page(document) -> None:
    document.add_page_break()


def _teacher_page(document, title: str, rows) -> None:
    _page(document)
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("TEACHER-ONLY")
    r.bold = True; r.font.size = Pt(30); r.font.color.rgb = RGBColor.from_string(PLUM)
    document.add_heading(title, level=1)
    add_branded_table(document, ["Item", "Best match / guidance"], rows)


def _cards(document, cards: list[str], prefix: str, columns: int = 2, per_page: int = 6) -> None:
    for start in range(0, len(cards), per_page):
        if start:
            _page(document)
        group = cards[start:start + per_page]
        rows = (len(group) + columns - 1) // columns
        table = document.add_table(rows=rows, cols=columns)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        _border(table, TEAL, 9)
        for i, cell in enumerate([c for row in table.rows for c in row.cells]):
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            cell.width = Inches(3.3)
            if i >= len(group):
                cell.text = ""
                continue
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(14)
            tag = p.add_run(f"{prefix}{start + i + 1}\n")
            tag.bold = True; tag.font.color.rgb = RGBColor.from_string(PLUM)
            body = p.add_run(group[i])
            body.font.size = Pt(13)
            p.paragraph_format.space_after = Pt(14)


def _lines(document, prompts: list[str], line_count: int = 2) -> None:
    table = document.add_table(rows=len(prompts), cols=1)
    _border(table, TEAL, 7)
    for prompt, cell in zip(prompts, table.column_cells(0)):
        p = cell.paragraphs[0]
        p.add_run(prompt).bold = True
        for _ in range(line_count):
            cell.add_paragraph("________________________________________________________________")


def add_y01_four_corners(document) -> None:
    items = [
        ("HEAR", "We listen to Kṛṣṇa-kathā.", "icon-hear.png"),
        ("CHANT", "We say Kṛṣṇa's names.", "icon-chant.png"),
        ("SERVE", "We help with love.", "icon-serve.png"),
        ("RESPECT", "We use kind words and care.", "icon-respect.png"),
    ]
    for index, (heading, sentence, image) in enumerate(items):
        if index: _page(document)
        _title(document, "Y01", "Four Corners Sign")
        p = document.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Inches(.45)
        r = p.add_run(heading); r.bold = True; r.font.size = Pt(44); r.font.color.rgb = RGBColor.from_string(PLUM)
        pic = document.add_paragraph(); pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pic.add_run().add_picture(str(ASSETS / image), width=Inches(3.2))
        s = document.add_paragraph(); s.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rr = s.add_run(sentence); rr.bold = True; rr.font.size = Pt(22)


def add_y02_scenario_cards(document) -> None:
    _title(document, "Y02", "Four Corners Scenario Cards", "Cut on cell borders. Read each card aloud.")
    cards = [
        "The family listens to a Kṛṣṇa story.", "We sing Hare Kṛṣṇa together.",
        "You help put the activity supplies away.", "You ask before touching a friend's toy.",
        "You sit quietly while another child answers.", "You help bring scripture to the family reading place.",
        "You repeat the memory line with the teacher.", "You join the mahā-mantra softly.",
        "You help a younger child find crayons.", "You keep feet on the floor in the host home.",
        "You listen when a parent reads one verse.", "Your family chants three mahā-mantras at home.",
    ]
    _cards(document, cards, "CARD ", per_page=6)
    keys = ["HEAR", "CHANT", "SERVE", "RESPECT", "RESPECT", "SERVE", "HEAR", "CHANT", "SERVE", "RESPECT", "HEAR", "CHANT"]
    _teacher_page(document, "Y02 Answer Key", [(str(i), key) for i, key in enumerate(keys, 1)])


def add_y03_bhakti_garden(document) -> None:
    _title(document, "Y03", "Bhakti Garden", "Color the petals. Memory phrase is in the flower center.")
    p = document.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(ASSETS / "flower-bhakti-garden.png"), width=Inches(6.5))
    document.add_paragraph("This week our family will grow: __________________________________________")
    document.add_paragraph("☐ We will try our tiny home practice at least once.")


def add_y04_badge_bookmark(document) -> None:
    _title(document, "Y04", "Memory Badge and Bookmark", "Cut around the two outlines. Text is printed inside each shape.")
    table = document.add_table(rows=1, cols=2); _border(table, SAFFRON, 8)
    badge, bookmark = table.rows[0].cells
    badge.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    badge.paragraphs[0].add_run().add_picture(str(ASSETS / "badge-outline.png"), width=Inches(3.0))
    bookmark.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    bookmark.paragraphs[0].add_run().add_picture(str(ASSETS / "bookmark-strip.png"), height=Inches(5.0))
    icons = document.add_table(rows=1, cols=4)
    for cell, name in zip(icons.rows[0].cells, ("icon-hear.png", "icon-chant.png", "icon-serve.png", "icon-respect.png")):
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cell.paragraphs[0].add_run().add_picture(str(ASSETS / name), width=Inches(1.1))


def add_y05_house_rule_sort(document) -> None:
    _title(document, "Y05", "House-Rule Picture Sort", "Sort each card onto SAFE or NEEDS RESET. Use calm, no-shame language.")
    mats = document.add_table(rows=1, cols=2)
    for cell, image in zip(mats.rows[0].cells, ("mat-safe.png", "mat-reset.png")):
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cell.paragraphs[0].add_run().add_picture(str(ASSETS / image), width=Inches(3.1))
    _page(document)
    cards = ["Feet on the floor while listening.", "Jumping on the host couch.", "Asking before borrowing crayons.",
             "Throwing a block indoors.", "Helping clean crayons into the bin.", "Running into a private bedroom without asking.",
             "Using kind words while waiting for a turn.", "Teasing a friend about a wrong answer."]
    _cards(document, cards, "CARD ", per_page=8)
    _teacher_page(document, "Y05 Answer Key", [("SAFE", "1, 3, 5, 7"), ("NEEDS RESET", "2, 4, 6, 8")])


def add_o01_observation(document) -> None:
    _title(document, "O01", "ŚB 1.2.18 Observation Sheet")
    add_verse_card(document, "ŚB 1.2.18", "",
        "naṣṭa-prāyeṣv abhadreṣu nityaṁ bhāgavata-sevayā / bhagavaty uttama-śloke bhaktir bhavati naiṣṭhikī",
        "When we regularly hear and serve the Bhāgavata (book and devotee association), troubles in the heart are cleared and steady devotion to the Lord becomes established.",
        "https://vedabase.io/en/library/sb/1/2/18/", "Scripture display/source: VedaBase; teaching meaning: KUTUMBA-original.")
    _lines(document, ["Repeated-practice words I notice:", "What result is described?",
                      "Why might regular (nityaṁ) matter for a family?", "One question for the facilitator:"], 1)


def add_o02_is_is_not(document) -> None:
    _title(document, "O02", "KUTUMBA Is / Is Not Sort")
    mats = document.add_table(rows=1, cols=2); _border(mats, PLUM, 10)
    for cell, label in zip(mats.rows[0].cells, ("KUTUMBA IS", "KUTUMBA IS NOT")):
        p = cell.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(label); r.bold = True; r.font.size = Pt(24); r.font.color.rgb = RGBColor.from_string(PLUM)
    _page(document)
    yes = ["family formation", "supports temple life", "parents participate", "small home practice", "respectful correction", "source-based learning"]
    no = ["drop-off babysitting", "replacement temple", "competition", "public sādhana leaderboard", "initiation or certification", "place to invent philosophy"]
    cards = [item for pair in zip(yes, no) for item in pair]
    _cards(document, cards, "CARD ", per_page=12)
    _teacher_page(document, "O02 Answer Key", [("KUTUMBA IS", "; ".join(yes)), ("KUTUMBA IS NOT", "; ".join(no))])


def add_o03_six_purpose(document) -> None:
    _title(document, "O03", "Six-Purpose Challenge")
    _cards(document, PURPOSES, "PURPOSE ", per_page=6)
    _page(document); _title(document, "O03", "Example Cards", "Match each example to one purpose.")
    _cards(document, EXAMPLES, "E", per_page=6)
    _teacher_page(document, "O03 Answer Key", [(f"E{i}", f"{key}. {PURPOSES[key-1]}") for i, key in enumerate(EXAMPLE_KEYS, 1)])


def add_o04_scenarios(document) -> None:
    _title(document, "O04", "Scenario Challenge")
    scenarios = [
        "A family never misses Saturday but has no shared practice at home.",
        "A family chants at home but stops joining temple or devotee gatherings because “home is enough.”",
        "A child bumps a plant; another child calls names, and an adult joins the mockery.",
        "A teacher is unsure how to answer a deep question about the soul.",
    ]
    for i, scenario in enumerate(scenarios, 1):
        if i > 1: _page(document)
        document.add_heading(f"Scenario {i}", level=2); document.add_paragraph(scenario)
        _lines(document, ["What is good?", "What is missing?", "What should happen next?", "Which KUTUMBA principle applies?"], 2)
    keys = [
        "Rhythm is healthy; add a tiny home practice without shame.",
        "Home effort is healthy; restore association and service.",
        "Stop mockery; calm reminder, redirect, and repair.",
        "Use the deferral line and verify the source; do not speculate.",
    ]
    _teacher_page(document, "O04 Answer Key", [(str(i), key) for i, key in enumerate(keys, 1)])


def add_o05_compass(document) -> None:
    _title(document, "O05", "Family Compass", "How should our family grow?")
    p = document.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(ASSETS / "compass-family.png"), width=Inches(5.6))
    _lines(document, ["HEAR:", "PRACTICE:", "SERVE:", "ASSOCIATE:"], 1)


def add_o06_exit_ticket(document) -> None:
    _title(document, "O06", "Exit Ticket", "Name: __________________________  Date: ______________")
    _lines(document, ["KUTUMBA is…", "KUTUMBA is not…", "One rule I understand…",
                      "One thing our family can try…", "One question I still have…"], 2)


def add_p01(document) -> None:
    _title(document, "P01", "One-Year Spiritual-Home Reflection (Private)")
    add_callout(document, "SAFETY_PRIVACY", "Write privately. Sharing is optional. Do not collect or place completed sheets in Git.")
    p = document.add_paragraph("One year from now, what do I hope our home feels like spiritually?")
    p.runs[0].bold = True; p.runs[0].font.size = Pt(18)
    for _ in range(9): document.add_paragraph("________________________________________________________________")
    document.add_paragraph("One word I want to protect this year: ______________________________")


def add_p02(document) -> None:
    _title(document, "P02", "Six-Purpose Card Sort")
    _cards(document, PURPOSES, "PURPOSE ", per_page=6)
    _page(document); _title(document, "P02", "Example Cards")
    _cards(document, EXAMPLES, "E", per_page=6)
    _page(document); _lines(document, ["Purpose that feels most needed now:", "Our concrete family example:"], 3)
    _teacher_page(document, "P02 Facilitator Match Key", [(f"E{i}", f"{key}. {PURPOSES[key-1]}") for i, key in enumerate(EXAMPLE_KEYS, 1)])


def add_p03(document) -> None:
    _title(document, "P03", "Two-Family Case")
    add_branded_table(document, ["Family A", "Family B"], [[
        "Attends every Saturday and enjoys the group, but does nothing together at home.",
        "Reads or chants at home but gradually stops joining temple and devotee association because home practice “is enough.”"
    ]])
    _lines(document, ["What is healthy in each family?", "What is missing?",
                      "What is one nonjudgmental repair?", "How can KUTUMBA support rather than replace temple life?"], 3)


def add_p04(document) -> None:
    _title(document, "P04", "Family Operating Agreement (Private)")
    add_callout(document, "SAFETY_PRIVACY", "No ranking and no public reading required. Keep completed agreements private.")
    _lines(document, ["One support we need:", "One contribution we can reliably make:",
                      "One Saturday-protection habit:", "One host-home behavior we commit to:"], 3)
    document.add_paragraph("Family initials (optional): ____________________  Date: ______________")


def add_p05(document) -> None:
    _title(document, "P05", "Family Saṅkalpa Builder (Private)")
    add_callout(document, "KEY_IDEA", "specific action + frequency + trigger + minimum version")
    _lines(document, ["Specific action:", "Frequency:", "Trigger (when and where):",
                      "Minimum version for a hard day:", "Where we will place the reminder:"], 2)
    icons = document.add_table(rows=1, cols=4)
    for cell, name in zip(icons.rows[0].cells, ("icon-hear.png", "icon-chant.png", "icon-serve.png", "icon-respect.png")):
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cell.paragraphs[0].add_run().add_picture(str(ASSETS / name), width=Inches(.75))
    document.add_paragraph("No photo proof. No ranking. The minimum version counts as success.")


YOUNGER_BUILDERS = [add_y01_four_corners, add_y02_scenario_cards, add_y03_bhakti_garden, add_y04_badge_bookmark, add_y05_house_rule_sort]
OLDER_BUILDERS = [add_o01_observation, add_o02_is_is_not, add_o03_six_purpose, add_o04_scenarios, add_o05_compass, add_o06_exit_ticket]
PARENT_BUILDERS = [add_p01, add_p02, add_p03, add_p04, add_p05]
