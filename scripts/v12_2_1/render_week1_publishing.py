#!/usr/bin/env python3
"""Rebuild the seven C1-W1 packets with safe guides and real printables."""
from __future__ import annotations

import hashlib
import json
import os
import re
import sys
import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import Callable

import pymupdf as fitz
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

REPO = Path(__file__).resolve().parents[2]
for item in (REPO, REPO / "scripts" / "v12", Path(__file__).parent):
    sys.path.insert(0, str(item))
from kutumba_docx_styles import PLUM, add_branded_table, add_callout, add_cover, add_header_footer, apply_kutumba_styles  # noqa: E402
from render_publication_docs import convert_with_word  # noqa: E402
from printable_layouts import OLDER_BUILDERS, PARENT_BUILDERS, YOUNGER_BUILDERS, add_p05  # noqa: E402
from render_markdown_safe import render_markdown_safe  # noqa: E402

BASE = REPO / "11-weekly-program-library" / "first-six-months" / "c1-w1-what-is-kutumba-and-why-are-we-here"
LAUNCH = REPO / "launch"
OUTPUT = REPO / "exports" / "final" / "week1"
STATUS = "Internal founding-cohort teaching material — human/temple review EXTERNAL_OPEN"
SCHEDULE = [
    ("1:50–2:00", "Arrival"), ("2:00–2:10", "Opening mantras + welcome"),
    ("2:10–2:30", "All-family launch/orientation"), ("2:30–3:10", "Parallel tracks"),
    ("3:10–3:30", "Reunification + bhakti laboratory"), ("3:30–3:40", "Snack + water"),
    ("3:40–3:55", "Saṅkalpa + Cycle 1 project + home practice"), ("3:55–4:00", "Close"),
]


@dataclass(frozen=True)
class Packet:
    name: str
    title: str
    subtitle: str
    audience: str
    build: Callable


def _relative(path: Path) -> str:
    return path.resolve().relative_to(REPO.resolve()).as_posix()


def _text(path: Path) -> str:
    if not path.is_file():
        raise FileNotFoundError(path)
    text = path.read_text(encoding="utf-8-sig")
    return "\n".join(line for line in text.splitlines() if "2:35" not in line and "3:25" not in line)


def _source(document, path: Path, heading: str) -> None:
    document.add_heading(heading, level=1)
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    document.add_paragraph(f"Canonical source: {_relative(path)}  •  SHA-256: {digest}", style="KUTUMBA Rights")
    title_match = re.search(r"^#\s+(.+)$", _text(path), re.M)
    render_markdown_safe(document, _text(path), title_match.group(1).strip() if title_match else heading)


def _page_source(document, path: Path, heading: str) -> None:
    document.add_page_break(); _source(document, path, heading)


def _page_builder(document, builder) -> None:
    document.add_page_break(); builder(document)


def _schedule(document, heading: str = "Locked Saturday Schedule") -> None:
    document.add_heading(heading, level=1)
    add_branded_table(document, ["Time", "Block"], SCHEDULE)
    add_callout(document, "TIME_CUE", "Track handoff is at 2:30; all groups reunite at 3:10.")


def _new(packet: Packet):
    doc = Document(); apply_kutumba_styles(doc)
    add_cover(doc, packet.title, packet.subtitle, "C1-W1 • What Is KUTUMBA, and Why Are We Here?", version="V12.2.1", status=STATUS)
    add_header_footer(doc, "C1-W1 • V12.2.1", packet.audience)
    for section in doc.sections:
        section.orientation = WD_ORIENT.PORTRAIT; section.page_width = Inches(8.5); section.page_height = Inches(11)
    return doc


def build_start(document) -> None:
    _schedule(document); _page_source(document, REPO / "V12_2-WEEK1-START-HERE.md", "Operational Entry Point")


def build_owner(document) -> None:
    _schedule(document); _page_source(document, BASE / "OWNER-RUNBOOK-V12.2.md", "Owner Runbook")


def build_orientation(document) -> None:
    _schedule(document)
    for path, title in [
        (LAUNCH / "KUTUMBA-C1-FAMILY-ORIENTATION.md", "Family Orientation"),
        (LAUNCH / "FAMILY-COVENANT.md", "Family Covenant"),
        (LAUNCH / "FAMILY-COVENANT-ACKNOWLEDGEMENT-TEMPLATE.md", "Blank Covenant Acknowledgement"),
        (LAUNCH / "CHILD-HOUSE-RULES.md", "Child House Rules"),
    ]: _page_source(document, path, title)


def build_track(document, guide: Path, guide_title: str, builders, rules: bool = False) -> None:
    _schedule(document); _page_source(document, guide, guide_title)
    if rules: _page_source(document, LAUNCH / "CHILD-HOUSE-RULES.md", "Child House Rules")
    for builder in builders: _page_builder(document, builder)


def _teacher_divider(document) -> None:
    document.add_page_break()
    p = document.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before = Inches(2.3)
    r = p.add_run("TEACHER-ONLY"); r.bold = True; r.font.size = Pt(32); r.font.color.rgb = RGBColor.from_string(PLUM)
    note = document.add_paragraph("Remove teacher-only pages before distributing participant copies."); note.alignment = WD_ALIGN_PARAGRAPH.CENTER


def build_saturday(document) -> None:
    _schedule(document, "Quick Schedule")
    for path, title in [
        (LAUNCH / "OPENING-MANTRAS-HANDOUT.md", "Opening Mantras"),
        (LAUNCH / "KUTUMBA-C1-FAMILY-ORIENTATION.md", "Family Orientation"),
        (LAUNCH / "FIRST-SIX-MONTHS-CALENDAR.md", "First-Six-Month Roadmap"),
        (LAUNCH / "FAMILY-COVENANT.md", "Family Covenant Summary"),
        (LAUNCH / "FAMILY-COVENANT-ACKNOWLEDGEMENT-TEMPLATE.md", "Blank Acknowledgement"),
        (BASE / "teacher" / "PARENT-GUIDE.md", "Parent Track Run Sheet"),
    ]: _page_source(document, path, title)
    for builder in PARENT_BUILDERS: _page_builder(document, builder)
    _page_source(document, LAUNCH / "CHILD-HOUSE-RULES.md", "Younger Child Rules")
    for builder in YOUNGER_BUILDERS: _page_builder(document, builder)
    _page_source(document, LAUNCH / "CHILD-HOUSE-RULES.md", "Older Child Rules")
    for builder in OLDER_BUILDERS: _page_builder(document, builder)
    _page_builder(document, add_p05)
    _page_source(document, BASE / "project" / "MODULE-PROJECT-BRIEF.md", "Cycle 1 Project Introduction")
    document.add_page_break(); document.add_heading("Week-1 Home Practice", level=1)
    for line in ("1. Offer one prayer.", "2. Chant three mahā-mantras.", "3. Each person offers one appreciation."):
        document.add_paragraph(line)
    add_callout(document, "HOME_PRACTICE", "Minimum: 5 minutes, at least once before next Saturday. No photo proof. No ranking.")
    _teacher_divider(document)
    _page_source(document, BASE / "teacher" / "YOUNGER-TEACHER-GUIDE.md", "Younger Teacher Run Sheet")
    _page_source(document, BASE / "teacher" / "OLDER-TEACHER-GUIDE.md", "Older Teacher Run Sheet")
    _page_source(document, BASE / "teacher" / "MAIN-FACILITATOR-GUIDE-V12.md", "Main Facilitator Speaking Map")
    document.add_page_break(); document.add_heading("Room / Material / Snack / Cleanup Checklist", level=1)
    add_branded_table(document, ["Done", "Area", "Check"], [
        ("☐", "Room", "Parent circle, younger floor zone, older table zone, clear aisles"),
        ("☐", "Materials", "Labeled track folders, pencils, crayons, tape, timer"),
        ("☐", "Privacy", "Completed forms remain private and never enter Git"),
        ("☐", "Snack", "Snack and water staged for 3:30; allergy plan confirmed"),
        ("☐", "Safety", "Calm corner visible; private rooms closed; exits clear"),
        ("☐", "Cleanup", "Tables wiped and host home reset assigned"),
    ])


def packets() -> list[Packet]:
    return [
        Packet("KUTUMBA-C1-W1-START-HERE-V12.2", "Week 1 — Start Here", "Owner-facing operational entry point", "Program owner", build_start),
        Packet("KUTUMBA-C1-W1-OWNER-RUNBOOK-V12.2", "Week 1 — Owner Runbook", "Friday preparation and Saturday execution", "Program owner", build_owner),
        Packet("KUTUMBA-C1-W1-FAMILY-ORIENTATION-AND-COVENANT-V12.2", "Family Orientation and Covenant", "Welcome, commitments, acknowledgement, and house rules", "Families", build_orientation),
        Packet("KUTUMBA-C1-W1-PARENT-TRACK-V12.2", "Parent Track", "Facilitator guide and complete parent printables", "Parent facilitator",
               lambda d: build_track(d, BASE / "teacher" / "PARENT-GUIDE.md", "Parent Facilitator Run Sheet", PARENT_BUILDERS)),
        Packet("KUTUMBA-C1-W1-YOUNGER-TEACHER-PACK-V12.2", "Younger Teacher Pack", "K–2 run sheet, child rules, and complete activities", "Younger teacher",
               lambda d: build_track(d, BASE / "teacher" / "YOUNGER-TEACHER-GUIDE.md", "Younger Teacher Run Sheet", YOUNGER_BUILDERS, True)),
        Packet("KUTUMBA-C1-W1-OLDER-TEACHER-PACK-V12.2", "Older Teacher Pack", "Grades 4–5 run sheet and complete activities", "Older teacher",
               lambda d: build_track(d, BASE / "teacher" / "OLDER-TEACHER-GUIDE.md", "Older Teacher Run Sheet", OLDER_BUILDERS)),
        Packet("KUTUMBA-C1-W1-SATURDAY-PRINT-PACKET-V12.2", "Saturday Print Packet", "Complete ordered participant and teacher packet", "Families and facilitators", build_saturday),
    ]


def strip_trailing_empty(pdf_path: Path) -> int:
    doc = fitz.open(pdf_path); removed = 0
    while doc.page_count > 1:
        blocks = doc[-1].get_text("blocks")
        body = " ".join(block[4] for block in blocks if block[1] > 60 and block[3] < 740).strip()
        if body: break
        doc.delete_page(-1); removed += 1
    if removed:
        temp = pdf_path.with_suffix(".tmp.pdf"); doc.save(temp, deflate=True); doc.close(); os.replace(temp, pdf_path)
    else: doc.close()
    return removed


def _forbidden_check(pdf_path: Path) -> int:
    forbidden = ["<div", "</div>", "page-break-after:", "```", "[ ear /", "[ beads /", "[ helping hands ]", "[ heart / home ]",
                 "35. What is healthy", "103. What should happen", "109. Offer one prayer"] + list("┌┐└┘─│├┤┬┴┼")
    with fitz.open(pdf_path) as pdf:
        text = "\n".join(page.get_text() for page in pdf); pages = pdf.page_count
    found = [token for token in forbidden if token.lower() in text.lower()]
    if found: raise RuntimeError(f"{pdf_path.name} contains forbidden output: {found}")
    for token in ("2:30", "3:10"):
        if token not in text: raise RuntimeError(f"{pdf_path.name} missing {token}")
    for token in ("2:35", "3:25"):
        if token in text: raise RuntimeError(f"{pdf_path.name} contains obsolete {token}")
    return pages


def main() -> int:
    OUTPUT.mkdir(parents=True, exist_ok=True)
    results = []
    for packet in packets():
        docx = OUTPUT / f"{packet.name}.docx"; pdf = OUTPUT / f"{packet.name}.pdf"
        document = _new(packet); packet.build(document)
        add_callout(document, "RIGHTS_NOTE", "KUTUMBA-original framing; cited third-party material retains its own rights. Human/temple review remains EXTERNAL_OPEN.")
        document.save(docx)
        if pdf.exists(): pdf.unlink()
        reported = convert_with_word(docx, pdf)
        strip_trailing_empty(pdf)
        pages = _forbidden_check(pdf)
        results.append({"name": packet.name, "docx": _relative(docx), "docx_bytes": docx.stat().st_size,
                        "pdf": _relative(pdf), "pdf_bytes": pdf.stat().st_size, "pages": pages, "word_reported_pages": reported})
        print(f"{packet.name}: {pages} pages")
    manifest = OUTPUT / "week1-render-manifest.json"
    manifest.write_text(json.dumps({"version": "V12.2.1", "packets": results}, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print(f"Wrote {_relative(manifest)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
